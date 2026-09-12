# PDF → 格子（pdf_grid）の番人 ── 2026-09-12（設計 docs/DESIGN-20260912-PDFの帳票を読む.md）
#
# ★ 格子を作る所が製品の正直さを全部背負う（甘い格子 → 自信のある誤値・実測で誤報 14）。
#   だから Excel 側と同じ厳しさで縛る: 陽性・陰性対照と変異で赤くなる形。
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest

from ailine_core import pdf_grid
from ailine_core.field_record import grade, value
from ailine_core.form_read import read_pdf_book

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "forms" / "あかね商事_2026-08.pdf"


# ── ② 部首だけを正規化する ────────────────────────────────
def test_kangxi_radicals_become_the_real_kanji():
    """★ 実物の日本語 PDF は `請求⽇` の「⽇」が U+2F47（見た目は同じ・別の文字）。"""
    assert pdf_grid.normalize_radicals("請求⽇") == "請求日"
    assert pdf_grid.normalize_radicals("⼩計") == "小計"
    assert pdf_grid.normalize_radicals("請求日") == "請求日"


def test_full_width_digits_are_left_alone_because_they_prove_the_cell_was_text():
    """★★ NFKC を文字列全体にかけない ── `１，３２０，０００` は Excel なら**文字**のセル。
    半角に化かすと「金額の欄に文字が入っている」という拒否ができなくなる。"""
    s = "１，３２０，０００"
    assert pdf_grid.normalize_radicals(s) == s
    assert pdf_grid.recover(s) == (s, "")


# ── ③ 描画の逆算 ────────────────────────────────────────────
@pytest.mark.parametrize("text, want", [
    ("¥32,400-", (32400, '¥#,##0-')),
    ("330,000円", (330000, '#,##0"円"')),
    ("1,234", (1234, "#,##0")),
    ("1", (1, "#,##0")),                     # ★ 裸の数も戻す（数量の列が空に見えないように）
    ("12.5", (12.5, "#,##0")),
])
def test_drawn_numbers_are_recovered_with_their_format(text, want):
    assert pdf_grid.recover(text) == want


@pytest.mark.parametrize("text", ["000-0000", "00-0000-0000", "S-202504-001", "1個",
                                  "2025年4月30日", "T1234567890123", "-"])
def test_things_that_only_look_numeric_stay_text(text):
    assert pdf_grid.recover(text) == (text, "")


# ── ① 行と列 ────────────────────────────────────────────────
def _w(text, x0, x1, top):
    return {"text": text, "x0": x0, "x1": x1, "top": top}


def test_rows_by_y_and_columns_by_the_widest_gaps():
    """実測の形（Adobe の明細）: 見出しと明細の x はずれるが、同じ列に落ちること。"""
    words = [_w("品名", 144, 164, 302), _w("単価", 309, 329, 302), _w("金額", 498, 518, 302),
             _w("ロゴデザイン制作", 110, 190, 332), _w("100,000円", 280, 326, 332),
             _w("100,000円", 457, 503, 332)]
    g = pdf_grid.grid_from_words(words, page_no=2)
    head = {c.value: c.col for c in g.all_cells() if c.row == 1}
    body = {c.col: c.value for c in g.all_cells() if c.row == 2}
    assert head == {"品名": 1, "単価": 2, "金額": 3}
    assert body == {1: "ロゴデザイン制作", 2: 100000, 3: 100000}
    assert g.cell(2, 2).at == "2頁2行2列"          # ★ 人が原本を指で追える番地


def test_two_words_that_land_on_one_cell_are_not_glued():
    g = pdf_grid.grid_from_words([_w("a", 10, 20, 100), _w("b", 12, 22, 100)])
    assert sorted(c.value for c in g.all_cells()) == ["a", "b"]


def test_an_empty_page_is_an_empty_grid():
    assert pdf_grid.grid_from_words([]).all_cells() == []


# ── 実物（自分の内容を LibreOffice で PDF にした検体）────────────
def test_the_fixture_reads_all_five_fields_without_touching_the_rules():
    recs = read_pdf_book(FIXTURE)
    got = {f: value(r) for f, r in recs.items()}
    assert got == {"宛先": "ナギ商会株式会社", "請求元": "株式会社あかね商事", "請求額": 33000,
                   "請求日": dt.date(2026, 8, 31), "請求番号": "INV-2026-08-777"}, got


def test_spaces_inside_a_cell_are_kept_so_the_band_is_found():
    """★★ 「合　　計」の空白は同じセルの証拠 ── 捨てると帯の合計が永久に見つからない。"""
    import pdfplumber
    with pdfplumber.open(str(FIXTURE)) as pdf:
        kept = pdf_grid.grid_from_page(pdf.pages[0])
        dropped = pdf_grid.grid_from_words(pdf.pages[0].extract_words(), page_no=1)
    def band_labels(g):       # 空白を落として「合計」そのものになるマス（『合計金額』は別）
        return {"".join(str(c.value).split()) for c in g.all_cells()} & {"合計", "小計"}
    assert band_labels(kept) == {"合計", "小計"}, band_labels(kept)
    assert band_labels(dropped) == set(), {str(c.value) for c in dropped.all_cells()}
    assert "合" in {str(c.value) for c in dropped.all_cells()}      # ★ 割れて 1 文字になっている


def test_a_pdf_from_the_product_path_never_claims_corroboration():
    """★ D4: 式が無いので「裏が取れた」は出ない。値は出る。"""
    rec = read_pdf_book(FIXTURE)["請求額"]
    assert value(rec) == 33000
    assert grade(rec) == "単", grade(rec)
    assert rec.unconfirmable and any("式" in x for x in rec.unconfirmable)


# ── D7: テキスト層が無い ──────────────────────────────────
def empty_page_pdf(path: Path) -> Path:
    """文字の無い 1 頁の PDF（スキャン相当）。★ 依存なしで書ける最小の形。"""
    body = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\n"
            b"trailer<</Root 1 0 R>>\n%%EOF\n")
    path.write_bytes(body)
    return path


def test_a_pdf_without_a_text_layer_is_named_not_silently_empty(tmp_path):
    p = empty_page_pdf(tmp_path / "スキャン.pdf")
    with pytest.raises(pdf_grid.NoTextLayer) as ei:
        read_pdf_book(p)
    assert "テキスト層" in str(ei.value) and "OCR" in str(ei.value)


# ── 行は垂直の中心で束ねる（2026-09-12・実物 construction_bill）──────────
def test_a_big_amount_and_its_small_label_share_a_row():
    """★ `¥ -`（高さ 24pt）は同じ行のラベル（高さ 14pt）より top が 5pt 上に出る。
    top で束ねると別の行に割れ、ラベルの右が空に見える。中心は 0.1pt しか違わない。"""
    words = [{"text": "ご請求金額", "x0": 133, "x1": 203, "top": 364.4, "bottom": 378.2},
             {"text": "¥   -", "x0": 283, "x1": 545, "top": 359.5, "bottom": 383.3}]
    g = pdf_grid.grid_from_words(words)
    assert g.rows == 1, [(c.row, c.value) for c in g.all_cells()]
    assert g.cell(1, 1).value == "ご請求金額" and g.cell(1, 2).value == 0


def test_an_accounting_zero_is_recovered_but_a_bare_dash_is_not():
    """★ `¥   -` は Excel の会計書式が 0 を描いた形。裸の `-` は「該当なし」の文字のまま。"""
    assert pdf_grid.recover("¥   -") == (0, '¥#,##0;;"-"')
    assert pdf_grid.recover("￥ -") == (0, '¥#,##0;;"-"')
    assert pdf_grid.recover("-") == ("-", "")
    assert pdf_grid.recover("¥-1") == ("¥-1", "")


# ── 見出しはどの行にも属さない（2026-09-12・実測の誤報から）──────────
def test_a_tall_title_spanning_two_rows_does_not_join_either():
    """★★ 実測の誤報: `請求書番号：` の右が**帳票のタイトル**になり、
    `請求番号 = 請　求　書` を出した（Wondershare の別レイアウト）。

    タイトル（高さ 43.7pt・57.4〜101.1）は小さい 2 行にまたがり、中心（79.3）が
    `請求書番号：`（中心 79.6）と 0.3pt しか違わない。
    ★ 高さの閾値では分けられない ── `¥ -` はラベルの行に入るべきだから。
      違いは構造: **またぐ行が 1 つなら同じ行・2 つ以上なら見出し**。
    """
    words = [
        {"text": "発行日：", "x0": 91, "x1": 131, "top": 59.7, "bottom": 69.7},
        {"text": "請　求　書", "x0": 298, "x1": 488, "top": 57.4, "bottom": 101.1},
        {"text": "請求書番号：", "x0": 71, "x1": 131, "top": 74.6, "bottom": 84.5},
    ]
    g = pdf_grid.grid_from_words(words)
    at = {str(c.value): (c.row, c.col) for c in g.all_cells()}
    title_row = at["請　求　書"][0]
    assert title_row != at["請求書番号："][0], at      # ★ 番号のラベルと同じ行にしない
    assert title_row != at["発行日："][0], at
    # ★ どの小さい行からも右に見えない（right_of で拾われない）
    label = g.cell(*at["請求書番号："])
    assert [c.value for c in g.right_of(label, span=6)] == [], at


def test_a_tall_amount_spanning_one_row_still_joins_its_label():
    """★ 陰性対照 ── またぐ行が 1 つなら同じ行のまま（上の処置で壊さない）。"""
    words = [{"text": "ご請求金額", "x0": 133, "x1": 203, "top": 364.4, "bottom": 378.2},
             {"text": "¥   -", "x0": 283, "x1": 545, "top": 359.5, "bottom": 383.3}]
    g = pdf_grid.grid_from_words(words)
    assert g.cell(1, 1).value == "ご請求金額" and g.cell(1, 2).value == 0


def test_the_grid_declares_that_its_columns_were_reconstructed():
    """★★ 格子が自分で申告する ── 表を歩く側（明細の掃き出し）はこれを見て降りる。

    ★ 実測: PDF で出た「明細と帯が合いません」39 件のうち **34 件が偽**（Excel では
      裏が取れて正しい値）。「空欄は誤値より安い」は「偽の疑いも安い」を意味しない。
    """
    g = pdf_grid.grid_from_words([{"text": "a", "x0": 1, "x1": 2, "top": 1, "bottom": 2}])
    assert g.columns_reconstructed is True
    import openpyxl
    from ailine_core.form_grid import Grid
    wb = openpyxl.Workbook()
    wb.active["A1"] = "a"
    assert Grid.read(wb.active).columns_reconstructed is False


def test_a_pdf_does_not_cry_wolf_about_the_detail_table():
    """★ 列を組み直した格子では明細の掃き出しをしない（偽の食い違いを出さない）。"""
    from ailine_core.form_read import detail_amount_sum
    rec = read_pdf_book(FIXTURE)["請求額"]
    assert value(rec) == 33000 and grade(rec) == "単"
    assert any("組み直している" in x for x in rec.unconfirmable), rec.unconfirmable
    import pdfplumber
    with pdfplumber.open(str(FIXTURE)) as pdf:
        g = pdf_grid.grid_from_page(pdf.pages[0])
    assert detail_amount_sum(g, None) == (None, None, (), None, 0)


# ── ソフトが生成した PDF（2026-09-12・Namakoo 提供・社名は合成）─────────────
SOFTWARE_PDF = Path(__file__).resolve().parent / "fixtures" / "forms" / "ソフト発行_請求書.pdf"


def test_a_software_generated_invoice_pdf_yields_all_five_fields():
    """★★ 朝は「PDF が読めるかも分からない」所から始まり、ここまで来た記録。

    この検体は **プログラムが生成した PDF**（フォントが Liberation-Sans ＋ Noto-Sans-CJK-JP で
    Producer が空 ── Excel で作って印刷したものではない）。会計ソフトの出力に近い作りで、
    中身は記入済み・社名は合成（example.com）なので repo に入れてよい。
    ★ 2 頁あるが請求書らしい印は 1 頁目だけ 3 つ・2 頁目 1 つ ── 自分で 1 頁目を選ぶ。
    """
    import datetime as dt2
    recs = read_pdf_book(SOFTWARE_PDF)
    got = {f: value(r) for f, r in recs.items()}
    assert got == {"宛先": "株式会社ミライ商事",
                   "請求元": "株式会社テックイノベーション",
                   "請求額": 407000,
                   "請求日": dt2.date(2026, 9, 12),
                   "請求番号": "INV-202609-001"}, got
    # ★ PDF なので裏取りは名乗らない（D4）── 値は出す
    assert grade(recs["請求額"]) == "単"
    assert any("式" in x for x in recs["請求額"].unconfirmable)


def test_the_detail_header_of_that_pdf_is_not_found_and_that_is_recorded():
    """★★ 見つからない ── 見出しが `品名・明細` で、一覧の語と**完全一致しない**。

    ★ 「一覧の語で始まる」に緩める案は測って**却下**した（2026-09-12）:
      緩めると 18 冊（検体 11・束 7）で**見出し行が別の場所へ移る**（正しい D21/V21 を捨てて
      A19/O19 を拾う）。得はこの 1 通だけで、PDF は掃き出しを降りているので点数は 1 も動かない。
      ★ セル数（『合計』で始まるセル 389 件）は指標にならなかった ── 見出し行が動くかを測って決めた。
    ★ 発火条件: `品名・明細` 型の見出しが**別の出所で 2 件目**出た時、または PDF の列が直って
      掃き出しを再開する時（そこで初めて得が測れる）。
    """
    from ailine_core.form_read import detail_amount_sum, invoice_signals
    grids = [g for _i, g in pdf_grid.grids_of(SOFTWARE_PDF) if len(invoice_signals(g)) >= 2]
    assert len(grids) == 1
    assert detail_amount_sum(grids[0], None)[3] is None


CARRIED_PDF = Path(__file__).resolve().parent / "fixtures" / "forms" / "ソフト発行_繰越あり.pdf"


def test_corporate_ligatures_are_unfolded_for_matching_but_kept_in_the_value():
    """★★ 実物（ソフト生成）が `㈱ミライ商事 御中` と書いていて、**宛先も請求元も取れなかった**。

    `㈱`(U+3231) は NFKC が `(株)` に、`㍿`(U+337F) が `株式会社` にほどき、どちらも法人格の
    一覧に既に在る。★ 全体に NFKC をかけない ── 全角の数字が半角に化けて
    「金額の欄が文字」の拒否が消える（`_LIGATURES` は閉じた文字クラスだけ）。
    ★ 値は**原本の文字そのまま**（`㈱ミライ商事`）── 照合だけを均す。
    """
    from ailine_core.form_read import norm
    assert norm("㈱アルファ") == "(株)アルファ"
    assert norm("㍿テック") == "株式会社テック"
    assert norm("１，３２０，０００") == "１，３２０，０００"      # ★ 全角の数字は寄せない
    recs = read_pdf_book(CARRIED_PDF)
    assert value(recs["宛先"]) == "㈱ミライ商事", value(recs["宛先"])
    assert value(recs["請求元"]) == "㈱テックコンサルティング", value(recs["請求元"])


def test_a_carried_balance_invoice_refuses_to_pick_an_amount():
    """★ 繰越のある請求書は「合計」と「今回お支払いいただく額」が違う ── 決めない（設計どおり）。"""
    rec = read_pdf_book(CARRIED_PDF)["請求額"]
    assert grade(rec) == "割" and value(rec) is None
    assert "繰越" in rec.blank_reason, rec.blank_reason
