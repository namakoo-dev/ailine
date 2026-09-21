# -*- coding: utf-8 -*-
"""列名の内側にしか出てこないシート名を、言及と数えない（2026-09-21・盲検 5 体目）。

★★ 出所（買い手の言葉）:

> **間違った ⚠ で判定が ✓ → △ に落ちました。**「『発注』は存在しません」と言いますが、
> 発注シートは在ります。

  再現した依頼（`bench/blind/5体目` の凍結セット・実物の冊で確認）:

      請求シートで請求金額から発注金額を引いた差額の列を作って

      シート … 請求 / 発注 / 商品名・発注金額だけ
      列   … 請求番号 / 仕入先 / 商品名 / 請求金額 / **発注金額** / 差額 / 確認 …

  『発注』はこの文中で **列名『発注金額』の内側にしか現れない**。それを言及と数え、
  「変わっていない」と誤警報を出し、正しくできた仕事の ✓ が △ に落ちた。

★★ 根は 2 つ、どちらも「知っているのに使っていない」形:

  ① **覆う語彙に列名が入っていなかった**。畳む器官（`drop_names_covered_by_longer`・
     名前の**位置**で見る）は 08-24 から在って、判定規則も正しい。
     ★ 器官は在ったが、**語彙が届いていなかった** ── 規則は一切変えず、覆う候補に
       実在する列名を足すだけで解ける。

  ② 文面が「**存在しません／変更されていません**」と、まったく違う 2 つの事実を
     1 つに畳んでいた。機械は原本のシート一覧を持っているので**どちらか知っている**。
     ★ 直しは「シートでは存在しないは起こりえない」という**推論に寄りかからない** ──
       実際に見て言い分ける。前提が将来崩れても嘘にならない側に置く。

★ 陰性対照は実測で入った過去の事故から取る（緩めても厳しくしても、どれかが赤くなる）。
"""
from __future__ import annotations

import ast
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

import ailine  # noqa: E402
from _product_source import product_files  # noqa: E402
from ailine_core.target_sheet import (  # noqa: E402
    drop_names_covered_by_longer, sheet_names_mentioned_in)

BUYER_TASK = "請求シートで請求金額から発注金額を引いた差額の列を作って"
BUYER_SHEETS = ["請求", "発注"]
BUYER_HEADERS = {"請求": ["請求番号", "商品名", "請求金額", "発注金額", "差額", "確認"],
                 "発注": ["発注番号", "商品名", "数量", "単価", "発注金額"]}


def _mentioned(task, sheets, headers=None):
    meta = {"sheets": sheets, "headers": headers or {}}
    got = ailine.extract_task_mentions(task, sheets, ailine._header_names_of(meta))
    return sorted(got["sheets"])


# --- 事故そのもの -------------------------------------------------------------------

def test_a_sheet_named_only_inside_a_column_name_is_not_a_mention():
    """★★ 事故そのもの ── 『発注』は『発注金額』の内側にしかない。"""
    assert _mentioned(BUYER_TASK, BUYER_SHEETS, BUYER_HEADERS) == ["請求"]


def test_without_the_column_vocabulary_the_false_mention_comes_back():
    """★★ 対照 ── **効いているのは列名を渡したから**であることを凍結する。

    ★ これが無いと「たまたま消えた」のか「直したから消えた」のか分からない。
    """
    assert _mentioned(BUYER_TASK, BUYER_SHEETS, None) == ["発注", "請求"]


def test_the_false_warning_is_gone_end_to_end(tmp_path):
    """★★ 画面まで通して、『発注』の ★ が消えていること（＋対照つき）。"""
    p = tmp_path / "突合.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求"
    ws.append(BUYER_HEADERS["請求"])
    ws.append(["A-1", "ねじ", 100, 90, None, None])
    wb.create_sheet("発注").append(BUYER_HEADERS["発注"])
    wb.save(p)
    before = ailine.snapshot(p)

    wb = openpyxl.load_workbook(p)
    wb["請求"]["E2"] = 10           # 差額の列だけ埋める（請求シートは変わる）
    wb.save(p)
    after = ailine.snapshot(p)

    meta = {"sheets": BUYER_SHEETS, "headers": BUYER_HEADERS}
    got = ailine.build_advisories(BUYER_TASK, before, after, None, meta=meta)
    assert not [ln for ln in got if "発注" in ln and ln.startswith("★")], got

    # ★ 対照: 列名を知らせなければ（meta 無し）、従来どおり誤警報が出る。
    plain = ailine.build_advisories(BUYER_TASK, before, after, None)
    assert [ln for ln in plain if "『発注』" in ln], plain


# --- 文面: 在るものを「存在しません」と言わない --------------------------------------

def test_a_sheet_that_exists_is_said_to_be_unchanged(tmp_path):
    """★★ 在って変わらなかったシートは「**変更されていません**」と言うこと。"""
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "台帳"
    wb.active.append(["取引先", "金額"])
    wb.create_sheet("参照").append(["x", 1])
    wb.save(p)
    before = ailine.snapshot(p)
    wb = openpyxl.load_workbook(p)
    wb["台帳"]["A2"] = "a"
    wb.save(p)
    after = ailine.snapshot(p)

    got = ailine.mention_overlap_advisory(
        {"cols": set(), "digit_cols": set(), "rows": set(), "sheets": {"参照"}},
        before, after)
    assert got == ["★ 依頼で言及された『参照』は変更されていません"], got
    assert not any("存在しません" in ln for ln in got), (
        "在るシートに『存在しません』と言っている ── 買い手が ✓ を疑った所")


def test_a_name_that_is_really_absent_is_still_said_to_be_absent(tmp_path):
    """★★ 本当に無い名前には「存在しません」と言えること。

    ★ 「シートの枝では起こりえない」という**推論に寄りかからない** ── 実際に見て
      言い分けているなら、無い側も正しく言えるはず。言えなければ片側しか直っていない。
    """
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "台帳"
    wb.active.append(["取引先", "金額"])
    wb.save(p)
    before = ailine.snapshot(p)
    wb = openpyxl.load_workbook(p)
    wb["台帳"]["A2"] = "a"
    wb.save(p)
    after = ailine.snapshot(p)

    got = ailine.mention_overlap_advisory(
        {"cols": set(), "digit_cols": set(), "rows": set(), "sheets": {"幻のシート"}},
        before, after)
    assert got == ["★ 依頼で言及された『幻のシート』は存在しません"], got


# --- 陰性対照: 実測で入った過去の事故 -------------------------------------------------

@pytest.mark.parametrize("why, task, sheets, headers, want", [
    ("08-24 独立して書かれた短い名前は消さない",
     "売上シートと売上60以上シートを見比べて", ["売上", "売上60以上"],
     {"売上": ["金額", "売上"]}, ["売上", "売上60以上"]),
    ("08-24 長い名前に畳まれる方は消す",
     "売上60以上シートを集計して", ["売上", "売上60以上"],
     {"売上": ["金額"]}, ["売上60以上"]),
    ("シートを本当に名指した回は残す",
     "発注シートを見て発注金額を確認して", ["請求", "発注"],
     BUYER_HEADERS, ["発注"]),
])
def test_the_specimens_that_earned_the_rule_do_not_move(why, task, sheets, headers, want):
    """★★ 列名を足したせいで過去の事故が再発しないこと。

    ★ 上 2 件は 08-24 に**実測で**入った線（集合で畳むと独立の短い名前まで消える／
      畳まないと誤警報が出る）。どちらも位置で見ることでしか両立しない。
    """
    assert _mentioned(task, sheets, headers) == want, why


def test_a_sheet_name_that_is_also_a_column_name_still_counts():
    """★★ シート名と**同じ**列名は覆いにしない（長さが同じものは覆えない）。

    ★ ここを緩めると「金額シートを更新して」の『金額』が、列『金額』に食われて消える。
    """
    got = _mentioned("金額シートを更新して", ["売上データ", "金額"],
                     {"売上データ": ["取引先", "金額"]})
    assert got == ["金額"], got


# --- 配線: 同じ取り出しを 2 箇所に書かない --------------------------------------------

def test_every_caller_passes_the_column_vocabulary():
    """★★ 呼び出し側の**全部**が列名を渡していること ── 分母は AST から導く。

    ★ 片方だけ渡すと、片方の経路にだけ誤警報が残る（この repo が何度も踏んだ形）。
      一覧を手で書くと、呼び出しを足した日に静かにずれる。
    """
    # ★ 本体を**パスで決め打ちしない**（2026-09-21・台帳 test_guard_ledger が止めた）。
    #   分割で実装が動いたとき、パス読みの番人は黙って空振りする。
    calls = []
    for f in product_files():
        for n in ast.walk(ast.parse(f.read_bytes().decode("utf-8"))):
            if not isinstance(n, ast.Call):
                continue
            name = (n.func.id if isinstance(n.func, ast.Name)
                    else getattr(n.func, "attr", None))
            if name == "extract_task_mentions":
                calls.append((f.name, n.lineno, len(n.args) + len(n.keywords)))
    assert calls, "呼び出しが 1 つも見つからない（名前が変わった？）"
    thin = [(f, ln) for f, ln, k in calls if k < 3]
    assert not thin, f"列名を渡していない呼び出しが {len(thin)} 箇所: {thin}"


def test_the_column_vocabulary_comes_from_one_place():
    """★ 見出しの取り出しは `_header_names_of` 1 本（呼び出し側に書き写さない）。"""
    assert ailine._header_names_of({"headers": {"a": ["x", "y"], "b": ["y", "z"]}}) == (
        "x", "y", "z")
    assert ailine._header_names_of(None) == ()
    assert ailine._header_names_of({}) == ()


def test_the_judging_rule_itself_was_not_touched():
    """★★ 直したのは**語彙**であって**規則**ではないこと。

    ★ 位置で見る器官に、同じ入力を与えれば同じ答えが返る ── ここが動いていたら、
      それは「覆う候補を足した」ではなく「畳み方を変えた」になっている。
    """
    task = "売上シートと売上60以上シートを見比べて"
    names = sheet_names_mentioned_in(task, ["売上", "売上60以上"])
    assert drop_names_covered_by_longer(task, names) == ["売上", "売上60以上"]
