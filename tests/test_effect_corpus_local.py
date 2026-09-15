# -*- coding: utf-8 -*-
"""効果の検体 ── 「この依頼を出したら、ブックはこうなっているべき」を手で宣言して測る。

★★ なぜ在るか（2026-09-15・盲検の買い手役が初手で 2 件踏んだ後に道具の棚を数えた）:

    翻訳の battery 81      入口だけ（op と slot）          → ①③ は翻訳が正しかったので通る
    単体 4,500             部品の契約                       → 退行しか止めない
    ゴールデン 752         verify_dsl_args の入出力         → 連鎖も事後条件も範囲外
    効果の行列             飾りの生存（図形・VBA・書式）    → 値の話ではない
    実物 87 冊             帳票の**読み**の答え             → run の経路ではない

  **誰も「適用後のブックがどうなっているべきか」を宣言していなかった。**
  ①（集計の後の並べ替えが明細に落ちて ✓）も ③（自分が作った計算列で絞れない）も、
  翻訳は正しく部品も正しく、**組み上がった結果だけ**が違っていた ── その隙間で生きていた。

★ 設計の線（`bench/effect_corpus.json` の _about と対）:
  1. **翻訳は凍結する**（plan を検体が持つ）── これは LLM を測る道具ではない。
     今日のバグは全部 LLM の下流だった。固定すれば揺れが消え、走らせるのが安くなる
  2. **答えは手で書く**。道具の出力から作らない（恒真になる）
  3. **変わっていないことも宣言する**（unchanged）── 消えたものは差分に出ない
  4. 採点器は自分を試す（下の 3 本の自己検査は実機を使わない ── 常に走る）
"""
import json
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_golden_transcripts import _isolate, _run_main  # noqa: E402

CORPUS = json.loads((REPO / "bench" / "effect_corpus.json").read_text(encoding="utf-8"))
CASES = CORPUS["cases"]


def _build(book_spec: dict, path: Path) -> Path:
    """検体のブックを作る（宣言どおりに ── ここで答えを作らない）。"""
    wb = openpyxl.Workbook()
    first = True
    for name, spec in book_spec["sheets"].items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        for row in spec["rows"]:
            ws.append(list(row))
        for addr, formula in (spec.get("formulas") or {}).items():
            ws[addr] = formula
    wb.save(path)
    return path


def _snapshot(path: Path) -> dict:
    """シート名 → {『A1』: 値}（式のセルは式文字列のまま ── 「変わっていない」の判定用）。"""
    wb = openpyxl.load_workbook(path)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = {c.coordinate: c.value for row in ws.iter_rows() for c in row
                     if c.value is not None}
    return out


def compare(expect: dict, after_path: Path, before: dict, rc: int) -> list:
    """宣言と実物を突き合わせて、破れの一覧を返す（空なら合格）。

    ★ 採点器はここ 1 本。下の自己検査がこの関数を壊して赤を確かめる。
    """
    breaks = []
    wb = openpyxl.load_workbook(after_path, data_only=True)
    after = _snapshot(after_path)
    if "exit" in expect and rc != expect["exit"]:
        breaks.append(f"終了コードが {rc}（宣言は {expect['exit']}）")
    if "sheets" in expect and wb.sheetnames != list(expect["sheets"]):
        breaks.append(f"シートの顔ぶれが {wb.sheetnames}（宣言は {expect['sheets']}）")
    for sheet in expect.get("unchanged") or []:
        if sheet not in after:
            breaks.append(f"『{sheet}』が消えている（1 セルも変わらないはずだった）")
        elif after[sheet] != before.get(sheet):
            diff = [k for k in set(after[sheet]) | set(before.get(sheet, {}))
                    if after[sheet].get(k) != before.get(sheet, {}).get(k)]
            breaks.append(f"『{sheet}』が変わっている（{len(diff)} セル: {sorted(diff)[:6]}）")
    for addr, want in (expect.get("cells") or {}).items():
        sheet, cell = addr.split("!", 1)
        if sheet not in wb.sheetnames:
            breaks.append(f"『{sheet}』が無い（{addr} を確かめられない）")
            continue
        got = wb[sheet][cell].value
        if got != want:
            breaks.append(f"{addr} が {got!r}（宣言は {want!r}）")
    return breaks


@pytest.mark.local
@pytest.mark.parametrize("case", CASES, ids=[c["id"] for c in CASES])
def test_the_book_looks_as_declared(case, tmp_path, monkeypatch, capsys):
    """★ 実機（LibreOffice）で適用し、**手で宣言した姿**と突き合わせる。"""
    book = _build(CORPUS["books"][case["book"]], tmp_path / f"{case['id']}.xlsx")
    before = _snapshot(book)
    _isolate(monkeypatch, tmp_path)
    monkeypatch.setattr(ailine, "translate_task",
                        lambda model, task, book_meta, temperature=0.1: {"plan": case["plan"]})
    rc, out = _run_main(["run", str(book), case["task"], "--overwrite"], capsys)
    breaks = compare(case["expect"], book, before, rc)
    assert not breaks, (f"{case['id']}: {case['why']}\n  破れ: " + "\n        ".join(breaks)
                        + f"\n--- 画面 ---\n{out}")


# ── 採点器の自己検査（実機を使わない ── 常に走る）────────────────────────────────

def _tiny(tmp_path):
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["a", 1])
    wb.save(p)
    return p


def test_the_scorer_calls_a_correct_book_correct(tmp_path):
    """★ 陽性対照 ── 宣言どおりのブックに破れを出さない。"""
    p = _tiny(tmp_path)
    before = _snapshot(p)
    assert compare({"sheets": ["S"], "unchanged": ["S"], "cells": {"S!A1": "a"}, "exit": 0},
                   p, before, 0) == []


@pytest.mark.parametrize("expect,why", [
    ({"cells": {"S!A1": "ちがう"}}, "値の違い"),
    ({"cells": {"S!Z9": "何か"}}, "空のセルに値を宣言"),
    ({"sheets": ["S", "無い"]}, "シートの顔ぶれ"),
    ({"exit": 3}, "終了コード"),
    ({"cells": {"無い!A1": "x"}}, "無いシートの宣言"),
])
def test_the_scorer_can_fail(expect, why, tmp_path):
    """★ 壊した宣言で必ず破れが出ること ── 出ない検査は在っても鳴らない。"""
    p = _tiny(tmp_path)
    assert compare(expect, p, _snapshot(p), 0), why


def test_the_scorer_sees_a_sheet_that_should_not_have_moved(tmp_path):
    """★ 負の被覆: 「1 セルも変わらない」はずのシートが変わったら鳴ること。"""
    p = _tiny(tmp_path)
    before = _snapshot(p)
    wb = openpyxl.load_workbook(p)
    wb["S"]["B1"] = 999
    wb.save(p)
    breaks = compare({"unchanged": ["S"]}, p, before, 0)
    assert breaks and "変わっている" in breaks[0], breaks


def test_every_case_declares_why_and_a_hand_answer():
    """★ 検体は「なぜ在るか」と**手で出した答え**を必ず持つ（道具の出力を写した検体を入れない）。"""
    bad = [c["id"] for c in CASES
           if not c.get("why") or not c.get("_hand_math") or not (c.get("expect") or {}).get("cells")]
    assert not bad, f"why / _hand_math / expect.cells の無い検体: {bad}"
    assert len(CASES) >= 4, "検体が減っている"
