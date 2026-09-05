"""式が値に潰されたら言うこと（2026-09-06）。

★★ 出所（2026-09-05 に測って確定した穴）: 実物の請求書には式が 29 個あった
  （`=IF(I17="","",G17*I17)` が 21 個 ほか）。そこへ「金額を全部 0 にして」と頼むと
  **式が消えて、事後条件は pass** する ── 宣言（列を 0 にする）は真だからだ。
  だが表としては壊れている: 以後 数量 を直しても金額が動かない。

  ★ `check_set_column_value` の**署名に `source_book` が無い**（適用前を一度も見ない）
    ので、式の消失は**原理的に見えなかった**。入口側から数えたら、既存セルに値を書く
    op は 5 本、うち適用前を見るのは 3 本、**「潰れた式」を見ているのは 0 本**だった。

★ 2026-08-31 の事故（`row_identity` を生んだ「2 セルだけ入れ替えた」は真だが表が矛盾）と
  **同じ形**。処置も同じ ── **直さない・言う**。

★★ 何を数えるかを実測で決めた（総数では区別できない）:
      値で潰す   式 3/3 → 0/3   ★ 言う
      行を削除   式 3/3 → 2/2   黙る（割合は変わらない）
      行を追加   式 3/3 → 4/4   黙る（`FillFormulasFromNeighbour` が式を運ぶ）
  **列ごとの「式だったセルの割合」**で見れば、行の増減・並べ替えに影響されない。
  op の宣言を見ないので、**op が増えても配線が要らない**。
  ★ 8/27 の実測「式の文字列は行が動けば変わるのが正しい」に従い、式の中身は比べない。

★★ 実物に当てて 1 つ直した: 初版は**見出しを 1 行目と決め打ち**していたが、
  実物の請求書は**見出しが 16 行目**で 1 行目は空 ── 全列が無名になって**黙っていた**。
  合成検体では鳴っていたので、実物に当てるまで気づけなかった（検体は治具まで含めて仮説）。
"""
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from _product_source import product_text  # noqa: E402 ── ★ 本体は場所でなく芯から読む
from ailine_core.formula_health import (  # noqa: E402
    before_after_advisories, formula_loss_advisory)

HEADERS = ["品名", "数量", "単価", "金額"]
ROWS3 = [("ボルト", 10, 50), ("ナット", 5, 20), ("ワッシャ", 8, 15)]


def _book(path, rows, *, amount=None, header_row=1, formula=True):
    """金額列を式（既定）か、指定の値で埋めた表を作る。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    for _ in range(header_row - 1):
        ws.append([])
    ws.append(HEADERS)
    for i, (n, q, u) in enumerate(rows, start=header_row + 1):
        ws.append([n, q, u])
        ws.cell(i, 4).value = (f"=B{i}*C{i}" if formula and amount is None else amount)
    wb.save(path)
    return path


# --- ① 潰したら言う ---------------------------------------------------------

def test_it_speaks_when_formulas_are_replaced_by_values(tmp_path):
    before = _book(tmp_path / "b.xlsx", ROWS3)
    after = _book(tmp_path / "a.xlsx", ROWS3, amount=0, formula=False)
    got = formula_loss_advisory(before, after)
    assert got, "式が全部消えたのに黙っている"
    assert "金額" in got[0] and "3/3" in got[0] and "0/3" in got[0], got[0]
    assert "追随しません" in got[0] and "直していません" in got[0], got[0]


def test_it_speaks_even_when_only_some_are_lost(tmp_path):
    """★ 一部だけ潰れた回も言う（割合が下がっている）。"""
    before = _book(tmp_path / "b.xlsx", ROWS3)
    after = _book(tmp_path / "a.xlsx", ROWS3)
    wb = openpyxl.load_workbook(after)
    wb["明細"].cell(3, 4).value = 999          # 1 セルだけ直値に
    wb.save(after)
    wb.close()
    got = formula_loss_advisory(before, after)
    assert got and "2/3" in got[0], got


# --- ② 鳴らない側（★ 狼少年にしない ── 対で縛る）---------------------------

@pytest.mark.parametrize("label, rows, kw", [
    ("行を 1 本削除", ROWS3[:2], {}),
    ("行を 1 本追加", ROWS3 + [("ねじ", 3, 40)], {}),
    ("何もしない", ROWS3, {}),
])
def test_it_stays_quiet_when_nothing_was_lost(tmp_path, label, rows, kw):
    before = _book(tmp_path / "b.xlsx", ROWS3)
    after = _book(tmp_path / "a.xlsx", rows, **kw)
    assert formula_loss_advisory(before, after) == [], label


def test_it_stays_quiet_when_there_were_no_formulas(tmp_path):
    """★ 元から式が無い表（実務のほとんど）で鳴らないこと。"""
    before = _book(tmp_path / "b.xlsx", ROWS3, amount=100, formula=False)
    after = _book(tmp_path / "a.xlsx", ROWS3, amount=0, formula=False)
    assert formula_loss_advisory(before, after) == []


def test_it_stays_quiet_when_the_column_is_gone(tmp_path):
    """★ 列ごと消えた回は黙る（頼まれたとおり ── DELETE_COLUMN）。"""
    before = _book(tmp_path / "b.xlsx", ROWS3)
    after = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    ws.append(HEADERS[:3])
    for n, q, u in ROWS3:
        ws.append([n, q, u])
    wb.save(after)
    wb.close()
    assert formula_loss_advisory(before, after) == []


def test_an_unreadable_book_is_silent(tmp_path):
    """★ 読めない回は黙る（断定しない）。"""
    before = _book(tmp_path / "b.xlsx", ROWS3)
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not a workbook")
    assert formula_loss_advisory(before, broken) == []
    assert formula_loss_advisory(broken, before) == []


# --- ③ 見出しが 1 行目でない表（★ 実物で見つけた穴）------------------------

def test_it_finds_the_column_when_the_header_is_not_the_first_row(tmp_path):
    """★★ 実物の請求書は見出しが 16 行目。1 行目決め打ちだと**全列が無名**になり黙った。"""
    before = _book(tmp_path / "b.xlsx", ROWS3, header_row=16)
    after = _book(tmp_path / "a.xlsx", ROWS3, amount=0, formula=False, header_row=16)
    assert formula_loss_advisory(before, after) == [], "見出し行を渡さなければ黙る（既定は 1 行目）"
    got = formula_loss_advisory(before, after, {"明細": 16})
    assert got and "金額" in got[0], got


@pytest.mark.parametrize("path", [Path("C:/Dev/_fixtures/misoca_invoice_blackline.xlsx")])
def test_it_speaks_on_the_real_invoice(tmp_path, path):
    """★ 実物（他社の請求書・repo の外）で鳴ること。無ければ skip する。"""
    if not path.is_file():
        pytest.skip("実物の検体が手元に無い（C:/Dev/_fixtures）")
    before = tmp_path / "b.xlsx"
    after = tmp_path / "a.xlsx"
    shutil.copy2(path, before)
    shutil.copy2(path, after)
    wb = openpyxl.load_workbook(after)
    ws = wb["misoca_invoice"]
    killed = 0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 11).value
        if isinstance(v, str) and v.startswith("="):
            ws.cell(r, 11).value = 0
            killed += 1
    wb.save(after)
    wb.close()
    assert killed >= 20, killed
    got = formula_loss_advisory(before, after, {"misoca_invoice": 16})
    assert got and "金額" in got[0], got


# --- ④ 配線（★ 純関数だけで満足しない）--------------------------------------

def test_the_advisory_comes_out_of_the_single_door(tmp_path):
    """★ 入口（before_after_advisories）を通しても出ること。"""
    before = _book(tmp_path / "b.xlsx", ROWS3)
    after = _book(tmp_path / "a.xlsx", ROWS3, amount=0, formula=False)
    got = before_after_advisories(before, after, {"col": "金額"},
                                  cell_ref=lambda r, c: f"R{r}C{c}")
    assert any("式で計算されていました" in g for g in got), got


def test_the_callers_pass_the_header_row():
    """★ 見出し行が呼び出し側から届いていること（実物で唯一効いた引数）。

    ★ 本体を**場所で決め打ちしない**（`tests/test_guard_ledger.py` の掟）── 昨日
      500 行を関数ごと動かした repo なので、パス直読みの番人は分割で空振りする。
    """
    src = product_text()
    wired = src.count("header_rows=(book_meta") + src.count("header_rows=(current_meta")
    assert wired >= 4, (
        f"見出し行を渡している経路が {wired} 箇所（op が決まる 4 経路のはず）── "
        "自由生成の段は book_meta を持たないので渡さない")
