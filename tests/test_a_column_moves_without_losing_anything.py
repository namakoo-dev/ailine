# -*- coding: utf-8 -*-
"""既にある列を動かす（2026-09-08 に足した操作）── その番人。

★★ 出所（盲検の検品 #18）: 「取引先の列を**一番左に**持ってきて」が断られ、しかも
  理由が「列『取引先』がありません」という**誤った診断**だった（実際は別シートの話を
  していた）。実務で毎日やる操作なので、断りの文言を直すのではなく**できるようにした**。

★ 部品は既に在った ── `MoveColumnTo` は「新しい列を作ってから動かす」ために実戦投入
  済み。足りなかったのは**既存の列を動かす op** だけ。列を動かす実装は 1 本のまま。

★ ここで縛るのは 4 つ:
    ① 位置の言い回し（一番左／末尾／誰かの右）が**実表の見出し**で解けること
    ② ヘルパへ渡す番号の**換算**（右へ動かす回は最終位置より 1 大きい）
    ③ 事後条件が「その 1 列だけが動いた」を、宣言でなく**適用前の並びから**確かめること
    ④ 動かす必要が無い回・列が無い回は**断る**こと（黙って何かしない）
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

import ailine
from ailine_core.postconditions.move import check_move_column

HEADERS = ["取引先", "項目", "件数", "単価", "金額", "税込金額", "締め日", "担当"]


# --- ① 位置の言い回し ---------------------------------------------------------

@pytest.mark.parametrize("task, want_at", [
    ("担当の列を一番左に持ってきて", 1),
    ("担当を先頭に移して", 1),
    ("単価の列を末尾に移動して", len(HEADERS) + 1),
    ("単価を一番右へ", len(HEADERS) + 1),
    # ★★ 実測した嘘の診断: 正規表現が**いちばん早い開始位置**から伸びるので
    #   『締め日を金額』を丸ごと掴み、「そんな列はありません」と言っていた。
    ("締め日を金額の右に移して", 6),
    ("担当を件数の左に移して", 3),
])
def test_the_position_words_are_resolved_against_the_real_headers(task, want_at):
    at, note = ailine.resolve_col_anchor(task, HEADERS)
    assert at == want_at, (at, note)


def test_a_position_that_is_not_written_is_not_invented():
    """★ 位置の言い回しが無ければ**決めない**（黙って端へ寄せない）。"""
    assert ailine.resolve_col_anchor("担当の列を動かして", HEADERS) == (None, None)


# --- ② 目的地とヘルパ引数の換算 -----------------------------------------------

def _resolve(task, col):
    bm = {"sheets": ["表"], "headers": {"表": HEADERS}}
    resolved = {"col": col, "_target_sheet": "表"}
    got = ailine._verify_move_column(resolved, set(), bm, task)
    return got, resolved


@pytest.mark.parametrize("task, col, want_from, want_to", [
    ("担当の列を一番左に持ってきて", "担当", 7, 0),
    ("単価の列を末尾に移動して", "単価", 3, 7),
    ("締め日を金額の右に移して", "締め日", 6, 5),
])
def test_the_final_position_is_what_we_promise(task, col, want_from, want_to):
    got, resolved = _resolve(task, col)
    assert got is None, got
    assert (resolved["_move_from"], resolved["_move_to"]) == (want_from, want_to)


@pytest.mark.parametrize("task, col, want", [
    # ★ 右へ動かす回だけ +1（途中で 1 本広げた並びでの位置を渡すヘルパのため）
    ("単価の列を末尾に移動して", "単価", (3, 8)),
    ("担当の列を一番左に持ってきて", "担当", (7, 0)),
    ("締め日を金額の右に移して", "締め日", (6, 5)),
])
def test_the_number_handed_to_the_helper_is_converted(task, col, want):
    """★ 列を動かす `Call` を出す口は repo に 1 つだけ（横断層の wrap）なので、
       ここで縛るのは**その口へ渡す材料**。実際に出た Basic は下の実機試験が見る。"""
    _got, resolved = _resolve(task, col)
    assert (resolved["_new_col_from"], resolved["_move_new_col_to"]) == want, resolved


# --- ④ 断る回 -----------------------------------------------------------------

@pytest.mark.parametrize("task, col, why", [
    ("取引先の列を一番左に持ってきて", "取引先", "もうそこに在ります"),
    ("部門の列を一番左に持ってきて", "部門", "ありません"),
    ("担当の列を動かして", "担当", "読み取れません"),
])
def test_it_refuses_instead_of_moving_something_else(task, col, why):
    got, _resolved = _resolve(task, col)
    assert got is not None and got[0] is False, got
    assert why in got[3], got[3]


# --- ③ 事後条件（適用前の並びから独立に組み直して比べる）----------------------

def _book(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_the_check_is_built_from_the_before_file_not_from_the_declaration(tmp_path):
    before = tmp_path / "b.xlsx"
    after = tmp_path / "a.xlsx"
    _book(before, ["A", "B", "C"], [[1, 2, 3], [4, 5, 6]])
    _book(after, ["C", "A", "B"], [[3, 1, 2], [6, 4, 5]])
    st, note = check_move_column(after, {"col": "C", "_move_to": 0}, 1, before)
    assert st == "pass", note

    # ★ 別の列まで動いていたら通さない
    _book(after, ["C", "B", "A"], [[3, 2, 1], [6, 5, 4]])
    st, note = check_move_column(after, {"col": "C", "_move_to": 0}, 1, before)
    assert st == "fail", note


def test_a_value_that_quietly_changed_is_caught(tmp_path):
    before = tmp_path / "b2.xlsx"
    after = tmp_path / "a2.xlsx"
    _book(before, ["A", "B", "C"], [[1, 2, 3], [4, 5, 6]])
    _book(after, ["C", "A", "B"], [[3, 1, 2], [6, 4, 99]])   # ★ 5 が 99 に化けた
    st, note = check_move_column(after, {"col": "C", "_move_to": 0}, 1, before)
    assert st == "fail", note


# --- ⑤ 実機（本物の LibreOffice で、式が付いてくることまで見る）----------------

REPO = Path(__file__).resolve().parent.parent


@pytest.mark.local
def test_it_really_moves_on_real_libreoffice(tmp_path):
    """★ 実機 ── 列が動き、**式の参照が付いてくる**。

    ★ 値の書き写しで実装すると、並びは正しく見えるのに式が別の列を指す
      （`=E2*1.1` が金額でなく単価を指す等）。だから式の中身まで読む。
    """
    book = tmp_path / "seikyu.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求"
    ws.append(["取引先", "件数", "単価", "金額", "税込金額", "担当"])
    for i, (name, cnt, tanka) in enumerate(
            [("あかね商事", 3, 1000), ("いろは工業", 5, 2000)], start=2):
        ws.append([name, cnt, tanka, f"=B{i}*C{i}", f"=D{i}*1.1", "田中"])
    wb.save(book)

    r = subprocess.run(
        [sys.executable, "-m", "ailine", "run", str(book),
         "担当の列を一番左に持ってきて", "--copy", "--timeout", "300"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
        env={**os.environ, "PYTHONPATH": str(REPO / "src")})
    assert r.returncode == 0, r.stdout[-900:]
    assert "✓" in r.stdout, r.stdout[-900:]

    out = openpyxl.load_workbook(book.with_name(book.stem + ".out.xlsx"))["請求"]
    assert [out.cell(1, c).value for c in range(1, 7)] == [
        "担当", "取引先", "件数", "単価", "金額", "税込金額"]
    for row in (2, 3):
        # ★ 列が 1 本右へずれたので、式の参照も 1 つ右を指しているはず
        assert out.cell(row, 5).value == f"=C{row}*D{row}", out.cell(row, 5).value
        assert out.cell(row, 6).value == f"=E{row}*1.1", out.cell(row, 6).value
    # ★ 値も落ちていない（キャッシュ値で読み直す）
    vals = openpyxl.load_workbook(
        book.with_name(book.stem + ".out.xlsx"), data_only=True)["請求"]
    assert [vals.cell(r_, 5).value for r_ in (2, 3)] == [3000, 10000]


# --- ⑥ 語彙が増えると意味の近い op に奪われる（2026-09-08・Namakoo が名指しした現象）--

def test_a_position_phrase_alone_does_not_make_it_a_column_move():
    """★★ 実測した奪い合い: 列移動を語彙に入れた途端、「鈴木の右に東棟」（1 セル書換）が
      **列移動**に読まれ、効果の行列で 5 件が ✓ → ？ に落ちた。

    ★ 直しは語彙でなく**実表**。行き先が実在の列に接地しないなら列の話ではないので、
      その計画は「別の仕事だから読み直すな」と言う資格を失う（宣言 needs_col_anchor）。
    ★ ここでは判定の材料（行き先が解けるか）だけを縛る ── 読み直しが実際に働くことは
      下の実機試験が見る。
    """
    heads = ["氏名", "所属", "内線", "メモ"]
    # ★ 『鈴木』は行の値であって列ではない → 位置は解けない
    assert ailine.resolve_col_anchor("鈴木の右に東棟", heads)[0] is None
    # ★ 対の試験: 実在の列を指した回はちゃんと解ける（門を閉じすぎない）
    assert ailine.resolve_col_anchor("メモを所属の右に移して", heads)[0] == 3


def test_the_declaration_says_which_op_needs_a_column_anchor():
    """★ 門は op 名でなく**宣言**を読む（op が増えても配線が要らない）。"""
    assert ailine.OP_WRITE_TARGET["MOVE_COLUMN"].needs_col_anchor is True
    assert ailine.OP_WRITE_TARGET["SWAP"].needs_col_anchor is False


@pytest.mark.local
def test_a_cell_write_is_not_stolen_by_the_new_column_move(tmp_path):
    """★ 実機 ── 「鈴木の右に東棟」が 1 セル書換として通ること（奪われないこと）。"""
    book = tmp_path / "meibo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名簿"
    ws.append(["氏名", "所属", "内線", "メモ"])
    for r in [["山田", "営業", 101, None], ["鈴木", "経理", 202, None],
              ["高橋", "総務", 305, None]]:
        ws.append(r)
    wb.save(book)
    r = subprocess.run(
        [sys.executable, "-m", "ailine", "run", str(book), "鈴木の右に東棟",
         "--copy", "--sheet", "名簿", "--timeout", "300"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
        env={**os.environ, "PYTHONPATH": str(REPO / "src")})
    assert r.returncode == 0, r.stdout[-700:]
    out = openpyxl.load_workbook(book.with_name(book.stem + ".out.xlsx"))["名簿"]
    assert [out.cell(3, c).value for c in range(1, 4)] == ["鈴木", "東棟", 202], (
        r.stdout[-700:])
