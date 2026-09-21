# 塊③(2/2)・中核 op 致命2（2026-08-24 の盲検）── 転記が参照表の「B 列」を決め打ち。
#
# ★ 実測: マスタ = 商品 / 区分 / **単価**（C 列）に対して「単価を転記して」と頼むと
#     変更点: C2: (空)→'果物'   C3: (空)→'果物'   C4: (空)→'高級'
#     ✓ 機械検証済みの内容です
#   **単価の列に「果物/果物/高級」が入った。** 数値であるべき列に文字列が入って ✓。
#
# ★ 根: 書き手（helpers の VLookupFromTable）が `oLook.getCellByPosition(1, j)`＝
#   「参照表 列1=値」を決め打ちし、検算（check_lookup_fill）も同じく列1・列2 決め打ちで
#   期待値を作る。**やる側と見る側が同じ思い込みを共有している**ので必ず一致する ── 恒真。
#
# ★★ 2026-09-21（盲検 5 体目）で**処方が一段進んだ**:
#   初版の契約は「2 列目でなければ**開示して ✓ を降ろす**」だった。断らなかったのは、
#   列の位置だけでは事故（マスタ=[商品,区分,単価]）と正しい依頼（明細=[商品,数量,単価]）を
#   区別できないから ── これは今も正しい。
#   だが買い手（小売店の店長）の発注記録は **6 列**で、転記はそもそも**使えなかった**:
#     「転記は参照表が 2 列ちょうどでないと動きません」
#   ★★ そして位置を割り出す式は、**開示のために既に書かれていた** ── 正しい列を
#     計算しておきながら、それで書きに行かず「意図と違うかもしれません」と言うだけだった。
#     **器官は在るが配線が無い。**
#
# 契約（2026-09-21 改定）:
#   ① 参照表の値の列は**名前で**決める（3 列以上のマスタでも正しく転記する）
#   ② 名前で引けない時だけ、何が書かれるかを名指しして開示する（従来の線を残す）
#   ③ 2 列マスタ（キー・値）は従来どおり通る（誤爆しない）
#   ④ 見出しが読めない参照表は従来どおり（黙って通す・断りの根拠が無い）
#   ★ ⑤ 書き手と検算は**別の道**で同じ名前に辿り着くこと（book_meta / ファイルの見出し）
#     ── 同じ決め打ちを共有したら、また恒真になる。

import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402
from ailine_core.postconditions.shape import check_lookup_fill  # noqa: E402


def _meta(sheets, headers):
    return {"sheets": sheets, "headers": headers,
            "header_rows": {s: 1 for s in sheets}}


THREE_COL = _meta(["明細", "マスタ"],
                   {"明細": ["商品", "単価"], "マスタ": ["商品", "区分", "単価"]})
ARGS = {"target_sheet": "明細", "target_col": "単価",
        "key_col": "商品", "source_sheet": "マスタ"}


def _call_line(resolved, meta):
    code = ailine.codegen_dsl("LOOKUP_FILL", dict(resolved), meta, use_formula=False)
    hits = [ln for ln in code.splitlines() if "VLookupFromTable" in ln]
    assert len(hits) == 1, code
    return hits[0].rstrip()


def test_a_three_column_master_is_read_by_name():
    """①★ 事故そのもの ── 単価が 3 列目でも、**単価を読んで**書くこと。

    ★ 生成された Basic が参照表の**3 列目（0 起点で 2）**を値として渡していることを見る。
      「警告が消えた」だけを見ると、直したのか黙らせたのか区別できない。
    """
    ok, resolved, _inf, err = ailine.verify_dsl_args(
        "LOOKUP_FILL", dict(ARGS), THREE_COL, task="マスタから単価を転記して")
    assert ok, err
    line = _call_line(resolved, THREE_COL)
    # 引数: (oDoc, headerRow, keyCol, resultCol, "マスタ", lookupKeyCol, lookupValueCol)
    assert line.endswith('"マスタ", 0, 2)'), (
        "参照表の『単価』（3 列目＝0 起点で 2）を渡していない: " + line)
    assert not (resolved.get("_warnings") or []), (
        "正しく読めるのに警告を出している（オオカミ少年になる）")


def test_the_check_reads_the_same_column_by_a_different_road(tmp_path):
    """★★⑤ 検算が**ファイルの見出し**から同じ列に辿り着くこと（恒真にしない）。

    ★ 書き手は `book_meta` の見出しから、検算は**冊そのもの**から引く。
      位置の決め打ちを共有していた頃は、間違って書いても必ず一致していた。
    """
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    ws.append(["商品", "単価"])
    ws.append(["りんご", 120])          # ★ 正しい転記結果（マスタの『単価』）
    m = wb.create_sheet("マスタ")
    m.append(["商品", "区分", "単価"])
    m.append(["りんご", "果物", 120])
    wb.save(p)
    assert check_lookup_fill(p, dict(ARGS))[0] == "pass"

    # ★ 陰性対照: 『区分』の値（旧実装が書いていたもの）が入っていたら落ちること。
    ws.cell(row=2, column=2).value = "果物"
    wb.save(p)
    status, reason = check_lookup_fill(p, dict(ARGS))
    assert status == "fail", (status, reason)


def test_a_name_the_master_does_not_have_is_still_disclosed():
    """②★ 名前で引けない時は従来どおり開示する（線を残す）。

    ★ ここを消すと「見つからないので 2 列目に落ちた」回が無言になる。
    """
    meta = _meta(["明細", "マスタ"],
                 {"明細": ["商品", "原価"], "マスタ": ["商品", "区分", "単価"]})
    ok, resolved, _inf, err = ailine.verify_dsl_args(
        "LOOKUP_FILL",
        {"target_sheet": "明細", "target_col": "原価",
         "key_col": "商品", "source_sheet": "マスタ"},
        meta, task="マスタから原価を転記して")
    assert ok, f"開示で足りるのに断った: {err}"
    text = " ".join(resolved.get("_warnings") or [])
    assert "区分" in text, f"何が書かれるかを名指ししていない: {text}"


def test_two_column_master_still_works():
    """③ 誤爆しない: キー・値の 2 列なら従来どおり。"""
    meta = _meta(["明細", "マスタ"],
                 {"明細": ["商品", "単価"], "マスタ": ["商品", "単価"]})
    ok, resolved, _inf, err = ailine.verify_dsl_args(
        "LOOKUP_FILL",
        {"target_sheet": "明細", "target_col": "単価",
         "key_col": "商品", "source_sheet": "マスタ"},
        meta, task="マスタから単価を転記して")
    assert ok, f"正しい 2 列マスタを落とした: {err}"
    assert _call_line(resolved, meta).endswith('"マスタ", 0, 1)')


def test_unknown_master_headers_are_not_refused():
    """④ 見出しが読めない参照表は断らない（根拠が無い時に止めない）。

    ★ 位置も決められないので、従来どおり 0/1 に落ちる。
    """
    meta = _meta(["明細", "マスタ"], {"明細": ["商品", "単価"]})
    ok, resolved, _inf, err = ailine.verify_dsl_args(
        "LOOKUP_FILL", dict(ARGS), meta, task="マスタから単価を転記して")
    assert ok, f"根拠が無いのに断った: {err}"
    assert _call_line(resolved, meta).endswith('"マスタ", 0, 1)')


def test_the_helper_takes_the_columns_as_arguments():
    """★ ヘルパ側が受け口を持っていること（**省略可能**＝既存の呼び方を壊さない）。

    ★ カタログは 5 引数の呼び方を模型に見せている。自由生成がその形で書いても
      従来どおり動くこと（既定 0/1）が、この直しの前提。
    """
    bas = (REPO / "src" / "ailine" / "helpers" / "AiLineHelpers.bas").read_text(encoding="utf-8")
    sig = [ln for ln in bas.splitlines() if ln.startswith("Sub VLookupFromTable")]
    assert len(sig) == 1, sig
    assert "Optional lookupKeyCol" in sig[0] and "Optional lookupValueCol" in sig[0], sig[0]
    body = bas.split("Sub VLookupFromTable")[1].split(chr(10) + "End Sub")[0]
    assert "getCellByPosition(0, j)" not in body, "キーの列を決め打ちしたまま"
    assert "getCellByPosition(1, j)" not in body, "値の列を決め打ちしたまま"


# --- 実機（LibreOffice が要る） -------------------------------------------------------

@pytest.mark.local
@pytest.mark.parametrize("why, call", [
    ("カタログが模型に見せている従来の呼び方（引数 5 つ）",
     'Call VLookupFromTable(oDoc, 0, 0, 1, "単価表")'),
    ("新しい呼び方（参照表の列を渡す）",
     'Call VLookupFromTable(oDoc, 0, 0, 1, "単価表", 0, 1)'),
])
def test_both_call_shapes_run_on_the_real_machine(tmp_path, why, call):
    """★★ `Optional` が実機で効くこと ── **走った**ことを値で確かめる。

    ★ Basic は予約語との衝突や引数の食い違いで**モジュールごと黙って死ぬ**。
      「適用した」は「走った」ではない ── 書き込まれた値まで見る。
    ★ カタログは模型に**引数 5 つ**の呼び方を見せている。新しい引数を足したせいで
      そちらが死んだら、自由生成の経路が黙って壊れる。
    ★★ この検体は治具の誤りを 1 度経由している（2026-09-21）: 最初は
      `Sub DoWork` + `ThisComponent` で組んで「5 引数は失敗する」と出た。
      本番の包みは `Sub Run(oDoc As Object)` ── **測定器を先に疑う**。
    """
    src = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    for r in [["商品", "単価"], ["りんご", None], ["バナナ", None]]:
        ws.append(r)
    m = wb.create_sheet("単価表")
    for r in [["商品", "単価"], ["りんご", 120], ["バナナ", 80]]:
        m.append(r)
    wb.save(src)
    work = tmp_path / "w"
    work.mkdir()
    _cat, helpers = ailine.load_helpers(ailine.DEFAULT_HELPERS)
    code = ("Option VBASupport 1" + chr(10) + "Option Explicit" + chr(10) * 2
            + "Sub Run(oDoc As Object)" + chr(10) + "    " + call + chr(10) + "End Sub" + chr(10))
    ok, err, _raw = ailine.basrun_apply(src, code, work, helper_files=helpers, timeout=180)
    assert ok, f"{why}: 走らなかった ── {(err or '')[-200:]}"
    got = [c.value for c in openpyxl.load_workbook(src)["明細"]["B"]][1:]
    assert got == [120, 80], f"{why}: 走ったが値が入っていない {got}"
