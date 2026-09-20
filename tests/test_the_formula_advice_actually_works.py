# -*- coding: utf-8 -*-
"""照合の断りが言う「開いて保存すれば値が入る」が、**本当に通る**こと（2026-09-20）。

★★ 出所（盲検 3 体目・製造業の購買）: 照合が「列が決まりません」としか言わず、買い手は
  既に列名を書いていた。本当の理由は『金額』が**式のままで計算結果を持たない**ことで、
  道具はそれを持っていたのに言っていなかった。2026-09-17 に理由の 1 行を足した:

    ★『金額』は**式のままで計算結果が入っていない**ため、金額の列として使えません
      （Excel か LibreOffice で一度開いて保存すると値が入ります）。

★★ この試験が見るのは、その**助言が本当に効くか**。合格線の 3 条目は
  「通る道を示し、**その道が通ることを機械で歩いて確かめてある**」── 文言が在るだけでは
  条を満たさない。この repo が何度も踏んだ「示した道が通らない（path_fails）」を、
  ここで実機ごと潰す。

★ 導通の盤（A 群）はここを歩けない ── 道そのものが LibreOffice を要求するため。
  盤は `実機で歩く` と判定してこの試験を**名指し**する（見ていないと別の所で見たを混ぜない）。
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

import ailine  # noqa: E402


def _books(d: Path) -> tuple:
    """A は金額が**式のまま**（計算結果なし）、B は値つき。"""
    a = d / "A.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["品名", "金額"])
    ws.append(["ボルト", "=100*2"])
    ws.append(["ナット", "=40*2"])
    wb.save(a)
    wb.close()
    b = d / "B.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["品名", "金額"])
    ws.append(["ボルト", 150])
    ws.append(["ワッシャー", 300])
    wb.save(b)
    wb.close()
    return a, b


def _run(a: Path, b: Path, task: str, home: Path):
    return subprocess.run(
        [sys.executable, "-m", "ailine", "run", str(a), str(b), task],
        cwd=str(REPO), capture_output=True, text=True, encoding="utf-8",
        errors="replace", timeout=900,
        env={**os.environ, "PYTHONPATH": str(REPO / "src"), "AILINE_HOME": str(home)})


def _open_and_save(book: Path, out_dir: Path) -> Path:
    """助言どおり「LibreOffice で一度開いて保存」する（人がやることを機械で再現）。

    ★ soffice の探し方は製品と同じ（basrun の office_dir）── 探し方を 2 通り持たない。
    """
    from ailine import _find_basrun_path, _load_module_from_path
    basrun = _find_basrun_path()
    assert basrun is not None, "basrun が見つからない"
    mod = _load_module_from_path(basrun, "_ailine_basrun_resave")
    office = Path(mod.office_dir())
    soffice = office / ("soffice.exe" if os.name == "nt" else "soffice")
    assert soffice.exists(), f"soffice が無い: {soffice}"
    out_dir.mkdir(parents=True, exist_ok=True)
    r = subprocess.run([str(soffice), "--headless", "--norestore", "--convert-to", "xlsx",
                        "--outdir", str(out_dir), str(book)],
                       capture_output=True, text=True, timeout=300,
                       encoding="utf-8", errors="replace")
    saved = out_dir / book.name
    assert saved.exists(), f"開いて保存できていない: {r.stdout[-300:]} / {r.stderr[-300:]}"
    return saved


@pytest.mark.local
def test_the_refusal_names_the_formula_column(tmp_path):
    """★ 引き金 ── 式のままの列を**名指しで**断ること（理由を言わない旧版に戻さない）。"""
    a, b = _books(tmp_path)
    r = _run(a, b, "金額で突き合わせて", tmp_path / "home")
    assert r.returncode == 3, r.stdout[-500:]
    assert "式のままで計算結果が入っていない" in r.stdout, r.stdout[-500:]
    assert "金額" in r.stdout, r.stdout[-500:]


@pytest.mark.local
def test_opening_and_saving_really_makes_the_column_usable(tmp_path):
    """★★ 本体 ── 示した道を**歩いたら着く**こと。

    ★ 助言は「Excel か LibreOffice で一度開いて保存すると値が入ります」。
      そのとおりに開いて保存し、同じ依頼をもう一度打って**通る**ことを確かめる。
    ★ ここが赤くなったら、買い手に**通らない道**を案内している（path_fails・最悪の形）。
    """
    a, b = _books(tmp_path)
    first = _run(a, b, "金額で突き合わせて", tmp_path / "home")
    assert first.returncode == 3, first.stdout[-400:]

    saved = _open_and_save(a, tmp_path / "resaved")
    # ★ 測定器を疑う: 本当に値が入ったかを、製品とは別の目（openpyxl）で確かめる
    ws = openpyxl.load_workbook(saved, data_only=True)["表"]
    got = [ws.cell(i, 2).value for i in (2, 3)]
    assert got == [200, 80], f"★ 開いて保存しても値が入っていない: {got}"

    again = _run(saved, b, "金額で突き合わせて", tmp_path / "home")
    assert again.returncode == 0, (
        "★★ 断りが示した道が通らない（path_fails）:\n" + again.stdout[-600:])


@pytest.mark.local
def test_naming_the_column_resolves_the_other_match_refusal(tmp_path):
    """★ もう 1 つの照合の断り ── 「候補: …。依頼文に列名を含めて」が通ること。

    ★ こちらは LLM も LibreOffice も要らないので導通の盤（A 群）も歩いている。
      ここは**買い手と同じ打ち方**（別プロセス・素の argv）で二重に確かめる。
    """
    a = tmp_path / "A.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["品名", "備考", "金額"])
    ws.append(["ボルト", "至急", 120])
    ws.append(["ナット", "", 80])
    wb.save(a)
    wb.close()
    b = tmp_path / "B.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["品名", "備考", "金額"])
    ws.append(["ボルト", "", 150])
    ws.append(["ワッシャー", "", 300])
    wb.save(b)
    wb.close()

    first = _run(a, b, "突き合わせて", tmp_path / "home")
    assert first.returncode == 3, first.stdout[-400:]
    assert "候補" in first.stdout, first.stdout[-400:]

    again = _run(a, b, "品名をキーに突き合わせて", tmp_path / "home")
    assert again.returncode == 0, (
        "★★ 断りが示した道が通らない（path_fails）:\n" + again.stdout[-600:])
