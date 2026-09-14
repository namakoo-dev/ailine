# -*- coding: utf-8 -*-
"""`ailine split <book> --by <見出し> --out <フォルダ>` の検体 ── 需要⑤の出口（2026-09-12）。

契約（設計 docs/DESIGN-20260912-担当者別に分けて配る.md）:
  - 値ごとに 1 冊（見出し行 ＋ その人の行 ＋ 出所列）＋ `_検分.xlsx`
  - 書き手の印（creator）と条件（description）が焼かれている ── 次に自分の出力と分かる
  - 合計・小計の行はどの冊にも入らない／空欄・表記ゆれ・複数担当は名指しされる
  - 担当者の見出しが 2 列に当たる冊は分けない（exit 4・1 冊も作らない）
  - 配る先に人のファイルがあれば exit 7（stack / forms と同じ関所）
  - ★★ 証明（D5）は**書いた冊を読み戻して**数える ── 読み戻しを差し替えたら exit 5 になり、
    1 冊も置かれない（＝自分の記憶で数えていないことの実証）
  - 原本は 1 バイトも変えない

★ LLM も LibreOffice も要らない（読むだけ・openpyxl のみ）。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

from ailine_core import split_people   # noqa: E402
from ailine_core import stack as stack_core   # noqa: E402

_HEADERS = ["日付", "案件", "顧客", "担当者", "金額"]
_ROWS = [
    ["2026-06-01", "広告", "甲社", "山田", 12000],
    ["2026-06-02", "印刷", "乙社", "佐藤", 8000],
    ["2026-06-03", "保守", "丙社", "山田", 25000],
    ["2026-06-04", "広告", "丁社", "山田　", 18000],   # ★ 表記ゆれ（全角空白）
    ["2026-06-05", "保守", "戊社", None, 7000],        # ★ 空欄
    ["2026-06-06", "印刷", "己社", "山田/佐藤", 5000],  # ★ 複数担当
    ["合計", None, None, None, 75000],                  # ★ 分けない行
]


def _book(path: Path, rows=None, headers=None, title=("2026年6月度 売上一覧",)):
    """見出しが 1 行目でない一覧表（実物の形）。★ 式は使わない（LO を起動しない）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "売上"
    for text in title:
        ws.append([text])
    ws.append(list(headers or _HEADERS))
    for row in (rows if rows is not None else _ROWS):
        ws.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _split(book: Path, out: Path, *extra, by="担当者"):
    return subprocess.run(
        [sys.executable, "-m", "ailine", "split", str(book), "--by", by,
         "--out", str(out), *extra],
        capture_output=True, text=True, timeout=180, encoding="utf-8",
        errors="replace", cwd=str(REPO))


def _values(path: Path) -> list:
    wb = openpyxl.load_workbook(path)
    got = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    return got


def _marks(path: Path) -> tuple:
    wb = openpyxl.load_workbook(path)
    got = (wb.properties.creator, wb.properties.description)
    wb.close()
    return got


@pytest.fixture()
def book(tmp_path):
    return _book(tmp_path / "元" / "売上一覧.xlsx")


def test_one_book_per_person_with_the_provenance_columns(book, tmp_path):
    """★ 値ごとに 1 冊。見出し ＋ その人の行 ＋ 出所列（元ファイル・元行）。"""
    out = tmp_path / "配る"
    r = _split(book, out, "--amount", "金額")
    assert r.returncode == 0, f"落ちた: {r.stdout}\n{r.stderr}"
    made = sorted(p.name for p in out.glob("*.xlsx"))
    assert made == sorted(["山田.xlsx", "山田_.xlsx", "佐藤.xlsx", "_検分.xlsx"]), made
    head, *rows = _values(out / "山田.xlsx")
    assert head == _HEADERS + list(stack_core.PROVENANCE_HEADERS), head
    # ★ 末尾は本人の合計行（2026-09-14 に足した ── 出所が空・中身は下の専用の番人が見る）。
    *rows, last = rows
    assert split_people.is_own_total(last[-1], list(last)), last
    assert [row[-2:] for row in rows] == [["売上一覧.xlsx", 3], ["売上一覧.xlsx", 5]], rows
    assert [row[3] for row in rows] == ["山田", "山田"]


def test_the_writer_mark_and_the_condition_are_burned_in(book, tmp_path):
    """★ 印（creator）と条件（description）── 次に「自分の出力か」を機械で言えるように。"""
    out = tmp_path / "配る"
    assert _split(book, out).returncode == 0
    creator, description = _marks(out / "山田.xlsx")
    assert creator == split_people.CREATOR_MARK
    cond = json.loads(description)
    assert cond == {"tool": "ailine", "kind": "split", "by": "担当者", "value": "山田"}
    assert stack_core.own_output_mark(out / "山田.xlsx") == split_people.CREATOR_MARK
    assert stack_core.own_output_mark(out / "_検分.xlsx") == split_people.CREATOR_MARK, \
        "★ 検分の冊が自分の出力と認識されない（次回の関所が人のファイルと誤認する）"


def test_the_original_is_not_touched(book, tmp_path):
    before = hashlib.sha256(book.read_bytes()).hexdigest()
    _split(book, tmp_path / "配る", "--amount", "金額")
    assert hashlib.sha256(book.read_bytes()).hexdigest() == before, "★ 原本が変わった"


def test_a_total_row_never_lands_in_anyones_book(book, tmp_path):
    """★★ 最悪の混入（D2）── **元の表の**合計行が誰かの冊に入っていないこと。

    ★★ 2026-09-14: 冊の末尾に本人の合計行を足した日、この番人は「合計」の語だけを見ていたので
      自分が足した行に噛んだ。見る物を**出所**に変えた ── 元から来た行（元行が数）に合計の語が
      在ってはいけない、が本当の契約（自分の合計行は出所が空で、金額は 75,000 ではない）。
    """
    out = tmp_path / "配る"
    assert _split(book, out, "--amount", "金額").returncode == 0
    for path in out.glob("*.xlsx"):
        if path.name == "_検分.xlsx":
            continue
        for row in _values(path)[1:]:
            if split_people.is_own_total(row[-1], list(row)):
                assert row[-2] is None and 75000 not in row, f"{path.name}: {row}"
                continue
            assert "合計" not in [str(v) for v in row], f"{path.name} に合計行が混入: {row}"
    assert not (out / "合計.xlsx").exists()


def test_the_report_names_the_blank_the_lookalike_and_the_two_people(book, tmp_path):
    r = _split(book, tmp_path / "配る", "--amount", "金額")
    assert "空欄" in r.stdout and "表記ゆれ" in r.stdout and "複数担当" in r.stdout, r.stdout
    assert "✓ 部分の和 ＝ 全体" in r.stdout, f"証明を出していない: {r.stdout}"


def test_json_carries_the_whole_contract(book, tmp_path):
    """★ 機械可読の側にも同じ事実（採点器と自動化がここを読む）。"""
    r = _split(book, tmp_path / "配る", "--amount", "金額", "--json")
    payload = json.loads(r.stdout.strip().splitlines()[-1])
    for key in ("parts", "blank", "multi", "excluded", "lookalike", "refused",
                "proof", "unparsed"):
        assert key in payload, f"--json に {key} が無い: {sorted(payload)}"
    assert payload["refused"] is None
    assert payload["blank"] == [7] and [m[0] for m in payload["multi"]] == [8]
    assert payload["excluded"] == [9]
    assert sorted(sorted(p) for p in payload["lookalike"]) == [sorted(["山田", "山田　"])]
    rows = payload["proof"]["rows"]
    assert rows["whole"] == rows["parts"] + rows["blank"] + rows["excluded"] + rows["multi"]
    assert payload["proof"]["ok"] is True
    placed = sorted(r0 for part in payload["parts"].values() for r0 in part["rows"])
    assert placed == [3, 4, 5, 6], placed
    assert payload["parts"]["山田"]["amount"] == 37000


def test_the_inspection_book_carries_the_proof(book, tmp_path):
    out = tmp_path / "配る"
    assert _split(book, out, "--amount", "金額").returncode == 0
    head, *rows = _values(out / "_検分.xlsx")
    assert head == list(split_people.REPORT_HEADERS), head
    kinds = {str(row[0]) for row in rows}
    for kind in ("配った", "空欄", "複数担当", "分けない行", "表記ゆれ", "証明（行）",
                 "証明（金額）"):
        assert kind in kinds, f"検分に『{kind}』が無い: {sorted(kinds)}"


def test_two_matching_headers_refuse_and_write_nothing(tmp_path):
    """★★ 変異「見出しが 2 列でも分ける」── exit 4 で 1 冊も作らない（設計 D1）。"""
    book = _book(tmp_path / "元" / "実績.xlsx",
                 headers=["日付", "顧客", "担当", "営業担当", "金額"],
                 rows=[["2026-06-02", "甲社", "内藤", "山田", 30000],
                       ["2026-06-03", "乙社", "内藤", "佐藤", 18000]])
    out = tmp_path / "配る"
    r = _split(book, out, by="担当")
    assert r.returncode == 4, f"分けてしまった: {r.stdout}"
    assert "担当" in r.stdout and "営業担当" in r.stdout, r.stdout
    # ★ 2026-09-13（買い手役・会計）: 「見出しの文字で 1 つに」だけでは従えない ── 次の一手を言う。
    #   『担当』は『担当』の列にそのまま一致するので `--exact` を案内する（同じ文字が 2 列なら元の表を変える）。
    assert "--exact" in r.stdout, r.stdout
    assert not out.exists() or not list(out.glob("*.xlsx")), "1 冊も作らないはず"


def test_a_missing_header_refuses_by_name(book, tmp_path):
    r = _split(book, tmp_path / "配る", by="記入者")
    assert r.returncode == 4
    assert "記入者" in r.stdout and "担当者" in r.stdout, r.stdout


def test_an_existing_human_file_stops_the_write(book, tmp_path):
    """★ 配る先に人のファイルがあれば exit 7（stack / forms と同じ関所）。"""
    out = tmp_path / "配る"
    out.mkdir()
    human = out / "大事な表.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "人が書いた大事な表"
    wb.save(human)
    wb.close()
    before = human.read_bytes()
    r = _split(book, out)
    assert r.returncode == 7, f"人のファイルの隣に黙って配った: {r.stdout}"
    assert human.read_bytes() == before
    assert "--overwrite" in r.stdout, "次の手を言っていない"
    assert not (out / "山田.xlsx").exists(), "止めたのに配っている"
    assert _split(book, out, "--overwrite").returncode == 0, "承知の上なら通ること"


def test_its_own_previous_output_is_not_a_gate(book, tmp_path):
    """★ 前回配った冊の上へは、関所なしで配り直せる（自分の出力だと分かるから）。"""
    out = tmp_path / "配る"
    assert _split(book, out).returncode == 0
    again = _split(book, out)
    assert again.returncode == 0, f"自分の前回出力を人のファイルと誤認した: {again.stdout}"


def test_the_other_sheets_are_disclosed(tmp_path):
    """★ 見たのは 1 枚目だけ ── 同じ冊に隠れている古い表を黙って無視しない。"""
    book = _book(tmp_path / "元" / "記録.xlsx")
    wb = openpyxl.load_workbook(book)
    ws2 = wb.create_sheet("売上_旧")
    ws2.append(_HEADERS)
    ws2.append(["2026-05-01", "広告", "甲社", "山田", 19000])
    ws2.sheet_state = "hidden"
    wb.save(book)
    wb.close()
    r = _split(book, tmp_path / "配る")
    assert "売上_旧" in r.stdout, f"他のシートを開示していない: {r.stdout}"


def test_a_named_sheet_that_is_missing_is_refused_not_replaced(book, tmp_path):
    r = _split(book, tmp_path / "配る", "--sheet", "存在しない")
    assert r.returncode == 4 and "売上" in r.stdout, r.stdout


def test_a_clean_table_raises_no_findings(tmp_path):
    """★ 陰性対照: 何も仕込まない表で所見 0（偽の疑いを立てたら道具の負け）。"""
    book = _book(tmp_path / "元" / "集計.xlsx",
                 rows=[["2026-05-01", "広告", "甲社", "山田", 20000],
                       ["2026-05-02", "印刷", "乙社", "佐藤", 9000],
                       ["2026-05-05", "保守", "丙社", "鈴木", 8500]])
    r = _split(book, tmp_path / "配る", "--amount", "金額", "--json")
    payload = json.loads(r.stdout.strip().splitlines()[-1])
    assert payload["blank"] == [] and payload["lookalike"] == []
    assert payload["multi"] == [] and payload["excluded"] == [] and payload["unparsed"] == []
    assert payload["proof"]["ok"] is True and len(payload["parts"]) == 3


def test_verify_does_not_pretend_the_mark_is_missing(book, tmp_path):
    """★ `ailine verify` は split の出力に「印がありません」と言わない（正直に断る）。

    ★ 2026-09-12（③）: 独立の検算が入ったので、断り文は「まだ実装していません」から
      **通る形の名指し**に変わった ── 1 冊では和が閉じないのでフォルダで受ける。
      詳しい検算そのものの契約は `tests/test_verify_split.py`。
    """
    out = tmp_path / "配る"
    assert _split(book, out).returncode == 0
    r = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out / "山田.xlsx"),
                        str(book.parent)], capture_output=True, text=True, timeout=120,
                       encoding="utf-8", errors="replace", cwd=str(REPO))
    assert "印がありません" not in r.stdout, r.stdout
    assert "ailine verify <出力フォルダ> <元の冊>" in r.stdout, r.stdout


def test_the_proof_is_counted_from_the_written_books_not_from_memory(book, tmp_path,
                                                                     monkeypatch):
    """★★ 変異「証明を出力でなく自分の記憶で数える」がここで赤くなる。

    読み戻し（`xml_readback.read_grid`）の答えを 1 行だけ削って返すよう差し替える。
    記憶で数えているなら差し替えは効かず exit 0 のまま通る ── 出力から数えているなら
    等式が閉じず exit 5 になり、**1 冊も置かれない**。
    """
    import ailine

    real = ailine.xml_readback.read_grid

    def lying_read(path, sheet_name=None):
        data = real(path, sheet_name)
        # ★ 落とすのは最後の**明細**の行（最終列＝元行が数の行）。★★ 2026-09-14: 冊の末尾に
        #   本人の合計行（出所は空）を足した日に「最後の行」が合計行に変わり、この変異が
        #   素通りした ── 治具が測る場所を失っていた（直すのは治具の側）。
        cols = max((c for (_r, c) in data["grid"]), default=0)
        rows = [r for (r, c), v in data["grid"].items()
                if r > 1 and c == cols and isinstance(v, (int, float))]
        if rows:
            drop = max(rows)
            data["grid"] = {k: v for k, v in data["grid"].items() if k[0] != drop}
        return data

    monkeypatch.setattr(ailine.xml_readback, "read_grid", lying_read)
    out = tmp_path / "配る"
    args = argparse.Namespace(book=str(book), by="担当者", out=str(out), amount="金額",
                              sheet=None, overwrite=False, json=False)
    assert ailine.cmd_split(args) == 5, "★ 読み戻しが嘘をついても通った（記憶で数えている）"
    assert not out.exists() or not list(out.glob("*.xlsx")), \
        "★ 証明できていないのに配った"


# --- 配る先の残留と、上書きした冊（2026-09-13・買い手役の初見・会計）-----------------------
#
# ★★ 名前のゆれを直して同じ配る先へ配り直したら、前回の冊が残ったまま画面は ✓ と「配った冊 1 件」。
#   zip で送ると本人は 3 冊（古い切り方 2 冊）を受け取る。split の ✓ は「自分が書いた冊」の
#   保証で「配るフォルダ」の保証ではない ── 残留も上書きも名指しする。


def test_a_stale_book_left_in_the_folder_is_named(book, tmp_path):
    # ★ 前回の自分の冊は、本物の split の出力を別名で置く（印の中身を書き写さない）
    first = tmp_path / "前回"
    assert _split(book, first).returncode == 0
    out = tmp_path / "配る"
    out.mkdir()
    import shutil
    shutil.copy2(next(p for p in first.glob("*.xlsx") if not p.name.startswith("_")),
                 out / "退職者.xlsx")
    r = _split(book, out)
    assert r.returncode == 0, r.stdout
    assert "⚠ 配る先に前回の冊が 1 冊残っています: 退職者.xlsx" in r.stdout, r.stdout
    assert "zip" in r.stdout, r.stdout


def test_overwriting_its_own_previous_books_is_said_out_loud(book, tmp_path):
    out = tmp_path / "配る"
    first = _split(book, out)
    assert first.returncode == 0 and "上書きしました" not in first.stdout, first.stdout
    r = _split(book, out)
    assert r.returncode == 0, r.stdout
    assert "前回の自分の出力" in r.stdout and "上書きしました" in r.stdout, r.stdout
    assert "残っています" not in r.stdout, r.stdout          # ★ 全部書き直したなら残留は無い


def test_the_stale_warning_comes_before_the_check_mark(book, tmp_path):
    """★ 2 回目の買い手役: ✓ を読んで安心する向きに作用する ── 一番危ない ⚠ は ✓ より前に。"""
    first = tmp_path / "前回"
    assert _split(book, first).returncode == 0
    out = tmp_path / "配る"
    out.mkdir()
    import shutil
    shutil.copy2(next(p for p in first.glob("*.xlsx") if not p.name.startswith("_")),
                 out / "退職者.xlsx")
    r = _split(book, out)
    lines = r.stdout.splitlines()
    warn = next(i for i, ln in enumerate(lines) if ln.startswith("⚠ 配る先に前回の冊"))
    ok = next(i for i, ln in enumerate(lines) if ln.startswith("✓"))
    assert warn < ok, r.stdout


# --- 元の表の合計行と明細の和（2026-09-14・会計役の MISSING #3）------------------------------
#
# ★★ 合計行は「分けない行」として除外していたが、その値が明細の和と合っているかを言わなかった。
#   10 月分の一覧に 9 月の合計が持ち越されていても何も出ない ── 静かに壊れる側。


def test_the_original_total_row_is_reconciled_on_screen(book, tmp_path):
    """★ 合っていても言う（締めの根拠に使う 1 行）。"""
    r = _split(book, tmp_path / "配る", "--amount", "金額")
    assert r.returncode == 0, r.stdout
    assert "（元の表の合計行" in r.stdout and "明細の和が一致）" in r.stdout, r.stdout


def test_a_stale_total_row_is_named_with_the_difference(book, tmp_path):
    """★ 元の表の合計が古い（前月の持ち越し）回 ── 差を名指しする。分けた冊は明細のとおり。"""
    wb = openpyxl.load_workbook(book)
    ws = wb.active
    heads = {str(c.value): c.column for c in ws[2]}
    row = [None] * ws.max_column
    row[0] = "合計"
    row[heads["金額"] - 1] = 999999
    ws.append(row)
    wb.save(book)
    wb.close()
    r = _split(book, tmp_path / "配る", "--amount", "金額")
    assert r.returncode == 0, r.stdout
    assert "⚠ 元の表の合計行" in r.stdout and "合いません" in r.stdout, r.stdout
    assert "分けた冊は明細のとおりです" in r.stdout, r.stdout


# --- 配った冊に本人の合計（2026-09-14・会計役 MISSING #2）-------------------------------------
#
# ★★ 「山田太郎.xlsx は明細 2 行だけ。8800 は _検分.xlsx にしかない。本人に渡す冊なら合計が欲しい」
# ★ 足した行も**検算の対象**にする ── 出所（元行）が空の行を黙って飛ばすと「捏造した行」の穴になる。
#   判断は `split_people.is_own_total` 1 箇所（書く側・事後条件・独立検算が同じ関数を呼ぶ）。


def _books(out_dir):
    return {p.stem: p for p in sorted(out_dir.glob("*.xlsx")) if not p.name.startswith("_")}


def test_each_book_ends_with_its_own_total(book, tmp_path):
    out = tmp_path / "配る"
    r = _split(book, out, "--amount", "金額")
    assert r.returncode == 0, r.stdout
    checked = []
    for stem, path in _books(out).items():
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        head = [str(c.value) for c in ws[1]]
        ai = head.index("金額")
        last = rows[-1]
        assert last[head.index("担当者")] == "合計", (stem, last)
        assert last[-1] is None and last[-2] is None, f"{stem}: 合計行に出所が入っている: {last}"
        detail = sum(r0[ai] for r0 in rows[:-1] if isinstance(r0[ai], (int, float)))
        assert last[ai] == detail, (stem, last[ai], detail)
        assert ws.cell(row=ws.max_row, column=head.index("担当者") + 1).font.bold
        wb.close()
        checked.append(stem)
    assert len(checked) >= 2, checked


def test_the_proof_and_the_verifier_do_not_double_count_the_total(book, tmp_path):
    out = tmp_path / "配る"
    r = _split(book, out, "--amount", "金額")
    assert r.returncode == 0, r.stdout
    assert "✓ 部分の和 ＝ 全体" in r.stdout, r.stdout
    v = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(book),
                        "--amount", "金額"], cwd=str(REPO), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=300)
    assert v.returncode == 0 and "✓" in v.stdout, f"exit={v.returncode} / {v.stdout}"


def test_without_amount_no_total_row_is_added(book, tmp_path):
    """★ 陰性対照 ── 金額の列を指していない回は冊の形を変えない。"""
    out = tmp_path / "配る"
    assert _split(book, out).returncode == 0
    for stem, path in _books(out).items():
        wb = openpyxl.load_workbook(path)
        vals = [c.value for c in wb.active["A"]]
        wb.close()
        assert "合計" not in [str(v) for v in vals], stem


def test_a_tampered_total_is_a_break(book, tmp_path):
    """★★ 足した行は検算の対象 ── 合計を書き換えたら独立検算が破れる。"""
    out = tmp_path / "配る"
    assert _split(book, out, "--amount", "金額").returncode == 0
    path = next(iter(_books(out).values()))
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ai = [str(c.value) for c in ws[1]].index("金額") + 1
    ws.cell(row=ws.max_row, column=ai).value = 999999
    wb.save(path)
    wb.close()
    v = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(book),
                        "--amount", "金額"], cwd=str(REPO), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=300)
    assert v.returncode == 5, f"exit={v.returncode} / {v.stdout}"
    assert "★ 冊の合計が明細と合わない" in v.stdout, v.stdout


def test_a_row_without_provenance_and_without_the_word_is_still_a_break(book, tmp_path):
    """★★ 穴を開けていない証明 ── 出所の無い行は「合計」の語が無ければ今までどおり破れ。"""
    out = tmp_path / "配る"
    assert _split(book, out, "--amount", "金額").returncode == 0
    path = next(iter(_books(out).values()))
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    head = [str(c.value) for c in ws[1]]
    row = [None] * len(head)
    row[head.index("担当者")] = "こっそり足した行"
    row[head.index("金額")] = 1
    ws.append(row)
    wb.save(path)
    wb.close()
    v = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(book),
                        "--amount", "金額"], cwd=str(REPO), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=300)
    assert v.returncode == 5 and "元行が数でない" in v.stdout, f"exit={v.returncode} / {v.stdout}"


def test_a_person_named_like_a_total_is_never_handed_a_book(tmp_path):
    """★ 検体 S05 の家系（担当者の値そのものが合計の語）── **実測して予測を却下した**回。

    ★★ 2026-09-14: 「合計の語の人に合計行を足すと人と合計行が見分けられない」と読んで
      `own_total_row` に柵を置き、この e2e で鳴らすつもりだった。測ってみると上流の
      「分けない行」の判定が先に効き、『合計商事』の行は *誰の冊にも入らない* ── 冊自体が
      作られないので、その事故は CLI からは起こせない。柵は第二の柵として残すが、番人は
      `own_total_row` の単体試験（`test_split_people.py`）が持つ ── ★ 到達できない柵を
      e2e で「守っている」と名乗らない（到達できず＝未確認）。
    """
    b = _book(tmp_path / "元" / "一覧.xlsx",
              headers=["日付", "顧客", "担当者", "金額"],
              rows=[["2026-09-01", "甲社", "合計商事", 1000],
                    ["2026-09-02", "乙社", "内藤", 2000]])
    out = tmp_path / "配る"
    r = _split(b, out, "--amount", "金額")
    assert r.returncode == 0, r.stdout
    books = _books(out)
    assert "合計商事" not in books, f"合計の語の行に冊が出た: {sorted(books)}"
    assert "分けない行" in r.stdout, r.stdout
    wb = openpyxl.load_workbook(books["内藤"])
    assert wb.active.max_row == 3, "普通の人には足す"
    wb.close()


def test_a_wrong_total_never_reaches_the_person(book, tmp_path, monkeypatch):
    """★★ 足した行は**書いた直後に読み戻して**確かめる ── 嘘の合計を書いたら 1 冊も置かない。

    書き手（`own_total_row`）が 1 円ずれた合計を返すように差し替える。事後条件が合計行を
    見ていなければ exit 0 で配られてしまう ── 見ているなら exit 5 で誰にも届かない。
    ★ 「合計が明細と合わない」は独立検算にも在るが、あちらは**後から**の目。人に渡る前に
      止める目がこちら側にも要る（片方だけだと、配った後に気づく）。
    """
    import ailine

    real = ailine.split_people.own_total_row

    def off_by_one(out_headers, plan, value, row_nums, row_values):
        row = real(out_headers, plan, value, row_nums, row_values)
        if row is not None and plan.amount_column:
            row[plan.amount_column - 1] += 1
        return row

    monkeypatch.setattr(ailine.split_people, "own_total_row", off_by_one)
    out = tmp_path / "配る"
    args = argparse.Namespace(book=str(book), by="担当者", out=str(out), amount="金額",
                              sheet=None, overwrite=False, json=False)
    assert ailine.cmd_split(args) == 5, "★ 嘘の合計を書いても通った（合計行を確かめていない）"
    assert not out.exists() or not list(out.glob("*.xlsx")), "★ 証明できていないのに配った"
