# -*- coding: utf-8 -*-
"""道具は**自分が作ったシート**を知っている（2026-09-21・盲検 5 体目・出所追跡）。

★★ 出所（買い手の言葉）:

> ★ の一行がほぼ毎回出ます。**3 回目から読まなくなりました。大事な警告と見分けが
> つきません。**

  実測（凍結セッションの解釈行と依頼文＝どちらも run 時点の事実）: **11/18 = 61%**。
  そして発火した回で競合していたシートは ──
  『品番の重複除去』『集計』『商品名・発注金額だけ』『検分』。
  ★★ **どれも ailine 自身の置き土産**。自分が散らかしたものを理由に
    「どのシートか分かりません」と言っていた。

★★ なぜ形で推し量らないか（この直しの肝）:
  `write_precondition._looks_like_own_prior_output` は見出しの一致で「たぶん自分」を
  当てる代理指標で、docstring 自身が敵対検証（08-20）の結論をこう書いている ──
  「署名は**出所ではなく形**でしか無い」「**真の出所追跡は未実装**」。
  ★ その上に建てると、人の手作りシートを自分のものと誤認して**警告が要る回に黙る**。
    だから「作った時に本人が残した記録」を読む側を建てた。

★★ 失敗の向きを設計で決めてある: 記録が無い・冊を移した/複製した・履歴が読めない
  → **空** → 「全部 人のもの」→ **注記は出続ける**。黙って消える警告が一番高くつく。

★★ この番人は**本番の書き手を必ず通す**。
  `build_history_entry` のコメントに実例がある ── `out_sha` は作られていたのに
  キーの固定列挙から漏れ、履歴全行で欠落し、関所は**到達不能**だった。
  その時「番人が通した理由の方が重い: 検体が history の行を**手で書いて**いて、
  本番の書き手を一度も通っていなかった」。ここは同じ罠の前に立っている。
"""
from __future__ import annotations

import ast
import json
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

import ailine  # noqa: E402
from _product_source import product_files  # noqa: E402


def _folder_of_books(tmp_path: Path) -> Path:
    d = tmp_path / "店舗別"
    d.mkdir()
    for name, rows in (("本店", [["ねじ", 3, 100]]), ("駅前", [["くぎ", 2, 50]])):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["商品名", "数量", "売価"])
        for r in rows:
            ws.append(r)
        wb.save(d / f"売上_{name}.xlsx")
    return d


# --- 本番の書き手を通す（手で履歴を書かない） ----------------------------------------

def test_a_book_the_tool_made_is_recorded_by_the_real_writer(tmp_path, monkeypatch, capsys):
    """★★ `stack` が作った冊のシートが、**本番の経路で**台帳に載ること。

    ★ 買い手の `9月_全店売上.xlsx`（売上＋検分）は丸ごと `stack` の出力なのに、
      `stack` は履歴を **1 行も残していなかった**。
    """
    hist = tmp_path / "history.jsonl"
    monkeypatch.setattr(ailine, "HISTORY_FILE", hist)
    out = tmp_path / "全店.xlsx"
    rc = ailine.main(["stack", str(_folder_of_books(tmp_path)), "--out", str(out)])
    capsys.readouterr()
    assert rc == 0
    assert out.exists()

    made = ailine.sheets_ailine_made(out, path=hist)
    assert made, "作った冊の出所が台帳に無い（stack が記録していない）"
    assert made == set(openpyxl.load_workbook(out, read_only=True).sheetnames), made


def test_the_ledger_carries_the_made_sheets_key(tmp_path):
    """★★ `build_history_entry` は**キーを固定列挙**している ── そこから漏れると、
    作っていても台帳に載らず、読む側は永遠に空になる（`out_sha` の前例そのもの）。"""
    entry = ailine.build_history_entry(
        {"ok": True, "made_sheets": ["検分"]}, tmp_path / "b.xlsx", "t", "m", "none")
    assert "made_sheets" in entry, "台帳が made_sheets を写していない"
    assert entry["made_sheets"] == ["検分"]


def test_the_diff_and_the_provenance_are_decided_in_one_place():
    """★★ `result["changes"]` を直に書く場所がゼロであること（片配線にしない）。

    ★ 差分を書く経路は 4 つあった。そこへ「作ったシート」を足すと 5 つ目の片配線になる。
      判断は `record_diff` 1 本に畳み、呼び出し側には持たせない。
    """
    hits = []
    for f in product_files():
        for i, line in enumerate(f.read_bytes().decode("utf-8").splitlines(), 1):
            if 'result["changes"] =' in line and not line.strip().startswith("#"):
                hits.append((f.name, i, line.strip()))
    assert len(hits) == 1, f"差分を書く場所が {len(hits)} ある（1 本に畳むこと）: {hits}"


def test_record_diff_derives_the_made_sheets_from_the_snapshots():
    """★ 作成経路を**手で並べない** ── before/after の差分から機械的に出す。"""
    result = {}
    ailine.record_diff(result, {"sheets": ["元"]}, {"sheets": ["元", "検分", "集計"]}, ["x"])
    assert result["made_sheets"] == ["検分", "集計"], result
    assert result["changes"] == ["x"]


# --- 失敗の向き: 分からない時は黙らない -----------------------------------------------

def test_an_unknown_book_is_treated_as_the_users(tmp_path):
    """★★ 記録が無い冊は**空**（＝全部 人のもの）── 注記は出続ける。"""
    assert ailine.sheets_ailine_made(tmp_path / "見知らぬ.xlsx",
                                      path=tmp_path / "無い.jsonl") == set()


def test_an_unreadable_ledger_is_treated_as_the_users(tmp_path):
    """★ 台帳が読めない回も空 ── **見ていないことを根拠に黙らない**。"""
    bad = tmp_path / "history.jsonl"
    bad.write_bytes(b"\xff\xfe not json at all")
    assert ailine.sheets_ailine_made(tmp_path / "b.xlsx", path=bad) == set()


def test_a_copied_book_keeps_warning(tmp_path, monkeypatch, capsys):
    """★★ 冊を**複製**したら出所は付いてこない ── 複製側では注記が出続けること。

    ★ 買い手も `突合_試し.xlsx` という複製を作っていた。そこで黙る方が危ない。
    """
    hist = tmp_path / "history.jsonl"
    monkeypatch.setattr(ailine, "HISTORY_FILE", hist)
    out = tmp_path / "全店.xlsx"
    ailine.main(["stack", str(_folder_of_books(tmp_path)), "--out", str(out)])
    capsys.readouterr()
    copy = tmp_path / "全店_試し.xlsx"
    copy.write_bytes(out.read_bytes())
    assert ailine.sheets_ailine_made(out, path=hist)
    assert ailine.sheets_ailine_made(copy, path=hist) == set(), (
        "複製にまで出所が付いてきている ── 複製は別の冊")


# --- パスの同一性（実測で捕まえた死に方） ----------------------------------------------

@pytest.mark.parametrize("how", ["そのまま", "スラッシュを逆に", "大文字にする", "相対を混ぜる"])
def test_the_same_file_written_differently_is_the_same_file(tmp_path, monkeypatch, how):
    """★★ 区切りの向き・大小文字・相対で**別物にならない**こと。

    ★★ 実測で捕まえた死に方（2026-09-21）: 生の文字列で突き合わせていたら、
      `stack` が書いた `C:\\…\\全店.xlsx` と人が打った `C:/…/全店.xlsx` が別物になり、
      記録は在るのに**一度も当たらなかった**。
    ★★ 怖いのは**挙動に出なかった**こと ── 当たらない＝「人のもの」扱いで注記が
      出続けるだけなので、安全側に倒れて静かに死んでいた。
      **陽性側を実測しなければ気づけない**（「出ないことは信号でない」）。
    """
    hist = tmp_path / "history.jsonl"
    monkeypatch.setattr(ailine, "HISTORY_FILE", hist)
    real = tmp_path / "全店.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "売上"
    wb.active.append(["a"])
    wb.save(real)
    ailine.record_made_book(real, "stack")

    asked = {"そのまま": str(real),
             "スラッシュを逆に": str(real).replace("\\", "/"),
             "大文字にする": str(real).upper(),
             "相対を混ぜる": str(tmp_path / "." / "全店.xlsx")}[how]
    assert ailine.sheets_ailine_made(asked, path=hist) == {"売上"}, how


# --- 門: 人が作った枚数で数える --------------------------------------------------------

def test_the_gate_counts_only_the_sheets_the_person_made():
    """★★ 『複数シートのブックだけ』の**複数**を、人が作った枚数で数えること。"""
    got = ailine._subject_slots("SET_WHERE", {"_target_sheet": "売上"},
                                 ["売上", "検分"], "見出しを太字にして",
                                 sheets_we_made=("検分",))
    assert not [s for s in got if s.key == "_target_sheet"], (
        "自分の置き土産を『選択肢』に数えている")


def test_the_gate_still_speaks_when_two_sheets_are_really_the_persons():
    """★★ 陰性対照 ── 人のシートが 2 枚なら今までどおり言うこと。

    ★ 買い手の `突合.xlsx`（請求・発注）がこれ。ここで黙ったら直しではなく退行。
    """
    got = ailine._subject_slots("SET_WHERE", {"_target_sheet": "請求"},
                                 ["請求", "発注", "商品名・発注金額だけ"],
                                 "確認という列を追加して",
                                 sheets_we_made=("商品名・発注金額だけ",))
    assert [s for s in got if s.key == "_target_sheet"], (
        "人のシートが 2 枚あるのに黙っている")


def test_knowing_nothing_keeps_the_old_behaviour():
    """★ 出所が空なら従来どおり全部数える（既定が安全側）。"""
    got = ailine._subject_slots("SET_WHERE", {"_target_sheet": "売上"},
                                 ["売上", "検分"], "見出しを太字にして")
    assert [s for s in got if s.key == "_target_sheet"]


# --- 分母: 冊を作る入口が記録を落としていないか ----------------------------------------

def test_every_command_that_commits_a_book_records_what_it_made():
    """★★ 冊を仕上げる入口の**全部**が出所を残すこと ── 分母は AST から導く。

    ★ 判定: `cmd_*` の中で作業用の一時ファイルを出力先へ写している関数は、
      同じ関数の中で `record_made_book` も呼ぶ。一覧を手で書かない
      （入口を足した日に静かにずれる）。
    ★ 記録の仕方は 2 通りある（冊まるごと＝`record_made_book`／足したシート＝
      `record_diff`）。**どちらかを通っていること**を見る ── 手段は縛らない。

    ★★ この番人は初回から穴を 3 つ掴んだ（2026-09-21）: `cmd_run_folder`・
      `cmd_run_match`（買い手が「いちばん役に立った」と言った **2 冊並べ**の経路）と
      `cmd_demo` が、冊を仕上げているのに出所を一切残していなかった。
    """
    missing = []
    for f in product_files():
        for node in ast.walk(ast.parse(f.read_bytes().decode("utf-8"))):
            if not (isinstance(node, ast.FunctionDef) and node.name.startswith("cmd_")):
                continue
            names = {getattr(n.func, "attr", None) or getattr(n.func, "id", None)
                     for n in ast.walk(node) if isinstance(n, ast.Call)}
            if "copy2" in names and not (names & {"record_made_book", "record_diff"}):
                missing.append((f.name, node.name))
    assert not missing, f"冊を仕上げているのに出所を残していない入口: {missing}"
