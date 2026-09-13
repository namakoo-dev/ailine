# -*- coding: utf-8 -*-
"""`ailine accounts <今回> --past <過去…> --out <冊>` の検体 ── 需要③の出口（2026-09-13）。

契約（設計 docs/DESIGN-20260913-経費の勘定科目を先例から引く.md ★ §6 が §2 を置き換える）:
  - 出力は 1 冊・シート 2 枚（`候補` と `検分`）。原本は 1 バイトも触らない（読むだけ）
  - 書き手の印（creator）と条件（description）が焼かれている ── 次に自分の出力と分かる
  - 出力先に人のファイルがあれば exit 7（stack / forms / split と同じ関所）
  - 今回と過去が同じファイルなら断る（自己先例は『裏が取れた』の最短路）
  - 見出しが決まらない冊は断る（1 冊も作らない）
  - 文字コードは `csv_quarantine.detect_encoding` を通す（弥生の cp932 が読める）
  - ★ 画面が出す「あとから確かめる」の呼び方を**そのまま実行して exit 0**（案内が嘘でない）

★ LLM も LibreOffice も要らない（読むだけ・openpyxl のみ）。
"""
from __future__ import annotations

import csv
import hashlib
import json
import shlex
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

from ailine_core import accounts_core   # noqa: E402
from ailine_core import stack as stack_core   # noqa: E402

#: マネーフォワードの書き出しを模した見出し（実物の列名）。
MF = ["取引No", "取引日", "借方勘定科目", "借方補助科目", "借方取引先", "借方金額(円)",
      "貸方勘定科目", "貸方取引先", "摘要"]

_PAST = [
    ["1", "2026/06/03", "通信費", "携帯", "甲通信", "8800", "現金", "", "6月分 電話代"],
    ["2", "2026/07/03", "通信費", "", "甲通信", "8800", "現金", "", "7月分 電話代"],
    ["3", "2026/06/10", "消耗品費", "文具", "乙商店", "1200", "現金", "", "コピー用紙"],
    ["4", "2026/06/20", "旅費交通費", "", "丙タクシー", "2300", "現金", "", "タクシー代"],
    ["5", "2026/07/20", "会議費", "", "丙タクシー", "4000", "現金", "", "打合せ"],
]
_TODAY = [
    ["11", "2026/08/03", "", "携帯", "甲通信", "8800", "現金", "", "8月分 電話代"],   # 確
    ["12", "2026/08/10", "", "", "乙商店", "1500", "現金", "", "ボールペン"],          # 単
    ["13", "2026/08/20", "", "", "丙タクシー", "2500", "現金", "", "タクシー代"],      # 割
    ["14", "2026/08/25", "", "", "未知商会", "900", "現金", "", "よく分からない"],     # 無
    ["15", "2026/08/26", "旅費交通費", "", "丙タクシー", "1000", "現金", "", "電車"],  # 埋まっている
    ["", "", "", "", "", "", "未払金", "", "継続行"],                                  # 継続行
]


def _csv(path: Path, rows, headers=MF, encoding="utf-8-sig"):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding=encoding, newline="") as f:
        writer = csv.writer(f)
        if headers is not None:
            writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
    return path


def _xlsx(path: Path, rows, headers=MF):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "仕訳帳"
    if headers is not None:
        ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


def _run(*argv):
    return subprocess.run([sys.executable, "-m", "ailine", *argv],
                          capture_output=True, text=True, timeout=300,
                          encoding="utf-8", errors="replace", cwd=str(REPO))


def _accounts(today: Path, past, out: Path, *extra):
    past_args = [str(p) for p in (past if isinstance(past, (list, tuple)) else [past])]
    return _run("accounts", str(today), "--past", *past_args, "--out", str(out), *extra)


def _payload(result):
    return json.loads(result.stdout.strip().splitlines()[-1])


def _sheet(path: Path, name: str) -> list:
    wb = openpyxl.load_workbook(path)
    got = [list(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    return got


@pytest.fixture()
def books(tmp_path):
    """(今回, 過去) の 2 冊（MF 形式・utf-8-sig）。"""
    return (_csv(tmp_path / "元" / "今回.csv", _TODAY),
            _csv(tmp_path / "元" / "過去.csv", _PAST))


def test_the_book_has_the_two_sheets_and_the_candidate_columns(books, tmp_path):
    """★ 候補のシートは「今回の行 ＋ 出所列 ＋ 候補の科目・区分・根拠・先例の番地」。"""
    today, past = books
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, past, out)
    assert r.returncode == 0, f"落ちた: {r.stdout}\n{r.stderr}"
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [accounts_core.SHEET_NAME, accounts_core.REPORT_SHEET], \
        wb.sheetnames
    wb.close()
    head, *rows = _sheet(out, accounts_core.SHEET_NAME)
    assert head[:len(MF)] == MF, head
    assert head[len(MF):] == list(stack_core.PROVENANCE_HEADERS) \
        + list(accounts_core.OUTPUT_HEADERS), head[len(MF):]
    # ★ 今回の行は全部並ぶ（触らない行も ── 人が原本の代わりに読める形にする）。
    assert len(rows) == len(_TODAY), rows
    by_row = {row[len(MF) + 1]: row[len(MF) + 2:] for row in rows}
    assert by_row[2][0] == "通信費", by_row[2]
    assert by_row[4][0] is None, "★ 割 の行に値を出している"
    assert by_row[7][0] is None, "★ 継続行に候補を出している"


def test_the_writer_mark_and_the_condition_are_burned_in(books, tmp_path):
    """★ 印（creator）と条件（description）── 次に「自分の出力か」を機械で言えるように。"""
    today, past = books
    out = tmp_path / "候補.xlsx"
    assert _accounts(today, past, out).returncode == 0
    wb = openpyxl.load_workbook(out)
    creator, description = wb.properties.creator, wb.properties.description
    wb.close()
    assert creator == accounts_core.CREATOR_MARK
    cond = json.loads(description)
    assert cond["tool"] == "ailine" and cond["kind"] == accounts_core.KIND
    assert cond["today"] == today.name and cond["past"] == [past.name]
    assert stack_core.own_output_mark(out) == accounts_core.CREATOR_MARK, \
        "★ 自分の出力と認識されない（次回の関所が人のファイルと誤認する）"
    assert accounts_core.CREATOR_MARK in stack_core.CREATOR_MARKS


def test_the_inputs_are_not_touched(books, tmp_path):
    today, past = books
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (today, past)]
    assert _accounts(today, past, tmp_path / "候補.xlsx").returncode == 0
    after = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (today, past)]
    assert before == after, "★ 入力が変わった（読むだけの約束が破れている）"


def test_json_carries_the_whole_contract(books, tmp_path):
    """★ 機械可読の側にも同じ事実（採点器と自動化がここを読む）。"""
    today, past = books
    r = _accounts(today, past, tmp_path / "候補.xlsx", "--json")
    payload = _payload(r)
    for key in ("rows", "untouched", "lookalike", "refused", "grades", "keys_used",
                "notes", "原本が変わった", "encoding", "past", "file_written"):
        assert key in payload, f"--json に {key} が無い: {sorted(payload)}"
    assert payload["refused"] is None and payload["原本が変わった"] is False
    # ★ 行番号は**物理行**（見出しが 1 行目なのでデータは 2 行目から）。
    assert sorted(int(r0) for r0 in payload["rows"]) == [2, 3, 4, 5]
    assert [r0 for r0, _why in payload["untouched"]] == [6, 7]
    assert payload["rows"]["2"]["account"] == "通信費"
    assert payload["rows"]["4"]["account"] is None
    assert payload["rows"]["2"]["citations"], "★ 番地が機械可読で出ていない"
    assert payload["encoding"] == "utf-8-sig"


def test_the_report_names_the_blanks_and_the_grades(books, tmp_path):
    today, past = books
    r = _accounts(today, past, tmp_path / "候補.xlsx")
    for token in ("空欄", "区分", "候補を出す行", "貸方は鍵に入れていません"):
        assert token in r.stdout, f"{token} が画面に無い: {r.stdout}"


def test_the_inspection_sheet_separates_the_denominator_from_the_untouched(books, tmp_path):
    today, past = books
    out = tmp_path / "候補.xlsx"
    assert _accounts(today, past, out).returncode == 0
    head, *rows = _sheet(out, accounts_core.REPORT_SHEET)
    assert head == list(accounts_core.REPORT_HEADERS), head
    kinds = {str(row[0]) for row in rows}
    for kind in (accounts_core.BLANK_KIND, accounts_core.UNTOUCHED_KIND,
                 accounts_core.NOTE_KIND):
        assert kind in kinds, f"検分に『{kind}』が無い: {sorted(kinds)}"
    blanks = sorted(row[1] for row in rows if row[0] == accounts_core.BLANK_KIND)
    assert blanks == [4, 5], blanks


def test_the_verify_hint_printed_on_screen_actually_works(books, tmp_path):
    """★★ 案内した呼び方を**そのまま実行して exit 0**（知られない検算は無い検算）。"""
    today, past = books
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, past, out)
    hint = [ln for ln in r.stdout.splitlines() if ln.startswith("あとから確かめる: ")]
    assert hint, f"検算の呼び方が画面に無い: {r.stdout}"
    # ★ Windows のパスは `\` を含む ── posix=True で割ると区切りが消える（試験の治具の話）。
    argv = [tok.strip('"') for tok in shlex.split(hint[0].split(": ", 1)[1], posix=False)]
    assert argv[0] == "ailine", argv
    again = _run(*argv[1:])
    assert again.returncode == 0, f"案内した呼び方が通らない: {again.stdout}\n{again.stderr}"
    assert "破れはありません" in again.stdout, again.stdout


def test_an_existing_human_file_stops_the_write(books, tmp_path):
    """★ 出力先に人のファイルがあれば exit 7（stack / forms / split と同じ関所）。"""
    today, past = books
    out = tmp_path / "大事な表.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "人が書いた大事な表"
    wb.save(out)
    wb.close()
    before = out.read_bytes()
    r = _accounts(today, past, out)
    assert r.returncode == 7, f"人のファイルを黙って潰した: {r.stdout}"
    assert out.read_bytes() == before
    assert "--overwrite" in r.stdout, "次の手を言っていない"
    assert _accounts(today, past, out, "--overwrite").returncode == 0, \
        "承知の上なら通ること"


def test_its_own_previous_output_is_not_a_gate(books, tmp_path):
    today, past = books
    out = tmp_path / "候補.xlsx"
    assert _accounts(today, past, out).returncode == 0
    again = _accounts(today, past, out)
    assert again.returncode == 0, f"自分の前回出力を人のファイルと誤認した: {again.stdout}"


def test_the_same_file_on_both_sides_is_refused(books, tmp_path):
    """★★ 自己先例は『裏が取れた』の最短路 ── 同じファイルなら断る（設計 §6.3）。"""
    today, _past = books
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, today, out)
    assert r.returncode == 4, f"同じファイルを先例にした: {r.stdout}"
    assert today.name in r.stdout
    assert not out.exists(), "断ったのに冊を作っている"


def test_a_book_whose_columns_cannot_be_resolved_is_refused(tmp_path):
    """★ 列が決まらない冊は断る（見た見出しを並べる・1 冊も作らない）。"""
    today = _csv(tmp_path / "元" / "謎.csv", [["a", "b", "c"]], headers=["甲", "乙", "丙"])
    past = _csv(tmp_path / "元" / "過去.csv", _PAST)
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, past, out)
    assert r.returncode == 4, r.stdout
    assert accounts_core.DEBIT_ACCOUNT in r.stdout and "甲" in r.stdout, r.stdout
    assert not out.exists()


def test_an_xlsx_journal_is_read_the_same_way(tmp_path):
    """★ xlsx の仕訳帳（MF の列を Excel に貼った形）も同じ規則で読む。"""
    today = _xlsx(tmp_path / "元" / "今回.xlsx", _TODAY)
    past = _xlsx(tmp_path / "元" / "過去.xlsx", _PAST)
    r = _accounts(today, past, tmp_path / "候補.xlsx", "--json")
    payload = _payload(r)
    assert payload["refused"] is None, payload["refused"]
    assert payload["rows"]["2"]["account"] == "通信費"


def test_a_yayoi_journal_in_cp932_without_headers_is_read_by_position(tmp_path):
    """★★ 弥生（cp932・見出し行なし・25 列固定）── 列位置で受ける（設計 §6.2）。"""
    def row(date, account, sub, amount, memo):
        values = [""] * accounts_core.YAYOI_COLUMN_COUNT
        values[0] = "2000"            # 識別フラグ（★ 一次資料どおり 4 桁 ── 空の冊は検体の欠けだった）
        values[1], values[3], values[4] = "1", date, account
        values[5], values[8], values[10], values[16] = sub, amount, "現金", memo
        return values

    past = _csv(tmp_path / "元" / "弥生_過去.csv",
                [row("R08/06/04", "荷造運賃", "", "2031", "宅配便定期便"),
                 row("R08/06/20", "荷造運賃", "", "1500", "宅配便定期便")],
                headers=None, encoding="cp932")
    today = _csv(tmp_path / "元" / "弥生_今回.csv",
                 [row("R08/07/01", "", "", "2031", "宅配便定期便")],
                 headers=None, encoding="cp932")
    r = _accounts(today, past, tmp_path / "候補.xlsx", "--json")
    payload = _payload(r)
    assert payload["refused"] is None, payload["refused"]
    assert payload["encoding"] == "cp932", payload["encoding"]
    assert payload["header_row"] is None, "★ 弥生に見出し行は無い"
    assert payload["keys_used"] == ["借方補助科目", "摘要"], payload["keys_used"]
    assert payload["rows"]["1"]["account"] == "荷造運賃", payload["rows"]
    assert "R08/06/20" in payload["rows"]["1"]["reason"], \
        "★ 和暦略記を原文のまま運んでいない（日付は読まない側）"


def test_two_past_files_are_both_read(tmp_path):
    """★ 過去は複数渡せる（番地はファイル名つきで出る）。"""
    today = _csv(tmp_path / "元" / "今回.csv", _TODAY)
    one = _csv(tmp_path / "元" / "過去1.csv", _PAST[:2])
    two = _csv(tmp_path / "元" / "過去2.csv", _PAST[2:])
    r = _accounts(today, [one, two], tmp_path / "候補.xlsx", "--json")
    payload = _payload(r)
    assert payload["past"] == ["過去1.csv", "過去2.csv"], payload["past"]
    names = {c[1] for row in payload["rows"].values() for c in row["citations"]}
    assert names == {"過去1.csv", "過去2.csv"}, names


def test_a_folder_of_past_journals_is_accepted(tmp_path):
    """★ `--past` はフォルダでも受ける（直下だけ・help がそう名乗っている）。"""
    today = _csv(tmp_path / "元" / "今回.csv", _TODAY)
    folder = tmp_path / "過去たち"
    _csv(folder / "過去1.csv", _PAST[:2])
    _csv(folder / "過去2.csv", _PAST[2:])
    r = _accounts(today, folder, tmp_path / "候補.xlsx", "--json")
    payload = _payload(r)
    assert sorted(payload["past"]) == ["過去1.csv", "過去2.csv"], payload["past"]


def test_a_book_where_nothing_ever_hits_is_refused_not_quietly_empty(tmp_path):
    """★★ 全行『無』の静かな成功をしない（設計 §6.3）── 1 冊も作らない。"""
    today = _csv(tmp_path / "元" / "今回.csv",
                 [["11", "2026/08/03", "", "", "見知らぬ会社", "8800", "現金", "", "謎"]])
    past = _csv(tmp_path / "元" / "過去.csv", _PAST)
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, past, out)
    assert r.returncode == 4, f"黙って全行『無』で成功した: {r.stdout}"
    assert not out.exists()
    assert "文字コード" in r.stdout, f"疑う先を言っていない: {r.stdout}"


def test_the_route_is_declared_and_reachable():
    """★ ROUTE_KIND の等号の番人と、`ailine ops` から辿れること（到達できない機能は無い機能）。"""
    import ailine
    assert ailine.ROUTE_KIND["accounts"] == "multi"
    assert "accounts" in ailine.multi_file_routes()
    r = _run("ops")
    assert r.returncode == 0
    assert "ailine accounts" in r.stdout, r.stdout
