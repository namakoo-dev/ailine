# 宛先が決まらない冊で、請求元に**買い手の名前**を出さないこと ── 2026-09-12
#
# ★★ 事故の形（PDF 化した 87 冊で実測・14 件）:
#   請求書は「請求元」と「宛先」が**同じ形の 2 ブロック**で並ぶ。`read_issuer` は
#   「宛先の名前を含む候補を除く」で買い手を消しているが、**宛先が決まらないと
#   その除外が黙って無効**になる。すると買い手の名前が唯一の候補として残り、
#   候補 1 つ ＝ 単 ＝ **自信を持って買い手を請求元として出す**。
#   支払いの文脈で出しうる最悪の間違い。
#
# ★ `read_issuer` の docstring は以前から「宛先が 割 なので買い手を除けない ──
#   そこは決められないのが正直」と**宣言していた**のに、実体は値を出していた。
#   宣言と実体の食い違いは if では塞げない ── **項を渡して機械に確かめさせる**。
#
# ★ Excel でも同じ経路は在った。宛先が一度も失敗しなかったから発火しなかっただけで、
#   番人は「宛先が欠けたときの請求元」を一度も試していなかった（負の被覆）。
from __future__ import annotations

import openpyxl

from ailine_core.field_record import grade, value
from ailine_core.form_grid import Grid
from ailine_core.form_read import read_addressee, read_issuer


def _sheet(rows: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "請求書"
    for at, v in rows.items():
        ws[at] = v
    return ws


#: ★ 実物で起きた形（PDF 化した 87 冊で実測）: 敬称が名前の**右**のマスに在り、
#:   本当の請求元は屋号（法人格が無いので候補に上がらない）。
#:   ★ 敬称が名前の**下**なら `read_addressee` は「敬称の直上」で拾える ── 壊れるのは右の形。
DANGEROUS = {"B3": "ナギ商会株式会社", "C3": "御中",
             "G5": "あかね商事", "G6": "T1234567890123"}


def test_the_addressee_really_cannot_be_determined_in_this_fixture():
    """★ 検体が前提どおりであること（ここが崩れると下の 2 本は何も試していない）。"""
    rec = read_addressee(Grid.read(_sheet(DANGEROUS)))
    assert grade(rec) == "無", grade(rec)
    assert value(rec) is None


def test_the_buyers_name_is_never_reported_as_the_issuer():
    """★★ 本番: 買い手の名前を請求元として出さない（空欄＋理由）。"""
    g = Grid.read(_sheet(DANGEROUS))
    rec = read_issuer(g, read_addressee(g))
    assert value(rec) != "ナギ商会株式会社", "★ 買い手を請求元として出した"
    assert grade(rec) not in ("確", "単"), grade(rec)
    why = rec.blank_reason
    assert "宛先が決まらなかった" in why, why
    # ★ 見つけた名前は隠さない ── 人が判断できるよう番地つきで並べる
    assert "ナギ商会株式会社" in why and "B3" in why, why


def test_the_candidates_are_kept_as_rivals_so_nothing_is_silently_dropped():
    g = Grid.read(_sheet(DANGEROUS))
    rec = read_issuer(g, read_addressee(g))
    assert [r[1] for r in rec.rivals] == ["ナギ商会株式会社"], rec.rivals
    assert "見分けられない" in rec.rivals[0][2], rec.rivals


def test_nothing_changes_when_the_addressee_is_known():
    """★ 陰性対照 ── 宛先が取れている冊では今までどおり請求元を読む。"""
    g = Grid.read(_sheet({"B3": "ナギ商会株式会社　御中",
                          "G5": "株式会社あかね商事", "G6": "T1234567890123"}))
    a = read_addressee(g)
    assert value(a) == "ナギ商会株式会社"
    rec = read_issuer(g, a)
    assert grade(rec) == "単" and value(rec) == "株式会社あかね商事", (grade(rec), value(rec))


def test_a_book_with_no_org_name_at_all_still_says_what_it_looked_for():
    """★ 候補が 1 つも無い冊の理由は変えない（『法人格が付いた名前だけを探しています』）。

    ★ 新しい理由が古い理由を飲み込むと、屋号の冊で人が**探し方**を疑えなくなる。
    """
    g = Grid.read(_sheet({"B3": "ナギ商会", "B4": "御中", "G5": "あかね商事"}))
    rec = read_issuer(g, read_addressee(g))
    assert grade(rec) == "無"
    assert "法人格" in rec.blank_reason, rec.blank_reason
    assert "宛先が決まらなかった" not in rec.blank_reason, rec.blank_reason
