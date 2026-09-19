# -*- coding: utf-8 -*-
"""案 D を**本物の LibreOffice** で証明する（2026-09-19・Namakoo「GOだな」）。

★ 単体（tests/test_a_split_reading_is_not_acted_on.py）は下書き当てを偽物に差し替えている。
  ここは偽物を使わない ── 翻訳だけを固定し（割れ方を再現するため）、下書き当てと差分の
  比較は製品の本物を通す。

★★ 何を証明するか:
    ① 並べ替え→並べ替え と 並べ替え は、冊に起きることが**本当に同じ**で、断らずに通る
       （宣言で比べていた初版の誤りが効く場所 ── 畳まれない同一段）
    ② 行削除 と 列削除 は、冊に起きることが**本当に違う**ので、書く前に止まる
       （実機が拾った本物の揺れ ── 平らな削除依頼にモデルが 2 回に 1 回「列削除」を返す）
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402

pytestmark = pytest.mark.local


def _book(tmp_path: Path) -> Path:
    p = tmp_path / "在庫.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "在庫"
    ws.append(["品名", "棚", "数量"])
    for r in (["ボルト", "A-1", 120], ["ナット", "A-2", 80], ["ワッシャー", "B-1", 300]):
        ws.append(r)
    wb.save(p)
    wb.close()
    return p


def _fixed(answers: list):
    calls = []

    def fake(model, task, book_meta, temperature=0.1):
        calls.append(task)
        return answers[min(len(calls) - 1, len(answers) - 1)]
    return fake, calls


def test_same_book_change_reaches_even_when_the_plans_differ(monkeypatch, tmp_path, capsys):
    """① 並べ替え→並べ替え vs 並べ替え ── 本物の下書き当てで差分が一致し、到達する。

    ★ fold_identical_steps は**新しいシートを作る段しか畳まない**（「同じ並べ替えを 2 回の
      ような段は無害なので触らない」と宣言で絞ってある）。だから製品では 2 段のまま走り、
      2 回目は空振り ── 冊に起きることは 1 回と同じ。宣言（op の並び）で比べれば割れ、
      実体（差分）で比べれば同じ ── 案 D が効く場所そのもの。
    ★ 「同じ列を 2 回動かす」を検体にした初版は、2 段目の接地が『もうそこに在ります』と
      **正しく断る**ので割れて当然だった（検体が非現実的・D の判定は正しい）。
    """
    book = _book(tmp_path)
    monkeypatch.setattr(ailine, "HISTORY_FILE", tmp_path / "history.jsonl")
    monkeypatch.delenv("AILINE_SINGLE_READ", raising=False)
    srt = {"op": "SORT", "args": {"col": "数量", "order": "desc"}}
    two = {"plan": [srt, dict(srt)]}
    one = {"plan": [srt]}
    fake, calls = _fixed([two, one])
    monkeypatch.setattr(ailine, "translate_task", fake)
    rc = ailine.main(["run", str(book), "数量で降順に並べ替えて", "--copy",
                      "--sheet", "在庫", "--timeout", "120"])
    out = capsys.readouterr().out
    assert "読み方が分かれました" not in out, out
    assert rc == 0, out
    got = openpyxl.load_workbook(tmp_path / "在庫.out.xlsx", data_only=True)["在庫"]
    assert [got.cell(r, 3).value for r in range(2, 5)] == [300, 120, 80], (
        [got.cell(r, 3).value for r in range(2, 5)])


def test_different_book_changes_stop_before_writing(monkeypatch, tmp_path, capsys):
    """② 行削除 vs 列削除 ── 本物の下書き当てで差分が違い、書く前に止まる。"""
    book = _book(tmp_path)
    before = book.read_bytes()
    monkeypatch.setattr(ailine, "HISTORY_FILE", tmp_path / "history.jsonl")
    monkeypatch.delenv("AILINE_SINGLE_READ", raising=False)
    rows = {"plan": [{"op": "DELETE_ROWS", "args": {"at": "ナット"}}]}
    cols = {"plan": [{"op": "DELETE_COLUMN", "args": {"col": "棚"}}]}
    fake, calls = _fixed([rows, cols])
    monkeypatch.setattr(ailine, "translate_task", fake)
    rc = ailine.main(["run", str(book), "ナットを削除して", "--copy",
                      "--sheet", "在庫", "--timeout", "120"])
    out = capsys.readouterr().out
    assert rc == 3, out
    assert "読み方が分かれました" in out and "冊に起きること" in out, out
    assert book.read_bytes() == before, "★ 止まったのに原本が変わっている"
    assert not (tmp_path / "在庫.out.xlsx").exists(), "★ 止まったのに .out を書いた"
    # ★ 示した道が通ること（導通）── 候補の op で固定すれば到達する
    fake2, _ = _fixed([rows])
    monkeypatch.setattr(ailine, "translate_task", fake2)
    rc2 = ailine.main(["run", str(book), "ナットを削除して", "--copy", "--sheet", "在庫",
                       "--timeout", "120", "--op", "DELETE_ROWS"])
    assert rc2 == 0, capsys.readouterr().out
    got = openpyxl.load_workbook(tmp_path / "在庫.out.xlsx")["在庫"]
    assert "ナット" not in [got.cell(r, 1).value for r in range(2, 5)]
