# 理由が嘘をつかないこと ── 2026-09-12（実物の請求書 2 冊から出た欠陥）
#
# ★★ 事故の形: 実物の請求書（買い手と売り手に同じ社名が入っている形）で、器官は
#   法人格つきの名前を**見つけて、宛先と同じだから捨てて**いたのに、理由はこう言っていた:
#
#     「『株式会社』などの法人格が付いた名前だけを探しています ──
#       屋号や略称だけの請求書では見つかりません」
#
#   人は**探し方**を疑うが、本当に疑うべきは「宛先と同じ名前と判定された」の方だ。
#   ★ 捨てる判断は正しい。嘘をついていたのは理由。空欄の値段は「理由の正しさ」で決まる。
#
# ★ 捨てた候補は**全部**並べる ── 「宛先と同じ」と「雛形のまま」で分岐させたら、前者が
#   後者を隠して T05（発行者名が `株式会社 〇〇〇`）の名指しを消した（同じ日に踏んだ）。
from __future__ import annotations

import openpyxl

from ailine_core.field_record import grade, value
from ailine_core.form_grid import Grid
from ailine_core.form_read import read_addressee, read_issuer


def _sheet(rows: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "請求書"
    for at, v in rows.items():
        ws[at] = v
    return Grid.read(ws)


def _issuer(g):
    return read_issuer(g, read_addressee(g))


def test_when_the_only_name_matches_the_addressee_the_reason_says_so():
    """★★ 本番: 見つけて捨てたなら、そう言う（探し方を疑わせない）。"""
    g = _sheet({"A2": "サンプル株式会社　御中", "F5": "サンプル株式会社"})
    rec = _issuer(g)
    assert grade(rec) == "無" and value(rec) is None
    why = rec.blank_reason
    assert "宛先と同じ名前" in why, why
    assert "F5" in why and "サンプル株式会社" in why, why          # ★ 番地と名前を名指しする
    assert "屋号や略称だけの請求書では見つかりません" not in why, why   # ★ 嘘の方を言わない
    assert [r[0] for r in rec.rivals] == ["F5"], rec.rivals


def test_a_book_with_no_org_name_at_all_still_blames_the_search_not_the_data():
    """★ 陰性対照 ── 候補が 1 つも無い冊では、今までどおり「探し方」を言う
    （屋号だけの請求書はこちら ── 検体 T07 の名指し『屋号』がここに在る）。"""
    g = _sheet({"A2": "ナギ商会　御中", "F5": "あかね商店"})
    why = _issuer(g).blank_reason
    assert "法人格" in why and "屋号" in why, why
    assert "宛先と同じ名前" not in why, why


def test_a_placeholder_name_is_still_named_even_when_another_candidate_matched_the_addressee():
    """★★ 分岐で片方を隠さない ── 同じ日に踏んだ形（T05 の名指し『〇』が消えた）。"""
    g = _sheet({"A2": "サンプル株式会社　御中", "F5": "サンプル株式会社",
                "B4": "株式会社 〇〇〇"})
    why = _issuer(g).blank_reason
    assert "宛先と同じ名前" in why, why
    assert "〇" in why and "雛形のまま" in why, why
    assert len(_issuer(g).rivals) == 2, _issuer(g).rivals


def test_the_issuer_is_still_read_when_the_names_differ():
    """★ 陰性対照 ── ふつうの冊は今までどおり読む。"""
    g = _sheet({"A2": "ナギ商会株式会社　御中", "F5": "株式会社あかね商事"})
    rec = _issuer(g)
    assert grade(rec) == "単" and value(rec) == "株式会社あかね商事"
