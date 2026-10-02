# -*- coding: utf-8 -*-
"""盲検の欠陥の形 7「断り・警告の文言が偽」── 観測していないことを観測したように言わない（2026-10-02）。

★★ 出所: 盲検の欠陥の形 7（観測していないことを観測したように言う／ある場面でしか正しくない文を
  広く出す）の上位 4 件。地図（読むだけの子が作った）は 1 件も走らせていなかったので、
  **4 件とも先に再現してから**直した。この試験の「嘘になる場面」の検体が、その再現そのもの。

  S4  「原本は変更していません」を、言う時点で指紋を照合せず出口の前後関係だけで言っていた
  S1  「依頼は『X』と行番号の両方を指していますが…」が、X が**書き込む新しい値**でも鳴った
  S2  「（元の表はそのまま残っています）」が、実行した op が表の値を書き換えた回にも出た
  A2  「列X／行n は存在しません/変更されていません」が、在るのに変わらなかっただけの回にも出た

★ どの件も 2 本で縛る: **嘘になる場面で嘘を言わない**／**正しい場面では今までどおり言う**
  （後者が無いと、全部黙らせる直しも緑になる）。
"""
from __future__ import annotations

import argparse
import hashlib
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import ailine  # noqa: E402
from ailine_core import cli_render  # noqa: E402
from ailine_core.row_conflict import (  # noqa: E402
    value_not_in_the_named_row, written_values_in_declaration)
from test_golden_transcripts import _isolate, _run_main  # noqa: E402


@pytest.fixture(autouse=True)
def _no_real_libreoffice(monkeypatch):
    """⚠ の出た回は申し送りのシートを LibreOffice で書く（basrun_apply）── この試験は文言を見るので
    実機を起こさない。必要な試験は自分で basrun_apply を差し替える（後勝ち）。"""
    monkeypatch.setattr(ailine, "basrun_apply", lambda *a, **k: (True, None, "ok"))


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _book(path: Path, rows) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()
    return path


def _nb(tmp_path: Path) -> Path:
    return _book(tmp_path / "nb.xlsx", [["商品", "金額"], ["a", 100], ["b", 250]])


def _sorted_apply(out_book, code, workdir, helper_files=(), timeout=None):
    wb = openpyxl.load_workbook(out_book)
    ws = wb.active
    ws["A2"], ws["B2"], ws["A3"], ws["B3"] = "b", 250, "a", 100
    wb.save(out_book)
    return True, None, "ok"


def _sort_translate(monkeypatch):
    monkeypatch.setattr(
        ailine, "translate_task",
        lambda model, task, book_meta, temperature=0.1:
        {"op": "SORT", "args": {"col": "金額", "order": "desc"}})


def _damage(path: Path) -> None:
    """障害注入: 原本を別の（開ける）内容に変える ── 途中で書き換わって失敗した状況の再現。"""
    wb = openpyxl.Workbook()
    wb.active.append(["壊れた", 0])
    wb.save(path)


# =====================================================================================
# S4 ── 「原本は変更していません」は、指紋を照合した時だけ言う
# =====================================================================================

def test_the_claim_follows_the_fingerprint(tmp_path):
    book = _nb(tmp_path)
    before = ailine._file_digest(book)
    assert ailine._untouched_claim(book, before) == "原本は変更していません"
    assert ailine._untouched_claim(book, before, verb="触っていません") == "原本は触っていません"
    _damage(book)
    got = ailine._untouched_claim(book, before)
    assert "変更していません" not in got and "変わっている可能性があります" in got, got
    # 取れない（走る前の指紋が無い／今のファイルが読めない）なら、言い切らない
    assert "照合できませんでした" in ailine._untouched_claim(book, None)
    assert "照合できませんでした" in ailine._untouched_claim(tmp_path / "無い.xlsx", before)
    assert "変更していません" not in ailine._untouched_claim(tmp_path / "無い.xlsx", before)


def test_a_run_that_failed_after_breaking_the_book_does_not_claim_it_is_untouched(
        tmp_path, monkeypatch, capsys):
    """★ 嘘になる場面（障害注入で再現）: 置換が原本を壊した上で失敗を返す。
    旧: 「（原本 nb.xlsx は変更していません。作業結果は nb.out.xlsx に残っています）」。"""
    _isolate(monkeypatch, tmp_path)
    book = _nb(tmp_path)
    _sort_translate(monkeypatch)
    monkeypatch.setattr(ailine, "basrun_apply", _sorted_apply)

    def broken_replace(b, out, workdir, keep_backups=3):
        _damage(b)
        return False, "置換に失敗した"
    monkeypatch.setattr(ailine, "atomic_replace_inplace", broken_replace)
    rc, out = _run_main(["run", str(book), "金額で降順に並べ替えて"], capsys)
    assert rc != 0, out
    assert "は変更していません" not in out, out
    assert "変わっている可能性があります" in out and "作業結果は nb.out.xlsx に残っています" in out, out


def test_a_run_that_failed_without_touching_the_book_still_says_so(tmp_path, monkeypatch, capsys):
    """★ 正しい場面では今までどおり言う（照合して一致した）。"""
    _isolate(monkeypatch, tmp_path)
    book = _nb(tmp_path)
    before = book.read_bytes()
    _sort_translate(monkeypatch)
    monkeypatch.setattr(ailine, "basrun_apply", _sorted_apply)
    monkeypatch.setattr(ailine, "atomic_replace_inplace",
                        lambda *a, **k: (False, "バックアップに失敗したため中止した"))
    rc, out = _run_main(["run", str(book), "金額で降順に並べ替えて"], capsys)
    assert rc != 0, out
    assert book.read_bytes() == before
    assert "（原本 nb.xlsx は変更していません。作業結果は nb.out.xlsx に残っています）" in out, out


def test_a_copy_run_does_not_claim_the_original_is_untouched_when_it_changed(
        tmp_path, monkeypatch, capsys):
    """★ 嘘になる場面: --copy で原本に触らないはずの run の途中で、原本が変わった（別のプロセス・
    Excel の保存など）。旧: 成功の見出しの下に「（原本 nb.xlsx は変更していません）」。"""
    _isolate(monkeypatch, tmp_path)
    book = _nb(tmp_path)
    _sort_translate(monkeypatch)

    def apply_and_someone_edits_the_book(out_book, code, workdir, helper_files=(), timeout=None):
        _sorted_apply(out_book, code, workdir, helper_files, timeout)
        _damage(book)
        return True, None, "ok"
    monkeypatch.setattr(ailine, "basrun_apply", apply_and_someone_edits_the_book)
    rc, out = _run_main(["run", str(book), "金額で降順に並べ替えて", "--copy"], capsys)
    assert "（原本 nb.xlsx は変更していません）" not in out, out
    assert "変わっている可能性があります" in out, out


def test_a_copy_run_that_left_the_book_alone_still_says_so(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book = _nb(tmp_path)
    _sort_translate(monkeypatch)
    monkeypatch.setattr(ailine, "basrun_apply", _sorted_apply)
    rc, out = _run_main(["run", str(book), "金額で降順に並べ替えて", "--copy"], capsys)
    assert rc == 0, out
    assert "（原本 nb.xlsx は変更していません）" in out, out


def _make_one_generation_then_edit(book: Path) -> None:
    """undo の材料: 世代を 1 つ積んでから、原本を別の中身にする。"""
    ailine.make_backup(book)
    _book(book, [["商品", "金額"], ["z", 999]])


def _fail_copy_into(book: Path, monkeypatch, *, tear: bool) -> None:
    """book への書き込みだけを失敗させる。tear=True は copy2 が書き先を開いた瞬間に切り詰めて
    から失敗する形（実物の挙動）── 原本は半端になる。"""
    real = shutil.copy2

    def copy2(src, dst, *a, **k):
        if Path(dst).resolve() == book.resolve():
            if tear:
                book.write_bytes(book.read_bytes()[:20])
            raise OSError(28, "ディスクの空きがありません")
        return real(src, dst, *a, **k)
    monkeypatch.setattr(ailine.shutil, "copy2", copy2)


def test_undo_that_tore_the_book_does_not_claim_it_is_untouched(tmp_path, monkeypatch, capsys):
    """★ 嘘になる場面: restore_backup は原本へ直接 copy2 する。途中で失敗すると原本は半端なのに、
    旧: 「× 復元に失敗しました（…）。原本は変更していません。」"""
    _isolate(monkeypatch, tmp_path)
    book = _nb(tmp_path)
    _make_one_generation_then_edit(book)
    _fail_copy_into(book, monkeypatch, tear=True)
    rc, out = _run_main(["undo", str(book)], capsys)
    assert rc == 1, out
    assert "復元に失敗しました" in out, out
    assert "原本は変更していません" not in out, out
    assert "変わっている可能性があります" in out, out


def test_undo_that_failed_before_writing_still_says_the_book_is_untouched(
        tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book = _nb(tmp_path)
    _make_one_generation_then_edit(book)
    _fail_copy_into(book, monkeypatch, tear=False)
    rc, out = _run_main(["undo", str(book)], capsys)
    assert rc == 1, out
    assert "復元に失敗しました" in out and "原本は変更していません。" in out, out


def test_adopt_that_stops_on_a_changed_book_says_it_touched_nothing_only_after_checking(
        tmp_path, monkeypatch, capsys):
    """★ 正しい場面: 下書きを作った後に原本が変わっていて止まる ── 両方とも触っていない（照合済み）。"""
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "b.xlsx", [["商品", "売上"], ["りんご", 100]])
    draft = _book(tmp_path / "b（下書き）.xlsx", [["商品", "売上"], ["みかん", 100]])
    rc, out = _run_main(["adopt", str(draft), str(book), "--base-sha", "0" * 64], capsys)
    assert rc == 7, out
    assert "原本は触っていません。下書きは触っていません" in out, out


def test_csv_untouched_line_is_compared_with_the_file_now(tmp_path):
    """「原本 CSV: 変更なし（sha256 … 一致）」の「一致」は、今のファイルとの照合の結果だけ。"""
    csv = tmp_path / "a.csv"
    csv.write_bytes("商品,金額\na,100\n".encode("utf-8"))
    ev = argparse.Namespace(sha256=hashlib.sha256(csv.read_bytes()).hexdigest())
    assert ailine._csv_untouched_line(csv, ev) == f"原本 CSV: 変更なし（sha256 {ev.sha256} 一致）"
    csv.write_bytes("商品,金額\na,999\n".encode("utf-8"))
    got = ailine._csv_untouched_line(csv, ev)
    assert "変更なし" not in got and "変わっている可能性があります" in got, got


def _invoices(tmp_path: Path) -> Path:
    folder = tmp_path / "請求書"
    folder.mkdir()
    for n in ("a", "b"):
        _book(folder / f"{n}.xlsx", [["請求書"], ["請求番号", f"No-{n}"], ["合計", 1000]])
    return folder


def test_forms_flags_an_input_that_changed_while_it_ran(tmp_path, monkeypatch, capsys):
    """★ e2e（障害注入）: forms の途中で入力の冊が変わったら、「1 バイトも変えていません」と言わない。"""
    _isolate(monkeypatch, tmp_path)
    folder = _invoices(tmp_path)
    real = ailine.record_made_book

    def record_and_someone_edits_an_input(path, kind, *a, **k):
        _damage(folder / "a.xlsx")
        return real(path, kind, *a, **k)
    monkeypatch.setattr(ailine, "record_made_book", record_and_someone_edits_an_input)
    rc, out = _run_main(["forms", str(folder), "--out", str(tmp_path / "一覧.xlsx")], capsys)
    assert "1 バイトも変えていません" not in out, out
    assert "入力の指紋が前後で違います: a.xlsx" in out, out


def test_forms_that_left_the_inputs_alone_still_says_so(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    folder = _invoices(tmp_path)
    rc, out = _run_main(["forms", str(folder), "--out", str(tmp_path / "一覧.xlsx")], capsys)
    assert rc == 0, out
    assert "元の請求書は 1 バイトも変えていません" in out, out
    assert "入力の指紋が前後で違います" not in out, out


def test_split_flags_an_input_that_changed_while_it_ran(tmp_path, monkeypatch, capsys):
    """★ e2e（障害注入）: split の途中で入力の冊が変わったら、「1 バイトも変えていません」と言わない。"""
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "list.xlsx", [["担当", "金額"], ["佐藤", 100], ["鈴木", 200], ["佐藤", 50]])
    out_dir = tmp_path / "out"
    real = ailine.record_made_book

    def record_and_someone_edits_the_input(path, kind, *a, **k):
        _damage(book)
        return real(path, kind, *a, **k)
    monkeypatch.setattr(ailine, "record_made_book", record_and_someone_edits_the_input)
    rc, out = _run_main(["split", str(book), "--by", "担当", "--out", str(out_dir)], capsys)
    assert "1 バイトも変えていません" not in out, out
    assert "入力の指紋が前後で違います" in out, out


def test_split_that_left_the_input_alone_still_says_so(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "list.xlsx", [["担当", "金額"], ["佐藤", 100], ["鈴木", 200], ["佐藤", 50]])
    rc, out = _run_main(["split", str(book), "--by", "担当", "--out", str(tmp_path / "out")], capsys)
    assert rc == 0, out
    assert "元の表は 1 バイトも変えていません" in out, out
    assert "入力の指紋が前後で違います" not in out, out


# =====================================================================================
# S1 ── 書き込む新しい値は、行の名指しではない
# =====================================================================================

NAMES = ["佐藤", "鈴木", "田中", "高橋", "伊藤"]


def _people(tmp_path: Path) -> Path:
    return _book(tmp_path / "t.xlsx", [["商品", "担当"]] + [[f"品{i}", n] for i, n in enumerate(NAMES)])


def _run_set_cell(tmp_path, monkeypatch, capsys, task, value):
    """3 行目（＝鈴木の行）の担当を書き換える run（翻訳と LibreOffice だけ差し替える）。"""
    _isolate(monkeypatch, tmp_path)
    book = _people(tmp_path)
    monkeypatch.setattr(
        ailine, "translate_task",
        lambda model, task_, book_meta, temperature=0.1:
        {"op": "SET_CELL_VALUE",
         "args": {"row": "", "row_number": 3, "col": "担当", "value": value}})

    def apply(out_book, code, workdir, helper_files=(), timeout=None):
        wb = openpyxl.load_workbook(out_book)
        wb.active["B3"] = value
        wb.save(out_book)
        return True, None, "ok"
    monkeypatch.setattr(ailine, "basrun_apply", apply)
    return _run_main(["run", str(book), task, "--copy"], capsys)


def test_a_new_value_that_exists_in_another_row_is_not_a_conflicting_row_name(
        tmp_path, monkeypatch, capsys):
    """★ 嘘になる場面（再現）: 「3行目の担当を『佐藤』に」── 佐藤は 2 行目に在るが、これは書く新しい値。
    旧: ⚠ 依頼は『佐藤』と行番号の両方を指していますが、その行に『佐藤』はありません ＋ ✓ が △ に降格。"""
    rc, out = _run_set_cell(tmp_path, monkeypatch, capsys, "3行目の担当を『佐藤』に", "佐藤")
    assert rc == 0, out
    assert "両方を指しています" not in out, out
    assert "✓" in out, out


def test_a_named_value_that_is_not_the_written_one_still_conflicts(tmp_path, monkeypatch, capsys):
    """★ 対: 値を書く依頼でも、**行の名指し**（高橋は 5 行目）が 3 行目と食い違えば今までどおり鳴る。"""
    rc, out = _run_set_cell(tmp_path, monkeypatch, capsys, "3行目の高橋の担当を『佐藤』に", "佐藤")
    assert "『高橋』と行番号の両方を指しています" in out, out
    assert "✓" not in out, out


@pytest.mark.parametrize("scope, want", [
    ("操作:1セル書換 対象の行:品1 対象列:担当 書き込む値:佐藤 入れる位置:3行目（依頼文の行番号）", {"佐藤"}),
    ("操作:条件つき書換 書き込む列:備考 書き込む値:済 条件を見る列:金額", {"済"}),
    ("操作:一括書換 対象列:備考 値:未定", {"未定"}),
    ("操作:行追加 挿入位置:2 入れる値:品名=棚／数量=3", {"棚", "3"}),
    # ★ 抽出の「値:」は条件の値で、書き込みではない
    ("操作:抽出 対象列:取引先 条件:等しい 値:ヤマノ食品", set()),
    ("", set()),
])
def test_written_values_are_read_from_the_declaration_only(scope, want):
    assert written_values_in_declaration(scope) == want


def test_the_pure_check_skips_only_the_written_values(tmp_path):
    book = _people(tmp_path)
    task = "3行目の担当を佐藤にして"
    assert value_not_in_the_named_row(task, 3, book) == "佐藤"                       # 対照（宣言なし）
    assert value_not_in_the_named_row(task, 3, book, written={"佐藤"}) is None       # 書く値は数えない
    assert value_not_in_the_named_row("3行目の高橋を佐藤にして", 3, book,
                                      written={"佐藤"}) == "高橋"                    # 名指しは数える


# =====================================================================================
# S2 ── 「（元の表はそのまま残っています）」は、書き換えていない時だけ
# =====================================================================================

def _finish(tmp_path, task, scope, op, capsys, *, rewrite_out=False, name="s2"):
    book = _book(tmp_path / f"{name}.xlsx", [["取引先", "金額"], ["丸和物流", 1000], ["山野", 500]])
    out = tmp_path / f"{name}.out.xlsx"
    shutil.copy2(book, out)
    if rewrite_out:                                    # 実行した op が既存の値を書き換えた
        wb = openpyxl.load_workbook(out)
        wb.active["B2"] = 0
        wb.save(out)
    work = tmp_path / f"w{name}"
    work.mkdir(exist_ok=True)
    a = argparse.Namespace(inplace=False, task=task, json=False, model="m",
                           keep_backups=3, copy=True)
    ailine._finish_apply(a, book, out, work, {"op": op}, machine_verified=True,
                         scope=scope, scope_note="", warning_count=0)
    return capsys.readouterr().out


def test_the_original_is_not_said_to_be_untouched_when_the_op_rewrote_it(tmp_path, capsys):
    """★ 嘘になる場面（再現）: 「行を削除して」に一括書換が走り、金額が書き換わった。
    旧: 「実行した操作は行や列を取り除きません（元の表はそのまま残っています）」。"""
    shown = _finish(tmp_path, "丸和物流の行を削除して", "操作:一括書換 対象列:金額 値:0",
                    "SET_COLUMN_VALUE", capsys, rewrite_out=True)
    assert "行や列を取り除きません" in shown, shown
    assert "元の表はそのまま" not in shown, shown


def test_the_original_is_said_to_be_untouched_when_the_op_only_made_a_new_sheet(tmp_path, capsys):
    """★ 正しい場面: 抽出は新しいシートを作るだけ（宣言）で、原本の中身も 1 つも違わない（実体）。"""
    shown = _finish(tmp_path, "丸和物流の行を削除して",
                    "操作:抽出 対象列:取引先 条件:等しい 値:丸和物流", "EXTRACT", capsys, name="ok")
    assert "行や列を取り除きません（元の表はそのまま残っています）" in shown, shown


def test_the_declaration_alone_is_not_enough_the_real_diff_is_read_too(tmp_path, capsys):
    """★ 宣言は「新しいシートを作るだけ」でも、実体で既存の値が変わっていたら言わない
    （宣言だけに寄りかかると、宣言が嘘の回に嘘を言う）。"""
    shown = _finish(tmp_path, "丸和物流の行を削除して",
                    "操作:抽出 対象列:取引先 条件:等しい 値:丸和物流", "EXTRACT", capsys,
                    rewrite_out=True, name="diff")
    assert "行や列を取り除きません" in shown, shown
    assert "元の表はそのまま" not in shown, shown


# =====================================================================================
# A2 ── 「存在しません」と「変更されていません」を言い分ける（列・行の枝）
# =====================================================================================

def _snap(rows: int = 2, cols: int = 3) -> dict:
    cells = {f"Sheet!{r},{c}": (f"v{r}{c}", "General", None, False, None, None)
             for r in range(1, rows + 1) for c in range(1, cols + 1)}
    return {"sheets": ["Sheet"], "charts": 0, "chart_counts": {"Sheet": 0}, "cells": cells,
            "merges": {"Sheet": []}, "colw": {"Sheet": {}}, "rowh": {"Sheet": {}},
            "truncated": False, "true_rows": {"Sheet": rows}}


def _after_with_c_changed(before: dict) -> dict:
    after = dict(before, cells=dict(before["cells"]))
    after["cells"]["Sheet!2,3"] = ("y", "General", None, False, None, None)
    return after


def _mentions(**kw) -> dict:
    return {"cols": set(), "digit_cols": set(), "rows": set(), "sheets": set(), **kw}


def test_a_column_or_row_that_exists_but_did_not_change_is_not_said_to_be_absent():
    before = _snap()
    after = _after_with_c_changed(before)
    lines = ailine.mention_overlap_advisory(
        _mentions(cols={2}, digit_cols={1}, rows={1}), before, after)
    assert lines == ["★ 依頼で言及された『列B』は変更されていません",
                     "★ 依頼で言及された『列1』は変更されていません",
                     "★ 依頼で言及された『行1』は変更されていません"], lines


def test_a_column_or_row_that_is_not_in_the_table_is_said_to_be_absent():
    before = _snap()
    after = _after_with_c_changed(before)
    lines = ailine.mention_overlap_advisory(
        _mentions(cols={26}, digit_cols={40}, rows={99}), before, after)
    assert lines == ["★ 依頼で言及された『列Z』は存在しません",
                     "★ 依頼で言及された『列40』は存在しません",
                     "★ 依頼で言及された『行99』は存在しません"], lines


def test_the_sheet_branch_still_says_it_apart_and_all_three_branches_share_one_wording():
    before = _snap()
    lines = ailine.mention_overlap_advisory(
        _mentions(sheets={"Sheet", "幻"}), before, before)
    assert lines == ["★ 依頼で言及された『Sheet』は変更されていません",
                     "★ 依頼で言及された『幻』は存在しません"], lines
    # 1 つの口を通る（枝ごとの書き写しが残っていない）
    assert ailine._mention_untouched_line("列B", True) == "★ 依頼で言及された『列B』は変更されていません"
    assert ailine._mention_untouched_line("列B", False) == "★ 依頼で言及された『列B』は存在しません"
    assert "存在しません/変更されていません" not in ailine._mention_untouched_line("x", True)
