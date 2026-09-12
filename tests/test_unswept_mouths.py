# 掃き出していない口（G3）── 明細と一致しても「裏が取れた」と言わない場合 ── 2026-09-12
#
# ★★ 明細の合計と一致しても、それは**明細という 1 つの口**を掃いただけだ。
#   器官の初版（34b3e90）が既に名指ししていた「開いていない第三の口」──
#   備考の金額の再掲・通貨が円でない・そもそも納品書か見積書か ── は別の口で、
#   どれも「一致しているのに本体と違う」を作る。
#
# ★★ この 3 冊（T28・T38・X03）が長く正しく見えていたのは、**一文字の取りこぼし**
#   （見出し `品番•品名` の `•` が U+2022 で照合に当たらない）が明細の掃き出し自体を
#   壊していたから ── 番人が間違った理由で効いていた。一文字を直すと 確 を over-claim する。
#   だから G1（中黒の同一視）と G3 は**一緒に**出す。
#
# ★ 語の有無では測れない: misoca のフッタに「無料のクラウド見積・納品・請求書サービス」が
#   在り、語で数えると偽陽性 29 冊（実測）。**構造**で測る（実測で偽陽性 0）。
from __future__ import annotations

import openpyxl
import pytest

from ailine_core.field_record import GRADES_WITH_VALUE, grade, value
from ailine_core.form_grid import Grid
from ailine_core.form_read import norm, read_form, unswept_mouths


def _read(ws):
    """★ 式のブックも渡す ── 渡さないと D4（写しを見分けられない）が先に効いて、
    この試験が見たい「掃き出していない口」の効果が測れない。この検体に式は無いので
    同じシートを渡すのが「式を見た（無かった）」の正直な写しになる。"""
    return read_form(ws, ws_formula=ws)


def _invoice(extra: dict | None = None):
    """明細と帯が整合した、ふつうの請求書（掃き出しが通る形）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求書"
    ws["B2"] = "請求書"
    ws["B3"] = "ナギ商会株式会社　御中"
    ws["G5"] = "株式会社あかね商事"
    ws["B11"] = "ご請求金額"
    ws["C11"] = 9900
    ws["B15"] = "品番•品名"          # ★ U+2022 の中黒（実物の形）
    ws["E15"] = "数量"
    ws["G15"] = "単価"
    ws["H15"] = "金額"
    ws["B16"] = "検査料"
    ws["E16"] = 1
    ws["G16"] = 7000
    ws["H16"] = 7000
    ws["B17"] = "梱包資材"
    ws["E17"] = 1
    ws["G17"] = 2000
    ws["H17"] = 2000
    ws["E37"] = "小計"
    ws["H37"] = 9000
    ws["E38"] = "消費税"
    ws["H38"] = 900
    ws["E39"] = "合計金額"
    ws["H39"] = 9900
    for at, v in (extra or {}).items():
        ws[at] = v
    return ws


def test_the_bullet_middle_dot_is_folded_so_the_detail_header_is_found():
    """★★ G1: `品番•品名`(U+2022) と `品番・品名`(U+30FB) は照合上 同じもの。

    ★ 実物のベンダー雛形に `•` が 10 セル・検体 87 冊に 31 セル在り、**NFKC でも寄らない**。
      寄せないと明細の見出しが 85 冊中 42 冊しか見つからず、掃き出しが半分の版面で黙る。
    """
    assert norm("品番•品名") == norm("品番・品名") == "品番・品名"
    assert norm("小･計") == "小・計"
    rec = _read(_invoice())["請求額"]
    assert grade(rec) == "確", grade(rec)          # 掃き出せている（陽性対照）
    assert value(rec) == 9900


def test_a_plain_invoice_has_no_unswept_mouth():
    """★ 陰性対照 ── ふつうの請求書で口を見つけたらオオカミ少年。"""
    assert unswept_mouths(Grid.read(_invoice()), 9900) == ()


@pytest.mark.parametrize("extra, want_in_reason", [
    ({"B44": "金額は USD 表示です"}, "通貨"),
    ({"B2": "納品書"}, "名乗って"),
    ({"B42": "合計金額 121,000 円（税込）をご請求申し上げます"}, "再掲"),
])
def test_each_mouth_stops_the_top_grade_but_keeps_the_value(extra, want_in_reason):
    """★ 値は出す・主張の強さだけ弱める（答えが 3 冊すべてで『単』を許している）。"""
    ws = _invoice(extra)
    rec = _read(ws)["請求額"]
    assert value(rec) == 9900, value(rec)
    assert grade(rec) in GRADES_WITH_VALUE and grade(rec) != "確", grade(rec)
    mouths = unswept_mouths(Grid.read(ws), 9900)
    assert mouths and any(want_in_reason in m for m in mouths), mouths


def test_the_currency_check_does_not_fire_on_the_word_for_issuer():
    """★★ `元` を通貨の印に入れない ── **『請求元』に当たる**（実測で誤爆した）。"""
    ws = _invoice({"G4": "請求元", "B45": "1 元 = 20 円"})
    assert not any("通貨" in m for m in unswept_mouths(Grid.read(ws), 9900))


def test_a_long_sentence_mentioning_another_form_is_not_a_document_title():
    """★ misoca のフッタ「無料のクラウド見積・納品・請求書サービス」で鳴らない
    （語で数えると偽陽性 29 冊だった）。**短いセル**だけを帳票名と見る。"""
    ws = _invoice({"B48": "無料のクラウド見積・納品・請求書サービス「Misoca」"})
    assert not any("名乗って" in m for m in unswept_mouths(Grid.read(ws), 9900))


def test_a_restatement_that_agrees_is_not_a_mouth():
    """★ 備考に再掲があっても**本体と同じ**なら口ではない（再掲の有無では測らない）。"""
    ws = _invoice({"B42": "合計金額 9,900 円（税込）をご請求申し上げます"})
    assert not any("再掲" in m for m in unswept_mouths(Grid.read(ws), 9900))


def test_a_discount_row_is_part_of_the_detail_not_the_start_of_the_band():
    """★★ G2: `値引き` は帯の開始語ではない ── 帯は 小計/合計/消費税 で開く。

    ★ 実測 T15（値引き行 −5,000 の陰性対照）: 明細から落ちて 110,000 になり、
      小計 105,000 と合わない**偽の食い違い**を出していた。
    """
    ws = _invoice()
    ws["B18"] = "値引き"
    ws["E18"] = 1
    ws["G18"] = -1000
    ws["H18"] = -1000
    ws["H37"] = 8000          # 小計 = 7000 + 2000 - 1000
    ws["H38"] = 800
    ws["H39"] = 8800
    ws["C11"] = 8800
    rec = _read(ws)["請求額"]
    assert value(rec) == 8800, (grade(rec), value(rec), rec.blank_reason[:120])
    assert grade(rec) == "確", (grade(rec), rec.blank_reason[:120])
