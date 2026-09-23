# -*- coding: utf-8 -*-
"""「〜を**除いて**」は対象から外す意味であって、削除ではない（2026-09-09）。

★★ 出所（盲検 C と D が**独立に**同じ所を指した ── 本命の証拠）:

    依頼   「合計行を**除いて**売上の多い順に並べ替えて」
    計画   1段目 行削除（合計行）＋ 2段目 並べ替え
    実物   ★ **合計行が消える**（判定は ⚠ なので ✓ は騙っていないが、undo が要る）

  ★ しかも並べ替え・抽出・条件つき書換は**元から合計行を外す**（`_skip_rows`・
    「データ行でないため並べ替えません」と画面に出る）。つまりこの削除段は
    「機械が既にやることを人が言い足しただけ」── 落としてよい。

★★ 計画だけでは区別できないと実測した（同じ 2 段が返り、順番すら安定しない）:

      「合計行を除いて…並べ替えて」      DELETE_ROWS + SORT
      「合計の行を消してから…並べ替えて」 DELETE_ROWS + SORT   ← ★ こちらは削除が正しい

  ★ だから**依頼文の語**で分ける。消す意図が明示された回は落とさない。

★ 盲検 C はこの回を「はい（頼んだ通り）」と誤判定し、俺も最初「false ⚠」と読み違えた。
  **行数を数えて**初めて分かった ── 数える対象を指定すると判定が締まる。
"""
from __future__ import annotations

import pytest

import ailine
from ailine_core.drop_redundant_delete import drop_delete_that_was_only_a_qualifier

DEL = {"op": "DELETE_ROWS", "args": {"at": 6}}
SORT = {"op": "SORT", "args": {"col": "売上", "order": "desc"}}


def _skips(op: str) -> bool:
    return op in ailine._OPS_THAT_SKIP_NON_DATA_ROWS


def _run(plan, task):
    return drop_delete_that_was_only_a_qualifier(plan, task, _skips)


@pytest.mark.parametrize("task", [
    "合計行を除いて売上の多い順に並べ替えて",
    "合計行以外を売上の多い順に並べ替えて",
    "合計を抜いて売上順にして",
])
def test_a_qualifier_drops_the_delete_step(task):
    got, note = _run([DEL, SORT], task)
    assert got == [SORT], got
    assert note and "外しました" in note


@pytest.mark.parametrize("task", [
    # ★ 消す意図が明示されている回は落とさない
    "合計の行を消してから売上順に並べ替えて",
    "合計行を削除して、売上順に並べ替えて",
    "合計行を取り除いてから並べ替えて",
])
def test_an_explicit_removal_is_kept(task):
    got, note = _run([DEL, SORT], task)
    assert got == [DEL, SORT], got
    assert note is None


def test_only_ops_that_skip_by_themselves_qualify():
    """★ 「元から対象外にする op」でなければ落とさない ── 行追加は合計行を外さない。"""
    got, note = _run([DEL, {"op": "ADD_ROW", "args": {}}], "合計行を除いて行を足して")
    assert len(got) == 2 and note is None


def test_the_declaration_matches_what_the_machine_actually_skips():
    """★ 宣言（_OPS_THAT_SKIP_NON_DATA_ROWS）が実体からずれていないか。

    ★ 実体の目印は「データ行でないため…ません」を出す解決関数。宣言だけが増えると、
      合計行を外さない op の削除段まで落として**本当に必要な削除を消す**。

    ★★ 2026-09-16: ここは名簿の中身を**手書きで凍結**していた
      （`== frozenset({"SORT","EXTRACT","SET_WHERE"})`）。宣言を書き写しただけなので
      **実体とのずれは原理的に見えない** ── 実際 AGGREGATE が名簿から抜けていて
      「合計行を除いて部署別に集計して」が利用者の合計行を消していたのに、ここは緑だった。
      ★ 手書きをやめ、**実装から導いた集合**と突き合わせる番人へ畳んだ
        （tests/test_op_completeness.py の⑤）。縛りは 1 本にして、こちらは委譲する。
    """
    from _product_source import count_in_product
    from test_op_completeness import ops_that_skip_non_data_rows_derived_from_codegen
    assert count_in_product("データ行でないため") >= 3, "実体側の目印が減っている"
    assert (ailine._OPS_THAT_SKIP_NON_DATA_ROWS
            == ops_that_skip_non_data_rows_derived_from_codegen())


# --- 実機（両方向）----------------------------------------------------------

@pytest.mark.local
@pytest.mark.parametrize("task, keeps_total", [
    ("合計行を除いて売上の多い順に並べ替えて", True),
    ("合計の行を消してから売上順に並べ替えて", False),
])
def test_it_behaves_on_real_libreoffice(tmp_path, task, keeps_total):
    import os
    import subprocess
    import sys
    from pathlib import Path

    import openpyxl

    repo = Path(ailine.__file__).resolve().parents[2]
    book = tmp_path / "u.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "売上表"
    ws.append(["商品", "件数", "単価", "売上"])
    for name, cnt, tanka in (("りんご", 12, 100), ("みかん", 8, 100), ("ぶどう", 5, 300)):
        ws.append([name, cnt, tanka, cnt * tanka])
    ws.append(["合計", None, None, "=SUM(D2:D4)"])
    wb.save(book)

    r = subprocess.run(
        [sys.executable, "-m", "ailine", "run", str(book), task, "--copy",
         "--sheet", "売上表", "--timeout", "300"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
        env={**os.environ, "PYTHONPATH": str(repo / "src")})
    # ★★ 2026-09-20: ここは `returncode == 0` だけを見ており、**製品が持っていない
    #   決定性**を暗に要求していた。実測: 同じ HEAD ・同じ冊・同じ文で 5 回振って
    #   **4 回は exit 0、1 回は exit 3**（読みが分かれた断り）だった。
    #   断った回の二つの読みは**本当に違う**:
    #       読み 1: 並べ替え → 行削除  （合計行は残る）
    #       読み 2: 行削除 → 並べ替え  （A5 '合計'→空）
    #   ★ つまりこれは**検出器が正しく働いた**回であって、退行ではない。
    #     黙って片方を選ぶのが一番悪い形 ── それを塞ぐために入れた番人だ。
    #
    # ★★ だから縛り方を分ける（緑にするために緩めない）:
    #   ① 硬い契約（どちらの回でも破ってはいけない）: **黙って合計行を消さない**。
    #   ② 断った回は、**両方の読みを見せて選ばせている**ことまで確かめる
    #     （合格線②「適切な断り」の形。ただ断って終わりならここは赤いまま）。
    if r.returncode == 3 and "読み方が分かれました" in r.stdout:
        assert "読み方 1:" in r.stdout and "読み方 2:" in r.stdout, r.stdout[-800:]
        assert "--op" in r.stdout, "★ 選び方（通る道）を示していない: " + r.stdout[-800:]
        assert not book.with_name(book.stem + ".out.xlsx").exists(),             "★ 断ったのに .out を残している（合格線⑤: 1 バイトも書かない）"
        return
    assert r.returncode == 0, r.stdout[-800:]
    out = openpyxl.load_workbook(book.with_name(book.stem + ".out.xlsx"))["売上表"]
    names = [out.cell(i, 1).value for i in range(1, out.max_row + 1)]
    assert ("合計" in names) is keeps_total, (task, names, r.stdout[-500:])
