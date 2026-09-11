# ラベルが**表示形式（number_format）に焼き込まれた**帳票の番人 ── 2026-09-12
#
# ★★ 実物の癖: 画面には「請求日： 2026年6月30日」と見えるのに、
#   セルの中身は日付型だけで、**『請求日』という文字はどのセルにも無い**。
#   ラベルは表示形式 `"請求日： "yyyy\年m\月d\日;@` の中に在る。
#   実測（2026-09-12）: 検体 87 冊のうち 13 冊・実物の雛形にも在る ── 珍しい癖ではない。
#
# ★ 拾いすぎの危険も測ってある: ラベル語を含む表示形式は、全 180 冊で
#   日付型／数値のセルにしか付いていなかった（`円` を焼き込んだ書式 506 件は別語）。
from __future__ import annotations

import datetime as dt

import openpyxl

from ailine_core.field_record import grade, value
from ailine_core.form_grid import Grid
from ailine_core.form_read import (format_label, read_invoice_number,
                                   read_issue_date)

DATE_FMT = '"請求日： "yyyy\\年m\\月d\\日;@'
ISSUED_FMT = '"発行日： "yyyy"年"m"月"d"日";@'     # ★ 実物 inv21 はこちらの書き方
NUMBER_FMT = '"請求書番号： "0_);[RED]\\(0\\)'


def _book(cells: dict):
    """cells: 番地 → (値, 表示形式)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "請求書"
    for at, (v, fmt) in cells.items():
        ws[at] = v
        if fmt:
            ws[at].number_format = fmt
    return ws


def test_the_literal_text_is_pulled_out_of_the_display_format():
    assert format_label(DATE_FMT) == "請求日： "
    assert format_label(ISSUED_FMT) == "発行日： 年月日"     # ★ 実物の書き方は区切りも引用される
    assert format_label(NUMBER_FMT) == "請求書番号： "
    assert format_label("General") == ""
    assert format_label(None) == ""
    assert format_label('#,##0"円"') == "円"


def test_the_grid_carries_the_display_format():
    """★ 器官の層（Grid）が表示形式を運ぶ ── ここが空だと上の規則は永久に発火しない。"""
    g = Grid.read(_book({"H20": (dt.date(2026, 6, 30), DATE_FMT)}))
    assert g.cell(20, 8).fmt == DATE_FMT


def test_a_merged_cell_takes_the_format_of_its_anchor():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "請求書"
    ws["H20"] = dt.date(2026, 6, 30)
    ws["H20"].number_format = DATE_FMT
    ws.merge_cells("H20:J20")
    g = Grid.read(ws)
    assert g.cell(20, 10).fmt == DATE_FMT, "★ 結合の右側が書式を失っている"


def test_a_date_whose_label_lives_in_the_format_is_read():
    rec = read_issue_date(Grid.read(_book({"H20": (dt.date(2026, 6, 30), DATE_FMT)})))
    assert grade(rec) == "単", grade(rec)
    assert value(rec) == dt.date(2026, 6, 30)
    assert "表示形式" in rec.evidences[0].how, rec.evidences[0].how


def test_two_baked_labels_that_agree_are_one_value():
    """★ 実物 inv21 は 発行日（H6）と 請求日（H20）を両方持ち、同じ日が入っている。"""
    rec = read_issue_date(Grid.read(_book({
        "H6": (dt.date(2026, 6, 30), ISSUED_FMT),
        "H20": (dt.date(2026, 6, 30), DATE_FMT)})))
    assert value(rec) == dt.date(2026, 6, 30)
    assert len(rec.evidences) == 2, rec.evidences


def test_two_baked_labels_that_disagree_are_left_blank():
    """★ 食い違うなら決めない ── 空欄＋理由（値を作らない側に倒す）。"""
    rec = read_issue_date(Grid.read(_book({
        "H6": (dt.date(2026, 5, 1), ISSUED_FMT),
        "H20": (dt.date(2026, 6, 30), DATE_FMT)})))
    assert grade(rec) == "割", grade(rec)
    assert value(rec) is None
    assert "食い違" in rec.blank_reason, rec.blank_reason


def test_a_bare_number_behind_a_baked_label_is_reported_as_an_identifier():
    """★ 請求番号は識別子であって数ではない ── 文字で出す（20260601.0 にしない）。"""
    rec = read_invoice_number(Grid.read(_book({"H19": (20260601, NUMBER_FMT)})))
    assert grade(rec) == "単"
    assert value(rec) == "20260601", repr(value(rec))
    rec2 = read_invoice_number(Grid.read(_book({"H19": (20260601.0, NUMBER_FMT)})))
    assert value(rec2) == "20260601", repr(value(rec2))


def test_a_format_without_a_label_is_not_mistaken_for_one():
    """★ 陰性対照 ── 金額の書式（`#,##0"円"`）や素の日付書式は拾わない。"""
    ws = _book({"C11": (3300, '#,##0"円"'), "H20": (dt.date(2026, 6, 30), "yyyy/m/d")})
    assert grade(read_issue_date(Grid.read(ws))) == "無"
    assert grade(read_invoice_number(Grid.read(ws))) == "無"


def test_a_cell_text_label_still_wins_its_own_cell():
    """★ 文字のラベルと焼き込みが同じセルに重なっても、証拠を二重に数えない。"""
    ws = _book({"G3": ("請求日：", None), "H3": (dt.date(2026, 8, 31), DATE_FMT)})
    rec = read_issue_date(Grid.read(ws))
    assert value(rec) == dt.date(2026, 8, 31)
    assert len(rec.evidences) == 1, [e.how for e in rec.evidences]
