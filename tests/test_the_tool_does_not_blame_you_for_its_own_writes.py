# -*- coding: utf-8 -*-
"""道具が自分でやったことを、人のせいにして止めない（2026-09-18・盲検 3 体目 ⑥）。

★★ 起きたこと（製造業の購買担当・初見・盲検）:

    ⚠ 出力先に書けません: …発注台帳.out.xlsx
    （…は ailine が作った物ですが、**そのあと変更されています**。作業内容が消えるので
      上書きしません ── 別の場所へ移すか削除してから、もう一度実行してください）
    EXIT=7

  買い手:「私は `.out` を**読んだだけ**です」「しかも直前の同じ状況では通っていた ──
  **同じ状況で通ったり止まったり**する」。

★★ 再現して分かった筋（実測・2026-09-18）:
    ① `--copy` で `<stem>.out.xlsx` を作る        → 履歴にその指紋を残す
    ② 原本直接で走る                              → **ailine 自身が .out を作業ファイルに
                                                     使って書き換える**。だが履歴には
                                                     **原本のパスしか残らない**
    ③ もう一度 原本直接で走る                      → .out の最新の指紋は①のままで、
                                                     いま在る物と合わない → 「人が変えた」
  ★ 変えたのは**この道具自身**。買い手は正しかった。
  ★「通ったり止まったり」も説明がつく ── 間に原本直接の run が挟まったかで変わる。

★★ 同じ誤りの 2 回目である: 2026-08-26 に「印が無い＝**人のファイル**です」と**断定**して
  間違えた（実際は ailine が数分前に作った物）。その時の処方が同じ関数の隣に書いてある ──
  **「見たものと、その解釈を分ける」**。分かるのは「指紋が違う」まで。誰が変えたかは断定しない。

★ 直しは文言を緩めるのではなく**実体を記録する**（Namakoo 決裁 A・「断りは仕方なく
  断るにすぎない」）── 止める必要が無いなら止めない。合格線でいえば「適切な断り」でなく
  「**到達**」へ移す直し方。
★ 履歴の行は増やさない ── 買い手が欲しいのは月次の証跡で、そこに雑音を足さない。
  その run の記録に scratch_out / scratch_out_sha の欄を足す。

★★ そして**守るべき線は残す**: 2026-08-25 に「`--copy` の成果物が原本反映 1 回で警告なしに
  消えた」事故があり、この指紋はその再発防止で入っている。**人が育てた `.out` は今までどおり
  止める**。分けたのは「道具自身の書き換え」と「人の手」。
"""
import contextlib
import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402

ROWS = [["取引先", "金額", "担当"], ["甲社", 1700, "高橋"],
        ["乙社", 2400, "田中"], ["丙社", 900, "伊藤"]]

pytestmark = pytest.mark.local   # ★ 実機（LibreOffice）を通す


@pytest.fixture
def book(tmp_path, monkeypatch):
    monkeypatch.setattr(ailine, "HISTORY_FILE", tmp_path / "history.jsonl")
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "売上"
    for r in ROWS:
        ws.append(list(r))
    wb.save(p)
    return p


def _run(book, *argv):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = ailine.main(["run", str(book), *argv])
    return rc, buf.getvalue()


def test_the_tool_does_not_blame_you_for_its_own_scratch_write(book):
    """★★ 事故そのもの: ①--copy → ②原本直接 → ③原本直接 が止まらないこと。"""
    rc, out = _run(book, "担当ごとに金額を集計して", "--copy")
    assert rc == 0, out
    rc, out = _run(book, "担当が高橋の行を抜き出して")
    assert rc == 0, out
    rc, out = _run(book, "金額の大きい順に並べ替えて")
    assert rc == 0, (
        "★ 道具が自分で書き換えた .out を『人が変えた』と言って止めている:\n" + out)
    assert "そのあと変更されています" not in out, out


def test_a_human_edit_is_still_refused(book):
    """★★ 守るべき線（2026-08-25 の事故の再発防止）は生きていること。

    ★ ここが無いと、上の直しが「関所を外した」だけになる ── 人が育てた `.out` が
      警告なしに消える形へ逆戻りする。
    """
    rc, out = _run(book, "担当ごとに金額を集計して", "--copy")
    assert rc == 0, out
    o = book.with_name(book.stem + ".out" + book.suffix)
    wb = openpyxl.load_workbook(o)
    wb.create_sheet("人が足したシート")["A1"] = "大事なメモ"
    wb.save(o)
    rc, out = _run(book, "金額の大きい順に並べ替えて")
    assert rc == 7, f"人が手を入れた .out を黙って上書きした:\n{out}"
    assert "そのあと変更されています" in out, out


def test_the_scratch_fingerprint_is_recorded(book, tmp_path):
    """★ 実体の項が本当に残ること（残らなければ上の直しは効かない）。"""
    _run(book, "担当ごとに金額を集計して", "--copy")
    _run(book, "担当が高橋の行を抜き出して")
    rows = [json.loads(ln) for ln in
            (tmp_path / "history.jsonl").read_text(encoding="utf-8").splitlines() if ln.strip()]
    stamped = [e for e in rows if e.get("scratch_out")]
    assert stamped, f"原本直接の run が作業に使った .out の指紋を残していない: {rows}"
    e = stamped[-1]
    assert e["scratch_out"].endswith(".out.xlsx"), e
    assert e["scratch_out_sha"], e
    o = Path(e["scratch_out"])
    assert e["scratch_out_sha"] == ailine._file_digest(o), (
        "記録した指紋が、残っている実体と合っていない")


def test_a_run_without_a_leftover_out_stamps_nothing(book, tmp_path):
    """★★ 陰性対照: `.out` が前から在なかった run は、指紋の欄を残さない。

    ★ 自分で作った `.out` は原子的置換のあとに消えるので、記録する物がそもそも無い。
      ここが無いと「常に残す」実装でも上の試験が通る ── そして常に残すと、
      **自分で作って消した物の指紋**が履歴に残り、次の run の判定材料を濁す。
    ★ 2026-09-18 の変異試験で、条件の片方（自分の作業か）が赤にならなかった
      （消える順序で既に保証されていた）。その枝は消し、**代わりにここで縛る**。
    """
    rc, out = _run(book, "担当が高橋の行を抜き出して")
    assert rc == 0, out
    rows = [json.loads(ln) for ln in
            (tmp_path / "history.jsonl").read_text(encoding="utf-8").splitlines() if ln.strip()]
    assert rows, "履歴が無い"
    assert not [e for e in rows if e.get("scratch_out")], (
        f"残っていない .out の指紋を記録している: {rows}")


def test_the_message_says_the_contents_changed(book):
    """★ 文言が実体に合っていること ──「消していません」だけだと中身も無事と読める。"""
    _run(book, "担当ごとに金額を集計して", "--copy")
    _rc, out = _run(book, "担当が高橋の行を抜き出して")
    assert "中身は今回の結果に変わっています" in out, (
        "『消していません』だけで、中身が入れ替わったことを言っていない:\n" + out)
