# 入れ替える 2 つは依頼文が名指ししたものに限る（2026-10-01・依頼の項の台帳で B だった項目）。
#
# ★ 形: a/b は「実在の列名/行の名前か」しか見ていなかった（読み直しの道だけが依頼文に在ることを
#   要求していた）。実走行の実例:「税込み金額の順番を逆にして」→ 税込み金額 ⇄ 締め日
#   （締め日は依頼に無い）で ✓ が出ていた。
#
# 契約:
#   ① 依頼文が 2 つとも名指ししていれば黙る（列名・行の値・「C列」・「3行目」）
#   ② 名指しの無い方があれば ⚠（_warnings ＝ ✓→△ の材料）で、その名前を出す
#   ③ 候補が実表で 2 つしか無い（列が 2 つ・データ行が 2 行）なら入れ替えは 1 通り ── 黙る
#   ④ セルの入れ替え（座標は依頼文と実表から機械が解く）は変えない
#   ⑤ 見出しが 1 行目でない表でも同じ

import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402
from ailine_core import argcheck  # noqa: E402

HEAD = ["品名", "棚", "数量", "締め日"]
DATA = [["ナット", "A-1", 7, "月末"], ["ボルト", "A-2", 6, "20日"], ["ワッシャ", "B-1", 9, "月末"]]


def _book(tmp_path, head=HEAD, data=DATA, header_row=1):
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "在庫"
    for i in range(header_row - 1):
        ws.append(["部品在庫"] if i == 0 else [])
    ws.append(head)
    for r in data:
        ws.append(r)
    wb.save(p)
    return p


def _swap(tmp_path, task, a, b, head=HEAD, data=DATA, header_row=1):
    meta = {"sheets": ["在庫"], "headers": {"在庫": list(head)},
            "header_rows": {"在庫": header_row},
            "path": str(_book(tmp_path, head, data, header_row))}
    ok, r, _i, err = ailine.verify_dsl_args("SWAP", {"a": a, "b": b}, meta, task=task)
    assert ok, err
    return r


def _swap_warnings(r):
    return [w for w in r.get("_warnings", []) if "を指す語が見当たりません" in w]


# --- ① 名指しされていれば黙る ------------------------------------------------------------

def test_named_columns_are_silent(tmp_path):
    r = _swap(tmp_path, "棚と数量の列を入れ替えて", "棚", "数量")
    assert r["_axis"] == "column" and not _swap_warnings(r)


def test_columns_named_by_letter_are_silent(tmp_path):
    r = _swap(tmp_path, "B列とC列を入れ替えて", "棚", "数量")
    assert r["_axis"] == "column" and not _swap_warnings(r)


def test_named_rows_are_silent(tmp_path):
    r = _swap(tmp_path, "ナットとボルトを入れ替えて", "ナット", "ボルト")
    assert r["_axis"] == "row" and not _swap_warnings(r)


def test_rows_named_by_another_cell_of_them_are_silent(tmp_path):
    r = _swap(tmp_path, "A-1の行とB-1の行を入れ替えて", "ナット", "ワッシャ")
    assert r["_axis"] == "row" and not _swap_warnings(r)


# --- ② 名指しの無い方は ⚠ -----------------------------------------------------------------

def test_the_column_the_request_did_not_name_is_disclosed(tmp_path):
    """実走行の実例の形: 依頼は 1 列しか言っていない。"""
    r = _swap(tmp_path, "数量の順番を逆にして", "数量", "締め日")
    ws = _swap_warnings(r)
    assert len(ws) == 1 and "『締め日』" in ws[0] and "列" in ws[0]


def test_the_row_the_request_did_not_name_is_disclosed(tmp_path):
    r = _swap(tmp_path, "ナットを上に入れ替えて", "ナット", "ワッシャ")
    ws = _swap_warnings(r)
    assert len(ws) == 1 and "『ワッシャ』" in ws[0] and "行" in ws[0]


# --- ③ 候補が 2 つしか無ければ黙る ----------------------------------------------------------

def test_a_two_column_table_has_only_one_swap(tmp_path):
    r = _swap(tmp_path, "列を入れ替えて", "品名", "棚", head=["品名", "棚"],
              data=[["ナット", "A-1"], ["ボルト", "A-2"]])
    assert r["_axis"] == "column" and not _swap_warnings(r)


def test_a_two_row_table_has_only_one_swap(tmp_path):
    r = _swap(tmp_path, "行を入れ替えて", "ナット", "ボルト", data=DATA[:2])
    assert r["_axis"] == "row" and not _swap_warnings(r)


# --- ⑤ 見出しが 1 行目でない表 -------------------------------------------------------------

def test_header_not_on_the_first_row(tmp_path):
    r = _swap(tmp_path, "ナットを上に入れ替えて", "ナット", "ワッシャ", header_row=3)
    assert r["_axis"] == "row" and r["_b_pos"] == 6 and _swap_warnings(r)
    r2 = _swap(tmp_path, "ナットとワッシャを入れ替えて", "ナット", "ワッシャ", header_row=3)
    assert not _swap_warnings(r2)


# --- 変異: 配線を外すと黙る ---------------------------------------------------------------

def test_without_the_wiring_the_guess_would_pass_silently(tmp_path, monkeypatch):
    monkeypatch.setattr(argcheck, "_swap_targets_unnamed", lambda *a, **k: [])
    r = _swap(tmp_path, "数量の順番を逆にして", "数量", "締め日")
    assert r["_axis"] == "column" and not _swap_warnings(r)     # ← 直す前の挙動
