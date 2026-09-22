"""「見出し」と言われたら、**機械が**見出し行を決める（2026-09-22）。

★★ 事故（盲検 6 体目を追って再現）: `--header-row 3` を渡した冊に
  「見出しを太字にして」と頼むと ──

    解釈: 操作:太字 対象:row:1
    検算しました（太字）: 6 セルが太字
    ⚠ …機械保証はありません          ← **exit 0**

  実物を開くと **1 行目（空のタイトル行）が太字**で、見出し行（3 行目）は素のまま。
  道具は「頼まれた見出し」を太字にしていない。

★★ この族は**一度踏んでいる**（`ailine_core/subject.py` の冒頭に記録がある）──
  「見出しを太字にして」が `col:数量*単価` に解決され、**見出し行は太字にならないまま
  ✓ が出た**。その時の処置は「⚠ で言う」止まりだった。
  ★ 道具は今回も気づいていた（`⚠ 依頼文が指しているのは: 見出し`）。
    **言えているのに直していない** ── 今日は機械が決める側へ進めた。

★ なぜ模型に決めさせないか: few-shot の例が `{"target": "row:1"}` と書いてあり、
  模型はそれを**写している**だけ。見出しが何行目かは表を見ないと決まらない。
  ★ `_OP_SCHEMA_NOTES` が他の op で宣言している分担と同じ ──
    「これは入れない、機械が決める」。

★ 黙って直さない: 出所を `_sources["target"]` に残し、画面の「解釈:」行に出す。

    解釈: 操作:太字 対象:row:3(推定)（依頼文の『見出し』と、この冊の見出し行（3行目）
    から機械が決めました（模型は row:1 と言っていました））
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import ailine  # noqa: E402


@pytest.mark.parametrize("task", [
    "見出しを太字にして",
    "ヘッダーを太字にして",
    "項目名を太字にして",
    "見出し行に背景色を付けて",
])
def test_the_machine_decides_when_the_task_says_header(task):
    """★★ 本体 ── 依頼が「見出し」と言い、見出しが 1 行目でないなら機械が決める。"""
    assert ailine.header_row_the_task_means(task, 3) == 3, task


@pytest.mark.parametrize("task, header_row", [
    ("見出しを太字にして", 1),      # ★ 見出しが 1 行目なら直すものが無い
    ("1行目を太字にして", 3),       # ★ 人が行番号を言っている ── 横取りしない
    ("3行目を太字にして", 3),
    ("金額で降順に並べ替えて", 3),  # ★ 「見出し」と言っていない
    ("", 3),
])
def test_it_does_not_take_over_other_requests(task, header_row):
    """★ 陰性対照 ── 人が行番号を言った回と、見出しが 1 行目の回は触らない。"""
    assert ailine.header_row_the_task_means(task, header_row) is None, (task, header_row)


def test_it_does_not_read_the_book_again():
    """★ 見出し行は**呼び出し側が既に知っている**値を使う（冊を二度読むと割れる）。

    ★ `resolve_header_rows` が 1 箇所で決めた値が渡ってくる ── ここで読み直すと、
      同じ冊について 2 つの答えを持つことになる（この repo が何度も踏んだ形）。
    """
    import inspect
    src = inspect.getsource(ailine.header_row_the_task_means)
    # ★★ 「無いこと」だけを見る assert は、探す場所が空でも通る（関数が消えても緑）。
    #   番人の台帳がそれを止めた ── だから**在ること**と対にする。
    assert "header_row" in src, "★ 見出し行を受け取っていない（探す場所が空）"
    assert "return header_row" in src, "★ 受け取った値を返していない"
    assert "book_meta" not in src, "冊を読み直している"
    assert "BookView" not in src and "load_workbook" not in src, "冊を開いている"


def test_the_words_are_a_list_and_that_is_on_purpose():
    """★ 語の列挙で正しい ── 判定しているのが「人がどの場所を言ったか」だから。

    ★ 漏れた時の壊れ方: 語が無ければ**今までどおり**（模型の言うとおりになる）。
      黙って別の行を触るのではない。
    """
    assert "見出し" in ailine._HEADER_WORDS
    assert ailine.header_row_the_task_means("そんな語は無い依頼", 3) is None


@pytest.mark.local
def test_the_real_run_bolds_the_header_not_the_title(tmp_path):
    """★★ 実機 ── 買い手の冊と同じ形で、**見出し行が**太字になること。"""
    import os
    import subprocess

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "精算"
    ws["B1"] = "2026年8月 経費精算書"
    for i, h in enumerate(["社員", "日付", "科目", "金額"]):
        ws.cell(row=3, column=2 + i).value = h
    for r, row in enumerate([("佐藤", "2026-08-02", "旅費交通費", 1280),
                             ("鈴木", "2026-08-07", "会議費", 3500)], start=4):
        for i, v in enumerate(row):
            ws.cell(row=r, column=2 + i).value = v
    p = tmp_path / "in.xlsx"
    wb.save(p)

    repo = Path(ailine.__file__).resolve().parents[2]
    out = tmp_path / "out.xlsx"
    subprocess.run(
        [sys.executable, "-m", "ailine", "run", str(p), "見出しを太字にして",
         "--header-row", "3", "--copy", "--out", str(out), "--timeout", "150"],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
        cwd=str(repo), env={**os.environ, "PYTHONPATH": str(repo / "src"),
                            "AILINE_HOME": str(tmp_path / "home")})
    got = openpyxl.load_workbook(out)["精算"]
    assert got.cell(row=3, column=2).font.bold, "見出し行が太字になっていない"
    assert not got.cell(row=1, column=2).font.bold, "タイトル行を太字にしている"
