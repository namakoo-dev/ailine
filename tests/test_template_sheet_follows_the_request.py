# 雛形のシートは依頼文と突き合わせる（2026-10-01・依頼の項の台帳で B だった 2 項目:
# REPORT_PER_ROW.template_sheet・FORMAT_MAP.template_sheet）。
#
# ★ 形: 雛形のシートは「実在するか・印（{{列名}}）が在るか」しか見ていなかった。雛形が 2 枚ある冊
#   （日本語と英文）で、依頼が名指ししない方を LLM が選んでも ✓ が出る。
#
# 契約:
#   ① 依頼文が名指ししていれば黙る（『雛形』と『雛形_英文』の断片は証拠にしない）
#   ② 名指しが無く、印を持つシートが 1 枚なら機械が決める ── LLM と違えば書き換え、出典を出す
#   ③ 名指しが無く、印を持つシートが 2 枚以上なら ⚠（_warnings ＝ ✓→△ の材料）で他の候補を名指しする
#   ④ 見出しが 1 行目でないデータでも同じ（判定はシート名と雛形の印だけを見る）

import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402
from ailine_core import argcheck  # noqa: E402

HEAD = ["取引先", "品目", "金額"]
DATA = [["山田商事", "ボルト", 1000], ["丸和物流", "ナット", 2000]]


def _book(tmp_path, templates=("雛形", "雛形_英文"), header_row=1, memo=True):
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求データ"
    for i in range(header_row - 1):
        ws.append(["8月の請求"] if i == 0 else [])
    ws.append(HEAD)
    for r in DATA:
        ws.append(r)
    for name in templates:
        t = wb.create_sheet(name)
        t["A1"], t["B1"] = "取引先", "金額"
        t["A2"], t["B2"] = "{{取引先}}", "{{金額}}"
    if memo:
        wb.create_sheet("メモ")["A1"] = "連絡事項"
    wb.save(p)
    return p, wb.sheetnames


def _run(tmp_path, op, task, template, templates=("雛形", "雛形_英文"), header_row=1):
    path, sheets = _book(tmp_path, templates, header_row)
    meta = {"sheets": sheets, "headers": {"請求データ": list(HEAD)},
            "header_rows": {"請求データ": header_row}, "path": str(path)}
    args = {"template_sheet": template}
    if op == "REPORT_PER_ROW":
        args["name_col"] = "取引先"
    ok, r, inferred, err = ailine.verify_dsl_args(op, args, meta, task=task,
                                                   target_sheet="請求データ")
    assert ok, err
    return r, inferred


def _tpl_warnings(r):
    return [w for w in r.get("_warnings", []) if "雛形のシート" in w]


OPS = ("REPORT_PER_ROW", "FORMAT_MAP")


# --- ① 名指しされていれば黙る ------------------------------------------------------------

def test_named_template_is_silent(tmp_path):
    for op in OPS:
        r, _i = _run(tmp_path, op, "雛形シートで取引先ごとに写して", "雛形")
        assert r["template_sheet"] == "雛形" and not _tpl_warnings(r), op
        r, _i = _run(tmp_path, op, "英文の雛形で取引先ごとに写して", "雛形_英文")
        assert r["template_sheet"] == "雛形_英文" and not _tpl_warnings(r), op


# --- ② 印を持つシートが 1 枚なら機械が決める ---------------------------------------------------

def test_the_only_template_is_decided_by_the_table(tmp_path):
    for op in OPS:
        r, inferred = _run(tmp_path, op, "取引先ごとに写して", "雛形", templates=("雛形",))
        assert r["template_sheet"] == "雛形" and not _tpl_warnings(r), op
        assert "template_sheet" not in r.get("_sources", {}), op


def test_the_only_template_wins_over_a_sheet_without_marks(tmp_path):
    """LLM が印の無いシートを選んだ ── 印を持つ唯一のシートへ機械が直し、出典を出す。"""
    for op in OPS:
        r, inferred = _run(tmp_path, op, "取引先ごとに写して", "メモ", templates=("雛形",))
        assert r["template_sheet"] == "雛形", op
        assert "『雛形』だけ" in r["_sources"]["template_sheet"], op
        assert not _tpl_warnings(r), op


# --- ③ 2 枚以上なら ⚠ -----------------------------------------------------------------------

def test_two_templates_without_a_name_are_disclosed(tmp_path):
    for op in OPS:
        r, _i = _run(tmp_path, op, "取引先ごとに写して", "雛形")
        ws = _tpl_warnings(r)
        assert len(ws) == 1 and "『雛形』" in ws[0] and "『雛形_英文』" in ws[0], op
        assert r["template_sheet"] == "雛形", op            # 書き換えない


# --- ④ 見出しが 1 行目でないデータ ----------------------------------------------------------

def test_header_not_on_the_first_row(tmp_path):
    r, _i = _run(tmp_path, "REPORT_PER_ROW", "取引先ごとに写して", "雛形", header_row=3)
    assert _tpl_warnings(r)
    r2, _i = _run(tmp_path, "REPORT_PER_ROW", "雛形シートで取引先ごとに写して", "雛形", header_row=3)
    assert not _tpl_warnings(r2)


# --- 変異: 配線を外すと黙る -----------------------------------------------------------------

def test_without_the_wiring_the_guess_would_pass_silently(tmp_path, monkeypatch):
    monkeypatch.setattr(argcheck, "_template_sheet_from_request", lambda *a, **k: None)
    for op in OPS:
        r, _i = _run(tmp_path, op, "取引先ごとに写して", "雛形")
        assert r["template_sheet"] == "雛形" and not _tpl_warnings(r), op   # ← 直す前の挙動
