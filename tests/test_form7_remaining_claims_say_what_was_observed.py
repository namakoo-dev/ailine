# -*- coding: utf-8 -*-
"""盲検の欠陥の形 7「断り・警告の文言が偽」の残り ── 観測していないことを観測したように言わない（2026-10-02）。

★★ 前便（tests/test_form7_claims_say_what_was_observed.py）の続き。地図は読んだだけだったので、
  **どの件も先に再現してから**直した。この試験の「嘘になる場面」の検体が、その再現そのもの。

  A1  「列の値をそのまま書き換える操作は今のところ対応していません」── 一覧に一括書換が在る
  A3  「列の取り違えか、別のソフトの書き出しが混ざっています」── 閉じた二択・見た見出しが無い
  A4  `ailine ops` の「未対応です」「照合できないため生成せず断ります」── 能力の否定
  A5  「この道具に『列追加』を打ち消す操作はありません」── 逆の『列削除』が在る
  A6  「（人が置いたファイルか、途中で失敗した run が残した作業結果のどちらかです）」── 閉じた二択
  A7  「ailine が作った物ですが、そのあと変更されています」── 見たのは指紋の不一致だけ
  A8  「『所属』は列の名前です（書き込む値ではありません）」── 同じ字なだけの語を値でないと断定
  A9  「曜日や『今日から見て』のような読み方はこの道具にはありません」── 言っていない依頼にも付く
  範囲  「列『X』がありません」── どのシートを探したかを言わない
  S3  「指す先の中身が変わる式があります」── 何も動かなかった回にも鳴る

★ どの件も 2 本で縛る: **嘘になる場面で嘘を言わない**／**正しい場面では今までどおり言う**
  （後者が無いと、全部黙らせる直しも緑になる）。
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402
from ailine_core import accounts_core, anchor, cellmap, cli_render, intent  # noqa: E402


def _book(tmp_path: Path, rows, name="b.xlsx", sheet="売上") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    p = tmp_path / name
    wb.save(p)
    wb.close()
    return p


def _meta(p: Path) -> dict:
    return ailine.build_book_meta(p)


# --- A1: 一覧に在る操作を「対応していません」と言わない -------------------------------------

def test_a1_a_plain_rewrite_is_pointed_at_the_op_that_exists(tmp_path):
    """再現: 「金額の列の値を変えて」が単列の倍率と読まれ、一括書換が在るのに『対応していません』。"""
    p = _book(tmp_path, [["商品", "金額"], ["a", 100], ["b", 250]])
    ok, _r, _i, err = ailine.verify_dsl_args(
        "COMPUTE_COLUMN", {"operands": ["金額"], "operator": "*"}, _meta(p),
        task="金額の列の値を変えて", vocab={})
    assert not ok and err
    assert "対応していません" not in err, err
    assert ailine.OP_META["SET_COLUMN_VALUE"]["label"] in err, err   # 在る操作の名前を出す


def test_a1_the_rate_path_is_untouched(tmp_path):
    """対: 率の手がかりが在る依頼は今までどおり通る（黙らせる直しでないこと）。"""
    p = _book(tmp_path, [["商品", "金額"], ["a", 100]])
    ok, r, _i, err = ailine.verify_dsl_args(
        "COMPUTE_COLUMN", {"operands": ["金額"], "operator": "*"}, _meta(p),
        task="金額に消費税10%を掛けて", vocab={})
    assert ok, err
    assert r["factor"] == pytest.approx(1.1)


# --- A4: `ailine ops` は能力を否定しない ----------------------------------------------------

def test_a4_the_ops_table_does_not_deny_a_capability_it_cannot_know():
    text = "\n".join(cli_render.render_ops_table(
        ailine.OP_META, ailine.OP_SCHEMA, ailine._CONFIRM_FIELDS))
    assert "未対応です" not in text and "照合できないため" not in text, text[-600:]
    assert "決められません" in text          # 未対応か言い方の問題かは決められない、と言う
    assert "生成せず断ります" in text         # 断る振る舞いそのものは今までどおり言う


# --- A5: 打ち消す操作が在るかを宣言から言う -------------------------------------------------

def test_a5_an_inverse_op_that_exists_is_named():
    """再現: 『列追加』を打ち消せと頼まれ、逆の『列削除』が在るのに『打ち消す操作はありません』。"""
    why = ailine.task_asks_to_undo_this_op("列追加を取り消して", "ADD_COLUMN")
    assert why, "打ち消しの依頼が素通りした"
    assert "ありません" not in why, why
    assert ailine.OP_LABELS["DELETE_COLUMN"] in why, why


def test_a5_a_declared_none_still_says_there_is_none():
    why = ailine.task_asks_to_undo_this_op("セル結合を取り消して", "MERGE")
    assert why and "打ち消す操作はありません" in why, why


def test_a5_an_undeclared_op_is_not_denied(monkeypatch):
    """宣言表に無い op には『ありません』と言わない（『決められませんでした』）。"""
    monkeypatch.delitem(intent.INVERSE_OP, "BOLD")
    why = ailine.task_asks_to_undo_this_op("太字を解除して", "BOLD")
    assert why and "ありません" not in why and "決められませんでした" in why, why


def test_a5_every_op_is_declared():
    """op を足した日に宣言を忘れると赤（宣言が無い op は『決められませんでした』になる）。"""
    assert set(intent.INVERSE_OP) == set(ailine.OP_META)
    for op, inv in intent.INVERSE_OP.items():
        assert inv is None or inv in ailine.OP_META, (op, inv)


# --- A9: 日付の読み方は、利用者が言った時だけ ------------------------------------------------

def _extract_err(tmp_path, task, value="未確認"):
    p = _book(tmp_path, [["名前", "メモ"], ["a", "済"], ["b", "済"]])
    ok, _r, _i, err = ailine.verify_dsl_args(
        "EXTRACT", {"col": "メモ", "cmp": "contains", "value": value}, _meta(p), task=task, vocab={})
    assert not ok, "0 行一致の抽出が通っている"
    return err


def test_a9_no_calendar_words_means_no_calendar_claim(tmp_path):
    err = _extract_err(tmp_path, "未確認が入っている行を抜き出して")
    assert "1 行もありません" in err, err           # 見たことは言う
    assert "日付の読み方" not in err and "今日から見て" not in err, err
    assert "どの列のどの文字か" in err, err         # 道は残す


def test_a9_a_calendar_word_still_gets_its_sentence(tmp_path):
    err = _extract_err(tmp_path, "土日に出勤してる人だけ抜き出して", value="土")
    assert "日付の読み方" in err and "土日" in err, err


# --- A3: 見たことと推し量りを分ける ---------------------------------------------------------

def test_a3_the_refusal_shows_what_was_seen_and_the_way_out():
    msg = accounts_core._no_key_column_refusal(["取引日", "科目", "金額"])
    assert "取り違えか" not in msg and "どちらか" not in msg, msg
    assert "『科目』" in msg and "見た見出し" in msg, msg          # 見たものを見せる
    assert "--column" in msg and "かもしれません" in msg, msg      # 道と、推し量りの形


def test_a3_without_headers_it_does_not_claim_to_have_seen_them():
    msg = accounts_core._no_key_column_refusal(None)
    assert "見た見出し" not in msg, msg


def test_a3_plan_accounts_uses_it():
    plan = accounts_core.plan_accounts([], {}, {}, seen_headers=["あ", "い"])
    assert plan.refused and "『あ』" in plan.refused, plan.refused


# --- A6・A7: 出力先を断る文 ------------------------------------------------------------------

def test_a6_the_unknown_owner_is_not_a_closed_choice(capsys, tmp_path):
    ailine._refuse_output_conflict(tmp_path / "x.out.xlsx", None)
    out = capsys.readouterr().out
    assert "この道具が書いた記録がありません" in out, out          # 見たこと（今までどおり）
    assert "のどちらかです" not in out, out
    assert "かもしれません" in out and "分かりません" in out, out   # 推し量りは推し量りの形で


def test_a7_the_edited_output_says_what_was_compared(capsys, tmp_path):
    ailine._refuse_edited_output(tmp_path / "x.out.xlsx")
    out = capsys.readouterr().out
    assert "そのあと変更されています" in out, out                   # 指紋の不一致は事実として言う
    assert "指紋" in out and "分かりません" in out, out             # 何を見たかと、見ていないこと
    assert "ailine が作った物ですが" not in out, out


# --- A8: 同じ字なだけの語を、値でないと断定しない ---------------------------------------------

@pytest.mark.parametrize("value", ["所属", "名簿", "太字にする"])
def test_a8_a_lookalike_word_is_not_declared_not_a_value(value):
    why = intent.why_not_a_value(value, ["氏名", "所属"], ["名簿"], ["太字", "並べ替え"])
    assert why and "書き込む値ではありません" not in why, why
    assert "決められません" in why, why                              # 決められない、と言う


def test_a8_a_real_value_still_passes():
    assert intent.why_not_a_value("確認済", ["氏名"], ["名簿"], ["太字"]) is None


def test_a8_an_eraser_word_is_still_refused_by_the_declared_table():
    assert "できません" in intent.why_not_a_value("空", [], [], [])


# --- 範囲を言わない不在: 探したシートを言う -----------------------------------------------------

def _two_sheet_book(tmp_path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "売上"
    a.append(["商品", "金額"])
    a.append(["りんご", 100])
    b = wb.create_sheet("原価")
    b.append(["商品", "原価"])
    b.append(["りんご", 60])
    p = tmp_path / "two.xlsx"
    wb.save(p)
    wb.close()
    return p


def _err(op, args, task, p, sheet="売上"):
    ok, _r, _i, err = ailine.verify_dsl_args(op, args, _meta(p), task=task, vocab={},
                                               target_sheet=sheet)
    assert not ok, (op, args)
    return err


@pytest.mark.parametrize("op, args, task", [
    ("DELETE_COLUMN", {"col": "原価"}, "原価の列を削除して"),
    ("MOVE_COLUMN", {"col": "原価"}, "原価の列を一番左に動かして"),
    ("SET_CELL_VALUE", {"col": "原価", "row": "りんご", "value": "x"}, "りんごの原価を「x」にして"),
    ("ADD_ROW", {"values": {"原価": 1}, "at": 2}, "2行目に原価1を追加して"),
])
def test_scope_the_missing_column_names_the_sheet_it_looked_in(tmp_path, op, args, task):
    """再現: 『原価』は別のシートに在る。『列がありません』だけでは、どこを探したか分からない。"""
    err = _err(op, args, task, _two_sheet_book(tmp_path))
    assert "探したシート: 『売上』" in err and "原価" in err and "ありません" in err, err


def test_scope_a_missing_column_in_compute_names_the_sheet(tmp_path):
    err = _err("COMPUTE_COLUMN", {"operands": ["原価"], "operator": "*", "factor": 1.1},
               "原価に1.1を掛けて", _two_sheet_book(tmp_path))
    assert "探したシート: 『売上』" in err, err


def test_scope_a_missing_row_names_the_sheet(tmp_path):
    p = _two_sheet_book(tmp_path)
    at, note = anchor.resolve_row_anchor("バナナの下に行を追加して", _meta(p), "売上")
    assert at is None and "探したシート: 『売上』" in note, note


def test_scope_resolve_col_ref_and_anchor_name_the_sheet():
    _v, _inf, err = anchor.resolve_col_ref("原価", ["商品", "金額"], sheet="売上")
    assert err and "探したシート: 『売上』" in err, err
    _at, note = anchor.resolve_col_anchor("原価の右に列を追加して", ["商品", "金額"], sheet="売上")
    assert note and "探したシート: 『売上』" in note, note


def test_scope_the_sheet_is_not_invented_when_unknown():
    """シートが分からない回に、作り話のシート名を言わない。"""
    assert anchor.searched(None) == ""
    _v, _inf, err = anchor.resolve_col_ref("原価", ["商品"])
    assert "シート『" not in err, err


def test_scope_a_present_column_is_not_called_missing(tmp_path):
    """対: 在る列は今までどおり通る。"""
    p = _two_sheet_book(tmp_path)
    ok, _r, _i, err = ailine.verify_dsl_args(
        "DELETE_COLUMN", {"col": "金額"}, _meta(p), task="金額の列を削除して", vocab={},
        target_sheet="売上")
    assert ok, err


# --- S3: 指す先の中身が変わる式は、実際に変わった時だけ ---------------------------------------

def _drift_book(tmp_path, ordered: bool):
    """金額で降順に並べ替える依頼。ordered=True は最初から降順（何も動かない）。"""
    rows = [("りんご", 300), ("みかん", 200), ("ぶどう", 100)] if ordered else \
           [("りんご", 100), ("みかん", 200), ("ぶどう", 300)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "売上"
    ws.append(["商品", "金額"])
    for r in rows:
        ws.append(list(r))
    ws["A5"], ws["B5"] = "合計", "=SUM(B2:B4)"        # 合計行が並べ替えの終わりを決める
    ws["A7"], ws["B7"] = "ぶどうの金額", "=B4"       # 範囲の外から特定の 1 行を指す式
    p = tmp_path / ("ord.xlsx" if ordered else "unord.xlsx")
    wb.save(p)
    wb.close()
    return p


def _sorted_copy(src: Path, dst: Path) -> Path:
    wb = openpyxl.load_workbook(src)
    ws = wb["売上"]
    rows = sorted(([ws.cell(r, 1).value, ws.cell(r, 2).value] for r in (2, 3, 4)),
                  key=lambda x: -x[1])
    for i, (a, b) in enumerate(rows, start=2):
        ws.cell(i, 1, a), ws.cell(i, 2, b)
    wb.save(dst)
    wb.close()
    return dst


def _advisories(tmp_path, ordered: bool):
    p = _drift_book(tmp_path, ordered)
    meta = _meta(p)
    ok, resolved, _i, err = ailine.verify_dsl_args(
        "SORT", {"col": "金額", "order": "desc"}, meta, task="金額の大きい順に並べ替えて", vocab={})
    assert ok, err
    after_p = _sorted_copy(p, tmp_path / "after.xlsx")
    from ailine_core.dsl_step import compose_dsl_step_advisories
    adv = compose_dsl_step_advisories(
        "structural", "SORT", resolved, meta, "金額の大きい順に並べ替えて",
        ailine.snapshot(p), ailine.snapshot(after_p), deps=ailine._make_dsl_step_deps())
    return resolved, adv


def test_s3_nothing_moved_means_nothing_is_said(tmp_path):
    """再現: 最初から降順の表。何も動かないのに「指す先の中身が変わる式」と鳴っていた。"""
    resolved, adv = _advisories(tmp_path, ordered=True)
    assert resolved.get("_drift"), "候補が控えられていない（対の側が死んでいる）"
    assert not any("指す先の中身" in a for a in adv), adv
    assert not any("指す先の中身" in w for w in resolved.get("_warnings", [])), resolved.get("_warnings")


def test_s3_a_real_move_is_still_named(tmp_path):
    resolved, adv = _advisories(tmp_path, ordered=False)
    hit = [a for a in adv if "指す先の中身が変わった式" in a]
    assert hit and hit[0].startswith("★"), adv                       # ★ 付き＝ ✓→△ の数に入る
    assert "『" in hit[0] and "から" in hit[0], hit[0]               # 前→後の中身まで言う
    assert "直していません" in hit[0], hit[0]
    assert not resolved.get("_warnings"), "適用の前にも鳴っている"


def test_s3_a_truncated_snapshot_does_not_pretend_to_have_checked():
    before = {"cells": {}, "truncated": True}
    after = {"cells": {}, "truncated": True}
    line = cellmap.reference_drift_observed([("S", "B1", "=A9999", 9999, 1)], "S", before, after)
    assert line and "確かめていません" in line, line


def test_s3_untruncated_unchanged_is_silent():
    snap = {"cells": {"S!2,1": ("x", "General", None, False, None, None)}, "truncated": False}
    assert cellmap.reference_drift_observed([("S", "B1", "=A2", 2, 1)], "S", snap, dict(snap)) is None
