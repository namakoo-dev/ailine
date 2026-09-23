"""画面に出した警告は、出来た冊の中にも残る（盲検 6 体目 ⑪・致命）。

★ 事故: 画面には警告が 4 つ出ているのに、出来た冊を後から開いた人には 1 つも見えない。
  買い手の言葉:「2 行しか入っていない仕訳帳が、**警告なしで**手元に残ります」。
  メールで転送された先では、画面はもう存在しない。

★★ 偽の basrun は**本物と同じ壊れ方**をする ── `Call WriteInspectionSheet(...)` を含む
  Basic にヘルパの .bas が渡っていなければ、**黙って何もせず成功を返す**。
  2026-09-23 の初版はここでヘルパを渡し忘れ、実機で「成功・シート無し」になった。
  素直な偽物（ヘルパを見ない）ではこの形の退行を捕まえられない。
"""
import ast
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_golden_transcripts import _isolate, _run_main  # noqa: E402
from lo_fake import apply_inspection_sheets  # noqa: E402
from _product_source import product_text  # noqa: E402

SHEET = ailine.NOTE_SHEET
DEDUP_SHEET = "取引先の重複除去"

# 1 列目が空の行で走査が止まる形（買い手の仕訳帳と同じ）── 画面に ⚠ が出る
ROWS_WARNED = [["取引先", "金額"], ["甲社", 100], ["甲社", 250], ["乙社", 200],
               [None, None], ["丙社", 50]]
ROWS_CLEAN = [["取引先", "金額"], ["甲社", 100], ["甲社", 250], ["乙社", 200]]


def _book(tmp_path, rows):
    tmp_path.mkdir(parents=True, exist_ok=True)
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    wb.save(p)
    return p


def _has_helper(helper_files) -> bool:
    return any("Sub WriteInspectionSheet" in Path(f).read_text(encoding="utf-8")
               for f in helper_files)


def _fake_basrun(monkeypatch, dedup_rows, calls, fail_inspection=False):
    """1 回目は重複除去、`WriteInspectionSheet` を呼ぶ回は検分シートを書く。"""
    def fake_apply(out_book, code, workdir, helper_files=(), timeout=None):
        calls.append({"code": code, "helper": _has_helper(helper_files)})
        wb = openpyxl.load_workbook(out_book)
        if "removeByName" in code and SHEET in code:       # 古い申し送りを消す回
            if SHEET in wb.sheetnames:
                del wb[SHEET]
        elif "Call WriteInspectionSheet" in code:
            if fail_inspection:
                return False, "LibreOffice が落ちました", ""
            if not _has_helper(helper_files):
                return True, None, "ok"          # ★ 本物と同じ: 黙って何もしない
            apply_inspection_sheets(wb, code)
        else:
            ws = wb.create_sheet(DEDUP_SHEET)
            for r in dedup_rows:
                ws.append(list(r))
        wb.save(out_book)
        return True, None, "ok"
    monkeypatch.setattr(ailine, "basrun_apply", fake_apply)


def _run(tmp_path, monkeypatch, capsys, rows, dedup_rows, fail_inspection=False,
         stale_note=False):
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path, rows)
    if stale_note:                      # 前回の run が残した申し送り
        wb = openpyxl.load_workbook(book)
        wb.create_sheet(SHEET).append(["判定", "", "", "", "⚠ 前回の警告"])
        wb.save(book)
    monkeypatch.setattr(
        ailine, "translate_task",
        lambda model, task, book_meta, temperature=0.1:
        {"op": "DEDUP", "args": {"keys": ["取引先"]}})
    calls: list = []
    _fake_basrun(monkeypatch, dedup_rows, calls, fail_inspection)
    rc, out = _run_main(["run", str(book), "取引先が同じ行を重複として除いて", "--copy"], capsys)
    return rc, out, book.with_name("b.out.xlsx"), calls


def _sheet_rows(path, name):
    ws = openpyxl.load_workbook(path)[name]
    return [[c.value for c in r] for r in ws.iter_rows()]


def test_every_warning_on_the_screen_is_in_the_book(tmp_path, monkeypatch, capsys):
    """画面の ⚠ の行は、**一字一句そのまま**冊の申し送りシートに在る。"""
    rc, out, made, calls = _run(tmp_path, monkeypatch, capsys, ROWS_WARNED,
                                [["取引先", "金額"], ["甲社", 100], ["乙社", 200]])
    screen_warnings = [ln.strip() for ln in out.splitlines() if ln.strip().startswith("⚠")]
    assert screen_warnings, f"前提が崩れた: 画面に ⚠ が出ていない（検体を見直す）\n{out}"
    wb = openpyxl.load_workbook(made)
    assert SHEET in wb.sheetnames, f"申し送りシートが無い: {wb.sheetnames}\n{out}"
    rows = _sheet_rows(made, SHEET)
    assert rows[0] == ["種類", "対象", "行数", "金額", "なぜこうなったか"], rows[0]
    in_book = [r[4] for r in rows[1:]]
    # ★ 画面の ⚠ のうち、警告（advisories）由来のものはすべて冊に在る。
    #   判定行（⚠ 機械保証はありません…）も「判定」として在る。
    missing = [w for w in screen_warnings if w not in in_book]
    assert not missing, f"画面に出たのに冊に無い警告:\n{missing}\n冊: {in_book}"
    assert any(r[0] == "判定" for r in rows[1:]), f"判定の行が無い: {rows}"


def test_the_inspection_basic_gets_the_helper(tmp_path, monkeypatch, capsys):
    """★ 退行の名指し: 検分を書く回に、ヘルパの .bas が渡っていること。
       渡っていなければ本物の basrun は**成功を返して何もしない**。"""
    _rc, out, _made, calls = _run(tmp_path, monkeypatch, capsys, ROWS_WARNED,
                                  [["取引先", "金額"], ["甲社", 100], ["乙社", 200]])
    insp = [c for c in calls if "Call WriteInspectionSheet" in c["code"]]
    assert insp, f"検分を書く呼び出しが 1 回も無い\n{out}"
    assert all(c["helper"] for c in insp), "検分の Basic にヘルパが渡っていない（黙って空振りする）"


def test_a_clean_run_leaves_the_book_as_it_was(tmp_path, monkeypatch, capsys):
    """警告が無い回は何もしない ── 普段の出力の形を変えない。"""
    rc, out, made, calls = _run(tmp_path, monkeypatch, capsys, ROWS_CLEAN,
                                [["取引先", "金額"], ["甲社", 100], ["乙社", 200]])
    assert rc == 0, out
    assert "⚠" not in out, f"前提が崩れた: 綺麗な検体で ⚠ が出た\n{out}"
    assert SHEET not in openpyxl.load_workbook(made).sheetnames
    assert not [c for c in calls if "Call WriteInspectionSheet" in c["code"]], \
        "警告が無いのに 2 回目の往復をした"


def test_a_failed_write_is_said_and_does_not_stop_the_run(tmp_path, monkeypatch, capsys):
    """書けなかったら画面に出す（黙らない）。ただし原本の反映は済んでいるので run は止めない。"""
    rc_ok, _o, _m, _c = _run(tmp_path / "a", monkeypatch, capsys, ROWS_WARNED,
                             [["取引先", "金額"], ["甲社", 100], ["乙社", 200]])
    rc, out, made, _calls = _run(tmp_path / "b", monkeypatch, capsys, ROWS_WARNED,
                                 [["取引先", "金額"], ["甲社", 100], ["乙社", 200]],
                                 fail_inspection=True)
    assert "申し送りのシートは書けませんでした" in out, out
    assert rc == rc_ok, f"添え物の失敗で終了コードが変わった: {rc_ok} → {rc}"
    assert DEDUP_SHEET in openpyxl.load_workbook(made).sheetnames


def test_a_stale_note_from_last_time_is_removed_on_a_clean_run(tmp_path, monkeypatch, capsys):
    """★ 前回の申し送りが冊に残っていたら、綺麗な回に消す ── 残すと前回の警告が
       今回の冊について嘘を言う。★ 消すのは読み戻しの**前**（✓ の行が無いシートを言わない）。"""
    rc, out, made, _calls = _run(tmp_path, monkeypatch, capsys, ROWS_CLEAN,
                                 [["取引先", "金額"], ["甲社", 100], ["乙社", 200]],
                                 stale_note=True)
    assert rc == 0, out
    assert SHEET not in openpyxl.load_workbook(made).sheetnames, \
        f"前回の申し送りが残っている（前回の警告が今回の冊について嘘を言う）\n{out}"
    assert SHEET not in out, f"消したシートを読み戻したと言っている\n{out}"


def test_only_a_warned_run_changes_the_shape_of_the_book():
    """書く回の線: ⚠ の行が在るか、機械保証が無い回だけ。△ や ★ だけの回は書かない
       （冊のシートの顔ぶれは、宣言の検査・複合計画の後段・下流の道具すべてに効く）。"""
    assert ailine.note_is_due(["⚠ 1 行は検算できていません"], machine_verified=True)
    assert ailine.note_is_due([], machine_verified=False)
    assert not ailine.note_is_due(["★ 疑わしい: 変更が元データの範囲外です（I1）",
                                   "（新規シート『集計』の作成は意図どおりです）"],
                                  machine_verified=True)
    assert not ailine.note_is_due([], machine_verified=True)


def test_every_finish_apply_caller_passes_the_helpers():
    """★ 1 本で全経路を縛る: `_finish_apply(...)` の呼び出しはすべて helper_files を渡す。
       分母は手で並べず、製品の AST から数える。"""
    tree = ast.parse(product_text())
    calls = [n for n in ast.walk(tree)
             if isinstance(n, ast.Call) and getattr(n.func, "id", None) == "_finish_apply"]
    assert len(calls) >= 4, f"呼び出しの数が想定より少ない（数え方を疑う）: {len(calls)}"
    bare = [n.lineno for n in calls if not any(k.arg == "helper_files" for k in n.keywords)]
    assert not bare, f"helper_files を渡していない _finish_apply の呼び出し: 行 {bare}"
