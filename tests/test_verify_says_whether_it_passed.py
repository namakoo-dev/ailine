# -*- coding: utf-8 -*-
"""`ailine verify` は「検算した」でなく「**合格した／しなかった**」を言う（2026-09-18・盲検 3 体目 ②）。

★★ 起きたこと（製造業の購買担当・初見・盲検）:

    $ ailine verify <照合の出力> <元A> <元B>
    ■ ailine verify（照合）  out=…  A=…  B=…
    Σ A: 126833 / Σ B: 120563
    EXIT=0

  ★ 買い手の言葉:「**✓ も × も無い。この 2 つの数は私には意味がありません。
    『検算した』のか『合格した』のか分からない**」。
  ★ README は「出した出力は `ailine verify` で**独立に検算**できます ──
    **取り逃し・二重・元に無い値**を見ます」と約束している。**看板と実物のずれ**。

★★ 追ったら、またこの repo の持病だった:
  検算の出口は 5 経路（分けた冊／科目の候補／帳票の一覧／stack・extract／**照合**）。
  前の 4 つは `render_independent_verify_report` と `_independent_verify_exit` を
  **共有**しているのに、照合だけが 3 本目の器と 4 本目の判定の写しを持っていた。
  ★ `_independent_verify_exit` の docstring には「3 経路が書き写していた ── 片配線の足場」と
    **既に書いてある**。3 経路は畳んだが、照合だけ畳み残していた。

★ 照合だけが落としていたもの 3 つ:
    ① ✓／⚠ の判定の行     ② 分母（`keys` は**計算済みで一度も表示していなかった**）
    ③ 空虚な合格の禁止

★★ この検体そのものが 1 度失敗している（2026-09-18・記録として残す）:
  初版は helper が facts / vacuous を**自分で組んで**いて、本物の `_verify_match` を
  通っていなかった ── 検算の側を壊す変異 2 本が**緑のまま**だった。
  棚の線「**同じ関数で作った分母は恒真**」そのもの。
  ★ だから分母・判定は**実物の経路**で測る（冊を作る → 照合を走らせる → その出力を
    verify_match_output に渡す）。組み立てた dict は**破れの文言**にだけ使う。
"""
import contextlib
import io
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402
from ailine_core import verify as V  # noqa: E402
from ailine_core.cli_render import render_independent_verify_report  # noqa: E402

OK_A = [["取引先", "金額"], ["甲社", 1000], ["乙社", 2000], ["丙社", 3000]]
OK_B = [["取引先", "金額"], ["甲社", 1000], ["乙社", 2000], ["丙社", 3000]]
TASK = "取引先をキーにして金額を突き合わせて"


def _books(tmp_path, rows_a, rows_b):
    out = []
    for name, rows in (("a.xlsx", rows_a), ("b.xlsx", rows_b)):
        p = tmp_path / name
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in rows:
            ws.append(list(r))
        wb.save(p)
        out.append(p)
    return out


def _make_match_output(tmp_path, rows_a=None, rows_b=None):
    """実物の経路で照合の出力を作る（LLM は通らない・規則だけの経路）。"""
    a, b = _books(tmp_path, rows_a or OK_A, rows_b or OK_B)
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = ailine.main(["run", str(a), str(b), TASK])
    assert rc == 0, f"照合が作れなかった: {buf.getvalue()}"
    out = next(p for p in tmp_path.glob("*照合*.xlsx"))
    return out, a, b


def _real(tmp_path, **kw):
    out, a, b = _make_match_output(tmp_path, **kw)
    return V.verify_match_output(out, a, b)


def _lines(result):
    return render_independent_verify_report("照合", "out.xlsx", "a.xlsx / b.xlsx", result)


# --- 実物の経路で測る（分母・判定・空虚） -------------------------------------------

def test_a_clean_run_says_it_passed(tmp_path):
    """★ 事故そのもの: 破れが無いなら「✓ 破れはありません」と**言う**。"""
    said = "\n".join(_lines(_real(tmp_path)))
    assert "✓" in said, f"判定が出ていない（検算したのか合格したのか分からない）: {said}"
    assert "破れはありません" in said, said


def test_the_denominator_is_shown(tmp_path):
    """★★ 分母を出す ── `keys` は前から計算していて、一度も画面に出していなかった。

    ★ 「上の分母で測りました」と言う以上、その分母が上に無ければ嘘になる。
    """
    r = _real(tmp_path)
    said = "\n".join(_lines(r))
    assert r.get("keys"), f"分母が計算されていない: {r}"
    assert f"照合したキー: {r['keys']} 件" in said, (said, r["keys"])


def test_an_emptied_output_is_not_a_pass(tmp_path):
    """★★ 出力の行を消されたら ✓ を出さない ── **取り逃し**として名指しする。

    ★★ 2026-09-18: 初版のこの検体は「行を消したら**空虚**（測るものが無い）になる」と
      想定して書き、**赤くなった**。製品のほうが正しい ── 元帳に在るキーが出力から
      消えたのだから、それは「測るものが無い」ではなく**破れ**（取り逃し 3 件・exit 5）。
      ★ 検体の想定が実装より粗かった回。直すのは検体の側（今日 2 度目）。

    ★ 到達の記録: 照合の `vacuous`（0 行）は**製品の経路では到達できない** ──
      0 行の照合出力は列解決が先に断って作れず（実測）、出力を空にすれば上のとおり
      取り逃しが立つ。元も空なら列解決で断られる。**未到達の防御枝**として残す
      （「到達できず＝未確認であって安全でない」の側に置く・偽の合格は作らない）。
    """
    out, a, b = _make_match_output(tmp_path)
    wb = openpyxl.load_workbook(out)
    ws = wb["照合"]
    ws.delete_rows(2, ws.max_row)          # 見出しだけ残す
    wb.save(out)
    r = V.verify_match_output(out, a, b)
    said = "\n".join(_lines(r))
    assert "✓" not in said, f"行が消されているのに合格を名乗った: {said}"
    assert r.get("mismatch"), f"取り逃しを見逃した: {r}"
    assert "取り逃し" in said, said
    assert ailine._independent_verify_exit(r) == 5, "破れが在るのに exit 5 でない"


def test_a_tampered_output_is_caught(tmp_path):
    """★★ README の約束（取り逃し・二重・元に無い値）が実際に鳴ること。

    ★ 出力のキーを 1 つ書き換えると「取り逃し」と「元に無い値」が同時に立つ。
    """
    out, a, b = _make_match_output(tmp_path)
    wb = openpyxl.load_workbook(out)
    ws = wb["照合"]
    ws.cell(row=2, column=1).value = "存在しない社"
    wb.save(out)
    r = V.verify_match_output(out, a, b)
    said = "\n".join(_lines(r))
    assert r.get("mismatch"), f"改竄を見逃した: {r}"
    assert "⚠ 破れ" in said, said
    assert "取り逃し" in said and "元に無い値" in said, said
    assert ailine._independent_verify_exit(r) == 5, "破れが在るのに exit 5 でない"


# --- 破れの文言（ここだけ組み立てた形で測る） ----------------------------------------

def _built(mismatches):
    """★ 破れの**文言**だけを測るための組み立て。分母・判定・空虚には使わない
       （使うと恒真になる ── この検体が 1 度それで失敗した）。"""
    return {"mismatch": True, "facts": {"照合したキー": "5 件"},
            "breaks": [V._match_break(m) for m in mismatches], "vacuous": None}


def test_every_kind_of_break_is_named():
    """★ 畳んだ時に言葉を痩せさせていないこと（5 種類とも人の言葉で出る）。"""
    said = "\n".join(_lines(_built([
        {"kind": "count", "side": "A", "key": "甲社", "expected": 2, "written": 1},
        {"kind": "sum", "side": "B", "key": "乙社", "expected": 1000.0, "written": 900.0},
        {"kind": "diff", "key": "丙社", "expected": 100.0, "written": 0.0},
        {"kind": "missing_key", "key": "丁社"},
        {"kind": "extra_key", "key": "戊社"},
    ])))
    assert "⚠ 破れ 5 件:" in said, said
    for word in ("A側の件数", "B側の合計", "差額", "取り逃し", "元に無い値"):
        assert word in said, f"『{word}』が消えている: {said}"
    assert "捏造の可能性" in said, "言葉が痩せている（元の器に在った警告が消えた）"


def test_the_exit_code_comes_from_the_shared_organ():
    """★★ 判定の出口は 1 つ ── `0 if ok else 5` の 4 本目の写しを残さない。"""
    from _product_source import count_in_product
    assert count_in_product('return 0 if result.get("ok") else 5') == 0, (
        "★ 判定の写しが残っている（共有の出口 _independent_verify_exit を使うこと）")


def test_the_third_renderer_is_gone():
    """★ 死んだ器を残さない（次に読む人が『照合には別の器がある』と読む）。

    ★★ 「無いこと」だけの assert は**探す場所が空でも通る**（tests/test_guard_ledger.py が
      2026-09-18 にこの検体を捕まえた）。だから先に「在ること」を確かめてから無いと言う。
    """
    for f in ("src/ailine_core/cli_render.py", "src/ailine/__init__.py"):
        src = (REPO / f).read_text(encoding="utf-8")
        assert "render_independent_verify_report" in src, (
            f"★ 探す場所が空か形が変わっている（この検体が空回りする）: {f}")
        assert "render_verify_match_report" not in src, f"照合専用の器が残っている: {f}"
