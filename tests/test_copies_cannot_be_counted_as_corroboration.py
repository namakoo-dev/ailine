# 写し合いを見分けられない入口で「裏が取れた」を名乗らないこと ── 2026-09-12（D4）
#
# ★★ 「別々の出所が一致した」を根拠に**裏が取れた**と言うには、その 2 つが
#   写し合いでないと言えなければならない。Excel では式（`=合計`）を見て
#   「この欄は帯の写しだ」と分かる（`_depends_on`）。式が読めない入口では言えない:
#
#       PDF                    式が存在しない
#       Excel（式のブック無し）  式を読み込んでいない
#
# ★ 観測された被害者（2026-09-12 実測）:
#   - PDF: Wondershare の雛形で請求額が **裏が取れた** を名乗った（上部の欄と帯の合計は
#     元の Excel なら式で繋がった同じ事実のはず）
#   - Excel: 式を渡さないと B01 の根拠が 2 → 3 に増える（写しを独立と数えている）。
#     ただし 87 冊で区分が変わる冊は 0 ── **壊れた経路は在るが被害者は居ない**状態だった
#
# ★ 「PDF だから」では書かない ── 判定は 1 箇所（`copies_are_indistinguishable`）。
from __future__ import annotations

import openpyxl
import pytest

from ailine_core.field_record import (GRADES_WITH_VALUE, Evidence, Record, describe,
                                      grade, grade_of, value)
from ailine_core.form_read import copies_are_indistinguishable, read_book

REPO_BOOK_SHEET = "misoca_invoice"


def _two_agreeing():
    return (Evidence(rule="上部の請求額欄", value=36300, at="C11", how="C11"),
            Evidence(rule="帯の合計", value=36300, at="H39", how="H39"))


def test_agreeing_sources_earn_the_top_grade_only_when_copies_can_be_told_apart():
    """★ 本番: 掃き出し済みで一致していても、写しを見分けられないなら最上位を出さない。"""
    top = grade_of(_two_agreeing(), swept=True)
    capped = grade_of(_two_agreeing(), swept=True, copies_indistinguishable=True)
    assert top in GRADES_WITH_VALUE and capped in GRADES_WITH_VALUE
    assert top != capped, "★ 写しを見分けられないのに同じ区分を出している"
    # ★ 値は変わらない ── 弱くするのは**主張の強さ**だけ（空欄にはしない）
    for flag in (False, True):
        r = Record("請求額", _two_agreeing(), swept=True, copies_indistinguishable=flag)
        assert value(r) == 36300


def test_the_reason_says_which_of_the_two_doubts_it_is():
    """★ 「裏が取れていない」の理由は 2 種類あり、人の次の手が変わる ── 混ぜない。"""
    swept_not_done = Record("請求額", _two_agreeing(), swept=False)
    blind = Record("請求額", _two_agreeing(), swept=True, copies_indistinguishable=True)
    assert "食い違う数字が無いか" in describe(swept_not_done), describe(swept_not_done)
    assert "写し" in describe(blind), describe(blind)
    assert "食い違う数字が無いか" not in describe(blind), describe(blind)


def test_a_declared_reason_is_shown_instead_of_the_default():
    r = Record("請求額", _two_agreeing(), swept=True,
               copies_indistinguishable=True, copies_why="この帳票には式がありません")
    assert "この帳票には式がありません" in describe(r)


@pytest.mark.parametrize("evid, kw", [
    ((), {}),                                                     # 無
    ((Evidence(rule="上部", value=1, at="A1", how="a"),), {}),      # 単（出所 1 つ）
    (_two_agreeing(), {"conflict": True}),                        # 割（食い違いが立っている）
])
def test_it_changes_nothing_outside_the_one_case_it_is_for(evid, kw):
    """★ 陰性対照 ── 出所 1 つ・食い違い・手がかり無し は、この処置で動かない。"""
    assert (grade_of(evid, **kw)
            == grade_of(evid, copies_indistinguishable=True, **kw))


def test_disagreeing_values_stay_split_even_when_copies_are_blind():
    """★ 写しを見分けられないことを、食い違いを黙らせる口実にしない。"""
    e = (Evidence(rule="上部", value=100, at="A1", how="a"),
         Evidence(rule="帯", value=200, at="B2", how="b"))
    assert grade_of(e, swept=True, copies_indistinguishable=True) == grade_of(e, swept=True)
    assert grade_of(e, swept=True) not in GRADES_WITH_VALUE


# ── 判定は 1 箇所（式が読めるか）────────────────────────────
def test_the_single_predicate_is_about_formulas_not_about_pdf():
    blind, why = copies_are_indistinguishable(None)
    assert blind is True and "式" in why, why
    assert copies_are_indistinguishable(object()) == (False, "")


def test_an_excel_book_read_without_its_formulas_does_not_claim_corroboration(tmp_path):
    """★★ 潜在欠陥の番人: 同じブックを「式あり／式なし」で読んで、区分が違うこと。

    ★ 式を渡さない経路でも `確` が出ていた（写しを独立と数えていた）。
      この検体では区分が変わらなかったので**点数には出ない** ── だから機械で縛る。
    """
    book = tmp_path / "invoice.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = REPO_BOOK_SHEET
    ws["B2"] = "請求書"
    ws["B3"] = "ナギ商会株式会社　御中"
    ws["G5"] = "株式会社あかね商事"
    ws["B11"] = " ご請求金額　"
    ws["C11"] = "=H39"                     # ★ 上部の欄は帯の写し（式）
    ws["B15"] = "品番・品名"
    ws["E15"] = "数量"
    ws["G15"] = "単価"
    ws["H15"] = "金額"
    ws["B16"] = "用紙代"
    ws["E16"] = 1
    ws["G16"] = 30000
    ws["H16"] = 30000
    ws["E37"] = "小計"
    ws["H37"] = 30000
    ws["E38"] = "消費税"
    ws["G38"] = 0.1
    ws["H38"] = 3000
    ws["E39"] = "合計金額"
    ws["H39"] = 33000
    wb.save(book)
    wb.close()
    # ★ 式のキャッシュが無いので data_only では C11 が None になる ── 値を焼いた写しを別に作る
    wb = openpyxl.load_workbook(book)
    wb[REPO_BOOK_SHEET]["C11"] = 33000
    baked = tmp_path / "baked.xlsx"
    wb.save(baked)
    wb.close()

    values = openpyxl.load_workbook(baked, data_only=True)
    formulas = openpyxl.load_workbook(book, data_only=False)      # ★ 式は元のブックから
    try:
        with_f = read_book(values, wb_formula=formulas)["請求額"]
        without = read_book(values)["請求額"]
    finally:
        values.close()
        formulas.close()
    assert value(with_f) == value(without) == 33000
    assert grade(with_f) != grade(without), (
        "★ 式を読めない経路で、写しを独立した根拠として数えている")
    assert "写し" in describe(without), describe(without)
