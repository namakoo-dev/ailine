# -*- coding: utf-8 -*-
"""`accounts` の列名の壁に、通る扉を付ける（2026-09-20・盲検 4 体目 ⑤）。

★★ 出所（買い手の冊は見出しが『勘定科目』だった）── 道具はこう言っていた:

    × 『借方勘定科目』に当たる列がありません（別名: 借方勘定科目）。
      見た見出し: 『取引日』／『勘定科目』／『金額』／『摘要』
      （弥生の形（25 列・1 列目が 4 桁の数字）でもありません）

  誤診が 3 つ重なっていた:
    ① **『勘定科目』は目の前に在る**のに、近いとも言わない（人は「在るじゃないか」で止まる）
    ② 「別名: 借方勘定科目」── 別名が役割名と同じで情報がゼロ
    ③ 「弥生の形でもありません」── 見出し行の在る 4 列の冊に、無関係な第二の診断を
      くっつけて誤導していた（買い手は弥生を使っていない）
  そして**道が無い** ── 列を教える手段が無かった。

★ 直した形: `--column 役割=見出し`（複数可）。断りは近い列を名指しし、その場で打てる
  `--column` を見せる。弥生の注記は**幅が弥生と同じ時だけ**出す。
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

from ailine_core import accounts_core as C  # noqa: E402

HEADERS = ["取引日", "勘定科目", "金額", "摘要"]
ROWS = [(1, list(HEADERS)), (2, ["2026-09-01", "", 12000, "コピー用紙"])]


def _refusal(rows=None, overrides=None) -> str:
    _h, _hh, _map, why = C.resolve_accounts_columns(rows or ROWS, overrides)
    return why or ""


def test_the_near_column_is_named():
    """★★ ① 目の前に在る近い列を**名指しする**こと。"""
    why = _refusal()
    assert "『勘定科目』" in why, why
    assert "名前が近い" in why, why


def test_the_refusal_shows_a_command_you_can_type():
    """★★ 道を示すこと ── そのまま打てる `--column` を出す（合格線の 3 条目）。"""
    why = _refusal()
    assert "--column 借方勘定科目=勘定科目" in why, why


def test_the_yayoi_note_is_not_tacked_onto_an_unrelated_book():
    """★★ ③ 無関係な第二の診断を付けないこと。

    ★ 見出し行が在る 4 列の冊に「弥生の形でもありません」と言うのは**誤導**。
      買い手はそこで 30 分溶かす側へ行く（4 体目の汚染 1 と同じ形）。
    ★ 逆に**幅が弥生と同じ**なら言う価値がある ── そこは残す（下の試験）。
    """
    assert "弥生" not in _refusal(), _refusal()


def test_the_yayoi_note_still_appears_when_the_width_matches():
    """★ 25 列なのに受けられない冊には、弥生として受けられない理由を言うこと。

    ★ 「出さない」だけにすると、弥生を使っている人への説明まで消える。
    """
    wide = [(1, ["x"] + [""] * 24)]        # ★ 25 列・1 列目が 4 桁の数字でない
    why = _refusal(wide)
    assert "弥生" in why and "4 桁" in why, why


def test_all_the_missing_columns_are_named_at_once():
    """★ 足りない列は**全部まとめて**言うこと（往復を必要列の数だけ増やさない）。

    ★ 照合の側は最初からそうしている（「決まらなかった役割を全部集めてから報告」）──
      同じ考えがこちらに配線されていなかった。
    """
    why = _refusal()
    assert "借方勘定科目" in why and "借方金額" in why, why


def test_telling_it_the_column_actually_works():
    """★★ 教えた列で**決まる**こと（扉が開く）。"""
    _h, _hh, header_map, why = C.resolve_accounts_columns(
        ROWS, {C.DEBIT_ACCOUNT: "勘定科目", C.DEBIT_AMOUNT: "金額"})
    assert not why, why
    assert header_map[C.DEBIT_ACCOUNT] == 2, header_map
    assert header_map[C.DEBIT_AMOUNT] == 3, header_map


def test_a_column_that_is_not_there_is_refused_by_name():
    """★ 在ると思って渡した人に、**何が見えているか**を返すこと（推測で進まない）。"""
    why = _refusal(overrides={C.DEBIT_ACCOUNT: "そんな列"})
    assert "そんな列" in why and "ありません" in why, why
    assert "『勘定科目』" in why, "★ 見えている見出しを返していない: " + why


@pytest.mark.parametrize("bad, want", [
    ("借方勘定科目", "役割=見出し"),          # ★ `=` が無い
    ("そんな役割=勘定科目", "という役割はありません"),
    ("借方勘定科目=", "見出しが空"),
])
def test_a_broken_option_says_what_is_wrong(bad, want):
    """★ 指定の形が違う回も、何が違うかを言うこと。"""
    got, why = C.parse_column_overrides([bad])
    assert got == {} and why and want in why, (bad, why)


def test_the_roles_come_from_the_declaration():
    """★ 役割の顔ぶれは `COLUMN_ALIASES` が唯一の出どころ（手で並べない）。

    ★ 手で並べると、役割を足した日に**指定できないまま**静かに残る。
    """
    got, why = C.parse_column_overrides([f"{r}=見出し" for r in C.COLUMN_ALIASES])
    assert not why, why
    assert set(got) == set(C.COLUMN_ALIASES)


def test_every_command_that_reads_a_journal_offers_the_door():
    """★★ 片配線を作らない ── 仕訳の冊を読むコマンド**全部**に `--column` が在ること。

    ★ `accounts` だけに足すと、次の段（`accounts-apply`）で同じ壁に当たる。
      `verify` も同じ冊を読み直すので、無ければ検算だけ断られる。
    ★ 登録は 1 箇所（`_add_column_option`）── 文面と選択肢をそこで決める。
    """
    import ailine
    parser = ailine.build_parser()
    import argparse
    subs = [a for a in parser._actions if isinstance(a, argparse._SubParsersAction)][0]
    for name in ("accounts", "accounts-apply", "verify"):
        opts = {o for act in subs.choices[name]._actions for o in act.option_strings}
        assert "--column" in opts, f"★ {name} に扉が無い（同じ冊を読むのに教えられない）"
    from _product_source import count_in_product
    assert count_in_product("--column\", action=\"append\"") == 1, (
        "★ 口の登録が 2 箇所以上ある（文面と選択肢がずれる）")


def test_the_whole_sequence_walks_with_the_door_open(tmp_path):
    """★★ 買い手が通る 3 段を**端から端まで歩く**（2026-09-20）。

    ★★ なぜ要るか: 初版の番人は「argparse に `--column` が在るか」しか見ておらず、
      `verify` 側は値を受け取る配線が抜けていた（`NameError: overrides`）── 既存の試験が
      拾ったから助かったが、**口が在ること**と**効くこと**は別だ（設定≠動く）。
    ★ 3 段（accounts → accounts-apply → verify）は同じ冊を読む。どれか 1 つが
      受け取らなければ、買い手はそこで止まる。
    """
    import os
    import subprocess
    import openpyxl

    def mk(p, headers, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "仕訳"
        ws.append(headers)
        for r in rows:
            ws.append(r)
        wb.save(p)
        wb.close()
        return p

    heads = ["取引日", "勘定科目", "金額", "摘要"]
    now = mk(tmp_path / "今回.xlsx", heads, [["2026-09-01", "", 12000, "コピー用紙"]])
    past = mk(tmp_path / "過去.xlsx", heads, [["2026-08-01", "消耗品費", 5000, "コピー用紙"]])
    cand = tmp_path / "候補.xlsx"
    env = {**os.environ, "PYTHONPATH": str(REPO / "src"),
           "AILINE_HOME": str(tmp_path / "home")}
    door = ["--column", "借方勘定科目=勘定科目", "--column", "借方金額=金額"]

    def run(*args):
        return subprocess.run([sys.executable, "-m", "ailine", *args], cwd=str(REPO),
                              capture_output=True, text=True, encoding="utf-8",
                              errors="replace", timeout=900, env=env)

    r1 = run("accounts", str(now), "--past", str(past), "--out", str(cand), *door)
    assert r1.returncode == 0, "★ 1 段目（accounts）で止まった: " + r1.stdout[-600:]

    # ★ 人が右端に『採用』の列を足して ○ を付ける（決めるのは人 ── 道具は作らない）
    from ailine_core import accounts_apply
    wb = openpyxl.load_workbook(cand)
    ws = wb["候補"]
    head = [r for r in range(1, ws.max_row + 1)
            if any(str(ws.cell(r, c).value or "").strip() == "候補の科目"
                   for c in range(1, ws.max_column + 1))]
    assert head, "★ 候補のシートに『候補の科目』の見出しが無い"
    mark_col = ws.max_column + 1
    ws.cell(head[0], mark_col).value = accounts_apply.ADOPT_HEADER
    ws.cell(head[0] + 1, mark_col).value = "○"
    wb.save(cand)
    wb.close()

    out = tmp_path / "取込用.xlsx"
    r2 = run("accounts-apply", str(cand), str(now), "--out", str(out), *door)
    assert r2.returncode == 0, "★ 2 段目（accounts-apply）で止まった: " + r2.stdout[-600:]

    r3 = run("verify", str(cand), str(now), str(past), *door)
    assert r3.returncode == 0, "★ 3 段目（verify）で止まった: " + r3.stdout[-600:]
