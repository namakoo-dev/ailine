# -*- coding: utf-8 -*-
"""2 回目の買い手役 3 体（2026-09-13）が指した所の番人 ── 「ブックだけ開く人」の軸。

★★ 経理役: 「月末、黒い画面をずっと見ているわけではない（翌朝ブックだけ開く）」。読めなかった冊・
  載せなかった冊・月の混在が**画面にしか無く**、ブックに 1 セルも残らなかった。しかも受領した
  請求書 1 枚が読めずに落ちているのに verify は ✓ exit 0 ── 「✓ を見た時点で読むのをやめる」。
★ 会計役: 同じ中身の過去の冊 2 本で根拠の件数が黙って倍／弥生の定型文が同じシートで食い違う。
★ 事務職: `--by 元ファイル` で冊名が `04月.xlsx.xlsx`。
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from test_forms_e2e import _invoice, _forms, _sheets, _not_an_invoice            # noqa: E402
from test_accounts_e2e import _csv, _accounts, MF                                 # noqa: E402
from ailine_core import split_people                                              # noqa: E402


def _verify(out, folder):
    return subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(folder)],
                          cwd=str(REPO), capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=300)


def _redate(path, text):
    wb = openpyxl.load_workbook(path)
    wb.active["H3"] = text
    wb.save(path)
    wb.close()


# ── 束の要約シート ──────────────────────────────────────────────

def test_the_summary_sheet_keeps_what_the_screen_said(tmp_path):
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    _invoice(folder / "b.xlsx", "いろは工業株式会社", 5500)
    _redate(folder / "b.xlsx", "2026/5/31")
    _not_an_invoice(folder / "稟議書.xlsx", "稟議書")
    (folder / "こわれた.xlsx").write_bytes(b"text")
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    got = _sheets(out)
    assert "束の要約" in got, list(got)
    rows = got["束の要約"][1:]
    kinds = {row[0] for row in rows}
    assert {"読んだ冊", "読めなかった冊", "一覧に載せていない冊", "請求日の月", "束の所見"} <= kinds, rows
    assert any(row[0] == "読めなかった冊" and row[1] == "こわれた.xlsx" and "保存し直して" in row[2]
               for row in rows), rows
    assert any(row[0] == "一覧に載せていない冊" and row[1] == "稟議書.xlsx" for row in rows), rows
    months = [row for row in rows if row[0] == "請求日の月"]
    assert any("2026年5月" in row[2] and "b.xlsx" in row[1] for row in months), months
    assert any("2026年8月" in row[2] and "a.xlsx" in row[1] for row in months), months


def test_a_single_month_puts_no_month_rows_in_the_summary(tmp_path):
    """★ 陰性対照 ── 月が 1 つなら月の行は無い（余計な行を足さない）。"""
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    _invoice(folder / "b.xlsx", "いろは工業株式会社", 5500)
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    rows = _sheets(out)["束の要約"][1:]
    assert not [row for row in rows if row[0] == "請求日の月"], rows


# ── verify: 読めなかった冊が在れば ✓ を出さない ────────────────────

def test_verify_does_not_pass_a_list_that_is_missing_an_unreadable_book(tmp_path):
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    (folder / "こわれた.xlsx").write_bytes(b"text")
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    v = _verify(out, folder)
    assert v.returncode == 4, f"exit={v.returncode} / {v.stdout}"
    assert "✓" not in v.stdout, v.stdout
    assert "一覧は完全ではありません" in v.stdout and "こわれた.xlsx" in v.stdout, v.stdout


def test_verify_still_passes_a_complete_list(tmp_path):
    """★ 陰性対照 ── 全部読めた束は今までどおり ✓ exit 0。"""
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    v = _verify(out, folder)
    assert v.returncode == 0 and "✓" in v.stdout, v.stdout


def test_a_total_row_added_later_is_named_not_called_a_break(tmp_path):
    """★ 経理役: `run` で足した合計行を「★ 出所のファイルが無い: 合計」と破れに数えていた。"""
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    wb = openpyxl.load_workbook(out)
    ws = wb["一覧"]
    ws.append(["合計", None, None, 3300, None, None])
    wb.save(out)
    wb.close()
    v = _verify(out, folder)
    assert v.returncode == 0, f"exit={v.returncode} / {v.stdout}"
    assert "合計行（検算の対象外）: 1 行（合計）" in v.stdout, v.stdout
    assert "出所のファイルが無い" not in v.stdout, v.stdout


# ── 画面: 項目ごとの区分 ──────────────────────────────────────────

def test_the_grades_are_broken_down_by_field_on_screen(tmp_path):
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    _invoice(folder / "b.xlsx", "いろは工業株式会社", 5500)
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert r.returncode == 0
    lines = [ln for ln in r.stdout.splitlines() if ln.startswith("  請求額: ") or ln.startswith("  請求元: ")]
    assert len(lines) == 2, r.stdout
    assert any("確 2" in ln for ln in lines if ln.startswith("  請求額")), lines


# ── accounts: 同じ中身の過去は 1 本・弥生の定型文 ───────────────────

_PAST = [["1", "2026/08/25", "通信費", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "8月分 回線"],
         ["2", "2026/07/25", "通信費", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "7月分 回線"]]
_TODAY = [["11", "2026/09/25", "", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "9月分 回線"]]


def test_two_identical_past_files_count_as_one(tmp_path):
    today = _csv(tmp_path / "今回.csv", _TODAY)
    past = tmp_path / "過去"
    past.mkdir()
    _csv(past / "仕訳帳_原本.csv", _PAST)
    (past / "仕訳帳_原本 - コピー.csv").write_bytes((past / "仕訳帳_原本.csv").read_bytes())
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, past, out)
    assert r.returncode == 0, r.stdout
    assert "中身が同じなので 1 本に数えました" in r.stdout, r.stdout
    wb = openpyxl.load_workbook(out)
    reasons = [c.value for row in wb.worksheets[0].iter_rows(min_row=2) for c in row
               if isinstance(c.value, str) and "先例" in c.value]
    wb.close()
    assert reasons and all("過去 4 件" not in why for why in reasons), reasons
    assert any("過去 2 件" in why for why in reasons), reasons


def test_the_credit_note_does_not_claim_a_key_the_book_does_not_have(tmp_path):
    """★ 会計役（2 回とも）: 貸方取引先の無い冊で「貸方取引先だけは鍵に入れています」と
    「その列は無い」が同じシートに並んだ。"""
    headers = ["取引No", "取引日", "借方勘定科目", "借方補助科目", "借方金額(円)", "貸方勘定科目", "摘要"]
    past = [["1", "2026/08/25", "通信費", "本社回線", "8800", "未払金", "8月分 回線"]]
    today = [["11", "2026/09/25", "", "本社回線", "8800", "未払金", "9月分 回線"]]
    t = _csv(tmp_path / "今回.csv", today, headers=headers)
    p = _csv(tmp_path / "過去.csv", past, headers=headers)
    out = tmp_path / "候補.xlsx"
    assert _accounts(t, p, out).returncode == 0
    wb = openpyxl.load_workbook(out)
    notes = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and "鍵に入れ" in c.value]
    wb.close()
    assert notes, "定型文が出ていない（分母が痩せている）"
    assert not any("貸方取引先だけは鍵に入れています" in n for n in notes), notes


# ── split: 完全一致の見出しは採る・何を採ったか言う（設計 D1 の線を 1 段動かした）──────

def _plan(headers, by, **kw):
    grid = [(1, list(headers)), (2, ["甲社", "内藤", "山田", 1000]), (3, ["乙社", "内藤", "佐藤", 2000])]
    return split_people.plan_split(grid, 1, by, None, **kw)


def test_by_default_two_matches_still_refuse_but_name_the_way_out():
    """★ 設計 D1 はそのまま（既定では絞らない）── ただし断りが「元の表を変えろ」で終わらない。"""
    plan = _plan(["顧客", "担当者", "営業担当者", "金額"], "担当者")
    assert plan.refused is not None and "2 つあります" in plan.refused, plan.refused
    assert "--exact" in plan.refused, plan.refused


def test_exact_picks_the_literal_header_and_says_so():
    plan = _plan(["顧客", "担当者", "営業担当者", "金額"], "担当者", exact=True)
    assert plan.refused is None, plan.refused
    assert plan.by_column == 2, plan.by_column
    assert "『担当者』の列で分けました" in plan.by_note and "『営業担当者』" in plan.by_note, plan.by_note


def test_exact_cannot_rescue_a_partial_only_or_a_duplicated_header():
    """★ 陰性対照 ── 『担当』（部分一致だけ）も、『担当者』が 2 列も、--exact でも決めない。"""
    plan = _plan(["顧客", "担当者", "営業担当者", "金額"], "担当", exact=True)
    assert plan.refused is not None and "--exact" not in plan.refused, plan.refused
    plan = _plan(["顧客", "担当者", "担当者", "金額"], "担当者", exact=True)
    assert plan.refused is not None and "元の表で片方の見出しを変えて" in plan.refused, plan.refused


# ── split: 冊名の二重拡張子 ──────────────────────────────────────

def test_a_value_that_ends_with_xlsx_does_not_double_the_suffix():
    names = split_people.safe_filenames(["04月.xlsx", "05月.XLSX", "山田"])
    assert names["04月.xlsx"] == "04月" and names["05月.XLSX"] == "05月" and names["山田"] == "山田", names
