# -*- coding: utf-8 -*-
"""3 回目の買い手役 3 体（2026-09-13）が指した所の番人。

★★ 事務職の致命 2 つはどちらも元からあった穴:
  ① `run 9月売上.csv` が、隣で育てた `9月売上.xlsx` を黙って上書きし undo でも戻らない
  ② `×` で止まった回の `.out.xlsx` が翌月の縦積みに 1 冊として積まれ、Σ 38 → 53
★ 経理: 画面の集計（空欄 16 件）が一覧に載せていない冊の 5 項目を含んでいた（俺の片配線）
★ 会計: CSV の継続行（借方金額 '0'）が候補行に数えられ 13 行 vs 12 行／検算コマンドが貼れない
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from test_forms_e2e import _invoice, _forms, _not_an_invoice                          # noqa: E402
from test_accounts_e2e import _csv, _accounts, MF                                     # noqa: E402
from ailine_core import cli_render, intent, stack as stack_core                       # noqa: E402


def _run(argv, cwd):
    return subprocess.run([sys.executable, "-m", "ailine", *argv], cwd=str(cwd),
                          capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=300)


# ── ① 検疫の出力に作業が乗っていたら作り直さない ────────────────────────

def test_a_quarantined_book_that_was_edited_since_is_not_rebuilt(tmp_path):
    src = tmp_path / "9月売上.csv"
    src.write_text("品番,取引先,数量,単価\n0012,丸山工業,3,1200\n", encoding="utf-8-sig")
    assert _run(["csv", str(src)], tmp_path).returncode == 0
    book = tmp_path / "9月売上.xlsx"
    wb = openpyxl.load_workbook(book)
    wb.active["E1"] = "金額"                       # ★ 人（や run）が育てた列
    wb.active["E2"] = 3600
    wb.save(book)
    wb.close()
    r = _run(["csv", str(src)], tmp_path)
    assert r.returncode == 7, f"exit={r.returncode} / {r.stdout}"
    assert "そのあと変更されています" in r.stdout, r.stdout
    wb = openpyxl.load_workbook(book)
    assert wb.active["E1"].value == "金額", "★ 育てた列が消えた"
    wb.close()


def test_an_untouched_quarantined_book_is_still_rebuilt_quietly(tmp_path):
    """★ 陰性対照 ── 俺が置いたままなら今までどおり作り直してよい（毎月同じ CSV を指し直す運用）。"""
    src = tmp_path / "9月売上.csv"
    src.write_text("品番,取引先\n0012,丸山工業\n", encoding="utf-8-sig")
    assert _run(["csv", str(src)], tmp_path).returncode == 0
    r = _run(["csv", str(src)], tmp_path)
    assert r.returncode == 0, r.stdout


# ── ② 作業結果（.out.xlsx）は入力に数えない ─────────────────────────────

def test_a_work_result_left_in_the_folder_is_not_stacked(tmp_path):
    folder = tmp_path / "月次"
    folder.mkdir()
    for name in ("2026-06_売上.xlsx", "2026-06_売上.out.xlsx", "2026-07_売上.xlsx"):
        wb = openpyxl.Workbook()
        wb.active.append(["品番", "数量"])
        wb.active.append(["0012", 3])
        wb.save(folder / name)
        wb.close()
    r = _run(["stack", str(folder), "--out", str(tmp_path / "縦積み.xlsx")], tmp_path)
    assert r.returncode == 0, r.stdout
    assert "2 ファイル中 2 積んだ" in r.stdout, r.stdout
    assert "Σ数量: 元 6 / 出力 6" in r.stdout, r.stdout
    assert "2026-06_売上.out.xlsx" in r.stdout and "除外" in r.stdout, r.stdout


def test_the_work_result_rule_is_named_in_one_place():
    kept, excluded = stack_core.split_own_outputs([Path("a.out.xlsx"), Path("b.xlsx")])
    assert excluded == ["a.out.xlsx"], excluded
    assert stack_core.WORK_RESULT_SUFFIX == ".out.xlsx"


# ── ③ 画面の集計は一覧の実物と 1 対 1 ─────────────────────────────────

def test_screen_tallies_count_only_listed_books(tmp_path):
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    _not_an_invoice(folder / "年間予算表.xlsx", "年間予算表")
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert r.returncode == 0, r.stdout
    assert "一覧の空欄 0 件＋ 一覧に載せていない冊の 5 項目（検分にだけ）" in r.stdout, r.stdout
    assert "無 5" not in r.stdout.split("項目の区分")[1].splitlines()[0], r.stdout
    assert "請求日の無い冊" not in r.stdout, r.stdout


# ── ④ 継続行の '0' は候補行にしない（CSV でも xlsx でも同じ数）─────────────

_PAST = [["1", "2026/08/25", "通信費", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "8月分"]]
_TODAY = [["11", "2026/09/25", "", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "9月分"],
          ["", "", "", "", "", "0", "未払金", "アスクル", "（継続行）"]]


def test_a_continuation_row_with_a_text_zero_is_left_alone(tmp_path):
    today = _csv(tmp_path / "今回.csv", _TODAY)
    past = _csv(tmp_path / "過去.csv", _PAST)
    r = _accounts(today, past, tmp_path / "候補.xlsx")
    assert r.returncode == 0, r.stdout
    assert "候補を出す行 1 行のうち 1 行に科目の候補が出ました（触らない行 1" in r.stdout, r.stdout


# ── ⑤ 検算コマンドはそのまま貼れる ─────────────────────────────────────

def test_the_verify_hint_quotes_arguments_the_shell_would_eat():
    line = cli_render.verify_hint("ailine split", "配る", "一覧.xlsx",
                                  f"--amount {cli_render._quoted('金額(円)')}")[0]
    assert '--amount "金額(円)"' in line, line
    assert cli_render._quoted("金額") == "金額"


# ── ⑥ 縦積みの verify に合否の 1 行 ──────────────────────────────────

def test_the_stack_verify_report_says_it_matched():
    lines = cli_render.render_verify_report("縦積み.xlsx", "月次",
                                            {"row_count": {"source": 8, "output": 8}, "sums": {},
                                             "mismatch": None, "mismatches": []})
    assert any(ln.startswith("✓ 一致") for ln in lines), lines
    lines = cli_render.render_verify_report("縦積み.xlsx", "月次",
                                            {"row_count": {"source": 8, "output": 8}, "sums": {},
                                             "mismatch": {"kind": "row_count", "source": 8, "output": 9},
                                             "mismatches": []})
    assert not any(ln.startswith("✓") for ln in lines), lines


# ── ⑦ 「過去に 1 件もありません」は列の話 ──────────────────────────────

def test_a_missing_precedent_names_the_column(tmp_path):
    past = [["1", "2026/08/25", "消耗品費", "", "アスクル", "3000", "未払金", "", "コピー用紙"],
            ["2", "2026/08/25", "通信費", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "", "8月分"]]
    today = [["11", "2026/09/25", "", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "", "9月分"],
             ["12", "2026/09/25", "", "", "", "3000", "未払金", "アスクル", "デスクライト"]]
    r = _accounts(_csv(tmp_path / "今回.csv", today), _csv(tmp_path / "過去.csv", past), tmp_path / "候補.xlsx")
    assert r.returncode == 0, r.stdout
    wb = openpyxl.load_workbook(tmp_path / "候補.xlsx")
    texts = [c.value for row in wb.worksheets[0].iter_rows(min_row=2) for c in row
             if isinstance(c.value, str) and "1 件もありません" in c.value]
    wb.close()
    assert texts and all("貸方取引先の列に『アスクル』" in t for t in texts), texts


# ── ⑧ 読む形 ─────────────────────────────────────────────────────

def test_the_candidate_sheet_is_set_up_for_reading(tmp_path):
    today = _csv(tmp_path / "今回.csv", _TODAY[:1])
    past = _csv(tmp_path / "過去.csv", _PAST)
    out = tmp_path / "候補.xlsx"
    assert _accounts(today, past, out).returncode == 0
    wb = openpyxl.load_workbook(out)
    ws = wb.worksheets[0]
    assert ws.freeze_panes == "A2", ws.freeze_panes
    assert ws.auto_filter.ref, "オートフィルタが無い"
    col = [i + 1 for i, c in enumerate(ws[1]) if c.value == "根拠"][0]
    assert ws.cell(row=2, column=col).alignment.wrap_text is True
    assert wb.worksheets[1]["B1"].value == "元ファイルの行", wb.worksheets[1]["B1"].value
    wb.close()


# ── ⑨ 断りの言葉 ───────────────────────────────────────────────────

def test_split_tells_a_csv_holder_about_ailine_csv(tmp_path):
    src = tmp_path / "一覧.csv"
    src.write_text("担当者,金額\n山田,100\n", encoding="utf-8-sig")
    r = _run(["split", str(src), "--by", "担当者", "--out", str(tmp_path / "配る")], tmp_path)
    assert r.returncode != 0
    assert "ailine csv" in r.stdout, r.stdout


def test_a_fully_filled_journal_says_so_in_the_headline(tmp_path):
    filled = [["11", "2026/09/25", "通信費", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "", "9月分"]]
    r = _accounts(_csv(tmp_path / "今回.csv", filled), _csv(tmp_path / "過去.csv", _PAST), tmp_path / "候補.xlsx")
    assert r.returncode == 0, r.stdout
    assert "× 候補を出す行がありません ── 借方勘定科目が全部埋まっています" in r.stdout, r.stdout


# ── ⑩ export-csv はシート名を当てさせない ───────────────────────────────

def test_export_csv_picks_the_only_sheet_or_lists_them(tmp_path):
    one = tmp_path / "一枚.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet"
    wb.active.append(["a", "b"])
    wb.save(one)
    wb.close()
    r = _run(["export-csv", str(one)], tmp_path)
    assert r.returncode == 0 and "『Sheet』を書き出します" in r.stdout, r.stdout
    two = tmp_path / "二枚.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "縦積み"
    wb.active.append(["a"])
    wb.create_sheet("集計").append(["b"])
    wb.save(two)
    wb.close()
    r = _run(["export-csv", str(two)], tmp_path)
    assert r.returncode != 0 and "縦積み" in r.stdout and "集計" in r.stdout, r.stdout


# ── ⑫ 画像・ロゴは最初から断る ────────────────────────────────────────

def test_images_and_logos_are_named_as_things_we_do_not_do():
    for task in ("会社のロゴ画像を右上に貼って", "H2 に画像を挿入して", "写真を入れて"):
        assert intent.asked_for_what_we_do_not_do(task) == "画像の挿入", task
    assert intent.asked_for_what_we_do_not_do("売上から原価を引いた利益の列を作って") is None
