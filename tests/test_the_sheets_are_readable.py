# -*- coding: utf-8 -*-
"""書いた冊は**読める形**になっている（2026-09-14・買い手役・会計が 2 回・¥20,000 の #2）。

★★ 「根拠 最大 336 字・折り返しなし・フィルタなし・枠固定なし（2 回目）」。
  `inspection.reading_aids` は 2026-09-13 に作ったのに **15 シート中 1 枚**にしか通して
  いなかった（実測）── しかも通っていないのは人が一番読む「検分」。この repo の再発する
  欠陥（片配線・`docs/開発手法.md` §13）。

★★ この番人は**1 本で全経路を縛る** ── 3 つのコマンドを走らせ、**書いた冊を分母にして**
  全シートを検査する。手書きの白名簿にしない（白名簿は必ずずれる ── `render_folder_routes`
  の轍）。新しい出力シートを足して読む支度を忘れたら、ここが赤くなる。

契約:
  - 見出し行が在るシートは枠が固定されている（スクロールしても見出しが見える）
  - 合計行が**無い**シートはオートフィルタが在る（区分で絞れる ── 会計役の一番の要求）
  - 合計行が**在る**シートには付けない（絞ると合計が消える／残る形が読み手を惑わす）
  - 長い文（`inspection.WRAP_OVER_CHARS` 字超）が在る列は折り返す
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from ailine_core import inspection, total_row   # noqa: E402
from test_accounts_e2e import _PAST, _TODAY, _accounts, _csv   # noqa: E402
from test_forms_e2e import _forms, _invoice   # noqa: E402
from test_split_e2e import _book, _split   # noqa: E402


def _has_total_row(ws) -> bool:
    """そのシートの**末尾**が合計行か。

    ★★ 初版は「どこかの行に合計語が在れば」で測って、split の検分を誤判定した ──
      あそこは散文で「『合計』── 明細の行ではないので配っていません」と**書いてある**だけ。
      合計行は必ず末尾に在り、ラベルのセルが語そのものになる形（`TOTAL_LABEL` /
      `OWN_TOTAL_LABEL`）。★ 物差しを直す側（製品でなく治具）。
    """
    if (ws.max_row or 0) < 2:
        return False
    last = [v for v in next(ws.iter_rows(min_row=ws.max_row, values_only=True)) if v is not None]
    return any(total_row.row_has_total_word([v]) == str(v).strip() for v in last
               if isinstance(v, str))


def _long_columns(ws) -> list:
    out = []
    for col in range(1, (ws.max_column or 0) + 1):
        if any(isinstance(ws.cell(row=r, column=col).value, str)
               and len(ws.cell(row=r, column=col).value) > inspection.WRAP_OVER_CHARS
               for r in range(2, (ws.max_row or 1) + 1)):
            out.append(col)
    return out


def _judge(path: Path) -> list:
    """1 冊を検査して (冊, シート, 事実) を返す。"""
    facts = []
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        if (ws.max_row or 0) < 1:
            continue
        facts.append({
            "book": path.name, "sheet": ws.title,
            "freeze": ws.freeze_panes, "filter": ws.auto_filter.ref,
            "total": _has_total_row(ws), "long": _long_columns(ws),
            "wrapped": [c for c in range(1, (ws.max_column or 0) + 1)
                        if all((ws.cell(row=r, column=c).alignment or None) is not None
                               and ws.cell(row=r, column=c).alignment.wrap_text
                               for r in range(2, (ws.max_row or 1) + 1))],
        })
    wb.close()
    return facts


@pytest.fixture(scope="module")
def written(tmp_path_factory):
    """3 つのコマンドが**実際に書いた冊**（これが分母 ── 一覧を手で持たない）。"""
    tmp = tmp_path_factory.mktemp("読む形")
    folder = tmp / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)
    _invoice(folder / "b.xlsx", "いろは工業株式会社", 5500)
    out = tmp / "出力"
    out.mkdir()
    assert _forms(folder, out / "一覧.xlsx").returncode == 0
    assert _forms(folder, out / "一覧8月.xlsx", "--month", "2026-08").returncode == 0
    today = _csv(tmp / "今回.csv", _TODAY)
    past = _csv(tmp / "過去.csv", _PAST)
    assert _accounts(today, past, out / "候補.xlsx").returncode == 0
    book = _book(tmp / "元" / "売上一覧.xlsx")
    assert _split(book, out / "配る", "--amount", "金額").returncode == 0
    facts = []
    for path in sorted(out.rglob("*.xlsx")):
        facts += _judge(path)
    return facts


def test_the_fixture_actually_measures_something(written):
    """★ 空虚な合格の禁止 ── 長い文の列と合計行の在るシートが検体に**在る**ことを先に示す。"""
    assert len(written) >= 8, written
    assert [f for f in written if f["long"]], "長い文の列が 1 つも無い（何も測っていない）"
    assert [f for f in written if f["total"]], "合計行の在るシートが無い（陰性側を測れない）"


def test_every_sheet_freezes_its_header(written):
    bad = [(f["book"], f["sheet"]) for f in written if not f["freeze"]]
    assert not bad, f"見出しが固定されていないシート: {bad}"


def test_sheets_without_a_total_row_can_be_filtered(written):
    """★ 会計役の一番の要求 ── 『区分』で絞れること。"""
    bad = [(f["book"], f["sheet"]) for f in written if not f["total"] and not f["filter"]]
    assert not bad, f"オートフィルタが無いシート: {bad}"


def test_sheets_with_a_total_row_are_not_filtered(written):
    """★ 陰性対照 ── 合計行の在るシートには付けない（絞ると合計の在り方が変わる）。"""
    bad = [(f["book"], f["sheet"]) for f in written if f["total"] and f["filter"]]
    assert not bad, f"合計行の在るシートにフィルタが付いている: {bad}"


def test_long_reasons_are_wrapped(written):
    bad = [(f["book"], f["sheet"], sorted(set(f["long"]) - set(f["wrapped"])))
           for f in written if set(f["long"]) - set(f["wrapped"])]
    assert not bad, f"長い文の列が折り返されていない: {bad}"
