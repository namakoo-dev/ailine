# -*- coding: utf-8 -*-
"""開示の欠落（形 4）の在り処は、書かれた形から数え、台帳と等号で縛る（2026-10-02）。

★★ 前提: 盲検の欠陥の形 4（警告や前提が画面だけで残らない・どこにも残らない・黙って外す・黙ってファイルを
  残す）は 9 件出て、毎回**見つけた 1 つ**を直していた。記録に書かれた根は「出口の一覧を宣言から導かず、
  見つけた出口を 1 つずつ塞いでいる」。同じ家系（関所で止まった run が `.out` を黙って残す）は 3 度直した後に
  **4 件目**（Ctrl-C・想定外の例外）が出た ── 出口が `return` でなく `raise` だったから、`return` を数える
  番人には見えなかった。
  ★ だから 4 軸（終わり方・⚠ の行き先・黙って外す・置き換える）を tests/disclosure_core.py が AST で数え、
    開示なしの在り処を理由つきの台帳（tests/disclosure_register.json）と**等号で**縛る:
      ・台帳に無い在り処が現れたら赤（新しい出口・⚠・外す所を書いたら、分類するまで通らない）
      ・台帳に在って実装に無いのも赤（直したら台帳からも消す ── 開示なしは減る向きだけ）
      ・件数が違っても赤

★ この試験自身を疑う:
    ① 陽性対照: 直す前の形（出力を作った後に履歴の口を通らない return・`raise` で出る入口・画面だけの ⚠・
       言わない上限・関所の無い置き換え）の**小さな一時ソース**で、走査が鳴る
    ② 陰性対照: 直した形（口を通る・投げ直す前に後始末・言う語・関所）では鳴らない
    ③ 枝の取り違え: 別の枝（`if dry: _finish_run(); return 0`）の呼び出しは、後ろの出口を救わない
    ④ 直した実物（Ctrl-C・複合計画の途中の聞き返し・csv 暗黙前段・export-pdf・export-csv・demo）が
       **実際に**画面と履歴へ残す ── 数えただけでなく動かす
"""
from __future__ import annotations

import functools
import json
import shutil
import sys
import textwrap
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import disclosure_core as core  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

import ailine  # noqa: E402


@functools.lru_cache(maxsize=1)
def _real_found() -> tuple:
    found = core.scan()
    return tuple(found)


def _real_tally() -> dict:
    return core.tally(list(_real_found()))


# --- 本体 -------------------------------------------------------------------------------------

def test_every_undisclosed_site_is_registered_with_its_reason():
    diff = core.compare(_real_tally(), core.load_register())
    assert not diff["unlisted"], (
        "開示なしの在り処が台帳に無い（新しい出口・⚠・外す所・置き換える所は、**言う／残す／関所を通る**ように書くか、"
        "言えない理由を tests/disclosure_register.json に書くこと。★ 数を合わせるために足さない）: "
        f"{diff['unlisted']}")
    assert not diff["wrong_count"], f"件数が台帳と違う（実装, 台帳）: {diff['wrong_count']}"
    assert not diff["stale"], f"台帳に在って実装に無い在り処（直したなら台帳からも消す）: {diff['stale']}"


def test_every_register_entry_has_a_reason_and_an_unlock():
    data = json.loads(core.REGISTER.read_bytes().decode("utf-8"))
    for e in data["undisclosed"]:
        assert len(e.get("reason", "")) >= 20, (e["axis"], e["func"])
        assert len(e.get("unlock", "")) >= 6, (e["axis"], e["func"])
        assert e["class"] in core.UNDISCLOSED[e["axis"]], e
        assert isinstance(e["count"], int) and e["count"] >= 1, e


def test_the_scan_is_not_empty_on_any_axis():
    """★ 走査が 1 軸でも空回りしたら赤（在っても鳴らない、を作らない）。"""
    s = core.summary(list(_real_found()))
    for axis in "abcd":
        assert sum(s.get(axis, {}).values()) > 0, f"軸 {axis} の在り処が 0 件（走査が空回りしている）"
    assert s["a"].get("exc_guarded", 0) >= 1, "Ctrl-C・例外の出口を守る入口が 1 つも見えない"
    assert s["a"].get("routed", 0) >= 1 and s["b"].get("carried", 0) >= 1


# --- ①②③ 陽性・陰性・枝の取り違え（小さな一時ソース）-------------------------------------------------

def _scan(tmp_path, source: str) -> list:
    f = tmp_path / "mutant.py"
    f.write_bytes(textwrap.dedent(source).encode("utf-8"))
    return core.scan([f])


def _classes(found: list, axis: str) -> dict:
    out: dict = {}
    for x in found:
        if x["axis"] == axis:
            out[x["class"]] = out.get(x["class"], 0) + 1
    return out


def test_a_return_after_the_output_without_the_history_mouth_is_flagged(tmp_path):
    found = _scan(tmp_path, '''
        import shutil
        def cmd_x(a):
            shutil.copy2(a.src, a.out)
            if a.bad:
                return 3
            _finish_run(a)
            return 0
        def build(p):
            p.set_defaults(func=cmd_x)
    ''')
    c = _classes(found, "a")
    assert c.get("unrouted") == 1, c              # `return 3` ── 作った後・口を通らない
    assert c.get("routed") == 1, c                # `return 0` ── `_finish_run` の後


def test_a_return_before_the_output_exists_is_not_flagged(tmp_path):
    found = _scan(tmp_path, '''
        import shutil
        def cmd_x(a):
            if a.bad:
                return 3
            shutil.copy2(a.src, a.out)
            _finish_run(a)
            return 0
        def build(p):
            p.set_defaults(func=cmd_x)
    ''')
    c = _classes(found, "a")
    assert c.get("no_out") == 1 and not c.get("unrouted"), c


def test_another_branchs_mouth_does_not_rescue_a_later_exit(tmp_path):
    """★ 行の順だけで見ると `if dry: _finish_run(); return 0` が後ろの `return 3` を救ってしまう。"""
    found = _scan(tmp_path, '''
        import shutil
        def cmd_x(a):
            shutil.copy2(a.src, a.out)
            if a.dry:
                _finish_run(a)
                return 0
            return 3
        def build(p):
            p.set_defaults(func=cmd_x)
    ''')
    c = _classes(found, "a")
    assert c.get("unrouted") == 1, c


def test_a_condition_that_finishes_before_returning_true_routes_its_exit(tmp_path):
    found = _scan(tmp_path, '''
        import shutil
        def _answer(a):
            _finish_run(a)
            return True
        def cmd_x(a):
            shutil.copy2(a.src, a.out)
            if _answer(a):
                return 3
            return 0
        def build(p):
            p.set_defaults(func=cmd_x)
    ''')
    c = _classes(found, "a")
    assert c.get("routed", 0) >= 1 and not c.get("unrouted"), c


_RAISE_HEAD = '''
    import shutil
    def cmd_x(a):
        shutil.copy2(a.src, a.out)
        return 0
    def build(p):
        p.set_defaults(func=cmd_x)
'''
_RAISE_FIXED = '''
    import shutil
    def _finish_aborted(a, e):
        _finish_run(a)
    def cmd_x(a):
        try:
            shutil.copy2(a.src, a.out)
            return 0
        except BaseException as e:
            _finish_aborted(a, e)
            raise
    def build(p):
        p.set_defaults(func=cmd_x)
'''


def test_an_entry_that_can_leave_the_output_when_an_exception_escapes_is_flagged(tmp_path):
    """★ 4 件目の形: `return` の出口だけ数えると、例外の出口は見えない（HEAD の `_cmd_run_body` がこれ）。"""
    assert _classes(_scan(tmp_path, _RAISE_HEAD), "a").get("exc_unguarded") == 1
    c = _classes(_scan(tmp_path, _RAISE_FIXED), "a")
    assert c.get("exc_guarded") == 1 and not c.get("exc_unguarded"), c


def test_a_catch_that_swallows_or_skips_the_mouth_is_not_a_guard(tmp_path):
    swallow = '''
        import shutil
        def cmd_x(a):
            try:
                shutil.copy2(a.src, a.out)
            except BaseException:
                pass
        def build(p):
            p.set_defaults(func=cmd_x)
    '''
    no_reraise = '''
        import shutil
        def cmd_x(a):
            try:
                shutil.copy2(a.src, a.out)
            except BaseException as e:
                _finish_run(a)
        def build(p):
            p.set_defaults(func=cmd_x)
    '''
    for src in (swallow, no_reraise):
        assert _classes(_scan(tmp_path, src), "a").get("exc_unguarded") == 1, src


def test_an_exit_that_is_an_expression_is_counted_too(tmp_path):
    """★ `return e.exit_code`（盲検 5 体目 ①の 3 件目）── 数の出口の一覧に出ない出口。数の出口を持つ関数の
    中の式の出口も数える。"""
    found = _scan(tmp_path, '''
        import shutil
        def cmd_x(a):
            shutil.copy2(a.src, a.out)
            if a.bad:
                return 3
            try:
                step(a)
            except Gate as e:
                return e.exit_code
            _finish_run(a)
            return 0
        def build(p):
            p.set_defaults(func=cmd_x)
    ''')
    c = _classes(found, "a")
    assert c.get("unrouted") == 2, c            # `return 3` と `return e.exit_code`
    fixed = _scan(tmp_path, '''
        import shutil
        def cmd_x(a):
            shutil.copy2(a.src, a.out)
            try:
                step(a)
            except Gate as e:
                return _finish_gated(a, e.exit_code)
            _finish_run(a)
            return 0
        def build(p):
            p.set_defaults(func=cmd_x)
    ''')
    assert not _classes(fixed, "a").get("unrouted"), _classes(fixed, "a")


def test_a_function_that_receives_the_output_as_an_argument_is_already_after_it(tmp_path):
    """★ `_finish_apply(a, book, out_book, …)` が直に print した ⚠ 8 種（a70b5d3 の前）── 出力を作ったのは
    呼び出し側なので、関数の中の行の順だけでは『作った後』と分からない。"""
    found = _scan(tmp_path, '''
        def finish(a, book, out_book):
            print("⚠ 画面にしか出ない警告です")
    ''')
    assert _classes(found, "b") == {"print_after_out": 1}
    fixed = _scan(tmp_path, '''
        def finish(a, book, out_book):
            lines = []
            def _say(m):
                lines.append(m)
            _say("⚠ 運ばれる警告です")
    ''')
    assert _classes(fixed, "b") == {"carried": 1}


def test_a_warning_printed_straight_after_the_output_is_flagged(tmp_path):
    found = _scan(tmp_path, '''
        import shutil
        def cmd_x(a):
            lines = []
            print("⚠ 警告です（前）")
            shutil.copy2(a.src, a.out)
            print("⚠ 警告です（後）")
            lines.append("⚠ 運ばれる警告です")
            return lines
    ''')
    c = _classes(found, "b")
    assert c == {"print_no_out": 1, "print_after_out": 1, "carried": 1}, c


def test_a_cap_without_a_word_that_says_so_is_flagged(tmp_path):
    silent = '''
        MAX_ROWS = 10
        def read(ws):
            return ws[:MAX_ROWS]
    '''
    said = '''
        MAX_ROWS = 10
        def read(ws):
            if len(ws) > MAX_ROWS:
                print("先頭 10 行だけ読みました（上限）")
            return ws[:MAX_ROWS]
    '''
    assert _classes(_scan(tmp_path, silent), "c") == {"silent": 1}
    assert _classes(_scan(tmp_path, said), "c") == {"said": 1}


def test_a_skip_that_never_says_what_it_skipped_is_flagged(tmp_path):
    silent = '''
        def f(rows):
            out = []
            for r in rows:
                if r.is_total:
                    continue
                out.append(r)
            return out
    '''
    said = '''
        def f(rows):
            out = []
            for r in rows:
                if r.is_total:
                    print("合計行は除いて数えました")
                    continue
                out.append(r)
            return out
    '''
    assert _classes(_scan(tmp_path, silent), "c") == {"silent": 1}
    assert _classes(_scan(tmp_path, said), "c") == {"said": 1}


def test_a_write_without_a_gate_is_flagged(tmp_path):
    bare = '''
        import shutil
        def put(a):
            shutil.copy2(a.src, a.out)
    '''
    gated = '''
        import shutil
        def put(a):
            if a.out.exists() and not a.overwrite:
                return 7
            shutil.copy2(a.src, a.out)
    '''
    assert _classes(_scan(tmp_path, bare), "d") == {"ungated": 1}
    assert _classes(_scan(tmp_path, gated), "d") == {"gated": 1}


def test_a_new_undisclosed_site_turns_the_comparison_red(tmp_path):
    """③ 変異: 台帳に無い在り処を足した一時ファイルの走査を、実物の集計に足すと台帳との比較が赤になる。"""
    base = dict(_real_tally())
    extra = core.tally(_scan(tmp_path, '''
        import shutil
        def cmd_new(a):
            shutil.copy2(a.src, a.out)
            return 3
        def build(p):
            p.set_defaults(func=cmd_new)
    '''))
    assert extra, "変異が走査に出ない（対照が死んでいる）"
    merged = {**base, **extra}
    diff = core.compare(merged, core.load_register())
    assert diff["unlisted"], "新しい在り処が台帳に無いのに赤くならない"
    # 件数が増える向きも赤
    key = next(iter(base))
    bumped = dict(base)
    bumped[key] = base[key] + 1
    assert core.compare(bumped, core.load_register())["wrong_count"]


# --- ④ 直した実物を動かす（実 LibreOffice・実 ollama は要らない）------------------------------------------

def _book(d: Path, name: str = "商品リスト.xlsx") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["品番", "品名", "数量"])
    for r in [["0012", "ガラス花瓶", 3], ["0003", "タンブラー", 7]]:
        ws.append(r)
    p = d / name
    wb.save(p)
    wb.close()
    return p


def _history() -> list:
    return list(reversed(ailine.read_history(max_n=50)))


@pytest.mark.parametrize("boom,kind", [(KeyboardInterrupt(), "interrupted"),
                                        (RuntimeError("想定外"), "aborted")])
def test_a_run_that_dies_by_exception_names_and_records_the_work_file(tmp_path, monkeypatch, capsys, boom, kind):
    """★★ 4 件目の実物: 適用の途中で例外（Ctrl-C を含む）が抜けても、`.out` を言い・履歴に自分が作ったと残し、
    次の run が『この道具が書いた記録がありません』と塞がれない。例外は握りつぶさず投げ直す。"""
    book = _book(tmp_path)
    out = book.with_name(book.stem + ".out.xlsx")

    def _dispatch(a, b, workdir):
        shutil.copy2(b, out)            # 作業結果を作ってから落ちる
        raise boom

    monkeypatch.setattr(ailine, "_cmd_run_dispatch", _dispatch)
    with pytest.raises(type(boom)):
        ailine.main(["run", str(book), "品番の小さい順に並べ替えて", "--copy"])
    screen = capsys.readouterr().out
    assert out.exists()
    assert out.name in screen, "★ `.out` を残したのに画面が黙っている:\n" + screen
    rows = [r for r in _history() if r.get("out") and out.name in str(r["out"])]
    assert rows, f"★ 自分が作った `.out` を履歴に残していない: {_history()}"
    assert rows[-1]["failure_kind"] == kind and rows[-1]["ok"] is False
    assert ailine.refuse_if_output_is_someone_elses(
        book, "", chosen=out) is None, "★ 1 分前に自分が作った物を、次の run が塞いでいる"


def test_an_exception_before_the_output_exists_says_nothing_extra(tmp_path, monkeypatch, capsys):
    book = _book(tmp_path)

    def _dispatch(a, b, workdir):
        raise RuntimeError("途中で落ちた（出力はまだ）")

    monkeypatch.setattr(ailine, "_cmd_run_dispatch", _dispatch)
    with pytest.raises(RuntimeError):
        ailine.main(["run", str(book), "並べ替えて", "--copy"])
    assert "途中で止まりました" not in capsys.readouterr().out
    assert not [r for r in _history() if r.get("failure_kind") in ("aborted", "interrupted")]


def test_an_untouched_earlier_work_file_is_not_claimed_as_this_runs(tmp_path, capsys):
    """前から在る `.out` に触れていない回は『残しました』と名乗らない（この run が作った物でない）。"""
    book = _book(tmp_path)
    out = book.with_name(book.stem + ".out.xlsx")
    shutil.copy2(book, out)
    a = SimpleNamespace(book=str(book), task="t", model="m", json=False, out=None, copy=True,
                        inplace=False)
    ailine._finish_aborted(a, book, KeyboardInterrupt(), ailine._file_digest(out))
    assert capsys.readouterr().out == ""
    assert not _history()


def test_a_mid_plan_clarify_that_stops_names_the_work_file(tmp_path, capsys):
    """複合計画の 2 段目で聞き返しに落ちる（1 段目は適用済み）と `.out` が残る ── 言い、履歴の out も `.out`。"""
    book = _book(tmp_path)
    out = book.with_name(book.stem + ".out.xlsx")
    shutil.copy2(book, out)
    a = SimpleNamespace(book=str(book), task="売上シートの品番を並べ替えて", model="m", json=False,
                        out=None, copy=True, inplace=False)
    ailine._stamp_book(book)
    assert ailine._answer_before_asking(a, book, {}, ["表"], out_book=out) is True
    screen = capsys.readouterr().out
    assert out.name in screen, screen
    rows = [r for r in _history() if r.get("failure_kind", "").endswith("sheet_missing")]
    assert rows and rows[-1]["out"] == str(out), rows
    # out_book を渡さない呼び出し（単発の聞き返し）は従来どおり原本のパス
    capsys.readouterr()
    assert ailine._answer_before_asking(a, book, {}, ["表"]) is True
    assert out.name not in capsys.readouterr().out


def _csv(d: Path, body: bytes) -> Path:
    p = d / "売上.csv"
    p.write_bytes(body)
    return p


def _prestage(p: Path) -> int:
    """`run x.csv` の暗黙前段だけを動かす（その先の `_cmd_run_body` は呼び出し側が差し替える）。"""
    return ailine._cmd_run_csv_prestage(SimpleNamespace(book=str(p), task="並べ替えて", model="m", json=False))


def test_the_csv_prestage_records_its_conversion_and_its_warnings(tmp_path, monkeypatch, capsys):
    """`run x.csv` の暗黙前段が書いた xlsx を履歴に残す（⚠ ごと）── 残さないと次の run が塞がれ、⚠ は画面と消える。"""
    monkeypatch.setattr(ailine, "_cmd_run_body", lambda a: 0)
    p = _csv(tmp_path, "品番,数量\n0012,3\n0003\x01x,7\n".encode("utf-8"))
    rc = _prestage(p)
    assert rc == 3
    screen = capsys.readouterr().out
    assert "制御文字" in screen
    rows = [r for r in _history() if r.get("path") == "csv"]
    assert rows and rows[-1]["failure_kind"] == "csv_needs_review", rows
    assert any("制御文字" in n for n in rows[-1]["disclosed"]), rows[-1]
    assert rows[-1]["out_sha"], "出力の指紋が無い（次の run が『俺が置いたまま』と分からない）"


def test_a_second_csv_run_does_not_silently_rebuild_the_work_done_on_the_first(tmp_path, monkeypatch, capsys):
    """★ 置き換え: `run x.csv` を 2 回打つと、2 回目が 1 回目の作業結果を**黙って**作り直していた
    （前段が履歴に残さないので `edited_since` が素通り・実機で再現）。残すようにしたので、手が入っていれば断る。"""
    monkeypatch.setattr(ailine, "_cmd_run_body", lambda a: 0)
    p = _csv(tmp_path, "品番,数量\n0012,3\n0003,7\n".encode("utf-8"))
    assert _prestage(p) == 0
    xl = p.with_suffix(".xlsx")
    assert xl.exists()
    capsys.readouterr()
    assert _prestage(p) == 0                                        # 手が入っていない ── 作り直してよい
    wb = openpyxl.load_workbook(xl)
    wb.active["C1"] = "1 回目の run が足した列"                     # run が育てた、を真似る
    wb.save(xl)
    wb.close()
    capsys.readouterr()
    assert _prestage(p) == 7
    screen = capsys.readouterr().out
    assert "そのあと変更されています" in screen and "直接指定して実行してください" in screen, screen
    wb = openpyxl.load_workbook(xl)
    assert wb.active["C1"].value == "1 回目の run が足した列", "★ 作業結果が黙って消された"
    wb.close()


def test_export_csv_with_a_warning_leaves_it_in_the_history(tmp_path, capsys):
    """△ の回（式で計算結果の無いセルを空欄で書き出した）が、履歴では ✓ と区別がつかなかった。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["品名", "金額"])
    ws.append(["a", "=1+2"])         # openpyxl は計算結果を持たない
    p = tmp_path / "見積.xlsx"
    wb.save(p)
    wb.close()
    rc = ailine.main(["export-csv", str(p)])
    screen = capsys.readouterr().out
    assert rc == 0 and "△" in screen, screen
    rows = [r for r in _history() if r.get("path") == "export-csv"]
    assert rows and rows[-1]["ok"] is True
    assert any("空欄で書き出しました" in n for n in rows[-1]["disclosed"]), rows[-1]


def test_export_pdf_leaves_what_it_made_in_the_history(tmp_path, monkeypatch, capsys):
    """export-csv は履歴に残るのに export-pdf は 1 行も残らなかった（2 冊照合 4a61518 の兄弟）。"""
    book = _book(tmp_path)

    def _fake_pdf(book_path, out_path, **kw):
        Path(out_path).write_bytes(b"%PDF-1.4")
        return True, ""

    monkeypatch.setattr(ailine, "_soffice_to_pdf", _fake_pdf)
    monkeypatch.setattr(ailine.pdf_export, "verify_values_in_pdf",
                        lambda *a, **k: SimpleNamespace(available=True, missing=[], checked=6))
    rc = ailine.main(["export-pdf", str(book)])
    assert rc == 0, capsys.readouterr().out
    rows = [r for r in _history() if r.get("path") == "export-pdf"]
    assert rows and rows[-1]["ok"] is True and rows[-1]["out"].endswith(".pdf"), rows
    assert rows[-1]["out_sha"], "PDF の指紋が無い"
    # 値が見つからない回（rc 3）も PDF は残る ── 失敗として履歴に残す
    monkeypatch.setattr(ailine.pdf_export, "verify_values_in_pdf",
                        lambda *a, **k: SimpleNamespace(available=True, missing=["x"], checked=6))
    rc = ailine.main(["export-pdf", str(book), "--overwrite"])
    assert rc == 3
    rows = [r for r in _history() if r.get("path") == "export-pdf"]
    assert rows[-1]["ok"] is False and rows[-1]["failure_kind"] == "pdf_values_missing", rows[-1]


def test_demo_that_stops_halfway_names_what_it_already_placed(tmp_path, monkeypatch, capsys):
    real = shutil.copy2
    calls = {"n": 0}

    def _flaky(src, dst, *a, **k):
        calls["n"] += 1
        if calls["n"] == 3:
            raise PermissionError("読み取り専用")
        return real(src, dst, *a, **k)

    monkeypatch.setattr(ailine.shutil, "copy2", _flaky)
    rc = ailine.main(["demo", "--out", str(tmp_path / "d")])
    screen = capsys.readouterr().out
    assert rc == ailine.EXIT_ENVIRONMENT
    assert "置けた分は残っています" in screen, screen
    placed = sorted(p.name for p in (tmp_path / "d").iterdir())
    assert len(placed) == 2 and all(n in screen for n in placed), (placed, screen)


# --- ④' 実機（LibreOffice）で本当に起きた形 ── `-m local` ----------------------------------------------------

def _run_real(argv: list, plan=None, boom_in_apply=None):
    """製品を実機で 1 回走らせ、(exit, 画面) を返す。翻訳は固定する（毎回引ける）。"""
    import builtins
    import contextlib
    import io
    buf = io.StringIO()
    real_t, real_input, real_apply = ailine.translate_task, builtins.input, ailine.apply_dsl_step
    if plan is not None:
        ailine.translate_task = lambda *a, **k: {"plan": plan}

    def _no_tty(*_a, **_k):
        raise EOFError

    builtins.input = _no_tty
    if boom_in_apply is not None:
        def _boom(*_a, **_k):
            raise boom_in_apply
        ailine.apply_dsl_step = _boom
    rc = None
    try:
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
            try:
                rc = ailine.main(argv)
            except SystemExit as e:
                rc = e.code
            except BaseException as e:   # noqa: BLE001 ── 投げ直されることを見る
                rc = type(e).__name__
    finally:
        ailine.translate_task, builtins.input, ailine.apply_dsl_step = real_t, real_input, real_apply
    return rc, buf.getvalue()


_SORT = {"op": "SORT", "args": {"col": "品番", "order": "asc"}}


@pytest.mark.local
@pytest.mark.parametrize("flags", [["--copy"], []])
def test_ctrl_c_during_the_apply_does_not_strand_the_work_file(tmp_path, flags):
    """★★ 実機: 適用の途中の Ctrl-C。`.out` を言い・履歴に残し・次の run が塞がれない（--copy も既定の原本直接も）。"""
    book = _book(tmp_path)
    out = book.with_name(book.stem + ".out.xlsx")
    rc, screen = _run_real(["run", str(book), "品番の小さい順に並べ替えて", *flags], plan=[_SORT],
                           boom_in_apply=KeyboardInterrupt())
    assert rc == "KeyboardInterrupt", (rc, screen[-600:])
    assert out.exists(), "★ 引き金が引けていない ── 適用の前に落ちて `.out` が無い: " + screen[-600:]
    assert out.name in screen, "★ `.out` を残したのに画面が黙っている: " + screen[-700:]
    assert any(out.name in str(r.get("out")) for r in _history()), "★ 履歴に自分が作った `.out` が無い"
    rc2, screen2 = _run_real(["run", str(book), "品番の小さい順に並べ替えて", *flags], plan=[_SORT])
    assert "書いた記録がありません" not in screen2, screen2[-700:]


@pytest.mark.local
def test_a_clarify_after_a_first_step_names_the_work_file(tmp_path):
    """★★ 実機: 複合計画の 2 段目が聞き返し（存在しないシートの名指し）で止まる。1 段目は適用済みで `.out` が残る。"""
    book = _book(tmp_path)
    out = book.with_name(book.stem + ".out.xlsx")
    plan = [_SORT, {"op": "CLARIFY", "question": "どのシートですか"}]
    rc, screen = _run_real(["run", str(book), "売上シートの品番を並べ替えて、あと何か", "--copy"], plan=plan)
    assert rc == 3, (rc, screen[-600:])
    assert out.exists(), "★ 引き金が引けていない: " + screen[-600:]
    assert out.name in screen, "★ `.out` を残したのに画面が黙っている: " + screen[-700:]
    outs = [str(r.get("out")) for r in _history() if r.get("out")]
    assert any(out.name in o for o in outs), f"★ 履歴の out が `.out` でない: {outs}"
