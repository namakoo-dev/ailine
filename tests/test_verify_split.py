# -*- coding: utf-8 -*-
"""③ 分けた冊を**後から・ファイルだけから**独立に検算する（2026-09-12）。

## 何を縛るか

書き手（`cmd_split`）は既に書いた冊を読み戻して検算している。★ しかしその分母は
**書き手自身の意図（plan）**だった。`ailine verify` 側にはこの段が無く、`unsupported` を
正直に返すだけだった ── 「正直に無い」は在ることにならない。

契約（`ailine_core/verify_split.py`）:
  - 規則（誰にどの行を渡すか）を**再現しない**。分母は元の冊から取る
  - 元の非空行は「配った」か「検分が理由つきで名指しした」のどちらかに**ちょうど 1 回**
  - 出力の各行の値が元の行と**1 セルずつ**一致（規則を知らなくても壊れは分かる）
  - 金額: 配った ＋ 空欄 ＋ 複数担当 ＝ 元（★ 分けない行は除く ── 合計行が入るため）
  - ★★ **検分に書かれた数は読まない** ── 書き手の主張を分母にしたら検算にならない。
    この 1 点は「検分の数を書き換えても判定が変わらない」ことで縛る（下）
  - 1 冊だけ渡された形・印の無いフォルダは、合格でも不合格でもなく exit 4

★ LLM も LibreOffice も要らない（読むだけ・openpyxl のみ）。
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from ailine_core import split_people, verify_split   # noqa: E402
from ailine_core import stack as stack_core   # noqa: E402
from test_split_e2e import _book, _split   # noqa: E402 ── 同じ検体を使い回す（二重定義を作らない）

AMOUNT = "金額"


def _verify(out: Path, source: Path, *extra):
    return subprocess.run(
        [sys.executable, "-m", "ailine", "verify", str(out), str(source), *extra],
        capture_output=True, text=True, timeout=180, encoding="utf-8",
        errors="replace", cwd=str(REPO))


def _edit(path: Path, fn):
    wb = openpyxl.load_workbook(path)
    fn(wb.worksheets[0])
    wb.save(path)
    wb.close()


def _last_row(ws) -> int:
    return max(r for r in range(2, ws.max_row + 1)
               if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)))


def _headers(ws) -> list:
    return [str(c.value or "").strip() for c in ws[1]]


@pytest.fixture()
def split_out(tmp_path):
    """検体を 1 冊分けて、(出力フォルダ, 元の冊) を返す。"""
    book = _book(tmp_path / "元" / "売上一覧.xlsx")
    out = tmp_path / "配る"
    r = _split(book, out, "--amount", AMOUNT)
    assert r.returncode == 0, f"分けられなかった: {r.stdout}\n{r.stderr}"
    return out, book


def test_a_clean_split_verifies_from_the_files_alone(split_out):
    """★ 陽性対照 ── 素の出力は破れ 0 で通る。ここが 0 でなければ検算器を先に疑う。"""
    out, book = split_out
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"
    assert "破れはありません" in r.stdout, r.stdout
    # ★ 「破れ 0」だけでは**何も測っていない 0** と区別が付かない ── 分母が出ていることも縛る。
    for token in ("元の非空行", "配られた行", "Σ元"):
        assert token in r.stdout, f"分母が出ていない（{token}）: {r.stdout}"


def test_a_lost_row_is_named_with_its_row_number(split_out):
    """★ 消えたものは diff に出ない ── 配られなかった行を**元の行番号で**名指しする。"""
    out, book = split_out
    _edit(out / "山田.xlsx", lambda ws: ws.delete_rows(_last_row(ws)))
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "取り逃し" in r.stdout and "行目" in r.stdout, r.stdout
    assert "金額の和が閉じない" in r.stdout, r.stdout


def test_a_row_handed_to_two_people_is_caught(split_out):
    """★ 二重配布 ── 行数と Σ だけを見ていると気づけない（両方増えるので）。"""
    out, book = split_out
    got = []
    _edit(out / "山田.xlsx", lambda ws: got.append([c.value for c in ws[_last_row(ws)]]))
    _edit(out / "佐藤.xlsx", lambda ws: ws.append(got[0]))
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "二重配布" in r.stdout, r.stdout


def test_a_changed_value_is_caught_without_knowing_the_rules(split_out):
    """★★ 一番強い縛り ── 誰に配るかの規則を知らなくても、値の書き換えは分かる。"""
    out, book = split_out

    def tamper(ws):
        ws.cell(_last_row(ws), 2).value = "書き換えました"
    _edit(out / "山田.xlsx", tamper)
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "値が元と違う" in r.stdout, r.stdout


def test_a_row_that_never_existed_in_the_source_is_caught(split_out):
    """★ 捏造の側 ── 元に無い行番号を配ったら鳴る（取り逃しの裏返し）。"""
    out, book = split_out

    def forge(ws):
        head = _headers(ws)
        vals = [c.value for c in ws[_last_row(ws)]]
        vals[head.index(stack_core.PROVENANCE_HEADERS[1])] = 9999
        ws.append(vals)
    _edit(out / "山田.xlsx", forge)
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "捏造" in r.stdout, r.stdout


def test_the_writers_own_numbers_are_not_the_denominator(split_out):
    """★★★ 独立性そのものの縛り ── 検分に書かれた**数**（行数・金額）を書き換えても
    判定は変わらない。変わるなら、それは書き手の主張を分母にしている（恒真）。"""
    out, book = split_out
    report = out / (split_people.REPORT_STEM + ".xlsx")

    def lie(ws):
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 3).value is not None:
                ws.cell(r, 3).value = 999        # 行数
            if ws.cell(r, 4).value is not None:
                ws.cell(r, 4).value = 123456789  # 金額
    _edit(report, lie)
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 0, f"検分の数を信じてしまっている: {r.stdout}\n{r.stderr}"


def test_the_naming_in_the_inspection_sheet_is_actually_read(split_out):
    """★ 逆向き ── 名指しを消すと、配らなかった行が「取り逃し」になる。
    （上の試験と対で「数は読まない／名指しは読む」を挟み込む。）"""
    out, book = split_out
    report = out / (split_people.REPORT_STEM + ".xlsx")

    def erase(ws):
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "") in verify_split.NOT_DISTRIBUTED:
                ws.cell(r, 2).value = "（消しました）"
    _edit(report, erase)
    r = _verify(out, book, "--amount", AMOUNT)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "取り逃し" in r.stdout, r.stdout


def test_one_book_alone_is_refused_with_the_shape_that_works(split_out):
    """★ 空虚な合格を名乗らない ── 1 冊では和が閉じないので断る。
    ★ 断る時は**通る形を名指しする**（「できません」だけでは人は動けない）。"""
    out, book = split_out
    r = _verify(out / "山田.xlsx", book.parent)
    assert r.returncode == 4, f"{r.stdout}\n{r.stderr}"
    assert "ailine verify <出力フォルダ> <元の冊>" in r.stdout, r.stdout


def test_a_folder_without_our_mark_is_refused(tmp_path):
    """★ 人のフォルダを「検算した」ことにしない。"""
    stranger = tmp_path / "他人"
    stranger.mkdir()
    _book(stranger / "だれかの一覧.xlsx")
    source = _book(tmp_path / "元" / "売上一覧.xlsx")
    r = _verify(stranger, source)
    assert r.returncode == 4, f"{r.stdout}\n{r.stderr}"
    assert "印がありません" in r.stdout, r.stdout


def test_not_measuring_the_amount_is_said_out_loud(split_out):
    """★ 出ないことを合格の証拠にしない ── `--amount` が無い回は「測っていません」と書く。"""
    out, book = split_out
    r = _verify(out, book)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"
    assert "測っていません" in r.stdout, r.stdout


def test_the_verifier_does_not_reproduce_the_rules():
    """★★ 恒真の防止を**機械で**縛る ── 検算器が規則（誰に配るか）を呼んだら赤。

    ★ ここを緩めると、書き手と検算器が同じ間違いをして一致してしまう
      （この repo で実測済み: 合計行を両方が同じ関数で落として exit 0 になった）。
    """
    import ast
    text = (REPO / "src" / "ailine_core" / "verify_split.py").read_text(encoding="utf-8")
    # ★ 文字列で探した初版は docstring の「`plan_split` を呼ばない」に反応して赤くなった。
    #   **説明は呼び出しではない** ── 物差しは AST（実際に使っている名前）で取る。
    used, defined = set(), set()
    for node in ast.walk(ast.parse(text)):
        if isinstance(node, ast.Attribute):
            used.add(node.attr)
        elif isinstance(node, ast.Name):
            used.add(node.id)
        elif isinstance(node, ast.FunctionDef):
            defined.add(node.name)
    # ★ 「無いこと」の assert は単独では通ってしまう（集合が空でも通る）。
    #   **見えていることを先に示す**。★ 定義した名前と使った名前は別の集合
    #   （初版は `verify_split_folder` を「使った名前」に探して赤くなった ── 定義は使用でない）。
    for expected in ("verify_split_folder", "named_rows"):
        assert expected in defined, f"物差しが読めていない（定義に {expected} が無い）"
    for expected in ("read_grid", "read_core_properties"):
        assert expected in used, f"物差しが読めていない（使用に {expected} が無い）"
    for forbidden in ("plan_split", "lookalike_pairs", "entity_core", "total_row",
                      "find_header_row"):
        assert forbidden not in used, f"検算器が規則を呼んでいる: {forbidden}"
