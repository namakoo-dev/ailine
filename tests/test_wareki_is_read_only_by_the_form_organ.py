# 和暦（令和8年8月31日）を読む範囲の番人 ── 2026-09-12
#
# ★★ この試験の主題は「読めること」より **どこまでに留めたか**:
#   帳票を読む器官（form_read）だけが和暦を受ける。依頼文の値を読む
#   `parse_date_literal` と、列の型を決める `classify_date_column` は**今までどおり**。
#   測ったのは帳票の側だけなので、`ailine run` の経路へ黙って広げない。
#   ★ 「2 つを 1 つに畳む」方向の変更は、ここが赤くなる形で止める。
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from ailine_core.date_compare import (classify_date_column, parse_date_literal,
                                      parse_wareki_literal)
from ailine_core.field_record import grade, value
from ailine_core.form_grid import Grid
from ailine_core.form_read import read_issue_date


@pytest.mark.parametrize("raw, want", [
    ("令和8年8月31日", dt.date(2026, 8, 31)),
    ("令和元年5月1日", dt.date(2019, 5, 1)),       # ★ 元年 = 1 年
    ("平成31年4月30日", dt.date(2019, 4, 30)),
    ("昭和64年1月7日", dt.date(1989, 1, 7)),
    ("令和8/8/31", dt.date(2026, 8, 31)),
    (" 令和 8 年 8 月 31 日 ", dt.date(2026, 8, 31)),
])
def test_the_era_calendar_is_exact_arithmetic(raw, want):
    """★ 元号は法で決まる閉じた一覧 ── 言い回しの列挙（足し続ける形）ではない。"""
    assert parse_wareki_literal(raw) == want


@pytest.mark.parametrize("raw", [
    "R8.8.31", "H31/4/1", "S64.1.7",     # ★ 略記は帳票 178 冊で 0 件・R8 は品番にも見える
    "令和ビル3-2-1",                      # 住所
    "令和6年度分",                        # 年度（月日が無い）
    "令和8年2月30日",                     # 存在しない日
    "令和0年1月1日",
    "2026/8/31", "", "100", None,
])
def test_what_the_era_calendar_refuses(raw):
    assert parse_wareki_literal(raw) is None


def test_a_date_object_is_not_this_functions_job():
    """★ 既に日付のものは None を返す ── 呼び分けを曖昧にしない（合成は呼び出し側）。"""
    assert parse_wareki_literal(dt.date(2026, 8, 31)) is None


# ── ★ ここから「広げていない」ことの番人（負の被覆） ─────────────────
def test_the_request_side_parser_still_refuses_the_era_calendar():
    """★★ `parse_date_literal` は依頼文の値を読み、列の型判定にも使われる。
    ここに和暦を混ぜると `ailine run` の経路へ黙って波及する ── 測った範囲は帳票だけ。
    ★ 広げたくなったら、先に run 側で測ること（この試験を消すのは測った後）。"""
    assert parse_date_literal("令和8年8月31日") is None
    assert parse_date_literal("2026/8/31") == dt.date(2026, 8, 31)


def test_a_column_of_era_dates_is_not_promoted_to_a_date_column():
    """★ 列の型判定も今までどおり ── 和暦の列を日付列と名乗ると、比較や並べ替えが動く。"""
    kind, _has_time = classify_date_column(["令和8年8月31日", "令和8年7月31日"])
    assert kind == "other", kind
    assert classify_date_column(["2026/8/31", "2026/7/31"])[0] == "text_date"


# ── 器官の側（帳票を読む） ────────────────────────────────
def _sheet(date_cell_value):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "請求書"
    ws["G3"] = "請求日："
    ws["H3"] = date_cell_value
    return ws


def test_the_form_organ_reads_an_era_date_at_its_label():
    rec = read_issue_date(Grid.read(_sheet("令和8年8月31日")))
    assert grade(rec) == "単", grade(rec)
    assert value(rec) == dt.date(2026, 8, 31)


def test_the_form_organ_still_refuses_a_word_that_only_looks_like_an_era():
    """★ 『令和6年度分』は日付ではない ── 値を作らず、何を探したかを言う。"""
    rec = read_issue_date(Grid.read(_sheet("令和6年度分")))
    assert grade(rec) == "無", grade(rec)
    assert value(rec) is None
    assert "和暦" in rec.blank_reason, rec.blank_reason
