# -*- coding: utf-8 -*-
"""`ailine forms <folder> --out <path>` の検体 ── 需要①の出口（2026-09-11）。

契約:
  - 帳票の山を読み、**1 冊 1 行**の一覧を新しいブックに書く（原本は読むだけ）
  - 確/単 は値を書き、**割/無 は空欄**（2026-09-11 に凍結した判断）
  - 空欄は**必ず**検分シートに理由つきで 1 行出る（空欄の数 == 理由の数）
  - 請求書でないものからは値を作らない（官公庁の実表・空の雛形）
  - 出力先に人のファイルがあれば exit 7 で止まる（stack と同じ関所）
  - 自分の前回出力は入力に数えない（二重計上を防ぐ）

★ LLM も LibreOffice も要らない（読むだけ・openpyxl のみ）。
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent


def _invoice(path: Path, issuer: str, total: int, *, addressee: str = "ナギ商会株式会社"):
    """最小の請求書（帳票の形）。見出しの位置も帯も実物の形に合わせる。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求書"
    ws["B2"] = "請求書"
    ws["B3"] = addressee
    ws["B4"] = "経理部　御中"
    ws["G5"] = issuer
    ws["G6"] = "〒000-0000"
    ws["G3"] = "請求日："
    ws["H3"] = "2026/8/31"
    ws["G4"] = "請求番号："
    ws["H4"] = f"INV-{total}"
    ws["B11"] = " ご請求金額　"
    ws["C11"] = total
    ws["B15"] = "品番・品名"
    ws["E15"] = "数量"
    ws["G15"] = "単価"
    ws["H15"] = "金額"
    ws["B16"] = "用紙代"
    ws["E16"] = 1
    ws["G16"] = total // 11 * 10
    ws["H16"] = total // 11 * 10
    ws["E37"] = "小計"
    ws["H37"] = total // 11 * 10
    ws["E38"] = "消費税"
    ws["G38"] = 0.1
    ws["H38"] = total - total // 11 * 10
    ws["E39"] = "合計金額"
    ws["H39"] = total
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def _forms(folder: Path, out: Path, *extra):
    return subprocess.run(
        [sys.executable, "-m", "ailine", "forms", str(folder), "--out", str(out), *extra],
        capture_output=True, text=True, timeout=180, encoding="utf-8",
        errors="replace", cwd=str(REPO))


def _sheets(out: Path) -> dict:
    wb = openpyxl.load_workbook(out)
    got = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    got["_creator"] = wb.properties.creator
    wb.close()
    return got


@pytest.fixture()
def folder(tmp_path):
    d = tmp_path / "受領"
    _invoice(d / "a.xlsx", "あかね商事株式会社", 3300)
    _invoice(d / "b.xlsx", "いろは工業株式会社", 5500)
    return d


def test_one_row_per_invoice_with_the_fields_filled(folder, tmp_path):
    """★ 帳票の山 → 1 冊 1 行の一覧。"""
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, f"落ちた: {r.stdout}\n{r.stderr}"
    got = _sheets(out)
    assert got["_creator"] == "ailine forms", "★ 書き手の印が無い（次回 自分の出力と分からない）"
    head, *rows = got["一覧"]
    assert head == ["元ファイル", "請求元", "宛先", "請求額(税込)", "請求日", "請求番号"], head
    assert len(rows) == 2, rows
    by_file = {r0[0]: r0 for r0 in rows}
    assert by_file["a.xlsx"][1] == "あかね商事株式会社"
    assert by_file["a.xlsx"][2] == "ナギ商会株式会社"
    assert by_file["a.xlsx"][3] == 3300
    assert by_file["b.xlsx"][3] == 5500


def test_the_originals_are_not_touched(folder, tmp_path):
    """★ 原本は 1 バイトも変わらない（読むだけの契約）。"""
    import hashlib
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.glob("*.xlsx")}
    _forms(folder, tmp_path / "一覧.xlsx")
    after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.glob("*.xlsx")}
    assert before == after, "★ 原本が変わった"


def test_a_book_that_is_not_an_invoice_yields_no_values(folder, tmp_path):
    """★★ 請求書でないものから値を作らない（G2' の必達『混入で沈黙』）。

    ★ 官公庁の統計表を混ぜても、値は 1 つも出ない。
    ★★ 2026-09-13（買い手役の初見・経理）: 初版は**全列が空の行**として一覧に残していた ──
      その一覧に `run "合計行を付けて"` を頼むと 3/3 で ×（事後条件の検証対象が 0 件）。
      自分で作った空行が次の道具を殺すので、**一覧には載せず**、検分に 5 項目ぶんの理由を残す。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "家計調査"
    ws["A1"] = "家計調査"
    ws["A3"] = "世帯数"
    ws["B3"] = 1234
    ws["A4"] = "金額"
    ws["B4"] = 5678
    wb.save(folder / "統計.xlsx")
    wb.close()

    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0
    got = _sheets(out)
    names = [row[0] for row in got["一覧"][1:]]
    assert "統計.xlsx" not in names, f"★ 請求書でない冊が一覧に行として残った: {names}"
    assert len(names) == 2, names                      # 本物の 2 冊はそのまま
    reasons = [row for row in got["検分"][1:] if row[0] == "統計.xlsx"]
    assert len(reasons) == 5, f"外した冊の理由が検分に無い: {reasons}"
    assert "3 ファイル中 3 冊を読みました" in r.stdout, r.stdout   # ★ 分母は動かさない
    assert "一覧には載せていません" in r.stdout, r.stdout


def test_the_list_never_carries_an_all_blank_row(folder, tmp_path):
    """★ 空行は次の道具（合計行を付ける run）の事後条件を 0 件にする ── 1 行も作らない。"""
    _not_an_invoice(folder / "送付状.xlsx", "送付状")
    _not_an_invoice(folder / "稟議書.xlsx", "稟議書")
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    rows = _sheets(out)["一覧"][1:]
    assert len(rows) >= 2, rows
    blank = [row for row in rows if all(v in (None, "") for v in row[1:])]
    assert blank == [], blank


def test_verify_counts_a_left_out_book_instead_of_calling_it_a_break(folder, tmp_path):
    """★ 独立検算の側も同じ線 ── 理由つきで外した冊は「取り逃し」でなく、数えて名指しする。"""
    _not_an_invoice(folder / "稟議書.xlsx", "稟議書")
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    v = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(folder)],
                       cwd=str(REPO), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=300)
    assert v.returncode == 0, f"exit={v.returncode} / {v.stdout}"
    assert "一覧に載っていない冊（検分に理由あり）: 1 件（稟議書.xlsx）" in v.stdout, v.stdout
    assert "フォルダに在るのに一覧に無い冊" not in v.stdout, v.stdout


def test_every_blank_has_a_reason_in_the_inspection_sheet(folder, tmp_path):
    """★★ 空欄の数 == 理由の数（凍結した判断・型でも守っているが、ここでも数える）。

    ★「型が守っているはず」は検算ではない ── 出力を読んで数える。
    """
    # ★ 2026-09-13: 「何かのメモ」1 枚（項目が 1 つも取れない冊）は一覧に載らなくなった
    #   （空行が次の道具を殺すため）。空欄の分母は**請求書なのに 3 項目が欠けた冊**で作る。
    _invoice(folder / "欠け.xlsx", "", 7700)
    wb = openpyxl.load_workbook(folder / "欠け.xlsx")
    ws = wb.active
    for at in ("G3", "H3", "G4", "H4"):        # 請求日・請求番号のラベルと値を消す
        ws[at] = None
    wb.save(folder / "欠け.xlsx")
    wb.close()

    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    got = _sheets(out)
    assert any(row[0] == "欠け.xlsx" for row in got["一覧"][1:]), "★ 金額の取れた冊が一覧から消えた"
    blanks = sum(1 for row in got["一覧"][1:] for v in row[1:] if v is None)
    inspect = {(x[0], x[1]) for x in got["検分"][1:]}
    for row in got["一覧"][1:]:
        for i, v in enumerate(row[1:]):
            if v is None:
                field = ["請求元", "宛先", "請求額", "請求日", "請求番号"][i]
                assert (row[0], field) in inspect, f"★ 空欄に理由が無い: {row[0]}/{field}"
    assert blanks >= 3, f"★ 空欄が出ていない（検体が弱い）: {blanks}"
    reasons = [x[3] for x in got["検分"][1:] if (x[0], x[1]) in inspect]
    assert all(str(t).strip() for t in reasons), "★ 理由が空の行がある"


def test_the_report_names_the_denominator_and_the_blanks(folder, tmp_path):
    """★ 買い手の信用条件: 分母・区分の内訳・空欄の数を人に見せる。"""
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert "2 ファイル中 2 冊を読みました" in r.stdout, r.stdout
    assert "項目の区分" in r.stdout, r.stdout


def test_an_existing_human_file_stops_the_write(folder, tmp_path):
    """★ 出力先に人のファイルがあれば exit 7（stack と同じ関所）。"""
    out = tmp_path / "一覧.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "人が書いた大事な表"
    wb.save(out)
    wb.close()
    before = out.read_bytes()
    r = _forms(folder, out)
    assert r.returncode == 7, f"★ 人のファイルを黙って上書きした: {r.stdout}"
    assert out.read_bytes() == before, "★ 原本が変わった"
    assert "--overwrite" in r.stdout, "★ 次の手を言っていない"


def test_its_own_previous_output_is_not_counted_as_an_input(folder, tmp_path):
    """★ 前回の一覧を同じフォルダに置いても、入力に数えない（二重計上を防ぐ）。"""
    out = folder / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    assert "2 ファイル中 2 冊を読みました" in r.stdout, r.stdout
    assert "自分の出力" in r.stdout, r.stdout


def test_json_carries_the_denominator_and_the_blank_accounting(folder, tmp_path):
    """★ 機械可読の側にも分母と空欄の会計を載せる（人向けと同じ事実）。"""
    r = _forms(folder, tmp_path / "一覧.xlsx", "--json")
    payload = json.loads(r.stdout.strip().splitlines()[-1])
    assert payload["denominator"] == 2 and payload["collected"] == 2
    assert payload["blanks"] == payload["blanks_with_reason"], payload
    assert payload["grades"], payload


def test_the_inspection_sheet_does_not_cry_wolf(folder, tmp_path):
    """★★ 検分に載るのは**人が手を動かすものだけ**（空欄）── 普通の状態を並べない。

    ★ 実測（検体 87 冊）: 初版は「単（裏が取れていない）」も 1 行ずつ出していて、
      単 201 行が 割 12 + 無 15 の **27 行を埋めていた**。
      単 は口が 1 つしかない普通の状態で、それを所見として並べるのは
      「オオカミ少年防止を謳う道具が自分でオオカミ少年になる」形。
    """
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    got = _sheets(out)
    grades = {x[2] for x in got["検分"][1:]}
    assert "単" not in grades, f"★ 普通の状態（単）を検分に並べている: {got['検分'][1:3]}"
    assert "確" not in grades, "★ 裏が取れたものまで検分に並べている"


def test_the_grades_are_still_reachable_for_automation(folder, tmp_path):
    """★ 検分から外した『単/確』は **消えていない** ── --json に 1 冊 1 項目で残る。

    ★ 画面から消すことと、機械可読から消すことは別 ── 出す先を分けただけにする。
    """
    r = _forms(folder, tmp_path / "一覧.xlsx", "--json")
    payload = json.loads(r.stdout.strip().splitlines()[-1])
    fg = payload["field_grades"]
    assert len(fg) == 2 * 5, fg          # 2 冊 × 5 項目
    assert {x["field"] for x in fg} == {"請求元", "宛先", "請求額", "請求日", "請求番号"}
    assert all(x["grade"] in ("確", "単", "割", "無") for x in fg), fg


def test_a_missing_record_counts_as_a_blank_without_a_reason():
    """★ 負の被覆: 記録そのものが無い項目は『理由の無い空欄』として数えられること。

    ★ 2026-09-11 に実際に生まれた形 ── read_book の一枝が自前の 3 つ組を持っていて、
      請求日・請求番号を足したとき記録が無い空欄ができた。型（Record）は記録が無ければ
      働かない。初版の blanks_have_reasons はこれを continue で飛ばしていて数えなかった。
    """
    sys.path.insert(0, str(REPO / "src"))
    from ailine_core import forms_collect
    blanks, with_reason, bad = forms_collect.blanks_have_reasons([("x.xlsx", {})])
    assert blanks == len(forms_collect.FIELDS) and with_reason == 0
    assert bad and all("記録が無い" in b for b in bad), bad


def test_verify_tells_the_truth_about_its_own_forms_output(folder, tmp_path):
    """★ 穴 A（2026-09-11）: `ailine verify` は forms の出力に自分の印を見つけながら
    「印がありません」と言っていた。

    ★ 2026-09-13（③）: 独立の検算が入ったので、ここは「対応外」ではなく**本当に検算して**
      返る。検算そのものの契約は `tests/test_verify_forms.py`（含有・空欄の理由・冊の数）。
      この試験が守るのは**印の読み違いをしないこと**だけ。
    """
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    from ailine_core import verify as multifile_verify
    got = multifile_verify.verify_output(out, folder)
    assert not got.get("unmarked"), "★ 自分の印を『印が無い』と言っている"
    assert not got.get("unsupported"), got
    assert got.get("breaks") == [], got
    r = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(folder)],
                       capture_output=True, text=True, timeout=180, encoding="utf-8",
                       errors="replace", cwd=str(REPO))
    assert r.returncode == 0, r.stdout
    assert "印がありません" not in r.stdout, r.stdout


def test_a_duplicated_invoice_shows_up_in_the_bundle_findings(folder, tmp_path):
    """★★ 需要①の売り物: 1 冊ずつは正常でも、束で見ると『同じ請求書が 2 通』が分かる。"""
    import shutil
    shutil.copyfile(folder / "a.xlsx", folder / "a (2).xlsx")
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    got = _sheets(out)
    head, *rows = got["束の所見"]
    assert head == ["種類", "関わる冊", "なぜ怪しいか"], head
    assert len(rows) == 1, rows
    assert rows[0][0] == "重複" and "a (2).xlsx" in rows[0][1] and "a.xlsx" in rows[0][1], rows
    assert "重複" in rows[0][2]
    assert "束で見て怪しいもの 1 件" in r.stdout, r.stdout
    # ★ 値は作っていない ── 一覧は 3 行、どの値も所見で書き換わっていない
    assert len(got["一覧"]) == 4
    assert {x[3] for x in got["一覧"][1:]} == {3300, 5500}


def test_a_normal_folder_has_an_empty_bundle_sheet_and_no_warning(folder, tmp_path):
    """★ 陰性対照 ── 仕込みの無いフォルダで所見を出したらオオカミ少年。"""
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    got = _sheets(out)
    assert got["束の所見"][1:] == [], got["束の所見"]
    assert "怪しいもの" not in r.stdout, r.stdout


def test_json_carries_the_bundle_findings(folder, tmp_path):
    import shutil
    shutil.copyfile(folder / "b.xlsx", folder / "b_copy.xlsx")
    r = _forms(folder, tmp_path / "一覧.xlsx", "--json")
    payload = json.loads(r.stdout.strip().splitlines()[-1])
    assert [s["種類"] for s in payload["suspicions"]] == ["重複"], payload["suspicions"]
    assert sorted(payload["suspicions"][0]["冊"]) == ["b.xlsx", "b_copy.xlsx"]


# ── PDF（2026-09-12・設計 D7/D9）──────────────────────────────
FIXTURE_PDF = REPO / "tests" / "fixtures" / "forms" / "あかね商事_2026-08.pdf"


def test_a_pdf_in_the_folder_gets_its_own_row_next_to_the_excel_books(folder, tmp_path):
    """★★ 受け取る請求書は PDF が本流 ── xlsx と pdf が混ざったフォルダで両方が一覧に載る。"""
    import shutil
    shutil.copyfile(FIXTURE_PDF, folder / "あかね_pdf.pdf")
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    assert "3 ファイル中 3 冊を読みました" in r.stdout, r.stdout
    got = _sheets(out)
    by_file = {row[0]: row for row in got["一覧"][1:]}
    assert by_file["あかね_pdf.pdf"][1:4] == ["株式会社あかね商事", "ナギ商会株式会社", 33000], by_file
    assert by_file["あかね_pdf.pdf"][5] == "INV-2026-08-777"
    # ★ 原本は 1 バイトも触らない（PDF も）
    assert (folder / "あかね_pdf.pdf").read_bytes() == FIXTURE_PDF.read_bytes()


def test_a_scanned_pdf_is_named_as_unreadable_not_counted_as_read(folder, tmp_path):
    """★ 「読めなかった」と「請求書ではなかった」を混ぜない（D7）── 名指しして分母に残す。"""
    from test_pdf_grid import empty_page_pdf
    empty_page_pdf(folder / "スキャン.pdf")
    r = _forms(folder, tmp_path / "一覧.xlsx", "--json")
    assert r.returncode == 0, r.stdout
    payload = json.loads(r.stdout.strip().splitlines()[-1])
    assert payload["denominator"] == 3 and payload["collected"] == 2, payload
    assert [u["name"] for u in payload["unreadable"]] == ["スキャン.pdf"], payload["unreadable"]
    assert "テキスト層" in payload["unreadable"][0]["reason"]


def test_the_report_no_longer_lists_the_keys_of_the_excluded_dict(folder, tmp_path):
    """★ 初版は「（対象外）temp」など dict の鍵を毎回 5 行並べていた。"""
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert "（対象外）temp" not in r.stdout and "（対象外）other_format_names" not in r.stdout, r.stdout


# --- 経理の目で読める形（2026-09-13・買い手の初見 C8 / C9）------------------------
#
# ★★ 金額が `General` で `40000` と出ていた ── 経理は 3 桁区切りでないと目で検算しない
#   （1 桁の見落としがそのまま支払いになる）。★ 値は数値のまま・見え方だけ変える。


def test_the_amount_column_is_readable_as_money(folder, tmp_path):
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    wb = openpyxl.load_workbook(out)
    ws = wb["一覧"]
    col = [i + 1 for i, h in enumerate(next(ws.iter_rows(max_row=1, values_only=True)))
           if h == "請求額(税込)"][0]
    money = [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]
    assert len(money) == 2, f"分母（金額のセル）が変わった: {len(money)}"
    for cell in money:
        assert isinstance(cell.value, (int, float)), f"数値でなくなった: {cell.value!r}"
    assert {c.number_format for c in money} == {"#,##0"}, {c.number_format for c in money}
    # ★ 日付は openpyxl 既定の書式が付く（ここを一緒に壊していないことも見る）
    dcol = [i + 1 for i, h in enumerate(next(ws.iter_rows(max_row=1, values_only=True)))
            if h == "請求日"][0]
    assert "yy" in ws.cell(row=2, column=dcol).number_format, ws.cell(row=2, column=dcol).number_format
    wb.close()


def test_no_sheet_shows_empty_brackets(folder, tmp_path):
    """★ 空の括弧（`（）`）は「何か出すつもりで失敗した」に読める ── どのシートにも出さない。
    ★ 分母を言う: 3 枚のシートの**全文字列セル**を見ている（0 件を見て緑にしない）。"""
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    got = _sheets(out)
    texts = [v for name, rows in got.items() if name != "_creator"
             for row in rows for v in row if isinstance(v, str)]
    assert len(texts) >= 12, f"文字列セルが {len(texts)} 個しか無い（分母が痩せている）"
    bad = [t for t in texts if "（）" in t or "「」" in t or "（ ）" in t]
    assert not bad, bad


# --- 請求書でない冊と、空虚な合格（2026-09-13・買い手の初見 B7 / B8）---------------
#
# ★★ 送付状・稟議書を受領フォルダに混ぜたら、一覧に**全列が空の行**が並ぶだけで画面には
#   何も出ず、そのあと `ailine verify` が「含有を確かめた値 0 件」で **✓ 破れはありません**
#   と出した（実測）。0 件照合で合格を名乗るのは、この repo が何度も潰してきた形。


def _not_an_invoice(path: Path, title: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws["A1"] = title
    ws["A3"] = "ナギ商会株式会社 御中"
    ws["A5"] = "下記のとおりご連絡します"
    wb.save(path)
    wb.close()
    return path


def test_a_book_that_is_not_an_invoice_is_named_on_screen(tmp_path):
    """★ 「読めなかった」と「請求書ではなかった」を混ぜない（設計 D7 と同じ線）。"""
    folder = tmp_path / "受領"
    _not_an_invoice(folder / "送付状.xlsx", "送付状")
    _not_an_invoice(folder / "稟議書.xlsx", "稟議書")
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert r.returncode == 0, r.stdout
    assert "項目が 1 つも取れなかった冊 2 件" in r.stdout, r.stdout
    assert "送付状.xlsx" in r.stdout and "稟議書.xlsx" in r.stdout, r.stdout


def test_a_normal_folder_does_not_claim_anything_was_unreadable(folder, tmp_path):
    """★ 陰性対照 ── 普通の請求書で「1 つも取れなかった」と言ったらオオカミ少年。"""
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert "項目が 1 つも取れなかった冊" not in r.stdout, r.stdout


def test_verify_does_not_pass_a_list_with_no_values(tmp_path):
    """★★ 空虚な合格の禁止 ── 測るものが 0 件なら ✓ を出さず、exit 4（合格でも不合格でもない）。"""
    folder = tmp_path / "受領"
    _not_an_invoice(folder / "送付状.xlsx", "送付状")
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    v = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(folder)],
                       cwd=str(REPO), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=300)
    assert "含有を確かめた値: 0" in v.stdout, v.stdout
    assert "✓" not in v.stdout, f"0 件照合で合格を名乗った: {v.stdout}"
    assert "合格でも不合格でもありません" in v.stdout, v.stdout
    assert v.returncode == 4, f"exit={v.returncode} / {v.stdout}"


def test_verify_still_passes_a_list_that_has_values(folder, tmp_path):
    """★ 陰性対照 ── 値が在る一覧では今までどおり ✓（exit 0）。"""
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    v = subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(folder)],
                       cwd=str(REPO), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=300)
    assert v.returncode == 0 and "✓" in v.stdout, f"exit={v.returncode} / {v.stdout}"


# --- 税込の取り違え（2026-09-13・買い手役 3 体の初見 ── 事務職が自作の請求書 5 枚で踏んだ）------
#
# ★★★ `合計(税抜) → 消費税 → 税込合計` と**上から**並ぶ普通の請求書で、`請求額(税込)` に
#   **税抜**が入り、割にもならず、⚠ も出ず、verify は ✓ を出した（静かに 10% 少ない金額）。
#   真因は帯の合計の語に優先順位が無く**最初の一致で止めていた**こと。決め手は語でなく算術。


def _invoice_with_band(path: Path, band: list, *, issuer="株式会社さくら商会") -> Path:
    """帯の並びを外から指定できる最小の請求書（明細 1 行 100,000）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求書"
    ws["B2"] = "請求書"
    ws["B3"] = "ナギ商会株式会社"
    ws["B4"] = "経理部　御中"
    ws["G3"] = "請求日："
    ws["H3"] = "2026/8/31"
    ws["G4"] = "請求番号："
    ws["H4"] = "T-02"
    ws["G5"] = issuer
    ws["B15"] = "品名"
    ws["E15"] = "数量"
    ws["G15"] = "単価"
    ws["H15"] = "金額"
    ws["B16"] = "作業一式"
    ws["E16"] = 1
    ws["G16"] = 100000
    ws["H16"] = 100000
    for i, (label, value) in enumerate(band):
        ws[f"E{37 + i}"] = label
        ws[f"H{37 + i}"] = value
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _amount_and_grade(out: Path) -> tuple:
    wb = openpyxl.load_workbook(out)
    amount = [c.value for c in list(wb["一覧"].iter_rows(min_row=2, max_row=2))[0]][3]
    reasons = [r for r in wb["検分"].iter_rows(min_row=2, values_only=True) if r[1] == "請求額"]
    wb.close()
    return amount, reasons


def test_the_tax_inclusive_total_wins_over_the_pre_tax_total(tmp_path):
    """★★★ 合計(税抜)が税込合計より**上**に在っても、税込の方を採る（算術で決める）。"""
    folder = tmp_path / "受領"
    _invoice_with_band(folder / "B_合計と税込合計.xlsx",
                       [("合計", 100000), ("消費税", 10000), ("税込合計", 110000)])
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0, r.stdout
    amount, reasons = _amount_and_grade(out)
    assert amount == 110000, f"税抜が入った: {amount}"
    assert reasons == [], reasons


def test_two_different_totals_without_a_tax_row_are_a_split_not_a_value(tmp_path):
    """★ 算術で決まらないなら値を出さない（割）── 両方の番地と金額を名指しする。"""
    folder = tmp_path / "受領"
    _invoice_with_band(folder / "B2_消費税なし.xlsx", [("合計", 100000), ("税込合計", 110000)])
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    amount, reasons = _amount_and_grade(out)
    assert amount is None, f"決められないのに値を出した: {amount}"
    assert len(reasons) == 1 and reasons[0][2] == "割", reasons
    why = reasons[0][3]
    assert "H37" in why and "100,000" in why and "H38" in why and "110,000" in why, why


def test_the_ordinary_band_is_unchanged(tmp_path):
    """★ 陰性対照 ── 小計／消費税／税込合計 の普通の帯は今までどおり 110,000。"""
    folder = tmp_path / "受領"
    _invoice_with_band(folder / "A_小計と税込合計.xlsx",
                       [("小計", 100000), ("消費税", 10000), ("税込合計", 110000)])
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out).returncode == 0
    amount, reasons = _amount_and_grade(out)
    assert amount == 110000 and reasons == [], (amount, reasons)


def test_overwriting_its_own_previous_output_is_said_out_loud(folder, tmp_path):
    """★ 2026-09-13（買い手役 2/3）: 2 回目の実行が 1 回目と一字も変わらず、上書きしたとも言わなかった。"""
    out = tmp_path / "一覧.xlsx"
    assert "上書きします" not in _forms(folder, out).stdout
    r = _forms(folder, out)
    assert r.returncode == 0 and "前回の自分の出力 一覧.xlsx を上書きします" in r.stdout, r.stdout


# --- 期間外の混入を内訳で言う（2026-09-13・買い手役の初見・経理）------------------------------
#
# ★ 「9月受領分」に 5〜8 月が 4 冊（142,000 円）黙って混ざっていた。疑いにはしない（前月分が遅れて
#   混ざるのは実務で普通・検体の設計 §9.2）── 月ごとの冊数を 1 行言う。


def _redate(path: Path, text: str) -> None:
    wb = openpyxl.load_workbook(path)
    wb.active["H3"] = text
    wb.save(path)
    wb.close()


def test_mixed_months_are_counted_out_loud(folder, tmp_path):
    _redate(folder / "b.xlsx", "2026/5/31")
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert r.returncode == 0, r.stdout
    assert "請求日の月が 2 つ混ざっています: 2026年5月 1／2026年8月 1" in r.stdout, r.stdout
    assert "束で見て怪しいもの" not in r.stdout, "★ 月の違いを疑いにした（商慣行）"


def test_a_single_month_says_nothing_about_months(folder, tmp_path):
    """★ 陰性対照 ── 同じ月だけなら余計な行を足さない。"""
    r = _forms(folder, tmp_path / "一覧.xlsx")
    assert "混ざっています" not in r.stdout, r.stdout
