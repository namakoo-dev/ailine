# -*- coding: utf-8 -*-
"""需要③ 科目の候補の冊を**後から・ファイルだけから**独立に検算する（2026-09-13）。
   設計 docs/DESIGN-20260913-経費の勘定科目を先例から引く.md §6.5

## 何を縛るか

**鍵こそが規則**なので、検算側で鍵を作り直したら恒真になる。だから測るのは:
  - 出力が書いた**先例の番地**のセルを**読むだけ**で、候補の科目が本当にそこに在るか
  - 分母は**今回の入力**から: 借方が空で金額が在る行 ＝ 候補が出た行 ＋ 検分が理由を書いた行
  - ★★ 検分が書いた**件数は読まない**（書き手の主張を分母にしたら検算にならない）──
    「件数を書き換えても判定が変わらない」と「名指しを消すと鳴る」で挟み込む
  - 規則（`plan_accounts` / `KEYS`）を呼んでいないことを **AST** で縛る

★ LLM も LibreOffice も要らない（読むだけ・openpyxl のみ）。
"""
from __future__ import annotations

import ast
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from ailine_core import accounts_core, field_record   # noqa: E402
from ailine_core import stack as stack_core   # noqa: E402
# ★ 同じ検体を使い回す（二重定義を作らない ── test_verify_split と同じ作法）。
from test_accounts_e2e import _PAST, _TODAY, _accounts, _csv   # noqa: E402


def _verify(out: Path, *sources):
    return subprocess.run(
        [sys.executable, "-m", "ailine", "verify", str(out), *[str(s) for s in sources]],
        capture_output=True, text=True, timeout=300, encoding="utf-8",
        errors="replace", cwd=str(REPO))


@pytest.fixture()
def made(tmp_path):
    """検体を 1 冊だけ作って (候補の冊, 今回, 過去) を返す。"""
    today = _csv(tmp_path / "元" / "今回.csv", _TODAY)
    past = _csv(tmp_path / "元" / "過去.csv", _PAST)
    out = tmp_path / "候補.xlsx"
    r = _accounts(today, past, out)
    assert r.returncode == 0, f"作れなかった: {r.stdout}\n{r.stderr}"
    return out, today, past


def _edit(path: Path, sheet: str, fn):
    wb = openpyxl.load_workbook(path)
    fn(wb[sheet])
    wb.save(path)
    wb.close()


def _column(ws, header: str) -> int:
    for cell in ws[1]:
        if str(cell.value or "").strip() == header:
            return cell.column
    raise AssertionError(f"見出し『{header}』が無い")


def _row_of(ws, source_row: int) -> int:
    col = _column(ws, stack_core.PROVENANCE_HEADERS[1])
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col).value == source_row:
            return r
    raise AssertionError(f"元行 {source_row} の行が無い")


def test_a_clean_book_verifies_from_the_files_alone(made):
    """★ 陽性対照 ── 素の出力は破れ 0 で通る。ここが 0 でなければ検算器を先に疑う。"""
    out, today, past = made
    r = _verify(out, today, past)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"
    assert "破れはありません" in r.stdout, r.stdout
    # ★ 「破れ 0」だけでは**何も測っていない 0** と区別が付かない ── 分母も縛る。
    for token in ("今回の入力（借方が空で金額が在る行）", "候補が出た行",
                  "番地を読んで科目が在ることを確かめた件数"):
        assert token in r.stdout, f"分母が出ていない（{token}）: {r.stdout}"


def test_a_tampered_citation_is_caught_by_reading_the_cell(made):
    """★★ 一番強い縛り ── 番地を別の行にすり替えると、そこのセルを読んで鳴る。"""
    out, today, past = made

    def tamper(ws):
        row = _row_of(ws, 2)          # 通信費 を出した行
        col = _column(ws, accounts_core.OUTPUT_HEADERS[3])
        # ★ 過去 4 行目（消耗品費）へ番地をすり替える ── 科目はそこに無い。
        ws.cell(row, col).value = accounts_core.format_citations(
            [("借方取引先", past.name, 4)])
    _edit(out, accounts_core.SHEET_NAME, tamper)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "先例の番地にその科目が在りません" in r.stdout, r.stdout


def test_a_citation_pointing_at_a_row_that_does_not_exist_is_caught(made):
    """★ 捏造の側 ── 在りもしない行を根拠にしたら鳴る。"""
    out, today, past = made

    def forge(ws):
        row = _row_of(ws, 2)
        col = _column(ws, accounts_core.OUTPUT_HEADERS[3])
        ws.cell(row, col).value = accounts_core.format_citations(
            [("借方取引先", past.name, 9999)])
    _edit(out, accounts_core.SHEET_NAME, forge)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "先例の番地に行が無い" in r.stdout, r.stdout


def test_a_changed_account_is_caught_without_reproducing_the_rules(made):
    """★ 科目を書き換えると、番地のセルと合わないので鳴る（規則を知らなくても分かる）。"""
    out, today, past = made

    def tamper(ws):
        row = _row_of(ws, 2)
        ws.cell(row, _column(ws, accounts_core.OUTPUT_HEADERS[0])).value = "交際費"
    _edit(out, accounts_core.SHEET_NAME, tamper)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "交際費" in r.stdout and "捏造" in r.stdout, r.stdout


def test_an_account_on_a_row_that_must_stay_blank_is_caught(made):
    """★ 値を出さない区分（割 / 無）の行に科目が在れば鳴る（凍結した判断の側から測る）。"""
    out, today, past = made

    def tamper(ws):
        row = _row_of(ws, 4)          # 割 の行
        ws.cell(row, _column(ws, accounts_core.OUTPUT_HEADERS[0])).value = "旅費交通費"
    _edit(out, accounts_core.SHEET_NAME, tamper)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "値を出さない区分なのに科目が在る" in r.stdout, r.stdout


def test_a_grade_upgraded_without_a_second_key_is_caught(made):
    """★ 区分と番地の辻褄 ── 裏が取れたと名乗るなら別々の鍵を 2 本以上引いていること。"""
    out, today, past = made

    def tamper(ws):
        row = _row_of(ws, 3)          # 単 の行（鍵は 1 本）
        ws.cell(row, _column(ws, accounts_core.OUTPUT_HEADERS[1])).value = \
            field_record.CONFIRMED
    _edit(out, accounts_core.SHEET_NAME, tamper)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "区分と番地が食い違う" in r.stdout, r.stdout


def test_the_writers_own_counts_are_not_the_denominator(made):
    """★★★ 独立性そのものの縛り ── 検分に書かれた**件数**を書き換えても判定は変わらない。

    変わるなら、それは書き手の主張を分母にしている（恒真）。
    """
    out, today, past = made

    def lie(ws):
        col = list(accounts_core.REPORT_HEADERS).index("件数") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, col).value not in (None, ""):
                ws.cell(r, col).value = 999
    _edit(out, accounts_core.REPORT_SHEET, lie)
    r = _verify(out, today, past)
    assert r.returncode == 0, f"検分の件数を信じてしまっている: {r.stdout}\n{r.stderr}"


def test_the_naming_in_the_inspection_sheet_is_actually_read(made):
    """★ 逆向き ── 名指し（行番号）を消すと、空欄の行が「候補も理由も無い」で鳴る。

    （上の試験と対で「件数は読まない／名指しは読む」を挟み込む。）
    """
    out, today, past = made

    def erase(ws):
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "") == accounts_core.BLANK_KIND:
                ws.cell(r, 2).value = None
    _edit(out, accounts_core.REPORT_SHEET, erase)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "候補も理由も無い行" in r.stdout, r.stdout


def test_a_row_outside_the_declared_denominator_is_caught(made):
    """★ 宣言外の行（触らない行）に理由を立てたら鳴る（負の被覆の側）。"""
    out, today, past = made

    def add(ws):
        ws.append([accounts_core.BLANK_KIND, 7, "", "触らない行に理由を立てた"])
    _edit(out, accounts_core.REPORT_SHEET, add)
    r = _verify(out, today, past)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "宣言外の行" in r.stdout, r.stdout


def test_the_past_book_must_be_handed_over(made):
    """★ 出ないことを合格の証拠にしない ── 番地の冊が渡されなければ「確かめられなかった」。"""
    out, today, _past = made
    other = today.parent / "別の過去.csv"
    _csv(other, _PAST)
    r = _verify(out, today, other)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"
    assert "確かめられなかった番地" in r.stdout, r.stdout


def test_one_source_alone_is_refused_with_the_shape_that_works(made):
    """★ 空虚な合格を名乗らない ── 今回だけでは番地を読めないので断る。

    ★ 断る時は**通る形を名指しする**（「できません」だけでは人は動けない）。
    """
    out, today, _past = made
    r = _verify(out, today)
    assert r.returncode == 4, f"{r.stdout}\n{r.stderr}"
    assert "ailine verify <候補の冊> <今回の仕訳> <過去の仕訳…>" in r.stdout, r.stdout


def test_a_book_without_our_mark_is_refused(tmp_path):
    """★ 人のファイルを「検算した」ことにしない。"""
    stranger = tmp_path / "他人の表.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = accounts_core.SHEET_NAME
    wb.active.append(list(accounts_core.OUTPUT_HEADERS))
    wb.save(stranger)
    wb.close()
    today = _csv(tmp_path / "元" / "今回.csv", _TODAY)
    past = _csv(tmp_path / "元" / "過去.csv", _PAST)
    r = _verify(stranger, today, past)
    assert r.returncode == 4, f"{r.stdout}\n{r.stderr}"
    assert "印がありません" in r.stdout, r.stdout


def test_verify_from_the_folder_form_names_the_shape_that_works(made):
    """★ 元フォルダ 1 個の形（stack / forms）へ落ちても、誤診せず通る形を教える。"""
    out, today, _past = made
    r = _verify(out, today.parent)
    assert r.returncode == 4, f"{r.stdout}\n{r.stderr}"
    assert "印がありません" not in r.stdout, r.stdout
    assert "ailine verify <候補の冊> <今回の仕訳> <過去の仕訳…>" in r.stdout, r.stdout


def test_the_verifier_does_not_reproduce_the_rules():
    """★★ 恒真の防止を**機械で**縛る ── 検算器が規則（鍵）を呼んだら赤。

    ★ 文字列で探すと docstring の説明に反応する（**説明は呼び出しではない**）── 物差しは
      AST（実際に使っている名前）で取る。
    ★ 「無いこと」の assert は単独では通ってしまう（集合が空でも通る）── **見えている
      ことを先に示す**。定義した名前と使った名前は別の集合。
    """
    text = (REPO / "src" / "ailine_core" / "verify_accounts.py").read_text(encoding="utf-8")
    used, defined = set(), set()
    for node in ast.walk(ast.parse(text)):
        if isinstance(node, ast.Attribute):
            used.add(node.attr)
        elif isinstance(node, ast.Name):
            used.add(node.id)
        elif isinstance(node, ast.FunctionDef):
            defined.add(node.name)
    for expected in ("verify_accounts_book", "_blank_reasoned"):
        assert expected in defined, f"物差しが読めていない（定義に {expected} が無い）"
    for expected in ("read_grid", "read_core_properties", "parse_citations"):
        assert expected in used, f"物差しが読めていない（使用に {expected} が無い）"
    for forbidden in ("plan_accounts", "KEYS", "candidate_rows", "inspection_rows",
                      "reason_of", "lookalike_pairs", "_precedent_index"):
        assert forbidden not in used, f"検算器が規則を呼んでいる: {forbidden}"
