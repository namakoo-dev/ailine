# -*- coding: utf-8 -*-
"""空のシートは「見出しが決まらない」ではなく「**無い**」と言う（2026-09-21・盲検 5 体目 ②）。

★★ 出所（買い手が実際に打った 2 発・`bench/blind/5体目/argv.jsonl` に残っている）:

    09:56:07  run 空.xlsx 金額の合計を出して
    09:56:27  run 空.xlsx 金額の合計を出して --header-row 1

  1 打目 … 「見出しが何行目か分かりません。`--header-row 3` のように指定して再実行してください」
  2 打目 … 勧められたとおり打ったら、今度は **語彙外**（history.jsonl の failure_kind）

  ★★ **2 発とも原因は『冊が空』**なのに、画面が指した先は 2 回とも違うものだった。
    だから買い手は 2 回とも無駄打ちになった ── どちらの指摘に従っても永遠に直らない。

★★ これは新しい穴ではない ── **09-13 に塞いだ穴の、別の階層**。
  `multifile.nothing_to_read`（買い手役 3 体のうち 2 体・空のフォルダ）の docstring が
  既にこう書いている ── 「文書が無い」は 9（ENGINEERING.md の表）。
  同じ判断がシートの階層に配線されていなかった＝**片配線**。
  ★ 兄弟の番人: `tests/test_an_empty_folder_is_not_a_success.py`

★★ 線引き: 「**決まらない**」（材料はあるが一意に決まらない）は 3・
  「**無い**」（材料がゼロ）は 9。ENGINEERING.md の表が既にそう引いている。
  ★ ここを緩めると `--header-row` で本当に直る回まで 9 で止めてしまうので、
    陰性対照を 2 本置く（1 セルだけの冊／見出しが曖昧な冊）。

★ 塞いでも失う能力は無いと確かめた（実装前）: 1 セルに書く経路
  `resolve_cell_target_from_task` は実表の値を手がかりに位置を決めるので、
  手がかりが 0 の冊では元から None を返す（`if not rows or not heads: return None`）。
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

import ailine  # noqa: E402
from _product_source import window_around  # noqa: E402
from _run_argv import run_argv  # noqa: E402

#: ★ 買い手の依頼文そのまま（言い換えない ── 言い換えると再現でなく作文になる）。
BUYER_TASK = "金額の合計を出して"

#: ★ ENGINEERING.md の表「9 実行の前提が無い」。番号は下の試験が文書と突き合わせる。
MISSING_INPUT_EXIT = 9

#: ★ 見出しが**曖昧**な冊の StructDump（陰性対照用・既存の凍結試験と同じ形）。
AMBIGUOUS = {"sheets": {"Sheet": {"rows": {
    1: {"nonempty": 2, "str": 2, "bold": 0},
    2: {"nonempty": 2, "str": 1, "bold": 0},
    3: {"nonempty": 2, "str": 2, "bold": 0},
    4: {"nonempty": 2, "str": 1, "bold": 0},
}}}}


def _empty_book(tmp_path: Path, name: str = "空.xlsx") -> Path:
    """買い手の冊と同じもの ── 非空セル 0（実物も `_used_extent` が (0, 0)）。"""
    p = tmp_path / name
    openpyxl.Workbook().save(p)
    return p


def _book_with(tmp_path: Path, rows, name: str = "表.xlsx") -> Path:
    p = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(p)
    return p


def _run(tmp_path, monkeypatch, capsys, book, struct_dump=None, **over):
    """`ailine.main` を実際に回して (終了コード, 画面, 翻訳が呼ばれた回数) を返す。

    ★ 翻訳は**必ず爆弾**にする ── 「空だと気づく前に ollama を叩いた」回を緑にしない。
      買い手の 2 打目は翻訳まで行って『語彙外』と言われた。そこへ届く時点で負け。
    """
    monkeypatch.setattr(ailine, "HISTORY_FILE", tmp_path / "history.jsonl")
    monkeypatch.setattr(ailine, "build_struct_dump",
                        lambda book, workdir: (struct_dump if struct_dump is not None else {}))
    called = {"n": 0}

    def boom(*a, **k):
        called["n"] += 1
        return {"plan": [{"op": "FREEFORM", "args": {}}]}

    monkeypatch.setattr(ailine, "translate_task", boom)
    rc = ailine.main(run_argv(book=str(book), task=BUYER_TASK, **over))
    return rc, capsys.readouterr().out, called["n"]


# --- 事故そのもの ── 買い手が打った 2 発 -------------------------------------------

@pytest.mark.parametrize("over, shot", [
    ({}, "1 打目（素のまま）"),
    ({"header_row": 1}, "2 打目（勧められて --header-row 1 を足した）"),
])
def test_both_shots_the_buyer_typed_get_the_same_true_answer(
        tmp_path, monkeypatch, capsys, over, shot):
    """★★ **2 発とも同じ返事**になること ── 原因が同じなのだから返事も同じであるべき。

    ★ ここが 1 箇所に畳めていないと、2 打目だけ `--header-row` の分岐に吸われて
      翻訳まで行き、別の理由（語彙外）を言い出す＝買い手が踏んだ道。
    """
    rc, out, translated = _run(tmp_path, monkeypatch, capsys, _empty_book(tmp_path), **over)
    assert rc == MISSING_INPUT_EXIT, f"{shot}: 終了コードが {rc}（画面: {out!r}）"
    assert "表がありません" in out, f"{shot}: 空だと言っていない（{out!r}）"
    assert translated == 0, f"{shot}: 空だと気づく前に翻訳を呼んでいる（{translated} 回）"


def test_the_refusal_does_not_send_them_back_to_header_row(tmp_path, monkeypatch, capsys):
    """★★ **`--header-row` を勧めないこと** ── 買い手はそれを打って 2 打目を外した。

    ★ 「`--header-row` という文字を出すな」ではない（出さないと『じゃあどうすれば』になる）。
      **効かないと言い切る**のが要点。
    """
    _, out, _ = _run(tmp_path, monkeypatch, capsys, _empty_book(tmp_path))
    assert "では直せません" in out, f"`--header-row` が効かないと言っていない: {out!r}"
    assert "のように指定して再実行" not in out, f"まだ --header-row を勧めている: {out!r}"


def test_it_says_it_made_no_file(tmp_path, monkeypatch, capsys):
    """★ 断った回は「作っていない」と必ず言う（`nothing_to_read` と同じ作法）。

    ★ 09-13 の兄弟の事故は「作らず、作らなかったとも言わなかった」だった。
    """
    _, out, _ = _run(tmp_path, monkeypatch, capsys, _empty_book(tmp_path))
    assert "ファイルは作っていません" in out, out


# --- 陰性対照 ── 緩めるとここが赤くなる ------------------------------------------

def test_a_sheet_with_a_single_cell_is_not_empty(tmp_path, monkeypatch, capsys):
    """★★ 1 セルでも入っていれば「無い」ではない ── 従来の経路へ進むこと。

    ★ ここが赤くなる直し方＝「ほぼ空も空とみなす」。**材料がゼロ**の時だけが 9。
    """
    book = _book_with(tmp_path, [["金額"]])
    rc, out, _ = _run(tmp_path, monkeypatch, capsys, book)
    assert "表がありません" not in out, f"中身のある冊を空と言っている: {out!r}"
    assert rc != MISSING_INPUT_EXIT, f"中身のある冊を 9 で止めている（画面: {out!r}）"


def test_an_ambiguous_header_is_still_a_question(tmp_path, monkeypatch, capsys):
    """★★ 「決まらない」は今までどおり **3（聞き返し）** ── 9 に巻き込まないこと。

    ★ 曖昧な冊は `--header-row` で本当に直る。ここを 9 にすると直せる回を殺す。
    """
    book = _book_with(tmp_path, [["a", "b"], ["c", 1], ["d", "e"], ["f", 2]])
    rc, out, translated = _run(tmp_path, monkeypatch, capsys, book, struct_dump=AMBIGUOUS)
    assert rc == 3, f"曖昧な冊の終了コードが {rc}（画面: {out!r}）"
    assert "見出しが何行目か分かりません" in out, out
    assert "--header-row" in out, "曖昧な回では導線を出すべき"
    assert translated == 0


# --- 他のシートの案内は、中身のあるものだけ ------------------------------------------

def test_other_sheets_are_named_only_when_they_have_something(tmp_path):
    """★★ 全部空なのに `--sheet` を勧めたら、それこそ 3 打目の無駄打ちになる。"""
    wb = openpyxl.Workbook()
    wb.active.title = "空1"
    wb.create_sheet("空2")
    ws = wb.create_sheet("中身あり")
    ws.append(["金額"])
    p = tmp_path / "混在.xlsx"
    wb.save(p)

    name, others = ailine._sheet_with_nothing_in_it(p, "空1", ["空1", "空2", "中身あり"])
    assert name == "空1"
    assert others == ("中身あり",), f"空のシートまで勧めている: {others}"
    assert "--sheet" in ailine.nothing_in_the_sheet(name, others)

    wb2 = openpyxl.Workbook()
    wb2.active.title = "空1"
    wb2.create_sheet("空2")
    p2 = tmp_path / "全部空.xlsx"
    wb2.save(p2)
    name2, others2 = ailine._sheet_with_nothing_in_it(p2, "空1", ["空1", "空2"])
    assert (name2, others2) == ("空1", ()), (name2, others2)
    assert "--sheet" not in ailine.nothing_in_the_sheet(name2, others2), (
        "中身のあるシートが 1 枚も無いのに `--sheet` を勧めている")


def test_a_book_that_breaks_while_being_read_is_not_called_empty(tmp_path, monkeypatch):
    """★★ **開けたが途中で落ちた**回も『空』と言わない（2026-09-21・変異試験が見つけた）。

    ★ 初版は `load_workbook` が落ちる経路しか検体が無く、読み込めた**後**に落ちる経路を
      誰も見ていなかった ── 変異（そこを `return name, ()` に緩める）が**緑のまま通った**。
      ★ 「見ていないことを根拠にしない」は、両方の落ち方で守られて初めて守られている。
    """
    p = tmp_path / "途中で壊れる.xlsx"
    openpyxl.Workbook().save(p)

    def boom(ws):
        raise RuntimeError("読み取りの途中で落ちた")

    monkeypatch.setattr(ailine, "_used_extent", boom)
    assert ailine._sheet_with_nothing_in_it(p, "Sheet", ["Sheet"]) == (None, ())


def test_an_unreadable_book_is_not_called_empty(tmp_path):
    """★★ 読めなかった回を『空』と言わないこと ── **見ていないことを根拠にしない**。

    ★ secretscan と同じ線（見ていないので検出なしとは言えない）。
    """
    bad = tmp_path / "壊れ.xlsx"
    bad.write_bytes(b"not a zip")
    assert ailine._sheet_with_nothing_in_it(bad, "Sheet", ["Sheet"]) == (None, ())


# --- 契約の在り処 -------------------------------------------------------------------

def test_the_exit_code_is_the_one_the_table_documents():
    """★ 契約は文書側（`docs/ENGINEERING.md` の表）── 定数を読み合って恒真にしない。"""
    table = (REPO / "docs" / "ENGINEERING.md").read_text(encoding="utf-8")
    row = [ln for ln in table.splitlines() if ln.startswith(f"| {MISSING_INPUT_EXIT} |")]
    assert len(row) == 1, f"表に {MISSING_INPUT_EXIT} の行が {len(row)} 本ある"
    assert "シート" in row[0], (
        "表の 9 に『空のシート』が書かれていない ── 番号だけ揃えて文書に無いなら、"
        f"買い手も自動化も知りようがない: {row[0]}")


def test_the_check_uses_the_organ_that_already_exists():
    """★★ 空かどうかを**数え直さない** ── 判定は `_used_extent`（既存の器官）。

    ★ 同じ判断を 2 つ書くと、片方だけ直る日が来る（この穴自体がその形だった）。
    """
    seg = window_around("def _sheet_with_nothing_in_it", before=0, after=1600)
    assert "_used_extent" in seg, "空の判定を自前で数えている"
