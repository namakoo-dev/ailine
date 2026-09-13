# -*- coding: utf-8 -*-
"""20,000 円の動線（2026-09-13・買い手役 3 回目の 1 位）── 一往復で閉じる。

★ 経理: 「9 月分のフォルダを渡すと、9 月分だけの合計が出た一覧が返る」まで 1 本（`forms --month`）
★ 会計: 「採用」列に ○ を付けて渡し直すと、借方勘定科目を埋めた元の列順そのままの取込用ファイル
  （`accounts-apply`）── 毎月 200〜500 行の U→C 転記と末尾 6 列の削除が消える
"""
from __future__ import annotations

import csv
import io
import subprocess
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from test_forms_e2e import _invoice, _forms, _sheets, _not_an_invoice                # noqa: E402
from test_accounts_e2e import _csv, _accounts, MF                                   # noqa: E402


def _run(argv, cwd):
    return subprocess.run([sys.executable, "-m", "ailine", *argv], cwd=str(cwd),
                          capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=300)


def _redate(path, text):
    wb = openpyxl.load_workbook(path)
    wb.active["H3"] = text
    wb.save(path)
    wb.close()


# ── forms --month ─────────────────────────────────────────────────

def _month_folder(tmp_path):
    folder = tmp_path / "受領"
    _invoice(folder / "a.xlsx", "あかね商事株式会社", 3300)      # 2026/8/31
    _invoice(folder / "b.xlsx", "いろは工業株式会社", 5500)      # 2026/8/31
    _invoice(folder / "c.xlsx", "うめ物産株式会社", 7700)
    _redate(folder / "c.xlsx", "2026/5/31")                       # 別の月
    _invoice(folder / "d.xlsx", "えのき商店株式会社", 9900)
    _redate(folder / "d.xlsx", "")                                # 請求日なし
    _not_an_invoice(folder / "稟議書.xlsx", "稟議書")
    return folder


def test_month_keeps_only_that_month_and_adds_a_total(tmp_path):
    folder = _month_folder(tmp_path)
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out, "--month", "2026-08")
    assert r.returncode == 0, r.stdout
    assert "対象月 2026年8月: 一覧 2 冊／対象外 2 冊（別の月 1・請求日なし 1" in r.stdout, r.stdout
    assert "合計（対象月の請求額・一覧の末尾）: 8,800" in r.stdout, r.stdout
    got = _sheets(out)
    assert list(got)[:2] == ["一覧", "対象外"], list(got)
    names = [row[0] for row in got["一覧"][1:]]
    assert names == ["a.xlsx", "b.xlsx", "合計"], names
    assert got["一覧"][-1][3] == 8800, got["一覧"][-1]
    scope = {row[0]: row[6] for row in got["対象外"][1:]}
    assert "別の月の請求です（2026年5月）" in scope["c.xlsx"], scope
    assert "決められません" in scope["d.xlsx"], scope
    assert "稟議書.xlsx" not in scope and "稟議書.xlsx" not in names


def test_the_total_row_is_bold_and_formatted(tmp_path):
    folder = _month_folder(tmp_path)
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out, "--month", "2026-08").returncode == 0
    wb = openpyxl.load_workbook(out)
    ws = wb["一覧"]
    last = list(ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row))[0]
    assert last[0].value == "合計" and last[0].font.bold
    assert last[3].value == 8800 and last[3].number_format == "#,##0" and last[3].font.bold
    wb.close()


def test_verify_treats_the_out_of_scope_books_as_named_not_broken(tmp_path):
    folder = _month_folder(tmp_path)
    out = tmp_path / "一覧.xlsx"
    assert _forms(folder, out, "--month", "2026-08").returncode == 0
    v = _run(["verify", str(out), str(folder)], tmp_path)
    assert v.returncode == 0, f"exit={v.returncode} / {v.stdout}"
    assert "対象外の冊（--month で外した" in v.stdout and "c.xlsx" in v.stdout, v.stdout
    assert "合計行（検算の対象外）: 1 行" in v.stdout, v.stdout
    assert "✓" in v.stdout, v.stdout


def test_without_month_nothing_changes(tmp_path):
    """★ 陰性対照 ── 指定が無ければ今までどおり（対象外シートも合計行も無い）。"""
    folder = _month_folder(tmp_path)
    out = tmp_path / "一覧.xlsx"
    r = _forms(folder, out)
    assert r.returncode == 0 and "対象月" not in r.stdout, r.stdout
    got = _sheets(out)
    assert "対象外" not in got and got["一覧"][-1][0] != "合計", list(got)


def test_a_malformed_month_is_refused_in_words(tmp_path):
    folder = _month_folder(tmp_path)
    r = _forms(folder, tmp_path / "一覧.xlsx", "--month", "2026/8")
    assert r.returncode == 2 and "YYYY-MM" in r.stderr, r.stderr


# ── accounts-apply ────────────────────────────────────────────────

_PAST = [["1", "2026/08/25", "通信費", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "8月分"],
         ["2", "2026/08/26", "消耗品費", "", "アスクル", "3000", "未払金", "アスクル", "コピー用紙"]]
_TODAY = [["11", "2026/09/25", "", "本社回線", "ＮＴＴ西日本", "8800", "未払金", "ＮＴＴ西日本", "9月分"],
          ["12", "2026/09/26", "", "", "アスクル", "3000", "未払金", "アスクル", "コピー用紙"],
          ["13", "2026/09/27", "", "", "アマゾン", "1200", "未払金", "アマゾン", "電池"]]


def _candidates_with_adoption(tmp_path, encoding="utf-8-sig", preamble=()):
    today_path = tmp_path / "今回.csv"
    text = "\n".join(list(preamble) + [",".join(MF)] + [",".join(r) for r in _TODAY]) + "\n"
    today_path.write_bytes(text.encode(encoding))
    past = _csv(tmp_path / "過去.csv", _PAST)
    book = tmp_path / "候補.xlsx"
    assert _accounts(today_path, past, book).returncode == 0
    wb = openpyxl.load_workbook(book)
    ws = wb["候補"]
    col = ws.max_column + 1
    ws.cell(row=1, column=col, value="採用")
    cand = [i for i, c in enumerate(ws[1], start=1) if c.value == "候補の科目"][0]
    marks = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=cand).value and marks < 1:
            ws.cell(row=r, column=col, value="○")
            marks += 1
    wb.save(book)
    wb.close()
    return book, today_path, marks


def test_only_adopted_rows_are_written_into_the_original_column(tmp_path):
    book, today, marks = _candidates_with_adoption(tmp_path)
    assert marks == 1
    out = tmp_path / "取込用.csv"
    r = _run(["accounts-apply", str(book), str(today), "--out", str(out)], tmp_path)
    assert r.returncode == 0, r.stdout
    assert "採用 1 行" in r.stdout and "変わったセル 1 個" in r.stdout, r.stdout
    before = list(csv.reader(io.StringIO(today.read_text(encoding="utf-8-sig"))))
    after = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8-sig"))))
    diff = [(i, j) for i, (ra, rb) in enumerate(zip(before, after)) for j, (x, y) in enumerate(zip(ra, rb)) if x != y]
    assert diff == [(1, 2)], diff                       # ★ 1 行目のデータ・借方勘定科目だけ
    assert after[1][2] == "通信費" and after[2][2] == "" and after[3][2] == ""


def test_the_preamble_and_encoding_survive(tmp_path):
    """★ 弥生の形（説明行 2 行・cp932）でも、説明行と文字コードは元のまま。"""
    book, today, _ = _candidates_with_adoption(tmp_path, encoding="cp932",
                                               preamble=("仕訳日記帳", "会社名: ナギ商会"))
    out = tmp_path / "取込用.csv"
    r = _run(["accounts-apply", str(book), str(today), "--out", str(out)], tmp_path)
    assert r.returncode == 0, r.stdout
    raw = out.read_bytes()
    text = raw.decode("cp932")
    assert text.splitlines()[:2] == ["仕訳日記帳", "会社名: ナギ商会"], text[:60]
    assert "通信費" in text.splitlines()[3], text


def test_without_an_adoption_column_it_refuses_and_says_how(tmp_path):
    today = _csv(tmp_path / "今回.csv", _TODAY)
    past = _csv(tmp_path / "過去.csv", _PAST)
    book = tmp_path / "候補.xlsx"
    assert _accounts(today, past, book).returncode == 0
    r = _run(["accounts-apply", str(book), str(today), "--out", str(tmp_path / "取込用.csv")], tmp_path)
    assert r.returncode == 4 and "『採用』の列がありません" in r.stdout and "○" in r.stdout, r.stdout
    assert not (tmp_path / "取込用.csv").exists()


def test_the_output_must_keep_the_original_format(tmp_path):
    book, today, _ = _candidates_with_adoption(tmp_path)
    r = _run(["accounts-apply", str(book), str(today), "--out", str(tmp_path / "取込用.xlsx")], tmp_path)
    assert r.returncode == 4 and "同じ形式で" in r.stdout, r.stdout


def test_the_route_is_registered_everywhere():
    import ailine
    assert ailine.ROUTE_KIND["accounts-apply"] == "multi"
    assert ailine.NEEDS_MACHINE["accounts-apply"] is False
    assert "accounts-apply" in ailine.multi_file_routes()
