# -*- coding: utf-8 -*-
"""帳票を読む規則の番人（2026-09-11）。

★★ ここに在る試験は**全部、実測で踏んだ事故**から来ている。
  検体 v2（87 冊）と、一度も測っていない実物の雛形 26 冊で出たものだけを置く。
  ★ 想像で書いた「請求書はこうだろう」は 1 本も入れない ── それは俺の頭の中の
    請求書を守るだけで、買い手の帳票を守らない。

★ 検体そのものは repo に入らない（他社の帳票から作っているため）。
  だからここでは openpyxl でその場に**最小の帳票**を組み立てて縛る。
"""
from __future__ import annotations

import openpyxl
import pytest

from ailine_core.field_record import (CONFIRMED, NONE_FOUND, SINGLE, SPLIT,
                                       both_sides, describe, grade, value)
from ailine_core.form_grid import Grid
from ailine_core.form_read import clean_org_name, read_book, read_form


def book(rows: dict, merges=()):
    """{"B3": 値, …} から 1 枚の帳票を組む。戻り値: (値のシート, 式のシート)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for at, v in rows.items():
        ws[at] = v
    for m in merges:
        ws.merge_cells(m)
    return ws, ws


#: 使い回す最小の帳票 ── 上部の請求額・明細 2 行・帯（積み上げ型）。
def minimal(**over):
    rows = {
        "B2": "請求書",
        "B3": "ナギ商会株式会社",
        "B4": "経理部　御中",
        "G5": "あかね商事株式会社",
        "G6": "〒000-0000",
        "B11": " ご請求金額　", "C11": 3300, "D11": "（税込）",
        "B15": "品番・品名", "E15": "数量", "G15": "単価", "H15": "金額",
        "B16": "用紙代", "E16": 1, "G16": 1000, "H16": 1000,
        "B17": "運送費", "E17": 2, "G17": 1000, "H17": 2000,
        "E37": "小計", "H37": 3000,
        "E38": "消費税", "G38": 0.1, "H38": 300,
        "E39": "合計金額", "H39": 3300,
    }
    rows.update(over)
    return rows


def read(rows, merges=()):
    ws, wsf = book(rows, merges)
    return read_form(ws, ws_formula=wsf)


# ── ★ 名前を作る道は 1 本（片配線を塞ぐ） ──────────────────────
@pytest.mark.parametrize("field, cell", [("宛先", "B3"), ("請求元", "G5")])
def test_a_placeholder_name_is_refused_for_both_sides(field, cell):
    """★★ 1 本の試験で**宛先と請求元の両方**を縛る。

    ★ 2026-09-11、一度も測っていない実物の雛形 15 冊で踏んだ:
      プレースホルダ（`○○株式会社`）の番人を**請求元の側にだけ**書いていて、
      宛先は空の雛形に対して「宛先＝○○株式会社」を**値として出していた**。
    ★ 処方は「両方直す」ではなく「1 関数に畳んで呼び出し側に選ばせない」。
      だからこの試験も 1 本にする ── `clean_org_name` を壊せば**必ず両方赤くなる**。
    """
    recs = read(minimal(**{cell: "○○株式会社"}))
    r = recs[field]
    assert grade(r) == NONE_FOUND, f"★ 雛形のままの名前を {field} の値にした"
    assert "雛形" in r.blank_reason, f"★ 理由が雛形だと言っていない: {r.blank_reason}"
    assert "○○" in r.blank_reason, "★ 見つけたものを引用していない（人が直せない）"


def test_clean_org_name_is_the_only_door():
    """★ 名前を作る関数そのものの契約（両側がこれを通る）。"""
    assert clean_org_name("あかね商事株式会社") == ("あかね商事株式会社", "")
    assert clean_org_name("ナギ商会株式会社　御中")[0] == "ナギ商会株式会社"
    # ★ 1 セルに登録番号が改行で同居する（実物）
    assert clean_org_name("高梨産業株式会社\n(登録番号:T7010001234567)")[0] \
        == "高梨産業株式会社"
    # ★ 取次・担当の行は名前の行ではない
    assert clean_org_name("（取次）中央商事株式会社　担当 井上")[0] == ""
    assert clean_org_name("株式会社 〇〇〇")[0] == ""


# ── 宛先: 敬称の付いたセルが名前とは限らない ────────────────────
def test_the_honorific_cell_may_be_a_department_and_the_name_is_above():
    """★ `B3=社名` / `B4=経理部　御中` ── 御中が付くのは部署の側。

    ★ 2026-09-10 に「宛先 15/16 取れた」と誤報したのは、この御中セルを
      そのまま宛先として読んでいたから。
    """
    r = read(minimal())["宛先"]
    assert value(r) == "ナギ商会株式会社", f"★ 部署名を宛先にした: {value(r)!r}"


def test_the_honorific_may_live_in_the_name_cell_itself():
    """★ `B7=社名　御中` の骨もある。**両方の形が要る。**"""
    rows = minimal()
    del rows["B3"], rows["B4"]
    rows["B7"] = "ナギ商会株式会社　御中"
    r = read(rows)["宛先"]
    assert value(r) == "ナギ商会株式会社"


def test_a_gap_row_between_the_name_and_the_honorific_is_crossed():
    """★ 名前と御中のあいだが空くことがある（骨によって 1〜2 行）。"""
    rows = minimal()
    del rows["B4"]
    rows["B6"] = "総務課　御中"
    r = read(rows)["宛先"]
    assert value(r) == "ナギ商会株式会社"


# ── 請求額: ラベルの右をどう読むか ─────────────────────────────
def test_the_tax_rate_between_the_label_and_the_amount_is_not_the_amount():
    """★ `E38=消費税 | G38=0.1 | H38=300` ── いちばん近い右は**税率**。"""
    ws, wsf = book(minimal())
    g = Grid.read(ws)
    from ailine_core.form_read import _labelled_number, _LABEL_TAX
    got = _labelled_number(g, _LABEL_TAX, rows=(15, 60))
    assert got and got[0][1].value == 300, "★ 0.1（税率）を消費税額として拾った"


def test_a_money_amount_typed_as_text_is_not_silently_skipped():
    """★★ 答えの枠に文字で金額が入っていたら、**跨いで別の数を読まない**。

    ★ 2026-09-11 の実測: `¥1,320,000-` や `#REF!` を黙って跨ぎ、
      別のセルの数を「確」として出していた。
      黙って別のものを読むのが、この repo でいちばん高くつく失敗。
    """
    r = read(minimal(C11="¥1,320,000-"))["請求額"]
    assert grade(r) not in (CONFIRMED, SINGLE), "★ 別の数を黙って読んだ"
    assert "文字" in r.blank_reason or "読み取れ" in r.blank_reason


def test_a_broken_formula_in_the_amount_slot_is_reported():
    r = read(minimal(C11="#REF!"))["請求額"]
    assert grade(r) == NONE_FOUND
    assert "#REF!" in r.blank_reason


# ── ★ 一致だけで確と言わない・食い違いは区分に届く ────────────────
def test_the_detail_block_is_a_third_mouth():
    """明細の合計が帯と合っていれば、そこで初めて『確』が立つ。"""
    r = read(minimal())["請求額"]
    assert grade(r) == CONFIRMED
    assert value(r) == 3300
    assert "明細" in r.swept_how, f"★ 何を掃き出したか言っていない: {r.swept_how}"


def test_a_detail_that_disagrees_with_the_subtotal_blanks_the_amount():
    """★ 上部と帯が一致していても、明細と合わないなら値は出さない。

    ★ 集計範囲の外に明細が追記されている冊・SUMIF が行を落とす冊は、
      **上部と帯だけ見ていると一致していて気づけない**。
    """
    r = read(minimal(H17=9999))["請求額"]
    assert grade(r) == SPLIT, "★ 明細と小計の食い違いが区分に届いていない（片配線）"
    assert "明細" in r.blank_reason and "小計" in r.blank_reason
    assert value(r) is None


def test_a_row_whose_quantity_is_missing_is_caught_even_when_the_sum_agrees():
    """★★ 0 を足しても SUM は合う ── 満場一致で壊れた冊が通る形。

    ★ 数量の入力漏れで `単価 1,000 × 数量(空) = 金額 0`。
      小計・合計・上部の 3 つが揃って一致するので、行の中を見ないと分からない。
    """
    rows = minimal()
    rows["H17"] = 0
    del rows["E17"]                      # 数量が空
    rows["H37"] = 1000
    rows["H38"] = 100
    rows["H39"] = 1100
    rows["C11"] = 1100
    r = read(rows)["請求額"]
    assert grade(r) == SPLIT, "★ 数量が空で金額 0 の行を見逃した"
    assert "数量" in r.blank_reason, f"★ 何が悪いか言っていない: {r.blank_reason}"


def test_a_carried_forward_block_makes_the_amount_undecidable():
    """★ 繰越請求・源泉徴収がある帳票では、合計は**振り込む額ではない**。

    ★ 実物は `お振込金額　121,790 円` のように**ラベルと金額が同じセル**に入る。
    """
    r = read(minimal(B41="源泉所得税　▲10,210", B42="お振込金額　121,790 円"))["請求額"]
    assert grade(r) == SPLIT
    assert "源泉" in r.blank_reason, f"★ 見つけた欄を名指していない: {r.blank_reason}"


def test_the_workbooks_own_formula_is_not_an_independent_check():
    """★★ 合計が `=H37+H38` なら、小計＋消費税との一致は**検算ではない**。

    ★ 2026-09-10 の恒真（自分の写しと一致して裏が取れた）と同じ形。
    """
    ws = openpyxl.Workbook().active
    for at, v in minimal().items():
        ws[at] = v
    wsf = openpyxl.Workbook().active
    for at, v in minimal().items():
        wsf[at] = v
    wsf["H39"] = "=H37+H38"              # ★ 合計は小計と消費税の式そのもの
    recs = read_form(ws, ws_formula=wsf)
    r = recs["請求額"]
    ats = {e.at for e in r.evidences}
    assert "H37+H38" not in ats, "★ 帳票自身の式を 2 つ目の根拠として数えた（恒真）"
    assert any("検算になりません" in why for _at, _v, why in r.excluded), \
        "★ 数えなかったことを人に見せていない"


# ── 帯の形が 2 つある ────────────────────────────────────────
def test_the_band_can_be_a_table_where_the_tax_sits_in_another_column():
    """★ `C28=小計 | E28=金額` / `F26=消費税 | G26=…` ── 列が揃わない帯。

    ★ 初版は「消費税ラベルの右」を取り、**税区分の 10% 欄**を小計に足して
      偽の食い違いを出した（実測 R01/R02）。小計と同じ**行**の消費税列を見る。
    """
    rows = {
        "B1": "請求書",
        "B7": "ナギ商会株式会社　御中",
        "F2": "あかね商事株式会社", "F3": "〒000-0000",
        "B14": "ご請求金額", "E14": 153000,
        "C16": "品目", "E16": "単価", "F16": "数量", "G16": "金額",
        "B17": "用紙代", "E17": 90000, "F17": 1, "G17": 90000,
        "B18": "食品", "E18": 50000, "F18": 1, "G18": 50000,
        "C26": "10%対象", "D26": "対象額（税抜）", "E26": 90000,
        "F26": "消費税", "G26": 9000,
        "C27": "8%対象(※)", "E27": 50000, "G27": 4000,
        "C28": "小計", "E28": 140000, "G28": 13000,
    }
    r = read(rows)["請求額"]
    assert grade(r) != SPLIT, f"★ 表型の帯で偽の食い違いを出した: {r.blank_reason}"
    assert value(r) == 153000


def test_a_sheet_that_is_not_an_invoice_says_nothing():
    """★ 請求書でないものに答えない（実物の政府統計 6 冊で確認した振る舞い）。"""
    recs = read({"A1": "家計調査", "A3": "世帯数", "B3": 1234, "A4": "金額", "B4": 5678})
    for fld, r in recs.items():
        assert grade(r) == NONE_FOUND, f"★ 請求書でないのに {fld} を答えた: {value(r)!r}"
        assert r.blank_reason, "★ 空欄なのに理由が無い"


# ── ★ どのシートが請求書かも自分で決める ────────────────────────
def workbook(sheets: dict, hidden=()):
    """{"シート名": {"B3": 値, …}} から 1 冊を組む。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for at, v in rows.items():
            ws[at] = v
        if name in hidden:
            ws.sheet_state = "hidden"
    return wb


def test_the_explanation_sheet_is_not_the_invoice():
    """★ 1 枚目が『説明』で本体が 2 枚目（実物 12 冊がこの形）。

    ★★ 2026-09-11 まで、どのシートを読むかを**検体の答えから受け取っていた**。
      つまりこの問題を一度も解かずに満点を出していた ──
      買い手が持っていない手がかりで測っていた。
    """
    wb = workbook({"説明": {"A1": "この雛形の使い方", "A3": "① 社名を入れます"},
                   "適格請求書": minimal()})
    r = read_book(wb)["請求額"]
    assert value(r) == 3300, f"★ 説明シートを読んだ: {r.blank_reason}"


def test_two_invoice_sheets_in_one_book_cannot_be_decided():
    """★ 1 冊に請求書が 2 枚（8 月分・9 月分）── どちらか決められない。"""
    wb = workbook({"misoca_invoice": minimal(),
                   "請求書（9月分）": minimal(C11=4400, H39=4400)})
    for fld, r in read_book(wb).items():
        assert grade(r) == SPLIT, f"★ {fld}: 2 枚あるのに片方を選んだ"
        assert "シート" in r.blank_reason, f"★ 理由がシートを名指していない: {r.blank_reason}"
        assert value(r) is None


def test_a_found_conflict_is_never_reported_as_nothing_found():
    """★★ 食い違いを見つけたのに「何も無い」と言わない。

    ★ 初版は「根拠が 0 個なら 無」と逃げていて、請求書が 2 枚ある冊に対して
      「手がかりが見つかりません」と報告していた。人は探し方を疑う ── 嘘になる。
    """
    from ailine_core.field_record import grade_of
    assert grade_of((), conflict=True) == SPLIT, (
        "★ 食い違いを『無』と言った ── 決められないのであって、無いのではない")


def test_an_invoice_hidden_behind_an_empty_visible_sheet_is_reported_as_hidden():
    """★ 本体が非表示で、可視シートは空（実物にあった形）。

    ★ 読んでもよいが、**非表示だったことを人に言う**（黙って隠しシートを読まない）。
    """
    wb = workbook({"misoca_invoice": minimal(), "Sheet1": {}},
                  hidden=("misoca_invoice",))
    r = read_book(wb)["請求額"]
    assert value(r) == 3300
    assert "非表示" in (r.swept_how + r.blank_reason), "★ 隠れていたことを言っていない"


def test_a_book_with_no_invoice_sheet_names_what_it_looked_at():
    """★ どこにも請求書が無いなら、**見たシート名を並べて**空欄にする。"""
    wb = workbook({"家計調査": {"A1": "家計調査", "A3": "世帯数", "B3": 1234},
                   "注記": {"A1": "出典"}})
    for fld, r in read_book(wb).items():
        assert grade(r) == NONE_FOUND
        assert "家計調査" in r.blank_reason and "注記" in r.blank_reason, (
            f"★ 何を見たか言っていない: {r.blank_reason}")


def test_a_cover_letter_sheet_is_not_a_second_invoice():
    """★ 印が 1 つだけのシートを請求書と数えない。

    ★ 送付状には宛名（御中）が在る。そこだけ見て「請求書らしい」と数えると、
      送付状＋請求書の 1 冊で「請求書が 2 枚あります」と**偽の食い違い**が出る
      （正しい冊に印が立つ ── 誤報の次に悪い）。
    ★ 別の種類の印が 2 つ以上そろって初めて「らしい」とする。
    """
    wb = workbook({"送付状": {"B2": "ナギ商会株式会社　御中",
                              "B4": "平素は格別のお引き立てを賜り厚く御礼申し上げます。",
                              "B6": "下記の書類をお送りいたします。"},
                   "請求書": minimal()})
    r = read_book(wb)["請求額"]
    assert grade(r) != SPLIT, f"★ 送付状を 2 枚目の請求書と数えた: {r.blank_reason}"
    assert value(r) == 3300


def test_the_signals_are_of_different_kinds():
    """★ 印の数え方そのもの ── 同じ種類を 2 回数えて 2 にしない。"""
    from ailine_core.form_read import invoice_signals
    from ailine_core.form_grid import Grid
    ws, _ = book({"B2": "ナギ商会株式会社　御中", "B4": "経理部　御中", "B6": "総務課　御中"})
    sig = invoice_signals(Grid.read(ws))
    assert len(sig) == 1, f"★ 御中を 3 つ数えて 3 印にした: {sig}"


# ── ★ 請求元: 決められないなら決めない（2026-09-11・B′） ──────────
def test_the_addressee_mentioned_again_elsewhere_is_not_the_issuer():
    """★★ 別の行に再掲された**宛先**を、請求元として採らないこと。

    ★ 実測（検体 T10）: `D19=御請求先：ナギ商会株式会社／9 月分` が請求元の候補に残り、
      「右の列を優先」という**版面の癖**だけがそれを退けていた。癖を左優先に変えると
      この行が請求元として出る ── 検体が**禁止値**と名指ししている値だ。
    ★ 宛先の除外は「完全一致」では足りない。飾りが付いた再掲を**含んでいたら**除く。
    """
    rows = minimal()
    # ★ 再掲を請求元（G5）より**右**に置く ── 位置の癖（右優先）が外す側に倒れる配置。
    #   左に置くと癖がたまたま正解を拾って緑になり、**除外を一度も試さない検体**になる。
    rows["I41"] = "御請求先：ナギ商会株式会社／9 月分"
    r = read(rows)["請求元"]
    assert value(r) == "あかね商事株式会社", f"★ 宛先の再掲を請求元にした: {value(r)!r}"


def test_two_competing_organisations_are_not_decided_by_which_is_further_right():
    """★★ 発行元らしい名前が 2 つ残ったら、**位置で決めない**。

    ★ 「右のブロックほど発行者らしい」は実物の版面の癖であって根拠ではない。
      癖で当てにいくと、当たらない冊で**黙って間違った名前**を出す。
    ★ 区分の導出（grade_of）は、値の違う根拠が 2 つあれば 割 にする ── 新しい判断は要らない。
    """
    rows = minimal()
    rows["B30"] = "こだま産業株式会社"      # 左に別の組織名（宛先ではない）
    r = read(rows)["請求元"]
    assert grade(r) == SPLIT, f"★ 2 つ在るのに片方を選んだ: {value(r)!r}"
    assert value(r) is None
    both = {v for _at, v, _how in both_sides(r)}
    assert both == {"あかね商事株式会社", "こだま産業株式会社"}, both
    assert "あかね商事株式会社" in r.blank_reason and "こだま産業株式会社" in r.blank_reason


def test_a_registration_number_still_decides_between_two_organisations():
    """★ 登録番号は残す ── あれは「どちらのブロックが発行者か」の**役割**の証拠。

    ★ 名前の裏取り（値の証拠）には使わない ── 登録番号は名前を一言も言っていない。
      実測（2026-09-11）: 文書の属性を裏取りに使うと 90 冊で「弥生株式会社」を
      確として出す偽の裏取り装置になった。役割と値を混ぜない線はここで引く。
    """
    rows = minimal()
    rows["B30"] = "こだま産業株式会社"
    rows["G6"] = "T1234567890123"          # ★ あかね商事の側にだけ登録番号
    r = read(rows)["請求元"]
    assert value(r) == "あかね商事株式会社", f"★ 登録番号で絞れていない: {r.blank_reason}"
    assert "登録番号" in describe(r)


def test_a_single_candidate_still_passes():
    """★ 負の被覆 ── 候補が 1 つなら今までどおり値を出す（断りが広がりすぎない）。"""
    r = read(minimal())["請求元"]
    assert grade(r) == SINGLE and value(r) == "あかね商事株式会社"


# ── 請求日・請求番号（2026-09-11・束で疑う前に要る 4・5 つ目の項目） ──────────
import datetime as _dt


def test_issue_date_to_the_right_of_the_label_is_read_as_a_date():
    """★ `G3=請求日：` の右の値を日付として読む（文字の 3 形 と Excel の日付型）。"""
    for raw in ("2026/8/31", "2026-08-31", "2026年8月31日", _dt.datetime(2026, 8, 31)):
        r = read(minimal(G3="請求日：", H3=raw))["請求日"]
        assert value(r) == _dt.date(2026, 8, 31), f"★ {raw!r} を日付として読めない: {r.blank_reason}"


def test_issue_date_glued_to_its_label_in_one_cell():
    """★ 実物は `請求日：2026/8/31` と 1 セルに同居する形がある。ラベルを剥がしてから読む。"""
    r = read(minimal(G3="請求日：2026/8/31"))["請求日"]
    assert value(r) == _dt.date(2026, 8, 31), r.blank_reason


def test_a_placeholder_date_is_not_a_date():
    """★ 雛形の `××年1月1日` は日付ではない ── 値を作らず、理由で『請求日』を名指す。"""
    r = read(minimal(G3="請求日：", H3="××年1月1日"))["請求日"]
    assert grade(r) == NONE_FOUND, f"★ placeholder を日付にした: {value(r)!r}"
    assert "請求日" in r.blank_reason
    # ★ 区分はパーサだけでも 無 になる（「××年」は読めない）。この判定の役目は**文面** ──
    #   「日付として読めません」ではなく「雛形の埋め草のまま」と言うこと。ここを要求しないと
    #   判定を消しても緑のままで、番人が何も見ない（同日 4 回目の「重ねた守り」）。
    assert "雛形" in r.blank_reason, f"★ 埋め草だと言っていない: {r.blank_reason}"


def test_an_empty_date_slot_does_not_fall_through_to_the_cell_below():
    """★★ 事前測定で踏んだ罠: ラベルは縦に積まれ、値は右。右が空のとき「下」へ落ちると
       次のラベル（請求番号：）や**発行者名**を日付として拾う（probe で 77 件がそれだった）。
    """
    rows = minimal(G3="請求日：")             # 右（H3）は空。下（G4）は次のラベル、G5 は発行者名
    rows["G4"] = "請求番号："
    r = read(rows)["請求日"]
    assert grade(r) == NONE_FOUND, f"★ 下のセルを日付にした: {value(r)!r}"
    assert "請求日" in r.blank_reason


def test_invoice_number_is_read_right_of_its_label_and_never_from_below():
    """★ 請求番号も同じ線: 右だけ・同居も可・下には落ちない（下は発行者名 G5）。"""
    r = read(minimal(G4="請求番号：", H4="INV-2026-0831"))["請求番号"]
    assert value(r) == "INV-2026-0831", r.blank_reason
    r2 = read(minimal(G4="請求番号：INV-0001"))["請求番号"]
    assert value(r2) == "INV-0001", r2.blank_reason
    r3 = read(minimal(G4="請求番号："))["請求番号"]         # 右は空・下は 発行者名
    assert grade(r3) == NONE_FOUND, f"★ 発行者名を請求番号にした: {value(r3)!r}"


def test_missing_date_and_number_are_blank_with_a_reason():
    """★ ラベルそのものが無い冊 ── 無 ＋ 何を探したかの理由。"""
    recs = read(minimal())
    for field in ("請求日", "請求番号"):
        assert grade(recs[field]) == NONE_FOUND
        assert field in recs[field].blank_reason, recs[field].blank_reason
