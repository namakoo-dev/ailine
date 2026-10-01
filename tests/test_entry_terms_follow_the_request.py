# -*- coding: utf-8 -*-
"""run（31 op）以外の入口でも、依頼の項を依頼と突き合わせる（2026-10-01・入口の台帳）。

★★ 台帳 tests/entry_term_register.json で B（実表とだけ）・D（何とも突き合わせない）に
  分類された項目のうち、誤った答えを事実として出すものから塞いだ。型は ae4c732 と同じ
  （依頼から読む → 食い違えば依頼／機械が勝つ → 開示 → 読めない時は変えない）。

  ・run（フォルダ）の LLM の値（D）: 「金額が40000以上」に LLM の 4000 がそのまま通っていた
    （1 冊の抽出は 09-14 に threshold.ground で直してあった ── 兄弟間の片配線）
  ・run（フォルダ）の LLM の列（B）: 実在するかだけ見ていた → 名指しが無ければ ⚠
  ・split の式の値（D）: 式のままの列が配った冊で空欄・`--amount` の証明は「0 ＝ 0」で ✓（実測）
  ・verify の式の値（D）: 照合・分けるが LibreOffice の値で答えを作るのに、検算は原本を読み
    正しい出力に「破れ 9 件」（実測）
  ・2冊の照合のシート（D）: 依頼が別のシートを名指ししても黙って 1 枚目で照合していた
"""
import argparse
import contextlib
import io
import json
import shutil
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402

HDRS = ["請求番号", "担当", "税抜", "税込"]
ROWS = [("A-1", "山田", 1000), ("A-2", "佐藤", 2000), ("A-3", "山田", 500)]


def _sales(p: Path, *, cached: bool) -> Path:
    """税込が式（cached=False で計算結果なし）の一覧。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(HDRS)
    for i, (no, who, net) in enumerate(ROWS, start=2):
        ws.append([no, who, net, int(net * 1.1) if cached else f"=C{i}*1.1"])
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    wb.close()
    return p


def _fake_libreoffice(monkeypatch, filled: Path):
    """LibreOffice の代わりに、値の入った同じ表を返す（器官そのものは local の試験が持つ）。"""
    def fake(book, workdir, timeout=None):
        dst = Path(workdir) / "normalized.xlsx"
        shutil.copy(filled, dst)
        return dst
    monkeypatch.setattr(ailine, "normalize_book", fake)


def _no_libreoffice(monkeypatch):
    def boom(book, workdir, timeout=None):
        raise SystemExit(9)
    monkeypatch.setattr(ailine, "normalize_book", boom)


def _main(argv, capsys):
    rc = ailine.main(argv)
    return rc, capsys.readouterr().out


# --- run（フォルダ）: LLM の値と列 --------------------------------------------------------

def _folder(tmp_path) -> Path:
    folder = tmp_path / "src"
    folder.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["伝票", "取引先", "金額"])
    for r in (("J-1", "甲", 50000), ("J-2", "乙", 30000), ("J-3", "丙", 4500)):
        ws.append(list(r))
    wb.save(folder / "a.xlsx")
    return folder


def _translate(monkeypatch, args):
    monkeypatch.setattr(ailine, "translate_task",
                        lambda model, task, book_meta, temperature=0.1:
                        {"plan": [{"op": "EXTRACT", "args": args}]})


def test_folder_threshold_comes_from_the_request(tmp_path, monkeypatch, capsys):
    """★ 依頼は 40000、LLM は 4000 ── 依頼が勝ち、食い違いを言う（J-3 4500 を混ぜない）。"""
    _translate(monkeypatch, {"column": "金額", "cmp": "gte", "value": 4000})
    rc, out = _main(["run", str(_folder(tmp_path)), "金額が40000以上の行を抜き出して"], capsys)
    assert rc == 0, out
    assert "条件: 金額 40000" in out, out
    rows = [r[0].value for f in tmp_path.glob("*.xlsx")
            for r in openpyxl.load_workbook(f).active.iter_rows(min_row=2)]
    assert "J-1" in rows and "J-3" not in rows, rows


def test_positive_control_the_same_value_passes_silently(tmp_path, monkeypatch, capsys):
    """陽性対照: LLM が依頼どおりの数なら ⚠ は出ない（直しすぎない）。"""
    _translate(monkeypatch, {"column": "金額", "cmp": "gte", "value": 40000})
    rc, out = _main(["run", str(_folder(tmp_path)), "金額が40000以上の行を抜き出して"], capsys)
    assert rc == 0 and "⚠ 列" not in out and "LLM" not in out, out


def test_folder_threshold_without_a_number_is_refused(tmp_path, monkeypatch, capsys):
    """「少ない」の境目を機械が決めない（数が無ければ断る）。"""
    _translate(monkeypatch, {"column": "金額", "cmp": "lt", "value": 0})
    rc, out = _main(["run", str(_folder(tmp_path)), "金額が少ない行を抜き出して"], capsys)
    assert rc == 3, out
    assert not list(tmp_path.glob("*.xlsx")), "★ 断ったのに出力が在る"


def test_folder_column_not_named_is_disclosed(tmp_path, monkeypatch, capsys):
    """依頼が名指ししていない列を LLM が選んだら ⚠ で言う。"""
    _translate(monkeypatch, {"column": "金額", "cmp": "eq", "value": "甲"})
    rc, out = _main(["run", str(_folder(tmp_path)), "甲の行を抜き出して"], capsys)
    assert "⚠ 列『金額』は依頼文に名指しがありません" in out, out


def test_folder_value_not_in_the_request_is_disclosed(tmp_path, monkeypatch, capsys):
    _translate(monkeypatch, {"column": "取引先", "cmp": "eq", "value": "乙"})
    rc, out = _main(["run", str(_folder(tmp_path)), "取引先が甲の行を抜き出して"], capsys)
    assert "抽出する値『乙』は依頼文に見当たりません" in out, out


# --- split: 式の値 ---------------------------------------------------------------------------

def test_split_fills_formula_values_before_dealing(tmp_path, monkeypatch, capsys):
    """★ 配った冊に税込の値が入り、証明の金額が 0 ＝ 0 でない。"""
    book = _sales(tmp_path / "一覧.xlsx", cached=False)
    _fake_libreoffice(monkeypatch, _sales(tmp_path / "filled" / "x.xlsx", cached=True))
    out_dir = tmp_path / "配る"
    rc, out = _main(["split", str(book), "--by", "担当", "--amount", "税込",
                     "--out", str(out_dir)], capsys)
    assert rc == 0, out
    assert "LibreOffice で開いて計算させました" in out, out
    assert "金額 3850 ＝ 3850" in out, out
    ws = openpyxl.load_workbook(out_dir / "山田.xlsx").active
    assert [r[3] for r in ws.iter_rows(min_row=2, values_only=True)][:2] == [1100, 550]


def test_split_refuses_when_values_cannot_be_filled(tmp_path, monkeypatch, capsys):
    """★ 入れられなければ分けない（空欄の冊を配らない・証明の 0 ＝ 0 を出さない）。"""
    book = _sales(tmp_path / "一覧.xlsx", cached=False)
    _no_libreoffice(monkeypatch)
    out_dir = tmp_path / "配る"
    rc, out = _main(["split", str(book), "--by", "担当", "--amount", "税込",
                     "--out", str(out_dir)], capsys)
    assert rc == 4, out
    assert "計算結果の入っていない式のセルが 3 個" in out, out
    assert not out_dir.exists() or not list(out_dir.glob("*.xlsx")), "★ 断ったのに配っている"


def test_split_without_formulas_does_not_call_libreoffice(tmp_path, monkeypatch, capsys):
    """邪魔をしていない回に LibreOffice 往復を払わない。"""
    book = _sales(tmp_path / "一覧.xlsx", cached=True)
    called = []
    monkeypatch.setattr(ailine, "read_with_values_filled_in",
                        lambda *a, **k: called.append(1) or (None, "呼ぶな"))
    rc, out = _main(["split", str(book), "--by", "担当", "--out", str(tmp_path / "配る")], capsys)
    assert rc == 0 and not called, out


# --- verify: 検算は使ったデータを読む -------------------------------------------------------

def test_verify_split_reads_the_filled_values(tmp_path, monkeypatch, capsys):
    book = _sales(tmp_path / "一覧.xlsx", cached=False)
    _fake_libreoffice(monkeypatch, _sales(tmp_path / "filled" / "x.xlsx", cached=True))
    out_dir = tmp_path / "配る"
    rc, out = _main(["split", str(book), "--by", "担当", "--amount", "税込",
                     "--out", str(out_dir)], capsys)
    assert rc == 0, out
    rc, out = _main(["verify", str(out_dir), str(book), "--amount", "税込"], capsys)
    assert rc == 0, "★ 正しい出力を破れと言っている:\n" + out
    assert "LibreOffice で開いて計算させた値で読み直しました" in out, out
    _no_libreoffice(monkeypatch)
    rc, out = _main(["verify", str(out_dir), str(book), "--amount", "税込"], capsys)
    assert rc == 4 and "検算しません" in out, out


def _payments(p: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["請求番号", "入金額"])
    ws.append(["A-1", 1100])
    ws.append(["A-2", 2200])
    wb.save(p)
    wb.close()
    return p


def test_verify_match_reads_the_filled_values(tmp_path, monkeypatch, capsys):
    a = _sales(tmp_path / "請求.xlsx", cached=False)
    b = _payments(tmp_path / "入金.xlsx")
    _fake_libreoffice(monkeypatch, _sales(tmp_path / "filled" / "x.xlsx", cached=True))
    rc, out = _main(["run", str(a), str(b), "請求番号をキーに税込と入金額を突き合わせて", "--json"],
                    capsys)
    assert rc == 0, out
    produced = json.loads(out.strip().splitlines()[-1])["out"]
    assert json.loads(out.strip().splitlines()[-1])["amount_a"] == "税込"
    rc, out = _main(["verify", produced, str(a), str(b)], capsys)
    assert rc == 0, "★ 正しい照合を破れと言っている:\n" + out
    assert "LibreOffice で開いて計算させた値で読み直しました" in out, out


# --- 2冊の照合: シート -----------------------------------------------------------------------

def _two_sheets(p: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "今月"
    ws.append(["請求番号", "金額"])
    ws.append(["A-1", 100])
    ws2 = wb.create_sheet("先月")
    ws2.append(["請求番号", "金額"])
    ws2.append(["A-1", 999])
    wb.save(p)
    wb.close()
    return p


def _match(a, b, task, tmp_path, monkeypatch):
    monkeypatch.setenv("AILINE_HOME", str(tmp_path / "home"))
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = ailine.cmd_run_match(argparse.Namespace(json=False), a, b, task)
    return rc, buf.getvalue()


def test_match_refuses_a_sheet_it_does_not_read(tmp_path, monkeypatch):
    """★ 依頼が『先月』シートを名指ししたら、1 枚目（今月）で黙って照合しない。"""
    a = _two_sheets(tmp_path / "請求.xlsx")
    b = _payments(tmp_path / "入金.xlsx")
    rc, out = _match(a, b, "請求の先月シートの金額と入金額を請求番号をキーに突き合わせて",
                     tmp_path, monkeypatch)
    assert rc == 3, out
    assert "『先月』シートは読みません" in out, out


def test_folder_refuses_a_sheet_it_does_not_read(tmp_path, monkeypatch, capsys):
    """フォルダ抽出も同じ器官: 基準ファイルの別のシートを名指しされたら、1 枚目で黙って抜き出さない。"""
    folder = tmp_path / "src"
    folder.mkdir()
    _two_sheets(folder / "a.xlsx")
    _translate(monkeypatch, {"column": "金額", "cmp": "gte", "value": 100})
    rc, out = _main(["run", str(folder), "先月シートの金額が100以上の行を抜き出して"], capsys)
    assert rc == 3, out
    assert "『先月』シートは読みません" in out, out
    assert not list(tmp_path.glob("*.xlsx")), "★ 断ったのに出力が在る"


def test_match_discloses_which_sheet_it_read(tmp_path, monkeypatch):
    """名指しが無ければ 1 枚目で照合し、ほかのシートが在ることを言う。"""
    a = _two_sheets(tmp_path / "請求.xlsx")
    b = _payments(tmp_path / "入金.xlsx")
    rc, out = _match(a, b, "請求番号をキーに金額と入金額を突き合わせて", tmp_path, monkeypatch)
    assert rc == 0, out
    assert "シート『今月』を読みました ── 同じ冊に『先月』もあります" in out, out
