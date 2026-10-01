# 行番号で決まった行は、依頼文が名指しした行に限る（2026-10-01・依頼の項の台帳で B だった 2 項目:
# SET_CELL_VALUE.row_number・NUMBER_FORMAT.row_number）。
#
# ★ 形: 行番号は「表の範囲内か／見出しより下か」しか見ていなかった。依頼が行を言っていないのに
#   LLM が行番号を返すと、その行に書いて（書式を掛けて）✓ が出る。
# ★ 前の子が残した理由:「読み直しの道が機械で入れる行番号（A1 表記・端の語）と verify の段で
#   区別できない」。印を足すのでなく、**同じ依頼文を同じ読み手に通す**ことで区別を要らなくした
#   ── 機械が入れた行番号は、その読み手でもう一度読めば必ず同じ行に戻る。
#
# 契約:
#   ① 依頼文が行を指していれば黙る ──「7行目」・A1 のセル（E5）・表の端（最終行の〜）・その行の値
#   ② 指していなければ ⚠（_warnings ＝ ✓→△ の材料）で開示する
#   ③ データ行が 1 行しか無い表は実表で決まる ── 黙る
#   ④ 読み直しの道が機械で入れた行番号は鳴らない（同じ読み手で戻る）
#   ⑤ 見出しが 1 行目でない表でも同じ

import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
import ailine  # noqa: E402
from ailine_core import argcheck  # noqa: E402

HEAD = ["取引先", "伝票", "売上", "担当"]
DATA = [["山田商事", "INV-1", 1000, "佐藤"],
        ["丸和物流", "INV-2", 2000, "鈴木"],
        ["近江スチール", "INV-3", 3000, "高橋"],
        ["合計", None, 6000, None]]


def _book(tmp_path, header_row=1, data=DATA):
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "取引"
    for i in range(header_row - 1):
        ws.append(["取引一覧"] if i == 0 else [])
    ws.append(HEAD)
    for r in data:
        ws.append(r)
    wb.save(p)
    return p


def _meta(path, header_row=1):
    return {"sheets": ["取引"], "headers": {"取引": list(HEAD)},
            "header_rows": {"取引": header_row}, "path": str(path)}


def _cell(tmp_path, task, row_number, header_row=1, data=DATA, value="田中"):
    ok, r, _i, err = ailine.verify_dsl_args(
        "SET_CELL_VALUE", {"row_number": row_number, "col": "担当", "value": value},
        _meta(_book(tmp_path, header_row, data), header_row), task=task)
    assert ok, err
    return r


def _fmt(tmp_path, task, row_number, header_row=1):
    ok, r, _i, err = ailine.verify_dsl_args(
        "NUMBER_FORMAT", {"row_number": row_number, "style": "thousands"},
        _meta(_book(tmp_path, header_row), header_row), task=task)
    assert ok, err
    return r


def _row_warnings(r):
    return [w for w in r.get("_warnings", []) if "行目を指す語が見当たりません" in w]


# --- ① 1 セル書換: 依頼文が行を指していれば黙る ------------------------------------------

def test_row_number_in_the_request_is_silent(tmp_path):
    r = _cell(tmp_path, "3行目の担当を「田中」にして", 3)
    assert r["_row_index"] == 3 and not _row_warnings(r)


def test_a1_cell_in_the_request_is_silent(tmp_path):
    r = _cell(tmp_path, "D3に「田中」と入れて", 3)
    assert not _row_warnings(r)


def test_table_edge_in_the_request_is_silent(tmp_path):
    r = _cell(tmp_path, "最終行の担当を「田中」にして", 5)
    assert not _row_warnings(r)


def test_row_value_in_the_request_is_silent(tmp_path):
    r = _cell(tmp_path, "丸和物流の担当を「田中」にして", 3)
    assert not _row_warnings(r)


# --- ② 1 セル書換: 指していなければ ⚠ ----------------------------------------------------

def test_unnamed_row_number_is_disclosed(tmp_path):
    r = _cell(tmp_path, "担当を「田中」にして", 3)
    ws = _row_warnings(r)
    assert len(ws) == 1 and "3行目" in ws[0] and "丸和物流" in ws[0]


def test_a_different_row_named_by_the_request_is_disclosed(tmp_path):
    r = _cell(tmp_path, "近江スチールの担当を「田中」にして", 3)
    assert _row_warnings(r)
    r2 = _cell(tmp_path, "4行目の担当を「田中」にして", 3)
    assert _row_warnings(r2)


# --- ③ データ行が 1 行の表 ----------------------------------------------------------------

def test_a_one_row_table_decides_the_row(tmp_path):
    r = _cell(tmp_path, "担当を「田中」にして", 2, data=DATA[:1])
    assert r["_row_index"] == 2 and not _row_warnings(r)


# --- ④ 読み直しの道が機械で入れた行番号は鳴らない ------------------------------------------

def test_machine_inserted_row_numbers_come_back_to_the_same_row(tmp_path):
    """読み直しの 3 つの読み手（行番号・表の端・行の名前）で入れた行は、同じ依頼文で戻る。"""
    path = _book(tmp_path)
    meta = _meta(path)
    for task in ("3行目の担当を「田中」にして", "最終行の担当を「田中」にして",
                 "丸和物流の担当を「田中」にして"):
        row = (ailine.task_names_a_row_number(task)
               or (ailine.task_names_a_table_edge_row(task, meta, "取引") or (None,))[0]
               or ailine.resolve_cell_target_from_task(task, meta, "取引")[0])
        assert row, task
        ok, r, _i, err = ailine.verify_dsl_args(
            "SET_CELL_VALUE", {"row_number": row, "col": "担当", "value": "田中"}, meta, task=task)
        assert ok, err
        assert not _row_warnings(r), task


# --- ⑤ 見出しが 1 行目でない表 -------------------------------------------------------------

def test_header_not_on_the_first_row(tmp_path):
    r = _cell(tmp_path, "担当を「田中」にして", 5, header_row=3)
    assert r["_row_index"] == 5 and _row_warnings(r)
    r2 = _cell(tmp_path, "丸和物流の担当を「田中」にして", 5, header_row=3)
    assert not _row_warnings(r2)


# --- 数値書式の行番号 ----------------------------------------------------------------------

def test_format_row_named_by_number_or_total_word_is_silent(tmp_path):
    assert not _row_warnings(_fmt(tmp_path, "5行目を桁区切りにして", 5))
    assert not _row_warnings(_fmt(tmp_path, "合計を金額表示にして", 5))


def test_format_row_not_named_is_disclosed(tmp_path):
    r = _fmt(tmp_path, "桁区切りにして", 3)
    ws = _row_warnings(r)
    assert len(ws) == 1 and "3行目" in ws[0]


def test_format_row_with_header_not_on_the_first_row(tmp_path):
    assert _row_warnings(_fmt(tmp_path, "桁区切りにして", 5, header_row=3))
    assert not _row_warnings(_fmt(tmp_path, "合計を金額表示にして", 7, header_row=3))


def test_format_reread_total_row_comes_back(tmp_path):
    """読み直しが合計行へ付け替える道（number_format_target）の行は、同じ依頼文で戻る。"""
    path = _book(tmp_path)
    kind, row = ailine.number_format_target("合計を金額表示にして", _meta(path), "取引")
    assert (kind, row) == ("row", 5)
    assert not _row_warnings(_fmt(tmp_path, "合計を金額表示にして", row))


# --- 変異: 配線を外すと黙る -----------------------------------------------------------------

def test_without_the_wiring_the_guess_would_pass_silently(tmp_path, monkeypatch):
    monkeypatch.setattr(argcheck, "_row_number_named_by_the_request", lambda *a, **k: True)
    assert not _row_warnings(_cell(tmp_path, "担当を「田中」にして", 3))     # ← 直す前の挙動
    assert not _row_warnings(_fmt(tmp_path, "桁区切りにして", 3))
