# 下書きを原本に清書しても、`ailine undo` で戻せる（2026-09-24）
#
# ★★ 発端（原本へ書く口の台帳 tests/test_write_routes_ledger.py が「既知の欠陥」として載せていた）:
#   GUI の「原本に反映」は下書きを原本へ `shutil.copy2` で被せるだけで、ロックの関所も
#   実行ロックもバックアップも通らなかった。CLI は「原本は自動でバックアップされ、
#   `ailine undo` で戻せます」と約束しているのに、この経路だけ戻せない。
#   さらに、下書きを作った**後**に人が原本を Excel 等で直していたら、清書はそれを**黙って**上書きした。
# ★ 直し（Namakoo 決裁・案 A）: CLI に `ailine adopt <下書き> <原本> [--base-sha HEX]` を足し、
#   GUI はそれを子プロセスで叩くだけにする（GUI は本体を import しない）。
#
# 契約:
#   ① 清書すると原本が下書きの中身になり、`ailine undo` で**元の原本のバイト**に戻る
#   ② 下書きは消えない・変わらない
#   ③ `--base-sha` が今の原本と違えば断る ── 原本・下書きとも 1 バイトも変わらない
#   ④ Excel ロック・実行ロックがあれば断る ── 原本は無変更
#   ⑤ 無い／同じファイル／拡張子違いで断る ── 原本は無変更
#   ⑥ GUI の清書の枝は `adopt` を叩き、原本へ直に書く `shutil.copy2(..., book)` を持たない
#      （★ 無いことの assert だけにしない ── 叩いていることを大声で確かめる対と、実際に画面の口を叩く検体）

import hashlib
import importlib.util
import json
import re
import sys
import threading
import urllib.request
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

import ailine  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_golden_transcripts import _isolate, _run_main  # noqa: E402

GUI_DIR = REPO / "gui"


def _book(path: Path, first: str) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["商品", "売上"])
    ws.append([first, 100])
    wb.save(path)
    return path


def _first(p: Path):
    return openpyxl.load_workbook(p).active.cell(row=2, column=1).value


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _pair(tmp_path: Path):
    """原本（りんご）と、その複製を直した下書き（みかん）。"""
    book = _book(tmp_path / "b.xlsx", "りんご")
    draft = _book(tmp_path / "b（下書き）.xlsx", "みかん")
    return book, draft


# --- ①② 清書して、戻せる。下書きは残る ------------------------------------------------

def test_adopt_makes_the_book_the_draft_and_undo_brings_the_original_back(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    original, drafted = book.read_bytes(), draft.read_bytes()
    rc, out = _run_main(["adopt", str(draft), str(book), "--base-sha", _sha(book)], capsys)
    assert rc == 0, out
    assert book.read_bytes() == drafted, "原本が下書きの中身になっていない"
    assert "反映しました" in out and "下書きはそのまま残しています" in out, out
    assert f'ailine undo "{book}"' in out, f"戻し方を打てる形で言っていない: {out}"
    # ② 下書きは消えない・変わらない
    assert draft.exists(), "下書きが消えた（画面は『そのまま残しています』と言う）"
    assert draft.read_bytes() == drafted, "下書きの中身が変わった"
    # ① undo で元の原本のバイトに戻る
    rc, out = _run_main(["undo", str(book)], capsys)
    assert rc == 0, out
    assert book.read_bytes() == original, "undo で元の原本に戻らない（控えが取られていない）"
    assert draft.read_bytes() == drafted
    # 作業フォルダを残さない
    assert not (tmp_path / ".ailine_b").exists(), "作業フォルダが残った"


def test_without_base_sha_it_says_it_did_not_compare(tmp_path, monkeypatch, capsys):
    """★ 照合しなかったことを黙らない（出ないことは信号でない）。"""
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    rc, out = _run_main(["adopt", str(draft), str(book)], capsys)
    assert rc == 0, out
    assert "照合していません" in out, out
    assert _first(book) == "みかん"


def test_adopt_leaves_a_trace_of_the_sheets_the_draft_added(tmp_path, monkeypatch, capsys):
    """★ 出所を残す ── 下書きで増えたシートは道具の作業の結果、元からのシートは人のもの。"""
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    wb = openpyxl.load_workbook(draft)
    wb.create_sheet("検分")
    wb.save(draft)
    rc, out = _run_main(["adopt", str(draft), str(book)], capsys)
    assert rc == 0, out
    person = openpyxl.load_workbook(book, read_only=True).sheetnames[0]
    made = ailine.sheets_ailine_made(book, path=ailine.HISTORY_FILE)
    assert made == {"検分"}, f"清書の出所が台帳に無い／人のシートを道具のものにした: {made} ({person})"


# --- ③ 下書きを作った後に原本が変わっていたら止める -------------------------------------

def test_a_stale_base_sha_is_refused_and_nothing_changes(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    base = _sha(book)
    # 下書きを作った後に、人が原本を Excel で直した
    wb = openpyxl.load_workbook(book)
    wb.active.cell(row=2, column=2, value=999)
    wb.save(book)
    book_now, drafted = book.read_bytes(), draft.read_bytes()
    rc, out = _run_main(["adopt", str(draft), str(book), "--base-sha", base], capsys)
    assert rc == 7, f"原本が変わっているのに止まらない: exit={rc}\n{out}"
    assert book.read_bytes() == book_now, "原本に書いた（人の変更を黙って消した）"
    assert draft.read_bytes() == drafted, "下書きを変えた"
    assert "変わっています" in out, out
    # ★ 次の一手を打てる形で（上書きするなら --base-sha を付けずに）
    assert f'ailine adopt "{draft}" "{book}"' in out, out
    assert ailine.list_backups(book) == [], "断ったのに控えを積んだ"


def test_a_matching_base_sha_goes_through(tmp_path, monkeypatch, capsys):
    """陰性対照 ── 大文字の hex でも同じ中身なら通る（断りを広げすぎない）。"""
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    rc, out = _run_main(["adopt", str(draft), str(book), "--base-sha", _sha(book).upper()], capsys)
    assert rc == 0, out
    assert "照合していません" not in out, out


# --- ④ ロック -------------------------------------------------------------------------

def test_an_excel_lock_is_refused_and_the_book_is_untouched(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    (tmp_path / f"~${book.name}").write_bytes(b"excel lock")
    before, drafted = book.read_bytes(), draft.read_bytes()
    rc, out = _run_main(["adopt", str(draft), str(book), "--base-sha", _sha(book)], capsys)
    assert rc == ailine.EXIT_WRITE_BLOCKED, f"exit={rc}\n{out}"
    assert book.read_bytes() == before, "ロックを無視して上書きした"
    assert draft.read_bytes() == drafted
    assert "残骸" in out, "関所の文言でない（別の断りを書き写した）"


def test_a_busy_run_lock_is_refused_and_the_book_is_untouched(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book, draft = _pair(tmp_path)
    monkeypatch.setattr(ailine, "acquire_run_lock", lambda *a, **k: (False, "別の ailine が実行中です"))
    before = book.read_bytes()
    rc, out = _run_main(["adopt", str(draft), str(book)], capsys)
    assert rc == 6, f"exit={rc}\n{out}"
    assert book.read_bytes() == before


# --- ⑤ 受け取れない組み合わせ ------------------------------------------------------------

def test_a_missing_draft_is_refused(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "b.xlsx", "りんご")
    before = book.read_bytes()
    rc = ailine.main(["adopt", str(tmp_path / "無い下書き.xlsx"), str(book)])
    err = capsys.readouterr()
    assert rc == 9, err
    assert "無い下書き.xlsx" in err.err + err.out
    assert book.read_bytes() == before


def test_the_same_file_is_refused(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "b.xlsx", "りんご")
    before = book.read_bytes()
    rc, out = _run_main(["adopt", str(book), str(book)], capsys)
    assert rc == ailine.EXIT_APPLY_FAILED, out
    assert "同じファイル" in out, out
    assert book.read_bytes() == before
    assert ailine.list_backups(book) == []


def test_a_different_extension_is_refused(tmp_path, monkeypatch, capsys):
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "b.xlsm", "りんご")
    draft = _book(tmp_path / "b（下書き）.xlsx", "みかん")
    before, drafted = book.read_bytes(), draft.read_bytes()
    rc, out = _run_main(["adopt", str(draft), str(book)], capsys)
    assert rc == ailine.EXIT_APPLY_FAILED, out
    assert "形式が違います" in out, out
    assert book.read_bytes() == before and draft.read_bytes() == drafted


def test_a_draft_that_cannot_be_opened_is_not_adopted(tmp_path, monkeypatch, capsys):
    """★ 原本へ被せる前の最後の確認（run と同じ物差し）── 開けない下書きで原本を潰さない。"""
    _isolate(monkeypatch, tmp_path)
    book = _book(tmp_path / "b.xlsx", "りんご")
    draft = tmp_path / "b（下書き）.xlsx"
    draft.write_bytes(b"not an xlsx")
    before = book.read_bytes()
    rc, out = _run_main(["adopt", str(draft), str(book)], capsys)
    assert rc == ailine.EXIT_APPLY_FAILED, out
    assert book.read_bytes() == before
    assert draft.read_bytes() == b"not an xlsx"


# --- ⑥ GUI は adopt を叩く ----------------------------------------------------------------

def _code_only(py: str) -> str:
    return chr(10).join(ln.split("#")[0] for ln in py.split(chr(10)))


def test_the_gui_adopt_branch_calls_the_cli_and_never_copies_onto_the_book():
    src = (GUI_DIR / "server.py").read_text(encoding="utf-8")
    i = src.index('if u.path == "/api/run":')
    j = src.index("elif u.path ==", i)
    block = _code_only(src[i:j])
    # ★ 叩いていること（大声の側）
    assert '_ailine(_args)' in block and '["adopt", _draft, book]' in block, (
        "清書の枝が ailine adopt を叩いていない")
    assert '"--base-sha"' in block, "原本がその後に変わったかを本体に照合させていない"
    assert "_DRAFT_BASE[book] = hashlib.sha256(" in block, "下書きを作った時の原本を覚えていない"
    # 片付けは成功した時だけ
    k = block.index('["adopt", _draft, book]')
    tail = block[k:]
    assert tail.index("if rc == 0:") < tail.index("_DRAFTS.pop(book, None)"), (
        "断られた時にも下書きの記憶を消している")
    # ★ 無いこと（静かな側）── gui/ のどこにも、原本へ直に被せる copy2 が無い
    hits = []
    for py in sorted(GUI_DIR.rglob("*.py")):
        for n, ln in enumerate(_code_only(py.read_text(encoding="utf-8")).splitlines(), 1):
            if re.search(r"shutil\.copy2\([^)]*,\s*book\)", ln):
                hits.append(f"{py.name}:{n}: {ln.strip()}")
    assert not hits, f"GUI が原本へ直に書いている: {hits}"


def _load_server():
    spec = importlib.util.spec_from_file_location("_ailine_gui_server_adopt", GUI_DIR / "server.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def gui(tmp_path):
    """本物の画面サーバを 127.0.0.1 の空きポートで立てる（子プロセスは AILINE_HOME を継ぐ ── conftest）。"""
    server = _load_server()
    srv = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    th = threading.Thread(target=srv.serve_forever, daemon=True)
    th.start()
    try:
        yield server, f"http://127.0.0.1:{srv.server_address[1]}"
    finally:
        srv.shutdown()
        srv.server_close()


def _post(url: str, body: dict) -> dict:
    req = urllib.request.Request(url, data=json.dumps(body).encode("utf-8"),
                                  headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=300) as r:
        return json.loads(r.read().decode("utf-8"))


def test_the_gui_adopt_is_undoable_through_the_real_page_endpoint(tmp_path, gui):
    """★ 画面の口を実際に叩く ── 清書された原本が、画面の「もとに戻す」と同じ `ailine undo` で戻る。"""
    server, base = gui
    book, draft = _pair(tmp_path)
    original, drafted = book.read_bytes(), draft.read_bytes()
    server._DRAFTS[str(book)] = str(draft)
    server._DRAFT_BASE[str(book)] = _sha(book)
    res = _post(base + "/api/run", {"book": str(book), "task": ""})
    assert res["rc"] == 0, res["text"]
    assert set(res) >= {"rc", "text", "json", "before", "target"}, sorted(res)
    assert book.read_bytes() == drafted
    assert draft.read_bytes() == drafted, "下書きが変わった／消えた"
    assert str(book) not in server._DRAFTS, "成功したのに下書きの記憶が残っている"
    rc, out, _ = server._ailine(["undo", str(book)])
    assert rc == 0, out
    assert book.read_bytes() == original, "画面から清書した原本が undo で戻らない"


def test_the_gui_refusal_keeps_the_draft_in_memory(tmp_path, gui):
    """★ 原本がその後に変わっていたら、画面経由でも止まる ── 下書きの記憶は消さない。"""
    server, base = gui
    book, draft = _pair(tmp_path)
    server._DRAFTS[str(book)] = str(draft)
    server._DRAFT_BASE[str(book)] = "0" * 64      # 下書きを作った時の原本とは違う
    before = book.read_bytes()
    res = _post(base + "/api/run", {"book": str(book), "task": ""})
    assert res["rc"] == 7, res["text"]
    assert book.read_bytes() == before, "画面経由で原本に書いた"
    assert server._DRAFTS.get(str(book)) == str(draft), "断られたのに下書きの記憶を消した"
    assert server._DRAFT_BASE.get(str(book)) == "0" * 64


def test_the_screen_does_not_say_applied_when_it_was_refused():
    """★ 画面の見出し「（原本に反映）」は rc が 0 の時だけ（断られた清書で『反映』と言わない）。

    ★ 2026-09-24: 初版の見出しは `res.target` だけを見ていた ── 清書の枝は断られても
      target=原本 を返すので、何もしていないのに「（原本に反映）」と出た。
    """
    html = (REPO / "gui" / "index.html").read_bytes().decode("utf-8")
    i = html.index('"（原本に反映）"')
    window = html[max(0, i - 400): i + 80]
    assert "res.rc === 0" in window, "「（原本に反映）」の条件に rc を見ていない:\n" + window
    assert '"（原本は無変更）"' in window, "断られた時の見出しが近くに無い"
