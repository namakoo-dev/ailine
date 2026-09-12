# -*- coding: utf-8 -*-
"""③ 帳票の一覧を**後から・ファイルだけから**独立に検算する（2026-09-13）。

## 何を縛るか

`ailine verify` は一覧に対して `{"unsupported"}` を返すだけだった。
★ **正直に「無い」ことは、在ることにならない。**

契約（`ailine_core/verify_forms.py`）:
  - 抽出の規則（どのラベルの右を読むか）を**再現しない**。`form_read` を呼ばない
  - 一覧の各行が名指しする出所ファイルが実在する
  - フォルダの帳票が全部一覧に在る（★ こちらでも読めない冊は咎めない）
  - 取り出した値が、その出所ファイルの**中に実際に在る**（含有）── PDF はテキスト層で見る
  - 空欄には理由が在る（検分に 1 件ずつ）
  - ★ 日付が和暦などで書かれていて西暦の形で見つからない時は、**破れにも合格にもせず
    「確かめられなかった」と数えて名前を出す**

★ LLM も LibreOffice も要らない（読むだけ）。
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from ailine_core import forms_collect, verify_forms   # noqa: E402
from test_forms_e2e import _invoice   # noqa: E402 ── 同じ検体を使い回す（二重定義を作らない）

FIXTURES = REPO / "tests" / "fixtures" / "forms"


def _forms_cmd(folder: Path, out: Path):
    return subprocess.run([sys.executable, "-m", "ailine", "forms", str(folder),
                           "--out", str(out)], capture_output=True, text=True, timeout=300,
                          encoding="utf-8", errors="replace", cwd=str(REPO))


def _verify(out: Path, folder: Path):
    return subprocess.run([sys.executable, "-m", "ailine", "verify", str(out), str(folder)],
                          capture_output=True, text=True, timeout=300,
                          encoding="utf-8", errors="replace", cwd=str(REPO))


def _edit(path: Path, fn):
    wb = openpyxl.load_workbook(path)
    fn(wb)
    wb.save(path)
    wb.close()


@pytest.fixture()
def listed(tmp_path):
    """請求書 3 通を一覧にして (一覧, フォルダ) を返す。"""
    folder = tmp_path / "受け取った"
    folder.mkdir()
    for i, total in enumerate((11000, 22000, 33000), start=1):
        _invoice(folder / f"請求書{i}.xlsx", f"取引先{i}株式会社", total)
    out = tmp_path / "一覧.xlsx"
    r = _forms_cmd(folder, out)
    assert r.returncode == 0, f"一覧が作れなかった: {r.stdout}\n{r.stderr}"
    return out, folder


def test_a_clean_list_verifies_from_the_files_alone(listed):
    """★ 陽性対照 ── 素の一覧は破れ 0 で通る。ここが 0 でなければ検算器を先に疑う。"""
    out, folder = listed
    r = _verify(out, folder)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"
    assert "破れはありません" in r.stdout, r.stdout
    assert "含有を確かめた値" in r.stdout, r.stdout
    # ★★ 「破れ 0」だけでは、**何も確かめられていない 0** と区別が付かない。
    #   この検体の請求日は `2026/8/31` と西暦で書いてあるので、確かめられていないと嘘になる。
    #   ★ この 1 行が無いと「日付の書き方を 1 つも試さない」変異が黙って通った（実測）。
    assert "確かめられなかった" not in r.stdout, r.stdout


def test_a_value_that_is_not_in_the_source_is_caught(listed):
    """★★ 一番強い縛り ── どのセルを読むかの規則を知らなくても、無い値は分かる。"""
    out, folder = listed
    _edit(out, lambda wb: wb[forms_collect.SHEET_NAME].cell(2, 2, "存在しない会社株式会社"))
    r = _verify(out, folder)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "含有の破れ" in r.stdout, r.stdout


def test_a_blank_without_a_reason_is_caught(listed):
    """★ 空欄の値段は理由の正しさで決まる ── 理由の無い空欄は破れ。"""
    out, folder = listed
    # ★ `cell(row, col, None)` は openpyxl が**代入を飛ばす**（None は無視される）。
    #   最初これで書いて、一覧が変わらないまま「破れ 0」になり試験が嘘をついた。
    def blank(wb):
        wb[forms_collect.SHEET_NAME].cell(2, 2).value = None
    _edit(out, blank)
    r = _verify(out, folder)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "空欄に理由が無い" in r.stdout, r.stdout


def test_a_book_missing_from_the_list_is_caught(listed):
    """★ 消えたものは diff に出ない ── 分母は**フォルダ**から作る（一覧からではない）。"""
    out, folder = listed
    _edit(out, lambda wb: wb[forms_collect.SHEET_NAME].delete_rows(2))
    r = _verify(out, folder)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "一覧に無い冊" in r.stdout, r.stdout


def test_a_source_that_does_not_exist_is_caught(listed):
    out, folder = listed
    _edit(out, lambda wb: wb[forms_collect.SHEET_NAME].cell(2, 1, "存在しない.xlsx"))
    r = _verify(out, folder)
    assert r.returncode == 5, f"{r.stdout}\n{r.stderr}"
    assert "出所のファイルが無い" in r.stdout, r.stdout


def test_a_pdf_source_is_verified_through_its_text_layer(tmp_path):
    """★★ PDF でも含有で測れる ── テキスト層をそのまま読む（OCR ではない）。

    ★ これが効かないと、PDF の一覧は「作れるが確かめられない」ままになる
      （受け取る請求書は PDF が圧倒的に多い ── pyproject の決裁理由）。
    """
    folder = tmp_path / "受け取ったPDF"
    folder.mkdir()
    src = FIXTURES / "ソフト発行_請求書.pdf"
    (folder / src.name).write_bytes(src.read_bytes())
    out = tmp_path / "一覧.xlsx"
    assert _forms_cmd(folder, out).returncode == 0
    r = _verify(out, folder)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"
    assert "破れはありません" in r.stdout, r.stdout
    # ★ 陰性対照: PDF でも「無い値」は鳴る（テキスト層を本当に見ている証拠）
    _edit(out, lambda wb: wb[forms_collect.SHEET_NAME].cell(2, 2, "存在しない会社株式会社"))
    assert _verify(out, folder).returncode == 5


def test_a_date_written_only_in_the_japanese_era_is_counted_as_unchecked(tmp_path):
    """★★★ 確かめられないものを、破れにも合格にもしない。

    元が和暦（`令和8年8月31日`）だと西暦の形では見つからない。ここで False にすれば
    嘘の警報、True にすれば空虚な合格 ── **「確かめられなかった」と数えて名前を出す**。
    ★ 和暦の対応表を検算器に持たせない理由: それは「読む規則」の再現になる。
    """
    folder = tmp_path / "和暦"
    folder.mkdir()
    book = folder / "請求書.xlsx"
    _invoice(book, "和暦商事株式会社", 44000)
    _edit(book, lambda wb: wb["請求書"].__setitem__("H3", "令和8年8月31日"))
    out = tmp_path / "一覧.xlsx"
    assert _forms_cmd(folder, out).returncode == 0
    got = verify_forms.verify_forms_list(out, folder)
    assert not got["breaks"], got["breaks"]
    assert "含有を確かめられなかった値" in got["facts"], got["facts"]
    r = _verify(out, folder)
    assert r.returncode == 0 and "確かめられなかった" in r.stdout, r.stdout


def test_our_own_output_left_in_the_folder_is_not_counted_as_an_input(listed):
    """★ 分母が汚れない ── 一覧をフォルダの中に置いても「一覧に無い冊」にならない。"""
    out, folder = listed
    (folder / out.name).write_bytes(out.read_bytes())
    r = _verify(out, folder)
    assert r.returncode == 0, f"{r.stdout}\n{r.stderr}"


def test_the_verifier_does_not_reproduce_the_extraction_rules():
    """★★ 恒真の防止を機械で縛る ── 検算器が抽出の器官を呼んだら赤。"""
    import ast
    text = (REPO / "src" / "ailine_core" / "verify_forms.py").read_text(encoding="utf-8")
    used, defined = set(), set()
    for node in ast.walk(ast.parse(text)):
        if isinstance(node, ast.Attribute):
            used.add(node.attr)
        elif isinstance(node, ast.Name):
            used.add(node.id)
        elif isinstance(node, ast.FunctionDef):
            defined.add(node.name)
    for expected in ("verify_forms_list", "contained", "source_values"):
        assert expected in defined, f"物差しが読めていない（定義に {expected} が無い）"
    for expected in ("read_grid", "read_core_properties"):
        assert expected in used, f"物差しが読めていない（使用に {expected} が無い）"
    for forbidden in ("form_read", "read_form", "read_pdf_book", "read_form_grid",
                      "grade_of", "pdf_grid"):
        assert forbidden not in used, f"検算器が抽出の規則を呼んでいる: {forbidden}"
