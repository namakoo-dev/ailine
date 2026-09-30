import openpyxl
wb = openpyxl.load_workbook("orig/請求一覧.xlsx")
m = openpyxl.load_workbook("orig/取引先マスタ.xlsx").active
ws = wb.create_sheet("取引先マスタ")
for row in m.iter_rows():
    for c in row:
        n = ws.cell(c.row, c.column, c.value)
        n.number_format = c.number_format
wb.save("請求一覧_マスタ付.xlsx")
print("ok")
