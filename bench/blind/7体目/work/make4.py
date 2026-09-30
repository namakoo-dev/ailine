import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
wb = openpyxl.load_workbook("orig/請求一覧.xlsx"); ws = wb.active
ws.conditional_formatting.add("I5:I12", CellIsRule(operator="equal", formula=['"未入金"'], fill=PatternFill("solid", fgColor="FFC7CE")))
dv = DataValidation(type="list", formula1='"入金済,未入金,一部入金"'); ws.add_data_validation(dv); dv.add("I5:I12")
ws["A4"].comment = Comment("請求番号は会計ソフトと同じ", "山田")
wb.save("請求一覧_書式.xlsx")
print("ok")
