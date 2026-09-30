import openpyxl
from openpyxl.styles import Font
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BS"
ws["A1"] = "合計残高試算表（貸借対照表）"; ws.merge_cells("A1:E1"); ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "株式会社丸山工業　自 2026/09/01 至 2026/09/30　（単位：円）"; ws.merge_cells("A2:E2")
ws.append([])
ws.append(["勘定科目", "前月残高", "借方", "貸方", "当月残高"])
data = [
    ("【流動資産】", None, None, None, None),
    ("現金", 120000, 50000, 43000, 127000),
    ("普通預金", 3450000, 980000, 1210000, 3220000),
    ("売掛金", 1650000, 776050, 467500, 1958550),
    ("流動資産 計", None, None, None, None),
    ("【固定資産】", None, None, None, None),
    ("工具器具備品", 480000, 0, 0, 480000),
    ("固定資産 計", None, None, None, None),
    ("資産合計", None, None, None, None),
]
for d in data:
    ws.append(list(d))
# 小計・合計は式
ws["B9"] = "=SUM(B6:B8)"; ws["C9"] = "=SUM(C6:C8)"; ws["D9"] = "=SUM(D6:D8)"; ws["E9"] = "=SUM(E6:E8)"
for c in "BCDE":
    ws[f"{c}12"] = f"={c}11"
    ws[f"{c}13"] = f"={c}9+{c}12"
for r in (5, 10):
    ws.merge_cells(f"A{r}:E{r}")
for r in range(6, 14):
    for c in range(2, 6):
        ws.cell(r, c).number_format = "#,##0"
wb.save("試算表.xlsx")
print("ok")
