# -*- coding: utf-8 -*-
"""基準④: 抜き出しに怪しい点があるとき、印が立つか（2026-09-10 夜）。

★ わざと壊した冊を 4 通り作り、3 体の共通部分（三口検算＋R3）がどう答えるかを見る。
   壊し方は「実務で実際に起きる形」に限る:

    T1 合計を 1 円ずらす        （人が手で直して式を潰した）
    T2 明細を 1 行消す           （消したのに合計が再計算されていない）
    T3 連絡先ブロックを消す      （発行者の 〒/TEL/登録番号 が無い雛形）
    T4 御中を 2 つにする         （転送された請求書・宛先が 2 段）

★ 期待する答えは「正しい値を出す」ではなく「**怪しいと言う**」こと。
   壊れた冊で『確』が出たら失格（黙って間違える）。
"""
import json
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, r"C:\Dev\ailine\src")
import openpyxl
from openpyxl.utils import get_column_letter

SP = Path(__file__).resolve().parent   # 検体は隣（追跡しない）
ANS = {a["file"]: a for a in json.loads((SP / "答え_received40.json").read_text(encoding="utf-8"))}
CONTACT = ("〒", "TEL", "FAX", "E-Mail", "登録番号", "電話")
HEAD_MONEY = ("ご請求金額", "御請求金額", "今回請求額", "請求金額", "合計金額")
SUM_LABELS = ("小計", "消費税", "合計金額", "合計")


def norm(s):
    return re.sub(r"[\s　]", "", str(s)).strip()


def merged(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = v
    return m


def val(ws, mm, r, c):
    v = ws.cell(row=r, column=c).value
    return mm.get((r, c)) if v in (None, "") else v


def inspect(p, sheet):
    """3 体の共通部分: 三口検算（請求額）＋ R3（請求元）＋ 御中の数。区分と理由を返す。"""
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb[sheet]
    mm = merged(ws)
    R, C = min(ws.max_row, 60), min(ws.max_column, 24)
    cells, ochus, contacts = {}, [], []
    detail_row = amount_col = None
    for r in range(1, R + 1):
        for c in range(1, C + 1):
            v = val(ws, mm, r, c)
            if not isinstance(v, str) or not v.strip():
                continue
            s = norm(v)
            raw = ws.cell(row=r, column=c).value            # ★ 結合の左上だけが生の値を持つ
            if "御中" in s and raw not in (None, "") and (r, c) not in ochus:
                ochus.append((r, c))
            if any(s.startswith(k) or k in s[:6] for k in CONTACT) and raw not in (None, ""):
                contacts.append((r, c))
            if s in ("金額", "金額(税抜)", "金額（税抜）") and detail_row is None:
                detail_row, amount_col = r, c
            if 2 <= len(s) <= 40:
                cells[(r, c)] = s
    ochus = sorted(set(ochus))
    # 請求元 (R3)
    issuer, issuer_why = None, "連絡先の塊が見つからない"
    for (r, c) in contacts:
        for rr in (r - 1, r - 2):
            s = cells.get((rr, c))
            if s and not any(k in s for k in CONTACT) and (rr, c) not in ochus:
                issuer, issuer_why = s, f"{get_column_letter(c)}{rr}（連絡先の直上）"
                break
        if issuer:
            break
    issuer_grade = "単" if issuer else "無"
    if len(ochus) >= 2:
        issuer_grade, issuer_why = "無", f"御中が {len(ochus)} 箇所 ── 宛先が一つに決まらない"
    # 請求額（三口）
    mouths = {}
    for r in range(1, (detail_row or R)):
        for c in range(1, C + 1):
            s = norm(val(ws, mm, r, c) or "")
            if any(s.startswith(k) for k in HEAD_MONEY):
                for cc in range(c + 1, min(C, c + 6) + 1):
                    v = val(ws, mm, r, cc)
                    if isinstance(v, (int, float)):
                        mouths["口①上部"] = v
                        break
    if detail_row and amount_col:
        s, rr, n, blanks = 0.0, detail_row + 1, 0, 0
        while rr <= R:
            filled = sum(1 for c in range(1, C + 1) if val(ws, mm, rr, c) not in (None, ""))
            av = val(ws, mm, rr, amount_col)
            if filled >= 2 and isinstance(av, (int, float)):
                s += av; n += 1; blanks = 0
            elif filled == 0:
                blanks += 1
                if blanks >= 2 and n:
                    break
            elif n:
                break
            rr += 1
        band = {}
        for r in range(rr, R + 1):
            for c in range(1, C + 1):
                lab = norm(val(ws, mm, r, c) or "")
                for lb in SUM_LABELS:
                    if lab.startswith(lb) and lb not in band:
                        v = val(ws, mm, r, amount_col)
                        if isinstance(v, (int, float)):
                            band[lb] = v
        if n and "消費税" in band:
            mouths["口②和+税"] = s + band["消費税"]
        if band.get("合計金額") or band.get("合計"):
            mouths["口③帯合計"] = band.get("合計金額") or band.get("合計")
    vals = {round(v, 2) for v in mouths.values()}
    if len(mouths) >= 2 and len(vals) == 1:
        g, why = "確", f"{len(mouths)} 口が一致"
    elif len(mouths) >= 2:
        g, why = "割", "食い違い: " + " / ".join(f"{k}={v:,.0f}" for k, v in mouths.items())
    elif len(mouths) == 1:
        g, why = "単", f"{list(mouths)[0]} のみ"
    else:
        g, why = "無", "金額のラベルが無い"
    wb.close()
    return {"請求額": (next(iter(mouths.values()), None), g, why),
            "請求元": (issuer, issuer_grade, issuer_why)}


# ── 壊した冊を作る（openpyxl で直値を書く＝人が手で直した状態） ─────
work = SP / "_tamper"
if work.exists():
    shutil.rmtree(work)
work.mkdir(parents=True)
base = SP / "received40" / "received_01.xlsx"          # misoca の骨
a = ANS["received_01.xlsx"]
sheet = a["シート"]


def make(name, fn):
    dst = work / name
    shutil.copy(base, dst)
    wb = openpyxl.load_workbook(dst)     # ★ data_only=False: 式を直値で潰す
    ws = wb[sheet]
    fn(ws)
    wb.save(dst)
    wb.close()
    return dst


def t1(ws):   # 合計を 1 円ずらす（式を直値に置き換える）
    ws["H39"] = a["請求額(税込)"] + 1
    ws["C11"] = a["請求額(税込)"] + 1


def t2(ws):   # 明細を 1 行消す（合計は再計算されない＝式を直値で固定）
    for cell in ("C11", "H37", "H38", "H39"):
        wb_val = {"C11": a["請求額(税込)"], "H37": a["小計(税抜)"],
                  "H38": a["消費税(10%)"], "H39": a["請求額(税込)"]}[cell]
        ws[cell] = wb_val
    ws["B16"] = None; ws["E16"] = None; ws["G16"] = None; ws["H16"] = None


def t3(ws):   # 連絡先ブロックを消す
    for r in range(6, 13):
        ws[f"G{r}"] = None


def t4(ws):   # 御中を 2 つにする
    ws["B3"] = "ナギ商会株式会社 御中"


import subprocess
def recalc_all():
    soffice = r"C:\Program Files\LibreOffice\program\soffice.exe"
    prof = (SP / "lo-profile-mk").resolve().as_uri()
    out = work / "_out"; out.mkdir(exist_ok=True)
    files = [str(f) for f in sorted(work.glob("t*.xlsx"))]
    subprocess.run([soffice, "--headless", "--norestore", "--nologo",
                    f"-env:UserInstallation={prof}", "--convert-to", "xlsx",
                    "--outdir", str(out)] + files, capture_output=True, timeout=600)
    for f in out.glob("*.xlsx"):
        shutil.copy(f, work / f.name)

_made = [("T0 壊していない（対照）", make("t0.xlsx", lambda ws: None)),
         ("T1 合計を 1 円ずらす", make("t1.xlsx", t1)),
         ("T2 明細を 1 行消す", make("t2.xlsx", t2)),
         ("T3 連絡先ブロックを消す", make("t3.xlsx", t3)),
         ("T4 御中を 2 つにする", make("t4.xlsx", t4))]
recalc_all()
cases = _made

print("基準④: 壊した冊で印が立つか（期待 = 対照は確・壊した冊は確でない）\n")
fails = 0
for label, p in cases:
    r = inspect(p, sheet)
    v, g, why = r["請求額"]
    iv, ig, iwhy = r["請求元"]
    flag = ""
    if label.startswith("T0"):
        if g != "確": flag = "  ★ 対照で確が出ない＝測定器が壊れている"; fails += 1
    elif label.startswith(("T1", "T2")):
        if g == "確": flag = "  ★ 壊れているのに確 ＝ 黙って間違える"; fails += 1
    elif label.startswith(("T3", "T4")):
        if ig not in ("無", "単") or iv == a["請求元"] and ig == "確":
            flag = "  ★ 請求元で印が立たない"; fails += 1
    print(f"{label}")
    print(f"    請求額 {str(v):>8s}  [{g}]  {why}{flag if not label.startswith(('T3','T4')) else ''}")
    print(f"    請求元 {str(iv):>8s}  [{ig}]  {iwhy}{flag if label.startswith(('T3','T4')) else ''}")
print(f"\n印が立たなかった: {fails} 件（0 が合格）")
