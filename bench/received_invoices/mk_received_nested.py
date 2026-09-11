# -*- coding: utf-8 -*-
"""検体「入」（社名の入れ子・重なり）を作る（2026-09-11）。

mk_received_v2.py の作法を最小限まねる。★ 既存の mk_received_v2.py / specimens_v2.py /
答え_received_v2.json は 1 行も変更しない ── 別立て（received_nested/ 出力・別 answer）。

  ★ 骨は実物の雛形をそのまま使う（結合・帯・式が全部残る）
  ★ 答えは流し込んだ側が持つ（値と番地の両方）。道具の出力から作らない
  ★ 実物の金額は式。LibreOffice に開き直させてキャッシュを入れる（recalc）
  ★ 受け入れ条件: 書く前に各セルの値を予測し、LibreOffice の計算結果と 1 セルずつ突き合わせる

使い方:
    python mk_received_nested.py            # 全部
    python mk_received_nested.py 入01 入03  # id を指定（煙試験）
"""
import shutil
import subprocess
import sys
import json
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl                                   # noqa: E402

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

# ★ mk_received_v2 / skeletons_v2 / specimens_v2 は読み取り専用で import するだけ。
#   import しても既存ファイルは一切変更されない（main() は __main__ ガードの中）。
import mk_received_v2 as base                      # noqa: E402
from skeletons_v2 import SKELETONS                 # noqa: E402
from specimens_nested import SPECS                 # noqa: E402

OUT = HERE / "received_nested"
ANSWER = HERE / "答え_received_nested.json"


# ── 1 冊作る ───────────────────────────────────────────────────
def build_one(spec, idx):
    sid = spec["id"]
    dst = OUT / f"recv_{sid}.xlsx"

    sk = SKELETONS[spec["骨"]]
    shutil.copy(sk["src"], dst)
    wb = openpyxl.load_workbook(dst)          # ★ 式を残す（data_only を付けない）
    ws = wb[sk["sheet"]]
    inside = base.anchor_map(ws)
    d, sums = sk["detail"], sk["sums"]

    # ---- 雛形の見本を消す（先に消さないと答えと食い違う）----
    for row in range(d["first"], d["last"] + 1):
        for col in d["clear"]:
            base.put(ws, inside, f"{col}{row}", None)
    for addr in sk.get("note") or []:
        base.put(ws, inside, addr, None)
    if spec["骨"] == "misoca13":
        # ★ この骨は B1 に『見積書 ESTIMATE』が雛形のまま残っている（skeletons_v2 の quirk）。
        base.put(ws, inside, "B1", None)

    # ---- 請求元（発行者）── この検体群は全冊 issuer_name を明示する ----
    issuer = spec["issuer_name"]
    iss = sk["issuer"]
    issuer_at = base.put(ws, inside, iss["name"], issuer) if "name" in iss else None
    if iss.get("tel"):
        base.put(ws, inside, iss["tel"], "TEL：03-3000-%04d" % (idx % 10000))
    if iss.get("regno"):
        base.put(ws, inside, iss["regno"], "T%013d" % (1000000000000 + idx))

    # ---- 宛先（自社）── 骨によって「社名/御中」が別セルか同じセルかが違う ----
    to = sk["to"]
    SEPARATE = ("misoca256", "misoca12", "misoca13", "misoca16", "inv21")
    to_name = spec.get("to_name", base.JISHA if hasattr(base, "JISHA") else None)
    if to.get("zip"):
        base.put(ws, inside, to["zip"], base.JISHA_ZIP)
    if to.get("addr"):
        base.put(ws, inside, to["addr"], base.JISHA_ADDR)
    if spec["骨"] in SEPARATE:
        to_at = base.put(ws, inside, to["name"], to_name)
        if to.get("attn"):
            base.put(ws, inside, to["attn"], spec.get("to_attn", "経理部　御中"))
        to_value = to_name
    else:
        to_at = base.put(ws, inside, to["name"], to_name + "　御中")
        to_value = to_name + "　御中"

    # ---- 明細 ----
    lines = base.gen_lines(idx, spec.get("n", 2))
    if len(lines) > (d["last"] - d["first"] + 1):
        raise AssertionError(f"{sid}: 明細が雛形の枠を超える")
    rate = spec.get("rate", 0.1)
    for k, ln in enumerate(lines):
        row = d["first"] + k
        ln.setdefault("name", base.ITEMS[(idx + k) % len(base.ITEMS)])
        ln.setdefault("unitname", base.UNITS[(idx + k) % len(base.UNITS)])
        if d.get("date"):
            base.put(ws, inside, f"{d['date']}{row}", f"2026-08-{(idx % 28) + 1:02d}")
        base.put(ws, inside, f"{d['name']}{row}", ln["name"])
        if d.get("unitname"):
            base.put(ws, inside, f"{d['unitname']}{row}", ln["unitname"])
        if ln.get("qty") is not None:
            base.put(ws, inside, f"{d['qty']}{row}", ln["qty"])
        if ln.get("price") is not None:
            base.put(ws, inside, f"{d['price']}{row}", ln["price"])
        if d.get("mark") and ln.get("mark"):
            base.put(ws, inside, f"{d['mark']}{row}", ln["mark"])
        if d.get("rate"):
            base.put(ws, inside, f"{d['rate']}{row}", ln.get("rate", 0.1))
    if sk.get("rate_cell"):
        base.put(ws, inside, sk["rate_cell"], rate)

    # ---- 予測（★ 書く前に立てる）----
    pred = {}
    for k, ln in enumerate(lines):
        a = base.line_amount(sk, ln)
        if a is not None:
            pred[base.cell_at(ws, inside, f"{d['amount']}{d['first'] + k}")] = a
    pred.update({base.cell_at(ws, inside, k): v
                 for k, v in base.predict_sums(sk, lines, rate).items()})
    if issuer_at:
        pred[issuer_at] = issuer
    pred[to_at] = to_value

    # ---- 個別の細工（備考への差し込みなど）----
    for op in spec.get("post") or []:
        if op[0] == "set":
            base.put(ws, inside, op[1], op[2])
            pred[base.cell_at(ws, inside, op[1])] = op[2]
        elif op[0] == "clear":
            base.put(ws, inside, op[1], None)
            pred[base.cell_at(ws, inside, op[1])] = None

    wb.save(dst)
    wb.close()

    ans = dict(
        file=dst.name, id=sid, 群=spec["群"], 骨=sk["src"].name, シート=sk["sheet"],
        帳票種別=sk["kind"], 狙い=spec["狙い"], 出所=spec["出所"], 落とし方=spec["落とし方"],
        請求元=dict(値=issuer, 番地=issuer_at),
        宛先=dict(値=to_value, 番地=to_at),
        明細=[dict(品目=ln["name"], 数量=ln.get("qty"), 単価=ln.get("price"),
                   金額=base.line_amount(sk, ln)) for ln in lines],
        帯の番地=sums, 予測セル={k: v for k, v in pred.items()},
        雛形の癖=sk.get("quirk"), 期待=spec["期待"],
    )
    return ans


# ── LibreOffice に開き直させる（式のキャッシュを入れる）── 出力先はこの検体専用 ──
def recalc(files):
    work = HERE / "_recalc_nested"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    r = subprocess.run(
        [str(base.SOFFICE), "--headless", "--norestore", "--nologo",
         f"-env:UserInstallation={base.PROFILE}",
         "--convert-to", "xlsx", "--outdir", str(work)] + [str(f) for f in files],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=1800)
    made = list(work.glob("*.xlsx"))
    for f in made:
        shutil.copy(f, OUT / f.name)
    shutil.rmtree(work)
    return len(made), (r.stdout + r.stderr)[-400:]


# ── 受け入れ条件: 読み戻して 1 セルずつ突き合わせる ───────────────────
def verify(answers):
    bad = []
    for a in answers:
        if not a["予測セル"]:
            continue
        wb = openpyxl.load_workbook(OUT / a["file"], data_only=True)
        ws = wb[a["シート"]]
        for addr, want in a["予測セル"].items():
            got = ws[addr].value
            if not base.same(want, got):
                bad.append((a["id"], addr, want, got))
        wb.close()
    return bad


def main():
    only = set(sys.argv[1:])
    specs = [s for s in SPECS if not only or s["id"] in only]
    if OUT.exists() and not only:
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True, exist_ok=True)

    answers = []
    for i, spec in enumerate(SPECS):
        if only and spec["id"] not in only:
            continue
        answers.append(build_one(spec, i))
    print(f"作った: {len(answers)} 冊 → {OUT}")

    n, tail = recalc(sorted(OUT / a["file"] for a in answers))
    print(f"LibreOffice で開き直した: {n}/{len(answers)}")
    if n != len(answers):
        print("  ", tail)

    bad = verify(answers)
    ok_ids = {a["id"] for a in answers} - {b[0] for b in bad}
    print(f"\n受け入れ検査: 一致 {len(ok_ids)} 冊 / 食い違い {len(set(b[0] for b in bad))} 冊")
    for sid, addr, want, got in bad:
        print(f"  x {sid} {addr}: 予測={want!r} 実測={got!r}")

    for a in answers:
        a["採用"] = a["id"] in ok_ids

    ANSWER.write_text(json.dumps(answers, ensure_ascii=False, indent=2),
                      encoding="utf-8", newline="\n")
    print(f"答え → {ANSWER}")
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
