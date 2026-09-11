# -*- coding: utf-8 -*-
"""検体「束（フォルダ）」を作る（2026-09-11）。

mk_received_nested.py と同じ被せ方: mk_received_v2.py / skeletons_v2.py / specimens_v2.py は
1 行も変更せず `import mk_received_v2 as base` で読み取り専用に使う。出力は別立て
（received_bundle/<束>/*.xlsx・答え_received_bundle.json）。

この検体の単位は「1 冊」でなく「1 束（フォルダ）」── 同じ取引先が毎月出るフォルダを
複数作り、束の中でだけ見える怪しさ（重複・訂正再発行・年の誤り・桁違い・番号の飛び/
重なり・請求日が空）を仕込む。詳しい設計根拠は specimens_bundle.py の docstring を参照。

使い方:
    python mk_received_bundle.py            # 全部
    python mk_received_bundle.py B03 B03dup # id を指定（煙試験）
"""
import datetime
import json
import shutil
import subprocess
import sys
import zlib
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl                                   # noqa: E402

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import mk_received_v2 as base                      # noqa: E402
from skeletons_v2 import SKELETONS                  # noqa: E402
from specimens_bundle import (                      # noqa: E402
    BUNDLES, DATE_ADDR, INVNO_ADDR, MISOCA13_DATE_LABEL_CELL,
    MISOCA13_INVNO_LABEL_CELL, INV21_HAKKOBI_CELL, SEPARATE_TO, JISHA,
)

OUT = HERE / "received_bundle"
ANSWER = HERE / "答え_received_bundle.json"


# ── 請求日の値を「書式」から実際にセルへ書く値へ変換する ──────────────────
def date_text(y, m, d, fmt):
    """slash/dash/kanji/reiwa を文字列にする（date/空 はここでは使わない）。"""
    if fmt == "slash":
        return f"{y}/{m}/{d}"
    if fmt == "dash":
        return f"{y:04d}-{m:02d}-{d:02d}"
    if fmt == "kanji":
        return f"{y}年{m}月{d}日"
    if fmt == "reiwa":
        return f"令和{y - 2018}年{m}月{d}日"
    raise AssertionError(fmt)


def write_date(ws, inside, sk_id, book, pred):
    """請求日を書く。misoca13（値セルが無い）と inv21（表示形式にラベルが焼き込まれた
    日付型セル）は骨ごとの特別扱いが要る（specimens_bundle.py の docstring 参照）。
    戻り値: (番地, 答え用の値)
    """
    req = book["請求日"]
    fmt = req["書式"]

    if sk_id == "misoca13":
        # ★ 値セルが無い骨。ラベルのセル自体に「請求日：2026/8/31」と書き足す。
        addr = MISOCA13_DATE_LABEL_CELL
        if fmt == "空":
            raise AssertionError("misoca13 は値セルが無いので『空』の罠には使わない")
        text = "請求日：" + date_text(req["年"], req["月"], req["日"], fmt)
        at = base.put(ws, inside, addr, text)
        pred[at] = text
        return at, text

    addr = DATE_ADDR[sk_id]
    at = base.cell_at(ws, inside, addr)

    if fmt == "空":
        # ★ ラベルはある（骨の実物のまま）が、値セルを空のまま残す＝罠そのもの。
        #   constr は雛形に『××年1月1日』のプレースホルダが残っているので、
        #   明示的に None で消してはじめて「本当に空」になる。
        base.put(ws, inside, addr, None)
        pred[at] = None
        return at, None

    if fmt == "date" or sk_id == "inv21":
        # ★ inv21 は表示形式（例:「請求日： "yyyy"年"m"月"d"日"」）にラベルが焼き込まれた
        #   日付型セル。文字列を書くと書式が効かず生の文字列がそのまま出る＝壊れる。
        #   必ず datetime.date で書く（"date" 書式を指定していなくても inv21 は強制）。
        val = datetime.date(req["年"], req["月"], req["日"])
        base.put(ws, inside, addr, val)
        if sk_id != "inv21":
            ws[at].number_format = "yyyy/mm/dd"
        pred[at] = val
        return at, val.isoformat()

    text = date_text(req["年"], req["月"], req["日"], fmt)
    base.put(ws, inside, addr, text)
    pred[at] = text
    return at, text


def write_invno(ws, inside, sk_id, book, pred):
    """請求番号を書く。戻り値: (番地, 答え用の値)。"""
    no = book["請求番号"]
    if sk_id == "misoca13":
        addr = MISOCA13_INVNO_LABEL_CELL
        text = "請求番号：" + str(no)
        at = base.put(ws, inside, addr, text)
        pred[at] = text
        return at, text
    addr = INVNO_ADDR[sk_id]
    at = base.cell_at(ws, inside, addr)
    base.put(ws, inside, addr, no)
    pred[at] = no
    return at, no


# ── 1 冊作る ───────────────────────────────────────────────────
def build_one(book, idx):
    sid = book["id"]
    sk_id = book["骨"]
    sk = SKELETONS[sk_id]
    dst = OUT / book["束フォルダ"] / book["ファイル名"]
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(sk["src"], dst)
    wb = openpyxl.load_workbook(dst)          # ★ 式を残す（data_only を付けない）
    ws = wb[sk["sheet"]]
    inside = base.anchor_map(ws)
    d, sums = sk["detail"], sk["sums"]

    # ---- 雛形の見本を消す ----
    for row in range(d["first"], d["last"] + 1):
        for col in d["clear"]:
            base.put(ws, inside, f"{col}{row}", None)
    for addr in sk.get("note") or []:
        base.put(ws, inside, addr, None)
    if sk_id == "misoca13":
        base.put(ws, inside, "B1", None)      # ★『見積書 ESTIMATE』の残骸を消す

    # ---- 請求元（取引先）----
    # ★ TEL・登録番号は「取引先」に固定して振る（seed で振ると、seed だけ違う
    #   訂正再発行ペアで TEL・登録番号まで違って見えてしまい、罠が弱くなる。
    #   実物の同一取引先なら TEL・登録番号は月をまたいで同じはず）
    issuer = book["取引先"]
    vseed = zlib.crc32(issuer.encode("utf-8"))
    iss = sk["issuer"]
    issuer_at = base.put(ws, inside, iss["name"], issuer) if "name" in iss else None
    if iss.get("tel"):
        base.put(ws, inside, iss["tel"], "TEL：03-3000-%04d" % (vseed % 10000))
    if iss.get("regno"):
        base.put(ws, inside, iss["regno"], "T%013d" % (1000000000000 + vseed % 1000000000000))

    # ---- 宛先（自社）----
    to = sk["to"]
    if to.get("zip"):
        base.put(ws, inside, to["zip"], base.JISHA_ZIP)
    if to.get("addr"):
        base.put(ws, inside, to["addr"], base.JISHA_ADDR)
    if sk_id in SEPARATE_TO:
        to_at = base.put(ws, inside, to["name"], JISHA)
        if to.get("attn"):
            base.put(ws, inside, to["attn"], "経理部　御中")
        to_value = JISHA
    else:
        to_at = base.put(ws, inside, to["name"], JISHA + "　御中")
        to_value = JISHA + "　御中"

    # ---- 明細（seed で決定的に生成。金額倍率があれば桁を変える）----
    n = book.get("n", 3)
    lines = base.gen_lines(book["seed"], n)
    scale = book.get("金額倍率", 1.0)
    if scale != 1.0:
        for ln in lines:
            ln["price"] = round(ln["price"] * scale)
    if len(lines) > (d["last"] - d["first"] + 1):
        raise AssertionError(f"{sid}: 明細が雛形の枠を超える")
    for k, ln in enumerate(lines):
        row = d["first"] + k
        ln.setdefault("name", base.ITEMS[(book["seed"] + k) % len(base.ITEMS)])
        ln.setdefault("unitname", base.UNITS[(book["seed"] + k) % len(base.UNITS)])
        if d.get("date"):
            base.put(ws, inside, f"{d['date']}{row}",
                     f"2026-{(book['seed'] % 12) + 1:02d}-{(book['seed'] % 28) + 1:02d}")
        base.put(ws, inside, f"{d['name']}{row}", ln["name"])
        if d.get("unitname"):
            base.put(ws, inside, f"{d['unitname']}{row}", ln["unitname"])
        if ln.get("qty") is not None:
            base.put(ws, inside, f"{d['qty']}{row}", ln["qty"])
        if ln.get("price") is not None:
            base.put(ws, inside, f"{d['price']}{row}", ln["price"])
        if d.get("rate"):
            # ★ inv21（rate_column）は税率の列が空だと SUMIF が 1 行も拾わず
            #   小計が黙って 0 になる（skeletons_v2.py の quirk）。必ず書く。
            base.put(ws, inside, f"{d['rate']}{row}", ln.get("rate", 0.1))
    if sk.get("rate_cell"):
        base.put(ws, inside, sk["rate_cell"], 0.1)

    # ---- 予測（★ 書く前に立てる）----
    pred = {}
    for k, ln in enumerate(lines):
        a = base.line_amount(sk, ln)
        if a is not None:
            pred[base.cell_at(ws, inside, f"{d['amount']}{d['first'] + k}")] = a
    pred.update({base.cell_at(ws, inside, k): v
                 for k, v in base.predict_sums(sk, lines, 0.1).items()})
    if issuer_at:
        pred[issuer_at] = issuer
    pred[to_at] = to_value

    # ---- 請求日・請求番号 ----
    date_at, date_ans = write_date(ws, inside, sk_id, book, pred)
    invno_at, invno_ans = write_invno(ws, inside, sk_id, book, pred)

    # inv21 は「発行日」がもう 1 箇所ある（H6）── 揃えないと 1 冊に日付が 2 つ食い違う
    if sk_id == "inv21" and book["請求日"]["書式"] != "空":
        req = book["請求日"]
        hakko_at = base.cell_at(ws, inside, INV21_HAKKOBI_CELL)
        val = datetime.date(req["年"], req["月"], req["日"])
        base.put(ws, inside, INV21_HAKKOBI_CELL, val)
        pred[hakko_at] = val

    wb.save(dst)
    wb.close()

    請求額addr = sums.get("請求額")
    請求額at = base.cell_at(ws, inside, 請求額addr) if 請求額addr else None
    請求額val = pred.get(請求額at) if 請求額at else None

    ans = dict(
        file=book["ファイル名"], id=sid, 束=book["束"], フォルダ=book["束フォルダ"],
        骨=sk["src"].name, シート=sk["sheet"], 帳票種別=sk["kind"],
        請求元=dict(値=issuer, 番地=issuer_at),
        宛先=dict(値=to_value, 番地=to_at),
        請求額=dict(値=請求額val, 番地=請求額at),
        請求日=dict(値=date_ans, 書式=book["請求日"]["書式"], 番地=date_at),
        請求番号=dict(値=invno_ans, 番地=invno_at),
        明細=[dict(品目=ln["name"], 数量=ln.get("qty"), 単価=ln.get("price"),
                   金額=base.line_amount(sk, ln)) for ln in lines],
        帯の番地=sums, 雛形の癖=sk.get("quirk"),
        複製元=book.get("複製元"),
        予測セル={k: v for k, v in pred.items()},
    )
    return ans


# ── LibreOffice に開き直させる（式のキャッシュを入れる）── 出力先はこの検体専用 ──
def recalc(files):
    work = HERE / "_recalc_bundle"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    made_total = []
    tail = ""
    # ★ soffice --convert-to は同名ファイルを 1 つの --outdir で衝突させる
    #   （束をまたいで同じファイル名は無いが、束フォルダを保ったまま束ごとに変換する）
    by_dir = {}
    for f in files:
        by_dir.setdefault(f.parent, []).append(f)
    for folder, fs in by_dir.items():
        r = subprocess.run(
            [str(base.SOFFICE), "--headless", "--norestore", "--nologo",
             f"-env:UserInstallation={base.PROFILE}",
             "--convert-to", "xlsx", "--outdir", str(work)] + [str(f) for f in fs],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=1800)
        tail = (r.stdout + r.stderr)[-400:]
        made = list(work.glob("*.xlsx"))
        for f in made:
            shutil.copy(f, folder / f.name)
            f.unlink()
        made_total += made
    shutil.rmtree(work, ignore_errors=True)
    return len(made_total), tail


# ── 受け入れ条件: 読み戻して 1 セルずつ突き合わせる ───────────────────
def same_date_aware(pred, got):
    """base.same に日付比較を足す（datetime 型を素直に比較できない可能性への対処）。"""
    if isinstance(pred, (datetime.date, datetime.datetime)):
        if got is None:
            return False
        gy = got
        if isinstance(gy, str):
            return False    # 日付型を期待したのに文字列で返ってきたら不一致
        pd = pred.date() if isinstance(pred, datetime.datetime) else pred
        gd = gy.date() if isinstance(gy, datetime.datetime) else gy
        return pd == gd
    return base.same(pred, got)


def verify(answers):
    bad = []
    for a in answers:
        if not a["予測セル"]:
            continue
        path = OUT / a["フォルダ"] / a["file"]
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[a["シート"]]
        for addr, want in a["予測セル"].items():
            got = ws[addr].value
            if not same_date_aware(want, got):
                bad.append((a["id"], addr, want, got))
        wb.close()
    return bad


def main():
    only = set(sys.argv[1:])

    # ---- BUNDLES を平らな冊リストに展開（束フォルダ名を確定）----
    all_books = []
    for bd in BUNDLES:
        folder = bd["束"]
        for book in bd["冊"]:
            book = dict(book)
            book["束"] = bd["束"]
            book["束フォルダ"] = folder
            all_books.append(book)

    specs = [b for b in all_books if not only or b["id"] in only]
    if OUT.exists() and not only:
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True, exist_ok=True)

    answers = []
    for i, book in enumerate(specs):
        answers.append(build_one(book, i))
    print(f"作った: {len(answers)} 冊 → {OUT}")

    files = sorted(OUT / a["フォルダ"] / a["file"] for a in answers)
    n, tail = recalc(files)
    print(f"LibreOffice で開き直した: {n}/{len(answers)}")
    if n != len(answers):
        print("  ", tail)

    bad = verify(answers)
    ok_ids = {a["id"] for a in answers} - {b[0] for b in bad}
    print(f"\n受け入れ検査: 一致 {len(ok_ids)} 冊 / 食い違い {len(set(b[0] for b in bad))} 冊")
    for sid, addr, want, got in bad:
        print(f"  x {sid} {addr}: 予測={want!r} 実測={got!r}")
    for a in answers:
        a["採用"] = a["id"] in ok_ids

    # ---- 答え JSON を「束一覧」の形に組み直す ----
    by_bundle = {}
    for a in answers:
        by_bundle.setdefault(a["束"], []).append(a)

    束一覧 = []
    for bd in BUNDLES:
        if only and bd["束"] not in by_bundle:
            continue
        冊 = by_bundle.get(bd["束"], [])
        怪しい済み = set()
        for s in bd["疑い"]:
            怪しい済み.update(s["冊"])
        for m in bd["迷う"]:
            怪しい済み.add(m["冊"])
        怪しくない = [b["file"] for b in 冊 if b["file"] not in 怪しい済み]
        束一覧.append(dict(
            束=bd["束"], 陰性対照=bd.get("陰性対照", False), 狙い=bd.get("狙い"),
            冊=冊, 疑い=bd["疑い"], 怪しくない=怪しくない, 迷う=bd["迷う"],
        ))

    ANSWER.write_text(
        json.dumps(dict(束一覧=束一覧), ensure_ascii=False, indent=2, default=str),
        encoding="utf-8", newline="\n")
    print(f"答え → {ANSWER}")
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
