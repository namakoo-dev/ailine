# -*- coding: utf-8 -*-
"""検体 B: 骨は実物・肉は合成 ──「受け取った請求書 40 冊」を作る（2026-09-10 夜）。

★ 3 体の architect が揃って「STEP 0」と言ったもの。Namakoo の回答（Q1: 受け取った請求書）で
  流し込み方が確定した:

      宛先（御中）   = 自社 = **全冊 同じ値**
      請求元         = 40 社 = **冊ごとに違う**   ← 欲しいのはこっち

★ 骨は実物の雛形をそのまま使う（結合セル・帯構造・256 列・消費税行・**式**が全部残る）。
  合成が持ち込む想像は「値」だけ。今日 3 回外した「構造の想像」が入らない。

★ 実物の金額は式（misoca: C11=SUM(H37,H38) 等）。openpyxl で保存しただけでは
  キャッシュが無く data_only=True で全部 None になる。**LibreOffice に開かせて保存する**
  段が要る（人が雛形に入力して保存したのと同じ状態にする）。

★ 答えは流し込んだ側が持つ。値と番地の両方。道具の出力から作らない。
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

SP = Path(__file__).resolve().parent   # ★ 雛形と検体はこの隣（追跡しない）
OUT = SP / "received40"
BASRUN = Path(r"C:\Dev\basrun\basrun.py")

# ★ 自社（受け取る側）── 全冊で同じ
JISHA = "ナギ商会株式会社"

ISSUERS = [
    "あかね商事", "いろは工業", "うえだ物産", "エバラ機械", "大久保商店",
    "花菱建材", "菊池電機", "くまがい精工", "ケイアイ物流", "小坂運輸",
    "佐々木製作所", "しなの化成", "鈴村テック", "瀬戸内海運", "曽根田工務店",
    "高梨産業", "千葉見本市", "つばめ交通", "寺岡食品", "戸田金属",
]

# 骨（実物の雛形）と、そこへ書き込む番地。★ 実物を開いて目で確かめた位置。
SKELETONS = [
    {   # misoca
        "src": SP / "real14" / "a0cf7302-blackline.xlsx",
        "sheet": "misoca_invoice",
        "issuer": "G5",          # 発行者（右ブロックの先頭）
        "atesaki": "B4",         # 「〜御中」
        "atesaki_fmt": "{} 御中",
        "detail": {"start": 15, "name": "B", "qty": "E", "unit": "G", "rows": 22, "also": ("F",)},
        "answer": {"請求額": "C11", "小計": "H37", "消費税": "H38", "合計": "H39"},
    },
    {   # マネーフォワード（建設）
        "src": SP / "construction_bill.xlsx",
        "sheet": "適格請求書（インボイス）",
        "issuer": "F2",
        "atesaki": "B7",
        "atesaki_fmt": "{}　御中",
        "detail": {"start": 16, "name": "C", "qty": "F", "unit": "E", "rows": 9, "also": ("B", "D"), "date": "B"},
        "answer": {"請求額": "E14", "小計": "E28", "消費税": "G28"},
    },
    {   # マネーフォワード（別レイアウト）
        "src": SP / "inv21.xlsx",
        "sheet": "インボイス対応請求書",
        "issuer": "B11",
        "atesaki": "B13",
        "atesaki_fmt": "{}　御中",
        "detail": {"start": 32, "name": "D", "qty": "I", "unit": "H", "rows": 14, "also": ("B", "C", "E", "F", "G", "J"), "date": "B", "rate": "J"},
        "answer": {"請求額": "D22"},
    },
]

ITEMS = ["用紙代", "運送費", "保守料", "部材費", "作業代", "設計費", "検査料"]


def build():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    answers = []
    for i in range(40):
        sk = SKELETONS[i % len(SKELETONS)]
        issuer = ISSUERS[i % len(ISSUERS)]
        if i >= len(ISSUERS):
            issuer += "（西日本）"
        dst = OUT / f"received_{i + 1:02d}.xlsx"
        shutil.copy(sk["src"], dst)

        wb = openpyxl.load_workbook(dst)       # ★ 式を残す
        ws = wb[sk["sheet"]]
        ws[sk["issuer"]] = issuer
        ws[sk["atesaki"]] = sk["atesaki_fmt"].format(JISHA)

        d = sk["detail"]
        # ★ 雛形には見本の明細が入っている（misoca の「サンプル」4 行など）。
        #   先に消さないと、答えが「俺が書いた分」と食い違う（実際に 0/40 で外した）。
        # ★ 結合セルは左上以外に書けない（MergedCell は read-only）。左上だけ触る。
        anchors = {(r.min_row, r.min_col) for r in ws.merged_cells.ranges}
        inside = set()
        for rng in ws.merged_cells.ranges:
            for rr in range(rng.min_row, rng.max_row + 1):
                for cc in range(rng.min_col, rng.max_col + 1):
                    if (rr, cc) not in anchors:
                        inside.add((rr, cc))
        for row in range(d["start"] + 1, d["start"] + 1 + d["rows"]):
            for col in (d["name"], d["qty"], d["unit"], *d.get("also", ())):
                cell = ws[f"{col}{row}"]
                if (cell.row, cell.column) in inside:
                    continue
                cell.value = None
        n_lines = 1 + (i % 3)
        lines = []
        for k in range(n_lines):
            row = d["start"] + 1 + k
            qty = 1 + ((i + k) % 5)
            unit = 1000 * (1 + ((i + k) % 9))
            if d.get("date"):
                # ★ inv21/建設 の金額式は IF(B="","",…) ── 取引日が無いと金額が出ない
                ws[f"{d['date']}{row}"] = f"2026-08-{(i % 28) + 1:02d}"
            ws[f"{d['name']}{row}"] = ITEMS[(i + k) % len(ITEMS)]
            ws[f"{d['qty']}{row}"] = qty
            if d.get("rate"):
                # ★ 集計は SUMIF(税率列, 10%, 金額列) ── 税率が空だと合計が 0 になる
                ws[f"{d['rate']}{row}"] = 0.1
            ws[f"{d['unit']}{row}"] = unit
            lines.append({"品目": ITEMS[(i + k) % len(ITEMS)], "数量": qty, "単価": unit,
                          "金額": qty * unit})
        wb.save(dst)
        wb.close()

        shokei = sum(l["金額"] for l in lines)
        answers.append({
            "file": dst.name, "骨": sk["src"].name, "シート": sk["sheet"],
            "請求元": issuer, "請求元の番地": sk["issuer"],
            "宛先": JISHA, "宛先の番地": sk["atesaki"],
            "明細": lines, "小計(税抜)": shokei,
            "消費税(10%)": round(shokei * 0.1),
            "請求額(税込)": shokei + round(shokei * 0.1),
            "答えの番地": sk["answer"],
        })
    (SP / "答え_received40.json").write_text(
        json.dumps(answers, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    return answers


def recalc():
    """★ LibreOffice に開き直させて保存し、式のキャッシュを入れる。

    ★ basrun に recalc は無いので soffice の --convert-to を使う（開く→計算→保存）。
      専用プロファイルを使い、利用者の GUI を巻き込まない（basrun と同じ作法）。
    """
    soffice = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
    profile = (SP / "lo-profile-mk").resolve().as_uri()
    work = OUT.parent / "_recalc"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    files = sorted(OUT.glob("*.xlsx"))
    r = subprocess.run(
        [str(soffice), "--headless", "--norestore", "--nologo",
         f"-env:UserInstallation={profile}",
         "--convert-to", "xlsx", "--outdir", str(work)] + [str(f) for f in files],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900)
    made = list(work.glob("*.xlsx"))
    for f in made:
        shutil.copy(f, OUT / f.name)
    shutil.rmtree(work)
    return len(made), len(files) - len(made), (r.stdout + r.stderr)[-300:]


if __name__ == "__main__":
    a = build()
    print(f"作った: {len(a)} 冊 → {OUT}")
    print(f"  骨: {sorted({x['骨'] for x in a})}")
    print(f"  請求元: {len({x['請求元'] for x in a})} 社（★ 冊ごとに違う）")
    print(f"  宛先: {sorted({x['宛先'] for x in a})}（★ 全冊 同じ）")
    print(f"  請求額(税込)の合計: {sum(x['請求額(税込)'] for x in a):,}")
    ok, ng, tail = recalc()
    print("")
    print(f"LibreOffice で開き直した: 成功 {ok} / 失敗 {ng}")
    if ng:
        print("  ", tail)
