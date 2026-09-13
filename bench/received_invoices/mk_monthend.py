# -*- coding: utf-8 -*-
"""検体「月末の束」(monthend) を作る（2026-09-13）。

★ LibreOffice は起動しない。骨（real14 / vendors の実物雛形）も使わない ── この検体の主題は
  帳票の項目抽出（① 需要）ではなく束の所見（forms_suspect）なので、`ailine_core.form_read` が
  読める最小限の様式を openpyxl だけで**直接**書く（式は使わない・全部リテラル値）。
  レイアウトは事前に `python -m ailine forms` へ実際に通して確かめた（read_book は骨に
  依存しない一般規則で読むため・form_read.py の docstring ①〜⑥ 参照）。

使い方:
    python mk_monthend.py
"""
from __future__ import annotations

import datetime
import json
import sys
import shutil
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl                                      # noqa: E402

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from specimens_monthend import BUNDLES, BUYER          # noqa: E402

OUT = HERE / "monthend_books"
ANSWER = HERE / "答え_monthend.json"

_TAX_RATE = 0.1


def _write_book(path: Path, issuer: str, buyer: str, date, invno, amount: int) -> None:
    """1 冊を書く。式は一切使わない（合計・小計・消費税もリテラル値）。

    番地は `python -m ailine forms` に実際に通して確かめた最小限の様式
    （G3/G4 に請求日・請求番号のラベル、B3/B4 が宛先、G6 以下が発行元、
    B16 以下が明細、G37〜G39 が帯、B11/C11 が上部の請求額欄）。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求書"

    # ── 宛先（自社）── ①御中は部署の側に付く（会社名そのものと分ける）
    ws["B3"] = buyer
    ws["B4"] = "経理部　御中"

    # ── 請求日・請求番号 ──
    ws["G3"] = "請求日："
    if date is not None:
        ws["H3"] = date
    # date is None のときは値セルを書かない（ラベルはあるが空＝罠そのもの）
    ws["G4"] = "請求番号："
    ws["H4"] = invno

    # ── 発行元（取引先）── 登録番号が付くのは発行元の側だけ（買い手には付かない）
    ws["G6"] = issuer
    ws["G7"] = "〒100-0001"
    ws["G8"] = "東京都千代田区大手町1-1-1"
    ws["G9"] = "TEL：03-1234-5678"
    ws["G10"] = "T1234567890123"

    # ── 明細（1 行・数量 1・単価=金額）──
    ws["B16"] = "品番・品名"
    ws["E16"] = "数量"
    ws["F16"] = "単位"
    ws["G16"] = "単価"
    ws["H16"] = "金額"
    subtotal = round(amount / (1 + _TAX_RATE))
    tax = amount - subtotal
    ws["B17"] = "作業一式"
    ws["E17"] = 1
    ws["F17"] = "式"
    ws["G17"] = subtotal
    ws["H17"] = subtotal

    # ── 帯（小計・消費税・合計）── リテラル値。合計は小計+消費税の式ではない
    #   （独立した根拠として数えてもらうため。式にすると「検算のつもりが同じ計算」になる）。
    ws["G37"] = "小計"
    ws["H37"] = subtotal
    ws["G38"] = "消費税"
    ws["H38"] = tax
    ws["G39"] = "合計金額"
    ws["H39"] = amount

    # ── 上部の請求額欄（帯とは別の独立したセル）──
    ws["B11"] = "ご請求金額"
    ws["C11"] = amount

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def main() -> int:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    束一覧 = []
    total_books = 0
    for bd in BUNDLES:
        folder = OUT / bd["束"]
        for book in bd["冊"]:
            _write_book(folder / book["ファイル名"], book["取引先"], BUYER,
                        book["請求日"], book["請求番号"], book["請求額"])
            total_books += 1

        疑い済み = set()
        for w in bd["疑い"]:
            疑い済み.update(w["冊"])
        怪しくない = [b["ファイル名"] for b in bd["冊"] if b["ファイル名"] not in 疑い済み]

        束一覧.append(dict(
            束=bd["束"], 狙い=bd.get("狙い"),
            冊=[dict(file=b["ファイル名"], id=b["id"], 取引先=b["取引先"],
                     請求日=b["請求日"].isoformat() if isinstance(b["請求日"], datetime.date) else None,
                     請求番号=b["請求番号"], 請求額=b["請求額"]) for b in bd["冊"]],
            疑い=[{k: v for k, v in w.items() if k != "メモ"} for w in bd["疑い"]],
            怪しくない=怪しくない,
        ))

    ANSWER.write_bytes(
        json.dumps(dict(束一覧=束一覧), ensure_ascii=False, indent=2, default=str)
        .encode("utf-8"))
    print(f"作った: {total_books} 冊 / {len(BUNDLES)} 束 → {OUT}")
    print(f"答え → {ANSWER}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
