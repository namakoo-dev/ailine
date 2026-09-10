# -*- coding: utf-8 -*-
"""雛形をそのまま覗く（読むだけ）。値・式・結合を並べる。"""
import sys, warnings
from pathlib import Path
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import get_column_letter as gl

p = Path(sys.argv[1])
maxr = int(sys.argv[2]) if len(sys.argv) > 2 else 60
wbf = openpyxl.load_workbook(p, data_only=False)
wbv = openpyxl.load_workbook(p, data_only=True)
for sn in wbf.sheetnames:
    wf, wv = wbf[sn], wbv[sn]
    print(f"### {p.name} 『{sn}』 {wf.max_row}行 x {wf.max_column}列 結合{len(wf.merged_cells.ranges)}")
    mg = {}
    for r in wf.merged_cells.ranges:
        mg[(r.min_row, r.min_col)] = str(r)
    for r in range(1, min(wf.max_row, maxr) + 1):
        out = []
        for c in range(1, min(wf.max_column, 30) + 1):
            f = wf.cell(row=r, column=c).value
            v = wv.cell(row=r, column=c).value
            if f in (None, "") and v in (None, ""):
                continue
            a = f"{gl(c)}{r}"
            m = mg.get((r, c))
            s = f"{a}"
            if m and ":" in m:
                s += f"[{m.split(':')[1]}]"
            if isinstance(f, str) and f.startswith("="):
                s += f"= {f}  ->{v!r}"
            else:
                s += f"= {f!r}"
            out.append(s)
        if out:
            print("  " + " | ".join(out))
    print()
