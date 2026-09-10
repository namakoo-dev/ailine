# -*- coding: utf-8 -*-
"""検体の「罠が本当に入っているか」を、答えの JSON を見ずに成果物だけから確かめる。

★ なぜ別に要るか: mk_received_v2 の受け入れ検査は「俺の予測」と「LibreOffice の計算」
  を突き合わせるもので、**狙った罠が入っているか**は一言も言っていない。
  予測が全部当たったまま、罠だけが抜けていることが起こりうる（出ないことは信号でない）。
  ここは成果物のセルを直接読んで、罠の在り処を名指しで主張する。

  ★ 検算は本体（mk_received_v2）と別実装にしてある。ここでは骨の番地表 _sk.py すら
    使わず、シートを走査して「御中を含むセル」「数値の集合」から判定する。
"""
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl

OUT = Path(__file__).resolve().parent / "received_v2"


def sheets(fid, visible_only=False):
    wb = openpyxl.load_workbook(OUT / f"recv_{fid}.xlsx", data_only=True)
    out = [wb[n] for n in wb.sheetnames
           if not visible_only or wb[n].sheet_state == "visible"]
    return wb, out


def cells(x, maxr=100, maxc=45):
    """★ 1 冊は「シートの束」として扱う。construction_bill は 1 枚目が『説明』で、
    1 枚目だけ見ると罠を全部見落とす（この点検自体が最初それで空振りした）。"""
    for ws in (x if isinstance(x, list) else [x]):
        for r in range(1, min(ws.max_row or 1, maxr) + 1):
            for c in range(1, min(ws.max_column or 1, maxc) + 1):
                v = ws.cell(row=r, column=c).value
                if v not in (None, ""):
                    yield r, c, v


def main_title(x):
    ss = x if isinstance(x, list) else [x]
    return max(ss, key=lambda w: sum(1 for _ in cells(w))).title


def texts(ws):
    return [(r, c, v) for r, c, v in cells(ws) if isinstance(v, str)]


def nums(ws):
    return [(r, c, v) for r, c, v in cells(ws) if isinstance(v, (int, float))
            and not isinstance(v, bool)]


def has(ws, sub):
    return [(r, c, v) for r, c, v in texts(ws) if sub in v]


RESULTS = []


def check(fid, name, cond, detail=""):
    RESULTS.append((fid, name, bool(cond), detail))


def main():
    # ── 基礎群: 骨 x 行数 が直交しているか（成果物側から数える）─────────
    import collections
    grid = collections.Counter()
    for i in range(1, 37):
        wb, ss = sheets(f"B{i:02d}")
        ws = ss
        w0 = max(ss, key=lambda w: len(w.merged_cells.ranges))
        # 骨の指紋: シート名 + 列数 + 結合数（番地表を使わない）
        fp = (w0.title, w0.max_column, len(w0.merged_cells.ranges))
        # 明細の本数: 品目に使った語がいくつ在るか
        items = ["用紙代", "運送費", "保守料", "部材費", "作業代", "設計費",
                 "検査料", "梱包資材", "出張旅費", "外注工賃"]
        n = sum(1 for _, _, v in texts(ws) if v.strip() in items)
        grid[(fp, n)] += 1
        wb.close()
    fps = {fp for fp, _ in grid}
    ns = {n for _, n in grid}
    check("B*", "骨 6 種 x 行数 6 種の総当たりで各 1 冊",
          len(fps) == 6 and ns == {0, 1, 2, 3, 5, 8} and all(v == 1 for v in grid.values()),
          f"骨={len(fps)} 行数={sorted(ns)} 重複={[k for k, v in grid.items() if v != 1]}")

    # ── 欠陥 (2) の直し: 「御中/様のセル = 宛先社名」が全冊では真でない ──
    onchu_is_name = []
    for i in range(1, 37):
        wb, ss = sheets(f"B{i:02d}")
        hits = has(ss, "御中") + has(ss, "様")
        onchu_is_name.append(any("ナギ商会" in v for _, _, v in hits))
        wb.close()
    check("B*", "御中/様 のセルが宛先社名である骨と、そうでない骨が両方在る",
          any(onchu_is_name) and not all(onchu_is_name),
          f"真 {sum(onchu_is_name)}/36 冊")

    # ── 税率群 ────────────────────────────────────────────────────
    for fid, want8 in [("R01", True), ("R02", True), ("R03", True), ("R04", True)]:
        wb, ss = sheets(fid)
        ws = ss
        got = any(abs(v - 0.08) < 1e-9 for _, _, v in nums(ws)) or bool(has(ws, "※"))
        check(fid, "8% の痕跡（税率 0.08 か ※）が在る", got == want8)
        wb.close()
    for fid in ("R05", "R06"):
        wb, ss = sheets(fid)
        ws = ss
        check(fid, "税率セルが 0.08",
              any(abs(v - 0.08) < 1e-9 for _, _, v in nums(ws)))
        wb.close()

    # ── 敵対群 ────────────────────────────────────────────────────
    def one(fid, visible_only=False):
        return sheets(fid, visible_only)

    wb, ws = one("T01")
    check("T01", "御中/様 のセルに自社名が入っていない",
          all("ナギ商会" not in v for _, _, v in has(ws, "様") + has(ws, "御中")),
          str(has(ws, "様"))[:60]); wb.close()

    wb, ws = one("T02")
    check("T02", "宛先ブロック側（左）にも TEL が在る",
          any(c <= 4 for r, c, v in has(ws, "TEL")) and
          any(c >= 5 for r, c, v in has(ws, "TEL"))); wb.close()

    wb, ws = one("T03")
    hit = [v for _, _, v in texts(ws) if "\n" in v and "登録番号" in v and "高梨産業" in v]
    check("T03", "発行者名と登録番号が 1 セルに改行で同居", hit, str(hit)[:70]); wb.close()

    wb, ws = one("T04")
    check("T04", "振込先に発行者と違う口座名義",
          has(ws, "チユウオウフアクタリング")); wb.close()

    wb, ws = one("T05")
    check("T05", "発行者名が伏せ字のまま", has(ws, "株式会社 〇〇〇")); wb.close()

    wb, ws = one("T06")
    check("T06", "担当欄にだけ社名、社名欄は空",
          has(ws, "担当：曽根田工務店") and not has(ws, "株式会社 〇〇〇")); wb.close()

    wb, ws = one("T07")
    check("T07", "氏名（屋号：〜）の形", has(ws, "（屋号：")); wb.close()

    for fid in ("T08", "T39"):
        wb, ws = one(fid)
        check(fid, "前株ありと前株なしの同一社名が別セルに在る",
              has(ws, "株式会社あかね商事") and
              any("あかね商事" in v and "株式会社" not in v for _, _, v in texts(ws)))
        wb.close()

    wb, ws = one("T09")
    check("T09", "別会社が 2 つ", has(ws, "しなの化成") and has(ws, "中央商事")); wb.close()

    wb, ws = one("T10")
    check("T10", "「御請求先：自社名」が本文に在る", has(ws, "御請求先：ナギ商会")); wb.close()

    wb, ws = one("T11")
    check("T11", "御中が 2 箇所", len(has(ws, "御中")) >= 2,
          f"{len(has(ws,'御中'))} 箇所"); wb.close()

    wb, ws = one("T12")
    check("T12", "発行者ブロックが空（〒/TEL/社名が右側に無い）",
          not any(c >= 6 for r, c, v in texts(ws) if "〒" in v or "TEL" in v)); wb.close()

    wb, ws = one("T13")
    ns_ = sorted({round(v, 2) for _, _, v in nums(ws) if v > 1000})
    check("T13", "税込合計と今回請求額が別の値として同居（繰越 80,000 の差）",
          any(abs((a - b) - 80000) < 1e-6 for a in ns_ for b in ns_), str(ns_)[:80]); wb.close()

    wb, ws = one("T14")
    check("T14", "源泉と振込金額が備考に在り、合計と違う",
          has(ws, "源泉所得税") and has(ws, "121,790")); wb.close()

    wb, ws = one("T15")
    check("T15", "マイナス金額の明細が在る",
          any(v < 0 for _, _, v in nums(ws))); wb.close()

    wb, ws = one("T16")
    wf = openpyxl.load_workbook(OUT / "recv_T16.xlsx")[main_title(ws)]
    formulas = [v for r, c, v in cells(wf) if isinstance(v, str) and v.startswith("=")
                and r >= 39]
    check("T16", "明細ゼロで合計が定数（合計行に式が無い）",
          132000 in [v for _, _, v in nums(ws)] and not formulas, str(formulas)[:60]); wb.close()

    # 罠の本体: Σ明細 と 小計 が食い違う 3 冊。★ 明細行の金額を骨に依らず拾う
    def detail_vs_sub(fid, band_row):
        wb, ws = one(fid)
        det = [(r, c, v) for r, c, v in nums(ws) if r < band_row and v >= 1000]
        wb.close()
        return det

    wb, ws = one("T17")
    check("T17", "集計範囲の外（25 行目）に 90,000 の明細が在る",
          any(r == 25 and v == 90000 for r, c, v in nums(ws))); wb.close()
    wb, ws = one("T18")
    check("T18", "税率が空の明細行が在る（15,000 の行）",
          any(v == 15000 for _, _, v in nums(ws))); wb.close()
    wb, ws = one("T19")
    check("T19", "軽減税率欄が全角スペースの行が在り、25,000 が小計から落ちている",
          any(v.strip() == "" and v != "" for _, _, v in texts(ws))
          and {40000, 25000, 20000, 60000, 66000} <= {v for _, _, v in nums(ws)}
          and 85000 not in {v for _, _, v in nums(ws)},   # Σ明細 85,000 が帳票に無い
          str(sorted({v for _, _, v in nums(ws) if v >= 10000}))); wb.close()

    wb, ws = one("T20")
    mx = max(v for _, _, v in nums(ws))
    check("T20", "シート内の最大値が請求番号（請求額ではない）",
          mx == 20260901, f"最大値={mx}"); wb.close()

    wb, ws = one("T21")
    check("T21", "請求額欄が文字列", has(ws, "¥1,320,000-")); wb.close()
    wb, ws = one("T34")
    check("T34", "請求額欄が全角数字の文字列", has(ws, "１，３２０，０００")); wb.close()

    wb, ws = one("T22")
    check("T22", "消費税 0（免税事業者）",
          any(v == 0 for _, _, v in nums(ws)) and has(ws, "免税事業者")); wb.close()

    wb, ws = one("T23")
    frac = [v for _, _, v in nums(ws) if isinstance(v, float) and abs(v - round(v)) > 1e-9
            and v > 1]
    check("T23", "金額に小数が出ている", frac, str(frac)[:60]); wb.close()

    wb, ws = one("T24")
    vs = {v for _, _, v in nums(ws)}
    check("T24", "小計 100,000 ＋ 消費税 10,000 なのに合計が 109,999（1 円の食い違い）",
          {100000, 10000, 109999} <= vs, f"{sorted(v for v in vs if v > 5000)}"); wb.close()

    wb, ws = one("T25")
    zr = [r for r, c, v in texts(ws) if "無償" in v]
    below = [r for r, c, v in nums(ws) if zr and r > zr[0] and v >= 10000]
    check("T25", "0 円の明細行があり、その下にまだ明細が在る",
          zr and below, f"0円行={zr} その下の金額行={sorted(set(below))[:4]}"); wb.close()

    wb, ws = one("T26")
    rows = sorted({r for r, c, v in texts(ws)
                   if v.strip() in ("用紙代", "運送費", "保守料", "部材費", "作業代",
                                    "設計費", "検査料", "梱包資材", "出張旅費", "外注工賃")})
    check("T26", "明細行の間に空行が在る",
          len(rows) == 2 and rows[1] - rows[0] == 2, str(rows)); wb.close()

    wb, ws = one("T27")
    tot = max(v for _, _, v in nums(ws))
    check("T27", "備考の再掲金額が合計と一致",
          any(f"{int(tot):,}" in v for _, _, v in texts(ws)), f"合計={tot}"); wb.close()
    wb, ws = one("T28")
    tot = max(v for _, _, v in nums(ws))
    check("T28", "備考の再掲金額が合計と食い違う",
          has(ws, "121,000") and int(tot) != 121000, f"合計={tot}"); wb.close()

    wb, ws = one("T29")
    check("T29", "単価 80,000 が在るのに、その行の金額が出ていない",
          any(v == 80000 for _, _, v in nums(ws)) and
          not any(v == 80000 * 1 for r, c, v in nums(ws) if c >= 8)); wb.close()

    wbf = openpyxl.load_workbook(OUT / "recv_T30.xlsx")
    wsf = wbf[wbf.sheetnames[0]]
    check("T30", "明細行 17 が非表示",
          wsf.row_dimensions[17].hidden is True); wbf.close()

    wbf = openpyxl.load_workbook(OUT / "recv_T31.xlsx")
    st = {n: wbf[n].sheet_state for n in wbf.sheetnames}
    check("T31", "請求書シートが非表示で、可視は空シートだけ",
          "hidden" in st.values() and "visible" in st.values(), str(st)); wbf.close()

    wbf = openpyxl.load_workbook(OUT / "recv_T32.xlsx")
    check("T32", "請求書のシートが 2 枚", len(wbf.sheetnames) == 2,
          str(wbf.sheetnames)); wbf.close()

    wb, ws = one("T33")
    check("T33", "請求額セルが #REF!", has(ws, "#REF!"), str(has(ws, "#REF!"))[:60]); wb.close()

    wb, ws = one("T35")
    # ★ 番地表を使わずに測る: ふつうの冊では合計の値が 2 箇所（合計行と上部の欄）に
    #   現れる。T35 は上部を空にしたので 1 箇所しか無いはず。
    def tops(fid):
        w, x = one(fid)
        vs = [v for _, _, v in nums(x)]
        m = max(vs)
        w.close()
        return vs.count(m)
    # ★ 対照は同じ骨（misoca16）の基礎冊。B27 は inv21 で、最大値が請求番号なので対照に
    #   ならなかった（最初それで空振りした）。
    check("T35", "合計の値がシート内に 1 箇所しか無い（上部の請求額欄が空）",
          tops("T35") == 1 and tops("B15") == 2,
          f"T35={tops('T35')} 対照B15(同じ骨)={tops('B15')}"); wb.close()

    wb, ws = one("T36")
    check("T36", "「請求元：」の右が担当者名",
          has(ws, "請求元：") and has(ws, "戸田 太郎")); wb.close()

    wb, ws = one("T37")
    check("T37", "明細の品目名に「合計金額」が入っている",
          any("合計金額調整費" == v.strip() for _, _, v in texts(ws))); wb.close()

    wb, ws = one("T38")
    check("T38", "通貨が USD と書いてある", has(ws, "USD")); wb.close()

    wb, ws = one("T40")
    items = ["用紙代", "運送費", "保守料", "部材費", "作業代", "設計費", "検査料",
             "梱包資材", "出張旅費", "外注工賃"]
    n = sum(1 for _, _, v in texts(ws) if v.strip() in items)
    check("T40", "明細が雛形の枠を使い切っている（13 行）", n == 13, f"{n} 行"); wb.close()

    # ── 混入群 ────────────────────────────────────────────────────
    wb, ws = one("X01")
    check("X01", "納品書である", "納品" in main_title(ws) or bool(has(ws, "納　品　書")),
          main_title(ws)); wb.close()
    wb, ws = one("X02")
    check("X02", "御中が 1 つも無く『様』が在る",
          not has(ws, "御中") and has(ws, "様"), main_title(ws)); wb.close()
    wb, ws = one("X03")
    check("X03", "『見積書 ESTIMATE』が残っている", has(ws, "見積書")); wb.close()
    for fid in ("X04", "X05"):
        wb, ws = one(fid)
        check(fid, "請求書の手がかりが無い（御中も請求金額も無い）",
              not has(ws, "御中") and not has(ws, "請求")); wb.close()

    # ── ★ 骨の癖が他の検体に相乗りしていないか（交絡の点検）──────────
    leak = []
    for f in sorted(OUT.glob("recv_*.xlsx")):
        fid = f.stem.replace("recv_", "")
        if fid == "X03":
            continue
        wb = openpyxl.load_workbook(f, data_only=True)
        for n in wb.sheetnames:
            if has(wb[n], "見積書"):
                leak.append(fid)
        wb.close()
    check("*", "『見積書』の文字が X03 以外に漏れていない", not leak, str(leak))

    ng = [x for x in RESULTS if not x[2]]
    for fid, name, ok, det in RESULTS:
        print(f"  {'OK ' if ok else '★NG'} {fid:4s} {name}" + (f"   [{det}]" if det else ""))
    print(f"\n罠の点検: {len(RESULTS) - len(ng)}/{len(RESULTS)} 合格")
    return 1 if ng else 0


if __name__ == "__main__":
    sys.exit(main())
