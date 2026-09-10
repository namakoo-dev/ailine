# -*- coding: utf-8 -*-
"""検体 v2 を作る（2026-09-11）。骨は実物の雛形・肉（値）だけ流し込む。

mk_received.py の作法をそのまま継ぐ:
  ★ 骨は実物の雛形をそのまま使う（結合・帯・式が全部残る）
  ★ 答えは流し込んだ側が持つ（値と番地の両方）。道具の出力から作らない
  ★ 実物の金額は式。LibreOffice に開き直させてキャッシュを入れる（recalc）
  ★ 結合セルは左上以外に書けない（MergedCell は read-only）→ アンカーに畳む

v2 で足したもの:
  ★ 受け入れ条件を機械に: 生成器は **書く前に各セルの値を予測** し、
    LibreOffice が計算した結果と 1 セルずつ突き合わせる。
    予測が外れた検体は採用しない（＝俺の雛形理解が間違っていたことが分かる）。
    ★ 予測は「検体を作る式」ではなく「雛形の式を読んで俺が立てた式」で、
      本体の抽出実装とは別口。読み戻しが恒真にならないのはこのため。

使い方:
    python mk_received_v2.py            # 全部
    python mk_received_v2.py T01 T17    # id を指定（煙試験）
"""
import json
import math
import shutil
import subprocess
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from skeletons_v2 import SKELETONS, SP            # noqa: E402
from specimens_v2 import SPECS, JISHA          # noqa: E402

OUT = HERE / "received_v2"
ANSWER = HERE / "答え_received_v2.json"
SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
PROFILE = (HERE / "lo-profile-adv").resolve().as_uri()

# 自社（受け取る側）── 全冊 同じ
JISHA_ZIP, JISHA_ADDR = "〒101-0021", "東京都千代田区外神田 1-1-1 ナギビル 5F"
ATTN_POOL = ["経理部　御中", "総務課　御中", "購買部　御中",
             "山田 太郎　様", "佐藤 花子　様", "経理部 鈴木　様"]

STEMS = ["あかね", "いろは", "うえだ", "エバラ", "大久保", "花菱", "菊池", "くまがい",
         "ケイアイ", "小坂", "佐々木", "しなの", "鈴村", "瀬戸内", "曽根田", "高梨",
         "千葉", "つばめ", "寺岡", "戸田", "永井", "西野", "沼田", "根岸", "野々村",
         "橋本", "檜山", "福住", "邊見", "堀之内", "槇野", "三隅", "武藤", "室伏",
         "茂木", "矢作", "湯浅", "横手", "吉良", "若槻", "六車", "和久井"]
TAILS = ["商事", "工業", "物産", "機械", "製作所", "運輸", "電機", "化成", "建材", "食品"]
FORMS = ["株式会社{}{}", "{}{}株式会社", "{}{}有限会社", "合同会社{}{}"]

ITEMS = ["用紙代", "運送費", "保守料", "部材費", "作業代", "設計費", "検査料",
         "梱包資材", "出張旅費", "外注工賃"]
UNITS = ["式", "個", "時間", "日", "ヶ月"]


def company(i):
    return FORMS[i % len(FORMS)].format(STEMS[i % len(STEMS)], TAILS[(i // 4) % len(TAILS)])


def xlround(x, d=0):
    """Excel の ROUND（0 から遠い方へ half up）。Python の round は銀行家丸め。"""
    m = 10 ** d
    return math.floor(abs(x) * m + 0.5) / m * (1 if x >= 0 else -1)


# ── 結合セル ───────────────────────────────────────────────────
def anchor_map(ws):
    inside = {}
    for rng in ws.merged_cells.ranges:
        a = (rng.min_row, rng.min_col)
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                if (rr, cc) != a:
                    inside[(rr, cc)] = a
    return inside


def put(ws, inside, addr, value):
    """★ 結合セルは左上にしか書けない。中を指されたらアンカーに畳む。"""
    col = "".join(ch for ch in addr if ch.isalpha())
    row = int("".join(ch for ch in addr if ch.isdigit()))
    key = (row, column_index_from_string(col))
    r, c = inside.get(key, key)
    ws.cell(row=r, column=c).value = value
    return f"{get_column_letter(c)}{r}"


def cell_at(ws, inside, addr):
    col = "".join(ch for ch in addr if ch.isalpha())
    row = int("".join(ch for ch in addr if ch.isdigit()))
    key = (row, column_index_from_string(col))
    r, c = inside.get(key, key)
    return f"{get_column_letter(c)}{r}"


# ── 明細 ───────────────────────────────────────────────────────
def gen_lines(idx, n):
    out = []
    for k in range(n):
        out.append(dict(name=ITEMS[(idx + k) % len(ITEMS)],
                        unitname=UNITS[(idx + k) % len(UNITS)],
                        qty=1 + ((idx * 5 + k) % 5),
                        price=1000 * (1 + ((idx * 7 + k * 3) % 9))))
    return out


def line_amount(sk, ln):
    """雛形の金額式を読んで、その行の金額セルがどうなるかを予測する。"""
    if ln is None:
        return None
    price, qty = ln.get("price"), ln.get("qty")
    style = sk["tax_style"]
    if style == "single_rate":          # misoca: =IF(単価="","",数量*単価)
        if price in (None, ""):
            return None
        return (qty or 0) * price
    if style == "mark_column":          # 建設: =IF(SUM(単価*数量),SUM(単価*数量),"")
        v = (price or 0) * (qty or 0)
        return v if v != 0 else None
    if style == "rate_column":          # inv21: =IF(取引日="","",単価*数量)
        return (price or 0) * (qty or 0)          # 取引日は必ず書く
    if style == "fixed10":              # spread: =数量*単価（IF なし）
        return (qty or 0) * (price or 0)
    raise AssertionError(style)


def predict_sums(sk, lines, rate):
    """帯（小計・消費税・合計・請求額）を予測する。★ 雛形の式から手で立てた式。"""
    style, s = sk["tax_style"], sk["sums"]
    amts = [line_amount(sk, ln) for ln in lines]
    if style == "single_rate":
        sub = sum(a for a in amts if a is not None)
        tax = sub * rate
        return {s["小計"]: sub, s["消費税"]: tax, s["合計"]: sub + tax,
                s["請求額"]: sub + tax}
    if style == "mark_column":
        b10 = sum(a for ln, a in zip(lines, amts)
                  if a is not None and (ln or {}).get("mark") in (None, ""))
        b8 = sum(a for ln, a in zip(lines, amts)
                 if a is not None and (ln or {}).get("mark") == "※")
        t10, t8 = xlround(b10 * 0.1, 0), xlround(b8 * 0.08, 0)
        return {s["対象10"]: b10, s["税10"]: t10, s["対象8"]: b8, s["税8"]: t8,
                s["小計"]: b10 + b8, s["消費税"]: t10 + t8,
                s["請求額"]: b10 + b8 + t10 + t8}
    if style == "rate_column":
        b10 = sum(a for ln, a in zip(lines, amts)
                  if a is not None and (ln or {}).get("rate", 0.1) == 0.1)
        b8 = sum(a for ln, a in zip(lines, amts)
                 if a is not None and (ln or {}).get("rate", 0.1) == 0.08)
        t10, t8 = xlround(b10 * 0.1, 1), xlround(b8 * 0.08, 1)
        return {s["対象10"]: b10, s["税10"]: t10, s["対象8"]: b8, s["税8"]: t8,
                s["小計"]: b10 + b8, s["消費税"]: t10 + t8,
                s["合計"]: b10 + b8 + t10 + t8, s["請求額"]: b10 + b8 + t10 + t8}
    if style == "fixed10":
        sub = sum(a or 0 for a in amts)
        tax = sub * 0.1
        return {s["小計"]: sub, s["消費税"]: tax, s["合計"]: sub + tax,
                s["請求額"]: sub + tax}
    raise AssertionError(style)


# ── 1 冊作る ───────────────────────────────────────────────────
def build_one(spec, idx):
    sid = spec["id"]
    dst = OUT / f"recv_{sid}.xlsx"

    # 官公庁の実表は無改造でコピーするだけ（請求書ではない冊）
    if spec["骨"] in ("gov1", "gov2"):
        src = SP / "gov" / ("jinsui_1.xlsx" if spec["骨"] == "gov1" else "kakei_fies_t2.xlsx")
        shutil.copy(src, dst)
        return dict(file=dst.name, id=sid, 群=spec["群"], 骨=src.name, シート=None,
                    狙い=spec["狙い"], 出所=spec["出所"], 落とし方=spec["落とし方"],
                    請求元=None, 宛先=None, 明細=[], 予測セル={},
                    期待={k: v for k, v in spec["期待"].items()}), {}

    sk = SKELETONS[spec["骨"]]
    shutil.copy(sk["src"], dst)
    wb = openpyxl.load_workbook(dst)          # ★ 式を残す（data_only を付けない）
    ws = wb[sk["sheet"]]
    inside = anchor_map(ws)
    d, sums = sk["detail"], sk["sums"]

    # ---- 雛形の見本を消す（先に消さないと答えと食い違う）----
    for row in range(d["first"], d["last"] + 1):
        for col in d["clear"]:
            put(ws, inside, f"{col}{row}", None)
    for addr in sk.get("note") or []:
        put(ws, inside, addr, None)
    if spec["骨"] == "misoca13" and not spec.get("keep_estimate_label"):
        # ★ この骨は B1 に『見積書 ESTIMATE』を雛形のまま持っている。
        #   X03 以外では消す ── 消さないと「見積書」の効果が骨に相乗りして交絡する。
        put(ws, inside, "B1", None)

    # ---- 請求元（発行者）----
    # ★ "AUTO"=社名を割り当てる / "KEEP"=雛形の伏せ字をそのまま残す / None=セルを空にする
    issuer = spec.get("issuer_name", "AUTO")
    iss = sk["issuer"]
    issuer_at = None
    if issuer == "KEEP":
        issuer_at = cell_at(ws, inside, iss["name"])
        issuer = ws[issuer_at].value
    else:
        if issuer == "AUTO":
            issuer = company(idx)
        if "name" in iss:
            issuer_at = put(ws, inside, iss["name"], issuer)
    issuer_plain = spec.get("issuer_plain", issuer)
    if iss.get("tel"):
        put(ws, inside, iss["tel"], "TEL：03-3000-%04d" % (idx % 10000))
    if iss.get("regno"):
        put(ws, inside, iss["regno"], "T%013d" % (1000000000000 + idx))

    # ---- 宛先（自社）★ 実物の置き場に合わせる（旧検体の欠陥 (2) の直し）----
    to = sk["to"]
    to_at = None
    if to.get("zip"):
        put(ws, inside, to["zip"], JISHA_ZIP)
    if to.get("addr"):
        put(ws, inside, to["addr"], JISHA_ADDR)
    if spec["骨"] in ("misoca256", "misoca12", "misoca13", "misoca16", "inv21", "irai"):
        # 社名セルと「御中/様」セルが分かれている骨
        to_at = put(ws, inside, to["name"], JISHA)
        put(ws, inside, to["attn"], spec.get("to_attn", ATTN_POOL[idx % len(ATTN_POOL)]))
        to_value = JISHA
    else:
        # 社名と「御中」が同じセルの骨（実物がそう）
        to_at = put(ws, inside, to["name"], JISHA + "　御中")
        to_value = JISHA + "　御中"

    # ---- 明細 ----
    if spec.get("lines") is not None:
        lines = [dict(ln) if ln else None for ln in spec["lines"]]
    else:
        lines = gen_lines(idx, spec.get("n", 0))
    if len(lines) > (d["last"] - d["first"] + 1):
        raise AssertionError(f"{sid}: 明細が雛形の枠を超える")
    rate = spec.get("rate", 0.1)
    for k, ln in enumerate(lines):
        row = d["first"] + k
        if ln is None:
            continue
        ln.setdefault("name", ITEMS[(idx + k) % len(ITEMS)])
        ln.setdefault("unitname", UNITS[(idx + k) % len(UNITS)])
        if d.get("date"):
            put(ws, inside, f"{d['date']}{row}", f"2026-08-{(idx % 28) + 1:02d}")
        put(ws, inside, f"{d['name']}{row}", ln["name"])
        if d.get("unitname"):
            put(ws, inside, f"{d['unitname']}{row}", ln["unitname"])
        if ln.get("qty") is not None:
            put(ws, inside, f"{d['qty']}{row}", ln["qty"])
        if ln.get("price") is not None:
            put(ws, inside, f"{d['price']}{row}", ln["price"])
        if d.get("mark") and ln.get("mark"):
            put(ws, inside, f"{d['mark']}{row}", ln["mark"])
        if d.get("rate"):
            put(ws, inside, f"{d['rate']}{row}", ln.get("rate", 0.1))
    if sk.get("rate_cell"):
        put(ws, inside, sk["rate_cell"], rate)

    # ---- 予測（★ 書く前に立てる）----
    pred = {}
    for k, ln in enumerate(lines):
        if ln is None:
            continue
        a = line_amount(sk, ln)
        if a is not None:
            pred[cell_at(ws, inside, f"{d['amount']}{d['first'] + k}")] = a
    pred.update({cell_at(ws, inside, k): v
                 for k, v in predict_sums(sk, lines, rate).items()})
    if issuer_at:
        pred[issuer_at] = issuer
    pred[to_at] = to_value

    # ---- 個別の細工 ----
    if spec.get("kurikoshi"):                       # spread5 繰越
        prev, paid = spec["kurikoshi"]
        put(ws, inside, sums["前回"], prev)
        put(ws, inside, sums["入金"], paid)
        pred[cell_at(ws, inside, sums["前回"])] = prev
        pred[cell_at(ws, inside, sums["入金"])] = paid
        pred[cell_at(ws, inside, sums["繰越"])] = prev - paid
        pred[cell_at(ws, inside, sums["今回請求額"])] = (
            prev - paid + pred[cell_at(ws, inside, sums["合計"])])
    if spec.get("extra_row"):                       # 集計範囲の外に追記
        e = spec["extra_row"]
        r = e["row"]
        put(ws, inside, f"{d['date']}{r}", e["date"])
        put(ws, inside, f"{d['name']}{r}", e["name"])
        put(ws, inside, f"{d['price']}{r}", e["price"])
        put(ws, inside, f"{d['qty']}{r}", e["qty"])
        put(ws, inside, f"{d['amount']}{r}", e["amount"])   # ★ 式を潰して手入力
        pred[cell_at(ws, inside, f"{d['amount']}{r}")] = e["amount"]
    if spec.get("const_delta"):
        cd = spec["const_delta"]
        base = pred[cell_at(ws, inside, sums[cd["cell"]])] + cd["delta"]
        put(ws, inside, sums[cd["cell"]], base)
        pred[cell_at(ws, inside, sums[cd["cell"]])] = base
        for other in spec.get("also_const", []):
            put(ws, inside, sums[other], base)
            pred[cell_at(ws, inside, sums[other])] = base
    for addr, v in (spec.get("const") or {}).items():
        put(ws, inside, addr, v)
        pred[cell_at(ws, inside, addr)] = v
    for addr in spec.get("clear_cells") or []:
        put(ws, inside, addr, None)
        pred.pop(cell_at(ws, inside, addr), None)
        pred[cell_at(ws, inside, addr)] = None
    if spec.get("ref_error"):
        # ★ 測って分かったこと: "=#REF!" と式で書いても、エラー値そのもの
        #   （data_type='e'）を置いても、LibreOffice は **#NAME?** に変えてしまう。
        #   本物の #REF! を残すには、実際に #REF! を返す式が要る。
        addr, formula, want = spec["ref_error"]
        put(ws, inside, addr, formula)
        pred[cell_at(ws, inside, addr)] = want
    if spec.get("note_recap"):
        total = pred[cell_at(ws, inside, sums["合計"])]
        txt = f"合計金額 {int(total):,} 円（税込）をご請求申し上げます"
        put(ws, inside, sk["note"][1], txt)
        pred[cell_at(ws, inside, sk["note"][1])] = txt
    for op in spec.get("post") or []:
        if op[0] == "set":
            put(ws, inside, op[1], op[2])
            pred[cell_at(ws, inside, op[1])] = op[2]
        elif op[0] == "clear":
            put(ws, inside, op[1], None)
            pred[cell_at(ws, inside, op[1])] = None
    for r in spec.get("hide_rows") or []:
        ws.row_dimensions[r].hidden = True
    if spec.get("dup_sheet"):
        cp = wb.copy_worksheet(ws)
        cp.title = spec["dup_sheet"]
    if spec.get("hide_sheet"):
        wb.create_sheet("Sheet1")           # 可視シートが 1 枚も無い bookは開けない
        ws.sheet_state = "hidden"

    wb.save(dst)
    wb.close()

    ans = dict(
        file=dst.name, id=sid, 群=spec["群"], 骨=sk["src"].name, シート=sk["sheet"],
        帳票種別=sk["kind"], 狙い=spec["狙い"], 出所=spec["出所"], 落とし方=spec["落とし方"],
        請求元=dict(値=issuer_plain, セルの値=issuer, 番地=issuer_at),
        宛先=dict(値=JISHA, セルの値=to_value, 番地=to_at),
        明細=[None if ln is None else
              dict(品目=ln["name"], 数量=ln.get("qty"), 単価=ln.get("price"),
                   金額=line_amount(sk, ln), 税率=ln.get("rate"), 軽減=ln.get("mark"))
              for ln in lines],
        帯の番地=sums, 税率セル=sk.get("rate_cell"),
        予測セル={k: v for k, v in pred.items()},
        雛形の癖=sk.get("quirk"),
        期待=spec["期待"], 同値=spec.get("同値"),
    )
    return ans, pred


# ── LibreOffice に開き直させる（式のキャッシュを入れる）──────────────
def recalc(files):
    work = HERE / "_recalc"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    r = subprocess.run(
        [str(SOFFICE), "--headless", "--norestore", "--nologo",
         f"-env:UserInstallation={PROFILE}",
         "--convert-to", "xlsx", "--outdir", str(work)] + [str(f) for f in files],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=1800)
    made = list(work.glob("*.xlsx"))
    for f in made:
        shutil.copy(f, OUT / f.name)
    shutil.rmtree(work)
    return len(made), (r.stdout + r.stderr)[-400:]


# ── 受け入れ条件: 読み戻して 1 セルずつ突き合わせる ───────────────────
def same(pred, got):
    if pred is None:
        return got in (None, "")
    if isinstance(pred, (int, float)) and isinstance(got, (int, float)):
        return abs(pred - got) < 1e-6
    return str(pred).replace("\r\n", "\n") == str(got).replace("\r\n", "\n")


def verify(answers):
    bad = []
    for a in answers:
        if not a["予測セル"]:
            continue
        wb = openpyxl.load_workbook(OUT / a["file"], data_only=True)
        ws = wb[a["シート"]]
        for addr, want in a["予測セル"].items():
            got = ws[addr].value
            if not same(want, got):
                bad.append((a["id"], addr, want, got))
        wb.close()
    return bad


def main():
    only = set(sys.argv[1:])
    specs = [s for s in SPECS if not only or s["id"] in only]
    if OUT.exists() and not only:
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True, exist_ok=True)

    answers = []
    for i, spec in enumerate(SPECS):
        if only and spec["id"] not in only:
            continue
        a, _ = build_one(spec, i)
        answers.append(a)
    print(f"作った: {len(answers)} 冊 → {OUT}")

    n, tail = recalc(sorted(OUT / a["file"] for a in answers))
    print(f"LibreOffice で開き直した: {n}/{len(answers)}")
    if n != len(answers):
        print("  ", tail)

    bad = verify(answers)
    ok_ids = {a["id"] for a in answers} - {b[0] for b in bad}
    print(f"\n受け入れ検査: 一致 {len(ok_ids)} 冊 / 食い違い {len(set(b[0] for b in bad))} 冊")
    for sid, addr, want, got in bad:
        print(f"  ✗ {sid} {addr}: 予測={want!r} 実測={got!r}")

    # 期待の "AUTO" を答えで埋める
    for a in answers:
        if a["予測セル"] and a.get("帯の番地"):
            addr = a["帯の番地"].get("請求額")
            a["請求額"] = dict(値=a["予測セル"].get(addr), 番地=addr)
        for key, exp in (a["期待"] or {}).items():
            if exp.get("値") == "AUTO":
                exp["値"] = {"請求元": a["請求元"] and a["請求元"]["値"],
                             "請求額": a.get("請求額") and a["請求額"]["値"],
                             "宛先": a["宛先"] and a["宛先"]["値"]}.get(key)
        a["採用"] = a["id"] in ok_ids

    ANSWER.write_text(json.dumps(answers, ensure_ascii=False, indent=2),
                      encoding="utf-8", newline="\n")
    print(f"答え → {ANSWER}")
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
