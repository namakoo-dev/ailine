# -*- coding: utf-8 -*-
"""検体「1 冊の一覧表を担当者ごとに別の冊へ分ける」を作る（2026-09-12）。

specimens_split.py の宣言的なデータ（列・行・仕込み）だけを見て、この生成器が
1) xlsx を書く 2) 書いた値をそのまま「予測セル」として控える 3) 担当者ごとの
グループ・空欄・複数担当・分けない行を機械的に拾って答え JSON に組む。

★ 式は一切使わない（合計・小計は数値をそのまま書く）。LibreOffice は起動しない
  （依頼書の制約: 「LO を使ってよい」と言われるまで起動禁止・別走行と衝突するため）。
  したがって受け入れ検査は openpyxl の読み戻しだけで完結する（data_only は不要 ──
  そもそもキャッシュを要る式が無い）。

使い方:
    python mk_split.py            # 6 冊を作り、読み戻して答えと突き合わせる
"""
import json
import re
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl                                   # noqa: E402
from openpyxl.utils import get_column_letter        # noqa: E402

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from specimens_split import BOOKS                   # noqa: E402

OUT = HERE / "split_books"
ANSWER = HERE / "答え_split.json"

MULTI_RE = re.compile(r"[/・]")   # 複数担当の目印（スラッシュ・中点）


def addr(col_idx, row):
    return f"{get_column_letter(col_idx + 1)}{row}"


# ── 1 冊を書く。戻り値: (Workbook, meta) ─────────────────────────
def build_book(spec):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]

    columns = spec["columns"]
    n_title = len(spec["title"])
    header_row = n_title + 1
    amount_idx = columns.index("金額")

    pred = {}

    for i, text in enumerate(spec["title"], start=1):
        c = addr(0, i)
        ws[c] = text
        pred[c] = text

    for ci, name in enumerate(columns):
        c = addr(ci, header_row)
        ws[c] = name
        pred[c] = name

    assignee = spec["assignee"]
    dual = isinstance(assignee, tuple)

    row = header_row + 1
    groups = {}            # 表記そのまま(未正規化) -> {"行": [...], "金額": 合計}
    blanks = []             # 担当者が空欄の行番号
    multi_rows = []          # (行番号, 表記) ── 複数担当
    dividing_rows = []       # 分けない行（空行・小計・合計・備考）
    running_total = 0        # 明細の数値の和（合計行の値そのものではない）
    since_subtotal = 0
    total_declared = None
    total_real = None

    for rec in spec["rows"]:
        kind = rec["kind"]

        if kind == "data":
            for ci, name in enumerate(columns):
                val = rec.get(name)
                c = addr(ci, row)
                ws[c] = val
                pred[c] = val
            amt = rec.get("金額")
            numeric_amt = amt if isinstance(amt, (int, float)) else None
            if numeric_amt is not None:
                running_total += numeric_amt
                since_subtotal += numeric_amt
            if not dual:
                name = rec.get(assignee)
                if name is None:
                    blanks.append(row)
                elif MULTI_RE.search(name):
                    multi_rows.append((row, name))
                else:
                    g = groups.setdefault(name, {"行": [], "金額": 0})
                    g["行"].append(row)
                    if numeric_amt is not None:
                        g["金額"] += numeric_amt
            row += 1

        elif kind == "blank":
            dividing_rows.append(row)
            row += 1

        elif kind == "subtotal":
            c_label = addr(0, row)
            ws[c_label] = rec["label"]
            pred[c_label] = rec["label"]
            c_amt = addr(amount_idx, row)
            ws[c_amt] = since_subtotal
            pred[c_amt] = since_subtotal
            dividing_rows.append(row)
            since_subtotal = 0
            row += 1

        elif kind == "total":
            c_label = addr(0, row)
            ws[c_label] = rec["label"]
            pred[c_label] = rec["label"]
            c_amt = addr(amount_idx, row)
            total_val = running_total + spec.get("合計ずれ", 0)
            ws[c_amt] = total_val
            pred[c_amt] = total_val
            dividing_rows.append(row)
            total_declared, total_real = total_val, running_total
            row += 1

        elif kind == "note":
            c = addr(0, row)
            ws[c] = rec["text"]
            pred[c] = rec["text"]
            dividing_rows.append(row)
            row += 1

        else:
            raise AssertionError(f"unknown row kind: {kind!r}")

    last_row = row - 1

    # ---- 表記ゆれの自己点検（宣言した表記が実際にどこかの行にあるか）----
    if spec.get("表記ゆれ") and not dual:
        assignee_values = {r.get(assignee) for r in spec["rows"] if r.get("kind") == "data"}
        for a, b in spec["表記ゆれ"]:
            if a not in assignee_values or b not in assignee_values:
                raise AssertionError(f"{spec['id']}: 表記ゆれ宣言 {(a, b)} が行データに見当たらない")

    # ---- 非表示シート ----
    hidden_info = None
    hidden_pred = {}
    if spec.get("非表示シート"):
        h = spec["非表示シート"]
        ws2 = wb.create_sheet(h["シート名"])
        for ci, name in enumerate(h["columns"]):
            c = ws2.cell(row=1, column=ci + 1, value=name)
            hidden_pred[c.coordinate] = name
        for ri, r in enumerate(h["rows"], start=2):
            for ci, name in enumerate(h["columns"]):
                c = ws2.cell(row=ri, column=ci + 1, value=r.get(name))
                hidden_pred[c.coordinate] = r.get(name)
        ws2.sheet_state = "hidden"
        hidden_info = dict(シート名=h["シート名"], 説明=h["説明"])

    meta = dict(
        header_row=header_row, columns=columns, assignee=assignee, dual=dual,
        groups=groups, blanks=blanks, multi_rows=multi_rows,
        dividing_rows=dividing_rows, last_row=last_row, pred=pred,
        hidden_info=hidden_info, hidden_pred=hidden_pred,
        total_declared=total_declared, total_real=total_real,
    )
    return wb, meta


# ── meta から答え JSON の 1 冊分を組む ───────────────────────────
def make_answer(spec, meta):
    columns, assignee, dual, header_row = (
        meta["columns"], meta["assignee"], meta["dual"], meta["header_row"])

    if dual:
        担当者の列 = None
    else:
        idx = columns.index(assignee)
        担当者の列 = dict(見出し=assignee, 番地=addr(idx, header_row))

    期待 = dict(
        担当者ごと=(None if dual else
                    {k: dict(行=v["行"], 金額=v["金額"]) for k, v in meta["groups"].items()}),
        空欄の行=(None if dual else meta["blanks"]),
        表記ゆれ=([list(p) for p in spec["表記ゆれ"]] if spec.get("表記ゆれ") else None),
        分けない行=meta["dividing_rows"],
        分けられない=(dict(理由に含む語="担当") if dual else None),
    )

    迷う = [
        dict(行=r, 理由=f"複数担当 {name} ── どちらの冊にも入れる／どちらにも入れず"
                        "名指し、どちらでも正解")
        for r, name in meta["multi_rows"]
    ]
    if spec.get("迷う_book"):
        迷う.append(dict(行=None, 理由=spec["迷う_book"]))

    合計行の食い違い = None
    if spec.get("合計ずれ", 0) and meta["total_declared"] is not None:
        合計行の食い違い = (
            f"合計行の金額 {meta['total_declared']} は明細の和 {meta['total_real']} と"
            f"違う（+{spec['合計ずれ']}）")

    return dict(
        id=spec["id"], file=spec["file"], シート=spec["sheet"],
        担当者の列=担当者の列, 見出しの行=header_row,
        期待=期待,
        怪しくない=bool(spec.get("陰性対照")),
        迷う=迷う,
        仕込み=spec.get("仕込み", []),
        非表示シート=meta["hidden_info"],
        合計行の食い違い=合計行の食い違い,
        予測セル=meta["pred"],
        非表示シートの予測セル=(meta["hidden_pred"] or None),
    )


# ── 受け入れ条件: 生成物を読み戻して予測セルと突き合わせる ──────────────
def verify_one(path, ans):
    bad = []
    wb = openpyxl.load_workbook(path)
    ws = wb[ans["シート"]]
    for c, want in ans["予測セル"].items():
        got = ws[c].value
        if got != want:
            bad.append((ans["id"], ans["シート"], c, want, got))
    if ans.get("非表示シートの予測セル"):
        hs = ans["非表示シート"]["シート名"]
        ws2 = wb[hs]
        if ws2.sheet_state != "hidden":
            bad.append((ans["id"], hs, "sheet_state", "hidden", ws2.sheet_state))
        for c, want in ans["非表示シートの予測セル"].items():
            got = ws2[c].value
            if got != want:
                bad.append((ans["id"], hs, c, want, got))
    wb.close()
    return bad


def main():
    if OUT.exists():
        for f in OUT.glob("*.xlsx"):
            f.unlink()
    OUT.mkdir(parents=True, exist_ok=True)

    answers = []
    all_bad = []
    for spec in BOOKS:
        wb, meta = build_book(spec)
        dst = OUT / spec["file"]
        wb.save(dst)
        wb.close()
        ans = make_answer(spec, meta)
        bad = verify_one(dst, ans)
        ans["採用"] = not bad
        answers.append(ans)
        all_bad += bad

    print(f"作った: {len(answers)} 冊 → {OUT}")
    ok = sum(1 for a in answers if a["採用"])
    print(f"受け入れ検査: 一致 {ok} 冊 / 食い違い {len(answers) - ok} 冊")
    for sid, sheet, c, want, got in all_bad:
        print(f"  x {sid} [{sheet}] {c}: 予測={want!r} 実測={got!r}")

    ANSWER.write_text(
        json.dumps(answers, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8", newline="\n")
    print(f"答え → {ANSWER}")
    return 0 if not all_bad else 1


if __name__ == "__main__":
    sys.exit(main())
