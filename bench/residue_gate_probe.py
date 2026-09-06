# -*- coding: utf-8 -*-
"""残差を「✓ を出す前の関所」にしたら、**正しい回でどれだけ鳴るか**を測る。

★★ なぜ測るか（2026-09-06）: 同じ日に 3 件の穴を直したが、どれも同じ形だった ──
  **依頼に在るものが宣言から落ち、宣言と実体は一致するので ✓ が出る**
  （2 条件の片落ち／見出しの語を行と誤読／式の消失）。
  事後条件は「宣言 vs 実体」しか見ないので、**依頼が痩せた分は原理的に見えない**。

  ★ 材料は在る: 残差（`residue.find_unconsumed_words`）が「依頼に在って、解決済みの
    引数に現れなかった語」を返す。今日それを 2 回使った（2 組目の条件・属性の登録）。

★★ だが**そのままでは狼少年**だった（履歴で 49% が鳴る）。絞りを入れて 5 形で分離した:

      鳴らす条件 = 残差のうち、**実表の列名**であるもの

    事故  2 条件の片落ち  → ['部門','営業']
    事故  誤読（行追加）  → ['金額']
    正常  1 条件 / 並べ替え / 罫線 → []

★ この道具が測るのは**誤爆率だけ**（正しい回で鳴ってはいけない）。検体は
  `bench/basic_ops_matrix.py` の凍結検体をそのまま使う ── あれは**正しい回ばかり**が
  入っているので、鳴ったらそれは誤爆だ。

★ **製品は 1 バイトも変えない**（計測のために製品を触ると、測っているものが変わる）。
  解決（`verify_dsl_args`）までを直接呼ぶ ── LibreOffice は起こさないので速い。

使い方:
    python bench/residue_gate_probe.py            # 全部
    python bench/residue_gate_probe.py --table 売上
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "bench"))

import ailine  # noqa: E402
from ailine_core import residue  # noqa: E402

import basic_ops_matrix as M  # noqa: E402


def _headers_of(book: Path, sheet: str | None, *, with_values: bool = False) -> set:
    """関所が鳴る語の集合。既定は**列名だけ**。

    ★ `with_values=True` は「列名 ∪ 実表に在るセルの値」に広げた版（2026-09-06 に測る）。
      感度の測定で分かった穴 ──「落ちたのが**値だけ**の時、関所は列名しか見ないので
      原理的に見えない」── を塞げるかを見るため。★ 広げれば誤爆が増えるはずなので、
      塞ぐ価値と釣り合うかを**先に測る**（実装してから測らない）。
    """
    try:
        bm = ailine.build_book_meta(book)
        sh = sheet or (bm.get("sheets") or [None])[0]
        out = {str(h) for h in ((bm.get("headers") or {}).get(sh) or []) if h}
    except Exception:
        return set()
    if with_values:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(book, data_only=True)
            ws = wb[sh] if sh in wb.sheetnames else wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 200), values_only=True):
                for v in row:
                    if isinstance(v, str) and len(v) >= 2:
                        out.add(v)
            wb.close()
        except Exception:
            pass
    return out


def probe_one(book: Path, task: str, heads: set, *,
              want_op: str | None = None) -> tuple:
    """(鳴ったか, 鳴った語, 理由) を返す。理由は下の 5 種のどれか。

    ★ 翻訳（LLM）→ 解決（verify_dsl_args）までを本物どおりに通す。

    ★★ なぜ理由を割るか（2026-09-06・Namakoo の指摘「残り 90 はどうする？」）:
      初版はここで 4 つの違う事情を 1 つの False に潰していた。246 件のうち 156 件しか
      届かず、**落ちた 90 件を「測らなくてよかったもの」として黙って捨てていた**。
      ★ 除いて正しいのは `refused`（機械が断った回＝✓ の直前ではない）**だけ**で、
        残り 3 種は**翻訳の取りこぼし**── 本番には二段目が在るので実際は ✓ まで行く。
      ★ しかも偏りが悪い側だ: 一発で通らないのは長く複雑な依頼で、それは残差が
        いちばん鳴りやすい側。捨てたままだと誤爆率は**楽な側に偏った分母**で出る。

    ★ `want_op` を渡すと、一発目が DSL の形に届かなかった時に**本番と同じ二段目**
      （`translate_task_fixed_op` ── op を機械が固定して args だけ埋めさせる）を当てる。
      ★ 歪みの自覚: op を教えるので **op の選択は本番より易しい**。ただし関所が見るのは
        「依頼の語 vs 引数」であって op の当て方ではない ── むしろ *op が正しい時の
        args の痩せ* こそ、今日直した 3 件がまさにその形だった。
    """
    bm = ailine.build_book_meta(book)
    op, args, why = "", {}, ""
    try:
        tr = ailine.translate_task(ailine.DEFAULT_MODEL, task, bm, temperature=0.1)
    except SystemExit:
        tr, why = None, "translate_error"
    if tr is not None:
        plan = (tr or {}).get("plan") or ([tr] if (tr or {}).get("op") else [])
        if not plan or not isinstance(plan[0], dict):
            why = "no_plan"
        else:
            op = str(plan[0].get("op") or "")
            args = plan[0].get("args") or {}
            if op not in ailine.OP_SCHEMA:
                why = "not_dsl:" + (op or "?")
    if why and want_op in ailine.OP_SCHEMA:
        second = ailine.translate_task_fixed_op(ailine.DEFAULT_MODEL, want_op, task, bm)
        if second:
            op, args, why = want_op, second.get("args") or {}, ""
        else:
            return (False, [], why + "+二段目も落ちた")
    if why:
        return (False, [], why)
    ok, resolved, _inf, _err = ailine.verify_dsl_args(
        op, args, bm, task=task, vocab=ailine.load_vocab())
    if not ok:
        return (False, [], "refused")      # ★ 断った回は「✓ の直前」ではない
    # ★★ 測定器の欠陥を 1 つ直した（2026-09-06・17 件の誤爆が全部これだった）:
    #   `find_unconsumed_words` は **文字列の値しか消費しない**（実装が isinstance(v, str)）。
    #   `operands = ['数量','単価']` のような**リストの中身**が消費されず、残差に出ていた。
    #   ★ 設計を捨てる前に測定器を疑う ── 平らにしてから測る。
    #   ★ もう 1 つ: 対象は `col:数量` / `cell:2,3` のように**接頭辞つき**で入る
    #     （CENTER_ALIGN / BOLD / FILL_COLOR）。接頭辞を落とさないと列名が消費されない。
    def _bare(s: str) -> str:
        for pre in ("col:", "row:", "cell:", "sheet:"):
            if s.startswith(pre):
                return s[len(pre):]
        return s

    flat = {}
    for k, v in resolved.items():
        if isinstance(v, str):
            flat[k] = v
            if _bare(v) != v:
                flat[k + "_bare"] = _bare(v)
        elif isinstance(v, (list, tuple, set)):
            for i, one in enumerate(v):
                if isinstance(one, str):
                    flat[f"{k}{i}"] = one
        elif isinstance(v, dict):
            for kk, vv in v.items():
                if isinstance(kk, str):
                    flat[f"{k}k{kk}"] = kk
                if isinstance(vv, str):
                    flat[f"{k}v{kk}"] = vv
    left = residue.find_unconsumed_words(task, flat, ailine._op_match_pool(op))
    fired = [w for w in left if w in heads]
    return (bool(fired), fired, "ok")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--table", default=None)
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--values", action="store_true",
                    help="★ 関所を「列名 ∪ 実表の値」に広げた版で測る")
    ap.add_argument("--no-second", action="store_true",
                    help="二段目（op 固定の再翻訳）を当てない ── 初版と同じ痩せた分母")
    a = ap.parse_args()

    keys = [a.table] if a.table else list(M.TABLES)
    d = Path(tempfile.mkdtemp())
    total = resolved_n = fired_n = 0
    fired_rows = []
    misses = Counter()
    for key in keys:
        book = M._build(key, d / f"{key}.xlsx")
        heads = _headers_of(book, M.TABLES[key].get("sheet"), with_values=a.values)
        for want_op, task, _check in M._cases_for(key):
            if a.limit and total >= a.limit:
                break
            total += 1
            fired, words, why = probe_one(
                book, task, heads, want_op=None if a.no_second else want_op)
            if why == "ok":
                resolved_n += 1
            else:
                misses[why] += 1
            if fired:
                fired_n += 1
                fired_rows.append((key, task, words))
    print()
    print("★ 関所の幅:", "列名 ∪ 実表の値（広げた版）" if a.values else "列名だけ（既定）")
    print(f"検体 {total} 件 ／ 解決まで通った {resolved_n} 件")
    print(f"★ 関所が鳴った: {fired_n} 件"
          f"（解決した回の {fired_n * 100 // max(1, resolved_n)}%）── **全部が誤爆**")
    for key, task, words in fired_rows[:20]:
        print(f"    [{key}] {words}  ← {task[:44]}")
    if len(fired_rows) > 20:
        print(f"    …ほか {len(fired_rows) - 20} 件")
    print()
    if misses:
        print()
        print("★ 届かなかった回の内訳（★ `refused` だけが「除いて正しい」）:")
        for why, n in misses.most_common():
            print(f"    {n:3d}  {why}")
    print()
    print("★ この検体は**正しい回ばかり**なので、鳴った数がそのまま誤爆の数。")
    print("★ 0 に近くなければ、関所として置けない（狼少年は本物の赤を隠す）。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
