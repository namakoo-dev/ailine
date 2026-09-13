# -*- coding: utf-8 -*-
"""検体「経費の勘定科目を先例から引く」を作る（2026-09-13・設計レビュー後に作り直し）。

specimens_accounts.py の宣言（過去の事実 (鍵,値,科目) と今回のシナリオ）だけを見て、

  1) 4 社の形（mf/csv・弥生/csv・xlsx(MF形式)・freee/csv）で past と today の実ファイルを書く
  2) 採点器の契約どおりの区分の導出（鍵は 借方取引先／貸方取引先／借方補助科目／摘要・
     鍵ごと＝出所ごとに数える）を実装した `grade_row()` で、今回の各対象行の区分・科目・
     根拠の語を**機械的に**計算する（specimens には区分や科目を一切書いていない）
  3) 答え_accounts.json を組む（rules は廃止・先例に在る科目／触らない行／複合仕訳／
     貸方だけ取引先／文字コード を必須項目として書く）

★ LibreOffice は起動しない（xlsx は openpyxl で直接書く。式は使わない）。
★ CSV/JSON はすべて write_bytes で書く（write_text は Windows で改行を LF→CRLF に
  巻き込む事故があるため。改行は明示的に \\n のみ）。

使い方:
    python mk_accounts.py            # 10 ケースを作り、答え_accounts.json を書く
"""
from __future__ import annotations

import itertools
import json
import sys
from datetime import date, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from specimens_accounts import CASES  # noqa: E402

OUT = HERE / "accounts_books"
ANSWER = HERE / "答え_accounts.json"

# ── 4 本の鍵（採点器 docstring のとおり）───────────────────────────
CANON_FIELDS = ("借方取引先", "貸方取引先", "借方補助科目", "摘要")

# ── 3 社の列（設計 docs/DESIGN-20260913 §1）─────────────────────────
# 弥生は公式の列をそのまま（25 列・見出し行なし）。MF/freee は「抜粋」からの
# 復元なので、抜粋に無い列は埋め草（filler）── 実物と一致を主張しない。
HEADERS_MF = [
    "取引No", "取引日", "借方勘定科目", "借方補助科目", "借方部門", "借方取引先",
    "借方税区分", "借方インボイス", "借方金額(円)", "借方税額", "貸方勘定科目",
    "貸方補助科目", "貸方部門", "貸方取引先", "貸方税区分", "貸方インボイス",
    "貸方金額(円)", "貸方税額", "摘要", "仕訳メモ", "タグ", "MF仕訳タイプ",
    "決算整理仕訳", "作成日時", "借方部門コード", "貸方部門コード", "経費精算番号",
]  # 27 列
assert len(HEADERS_MF) == 27

HEADERS_YAYOI = [
    "識別フラグ", "伝票No.", "決算", "取引日付", "借方勘定科目", "借方補助科目",
    "借方部門", "借方税区分", "借方金額", "借方税金額", "貸方勘定科目", "貸方補助科目",
    "貸方部門", "貸方税区分", "貸方金額", "貸方税金額", "摘要", "番号", "期日",
    "タイプ", "生成元", "仕訳メモ", "付箋1", "付箋2", "調整",
]  # 25 列 ★ 見出し行なし（弥生インポート形式・取引先の列が無い）
assert len(HEADERS_YAYOI) == 25

HEADERS_FREEE = [
    "日付", "伝票番号", "借方勘定科目", "借方科目コード", "借方補助科目", "借方取引先",
    "借方部門", "借方品目", "借方メモタグ", "借方金額", "借方内税/外税", "借方税区分",
    "借方税額", "借方摘要", "貸方勘定科目", "貸方科目コード", "貸方補助科目", "貸方取引先",
    "貸方部門", "貸方品目", "貸方メモタグ", "貸方金額", "貸方内税/外税", "貸方税区分",
    "貸方税額", "貸方摘要", "仕訳メモ",
]

HEADERS = {"mf": HEADERS_MF, "xlsx": HEADERS_MF, "yayoi": HEADERS_YAYOI, "freee": HEADERS_FREEE}
# 各ソフトの物理列名 ── canonical な 4 本の鍵をどの列に書くか（弥生には取引先の列が無い）
PHYSICAL = {
    "mf": {"借方取引先": "借方取引先", "貸方取引先": "貸方取引先",
           "借方補助科目": "借方補助科目", "摘要": "摘要"},
    "xlsx": {"借方取引先": "借方取引先", "貸方取引先": "貸方取引先",
             "借方補助科目": "借方補助科目", "摘要": "摘要"},
    "freee": {"借方取引先": "借方取引先", "貸方取引先": "貸方取引先",
              "借方補助科目": "借方補助科目", "摘要": "借方摘要"},
    "yayoi": {"借方補助科目": "借方補助科目", "摘要": "摘要"},
}
DESC_PHYSICAL = {"mf": "摘要", "xlsx": "摘要", "freee": "借方摘要", "yayoi": "摘要"}
SEQ_FIELD = {"mf": "取引No", "xlsx": "取引No", "freee": "伝票番号", "yayoi": "伝票No."}

CREDIT_ACCOUNTS = ("現金", "未払金", "普通預金")
BASE_DATE = date(2026, 6, 1)


def _norm(s) -> str:
    """設計 D2 の norm（空白の除去のみ・道具側の全角半角統一は検体では要らない）。"""
    return "".join(str(s or "").split())


def _date_obj(i: int) -> date:
    return BASE_DATE + timedelta(days=i % 90)


def _fmt_date(software: str, d: date, i: int) -> str:
    if software == "freee":
        return d.strftime("%Y-%m-%d")
    if software == "yayoi":
        variant = i % 3
        if variant == 0:
            return f"R{d.year - 2018:02d}/{d.month:02d}/{d.day:02d}"     # 和暦略記
        if variant == 1:
            return f"{d.year}/{d.month}/{d.day}"                          # ゼロ埋め無し
        return d.strftime("%Y%m%d")                                       # 数字だけ
    return d.strftime("%Y/%m/%d")   # mf / xlsx


def _amount_for(i: int) -> int:
    """3 桁〜6 桁に収まる決定的な金額（乱数は使わない・再現性のため）。"""
    return 300 + (i * 577) % 999000


def _credit_for(i: int) -> str:
    return CREDIT_ACCOUNTS[i % len(CREDIT_ACCOUNTS)]


# ── レビュー後の区分の導出そのもの（手で書かない・宣言から機械的に決まる）──
def grade_row(past_index: dict, canon_values: dict):
    """鍵ごと＝出所ごとに数える。戻り値 (科目|None, 区分, 根拠に含む語|None)。"""
    hit_pairs = []       # [(鍵, 値, 科目), ...] ── 1 科目に決まった鍵だけ
    split = False
    for field in CANON_FIELDS:
        v = canon_values.get(field)
        if not v:
            continue
        accs = past_index.get(field, {}).get(_norm(v))
        if not accs:
            continue
        distinct = set(accs)
        if len(distinct) >= 2:
            split = True          # この鍵の内訳が2科目以上
            continue
        hit_pairs.append((field, v, next(iter(distinct))))
    if split:
        return None, "割", None
    if not hit_pairs:
        return None, "無", None
    accounts = {a for _, _, a in hit_pairs}
    if len(hit_pairs) == 1:
        _, v, a = hit_pairs[0]
        return a, "単", v                              # 当たった鍵が1つ ── その語を根拠にする
    if len(accounts) == 1:
        return hit_pairs[0][2], "確", None              # ★ どの鍵の語を挙げるかは道具の裁量にする
    return None, "割", None                              # 鍵どうしが違う科目を指す


# ── 1 行を組む（canonical な鍵の値 → 物理列。account=None なら借方勘定科目は空）──
def _build_row(software: str, seq, dstr: str, account, canon_values: dict, amount, credit):
    headers = HEADERS[software]
    row = {h: "" for h in headers}
    phys = PHYSICAL[software]
    for field, val in canon_values.items():
        col = phys.get(field)
        if col:
            row[col] = val

    row[SEQ_FIELD[software]] = seq
    if software in ("mf", "xlsx"):
        row["取引日"] = dstr
        row["借方勘定科目"] = account or ""
        row["借方税区分"] = "課税仕入10%" if account else ""
        row["借方インボイス"] = "適格" if account else ""
        row["借方金額(円)"] = amount
        row["借方税額"] = round(amount * 0.1) if account else 0
        row["貸方勘定科目"] = credit
        row["貸方税区分"] = "対象外"
        row["貸方金額(円)"] = amount
        row["貸方税額"] = 0
        row["MF仕訳タイプ"] = "通常仕訳"
        row["決算整理仕訳"] = "無"
        row["作成日時"] = dstr + " 10:00:00"
    elif software == "freee":
        row["日付"] = dstr
        row["借方勘定科目"] = account or ""
        row["借方内税/外税"] = "内税" if account else ""
        row["借方税区分"] = "課対仕入10%" if account else ""
        row["借方金額"] = amount
        row["借方税額"] = round(amount * 0.1) if account else 0
        row["貸方勘定科目"] = credit
        row["貸方内税/外税"] = "対象外"
        row["貸方税区分"] = "対象外"
        row["貸方金額"] = amount
        row["貸方税額"] = 0
        if row.get("借方摘要"):
            row["貸方摘要"] = row["借方摘要"]
    else:  # yayoi
        # ★ 識別フラグは一次資料どおり 4 桁（2000 = 通常の仕訳）。初版は空にしていて、
        #   読み手を検体に合わせて広げかける事故になった（2026-09-13）── 検体を実物に寄せる。
        row["識別フラグ"] = "2000"
        row["取引日付"] = dstr
        row["借方勘定科目"] = account or ""
        row["借方税区分"] = "課税仕入" if account else ""
        row["借方金額"] = amount
        row["借方税金額"] = round(amount * 0.1) if account else 0
        row["貸方勘定科目"] = credit
        row["貸方税区分"] = "対象外"
        row["貸方金額"] = amount
        row["貸方税金額"] = 0
        row["タイプ"] = "0"
        row["付箋1"] = "0"
        row["付箋2"] = "0"
    return row


def _self_check(spec):
    """specimens の宣言そのものの矛盾を検体作成前に落とす（split_people と同じ作法）。"""
    if spec["software"] == "yayoi":
        for sc in spec["scenarios"]:
            assert not sc.get("vendor"), f"{spec['id']}: 弥生には取引先の列が無い（{sc['label']}）"
    if spec.get("貸方だけ取引先"):
        assert any(sc.get("credit_side") for sc in spec["scenarios"]), \
            f"{spec['id']}: 貸方だけ取引先=True なのに credit_side な行が無い"
    if spec.get("複合仕訳"):
        assert any(sc.get("untouchable") == "continuation" for sc in spec["scenarios"]), \
            f"{spec['id']}: 複合仕訳=True なのに継続行が無い"
    assert 1 <= len(spec["past_files"]) <= 3, f"{spec['id']}: past は 1〜3 ファイルのはず"


# ── 1 ケースぶんの行を展開する ─────────────────────────────────────
def expand_case(spec):
    software = spec["software"]
    past_index: dict = {}
    past_account_universe = set()

    def register(field, value, account):
        past_index.setdefault(field, {}).setdefault(_norm(value), []).append(account)
        past_account_universe.add(account)

    n_files = len(spec["past_files"])
    past_rows_per_file = [[] for _ in range(n_files)]
    past_seq = [0] * n_files
    file_cycle = itertools.cycle(range(n_files))
    i = 0

    def _write_past_fact(field, value, account):
        nonlocal i
        register(field, value, account)
        i += 1
        d = _date_obj(i)
        fidx = next(file_cycle)
        past_seq[fidx] += 1
        row = _build_row(software, past_seq[fidx], _fmt_date(software, d, i), account,
                          {field: value}, _amount_for(i), _credit_for(i))
        past_rows_per_file[fidx].append(row)

    for field, value, account in spec.get("extra_past_facts", []):
        _write_past_fact(field, value, account)

    today_rows = []              # [{"row":dict,"kind":...,"acc":...,"grade":...,"word":...}, ...]
    fresh_seq = [0]
    last_seq = [0]

    for sc in spec["scenarios"]:
        for field, value, account in sc.get("past_facts", []):
            _write_past_fact(field, value, account)

        untouchable = sc.get("untouchable")
        for occ in range(sc.get("today_count", 0)):
            i += 1
            d = _date_obj(i)
            dstr = _fmt_date(software, d, i)
            amount, credit = _amount_for(i), _credit_for(i)

            canon_values = {}
            vendor = sc.get("vendor")
            if vendor:
                canon_values["貸方取引先" if sc.get("credit_side") else "借方取引先"] = vendor
            subsidiary = sc.get("subsidiary")
            if subsidiary:
                canon_values["借方補助科目"] = subsidiary
            if sc.get("desc_overrides") is not None:
                desc = sc["desc_overrides"][occ]
            else:
                base = sc.get("desc_base") or ""
                desc = base if sc.get("desc_fixed") else f"{base} {d.month}月分"
            if desc:
                canon_values["摘要"] = desc

            if untouchable == "filled":
                fresh_seq[0] += 1; last_seq[0] = fresh_seq[0]
                account = sc["filled_account"]
                row = _build_row(software, fresh_seq[0], dstr, account, canon_values, amount, credit)
                today_rows.append(dict(row=row, kind="untouchable"))
            elif untouchable == "continuation":
                seq = last_seq[0]     # 取引No/伝票No. を継続行として再利用（借方は正当に空）
                row = _build_row(software, seq, dstr, None, {}, amount, credit)
                row[DESC_PHYSICAL[software]] = desc
                today_rows.append(dict(row=row, kind="untouchable"))
            elif untouchable == "total":
                fresh_seq[0] += 1; last_seq[0] = fresh_seq[0]
                row = _build_row(software, fresh_seq[0], dstr, None, {}, amount, credit)
                row[DESC_PHYSICAL[software]] = "合計"
                today_rows.append(dict(row=row, kind="untouchable"))
            else:
                fresh_seq[0] += 1; last_seq[0] = fresh_seq[0]
                row = _build_row(software, fresh_seq[0], dstr, None, canon_values, amount, credit)
                acc, grade, word = grade_row(past_index, canon_values)
                today_rows.append(dict(row=row, kind="target", acc=acc, grade=grade, word=word))

    return past_rows_per_file, today_rows, past_account_universe


def _write_rows(path: Path, software: str, rows: list, with_header: bool):
    headers = HEADERS[software]
    if software == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "仕訳"
        r = 1
        if with_header:
            for ci, h in enumerate(headers, start=1):
                ws.cell(row=r, column=ci, value=h)
            r += 1
        for row in rows:
            for ci, h in enumerate(headers, start=1):
                ws.cell(row=r, column=ci, value=row.get(h, ""))
            r += 1
        wb.save(path)
        wb.close()
        return
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    if with_header:
        w.writerow(headers)
    for row in rows:
        w.writerow([row.get(h, "") for h in headers])
    encoding = "cp932" if software == "yayoi" else "utf-8-sig"
    path.write_bytes(buf.getvalue().encode(encoding))


def build_case(spec) -> dict:
    _self_check(spec)
    software = spec["software"]
    with_header = software != "yayoi"
    past_rows_per_file, today_rows, past_account_universe = expand_case(spec)

    for fidx, fname in enumerate(spec["past_files"]):
        _write_rows(OUT / fname, software, past_rows_per_file[fidx], with_header)

    header_offset = 1 if with_header else 0
    期待_行 = {}
    触らない行 = []
    plain_rows = []
    for idx, item in enumerate(today_rows, start=1):
        r = idx + header_offset
        plain_rows.append(item["row"])
        if item["kind"] == "target":
            期待_行[str(r)] = {"科目": item["acc"], "区分": item["grade"], "根拠に含む語": item["word"]}
        else:
            触らない行.append(r)
    _write_rows(OUT / spec["today_file"], software, plain_rows, with_header)

    return dict(
        id=spec["id"],
        software=software,
        today=spec["today_file"],
        past=list(spec["past_files"]),
        先例に在る科目=sorted(past_account_universe),
        触らない行=触らない行,
        複合仕訳=bool(spec.get("複合仕訳")),
        貸方だけ取引先=bool(spec.get("貸方だけ取引先")),
        文字コード=("cp932" if software == "yayoi" else "utf-8-sig"),
        期待={"行": 期待_行},
        怪しくない=bool(spec.get("怪しくない")),
        会社=spec["company"],   # ★ 契約外（人が読む用の参考情報。score_accounts.py は見ない）
    )


def main() -> int:
    if OUT.exists():
        for f in list(OUT.glob("*.csv")) + list(OUT.glob("*.xlsx")):
            f.unlink()
    OUT.mkdir(parents=True, exist_ok=True)

    answers = []
    tally = {"確": 0, "単": 0, "割": 0, "無": 0}
    for spec in CASES:
        ans = build_case(spec)
        answers.append(ans)
        for v in ans["期待"]["行"].values():
            tally[v["区分"]] += 1
    total = sum(tally.values())

    payload = json.dumps(answers, ensure_ascii=False, indent=2)
    ANSWER.write_bytes((payload + "\n").encode("utf-8"))

    softs = sorted({a["software"] for a in answers})
    print(f"作った: {len(answers)} ケース（{softs}）→ {OUT}")
    print(f"today 対象合計 {total} 行 ── 確 {tally['確']} / 単 {tally['単']} "
          f"/ 割 {tally['割']} / 無 {tally['無']}")
    print(f"答え → {ANSWER}")

    bad = []
    if not (150 <= total <= 200):
        bad.append(f"today 合計行数が範囲外（150〜200 のはずが {total}）")
    for g, n in tally.items():
        if n < 8:
            bad.append(f"区分 {g} が {n} 行（8 以上のはず）")
    if bad:
        print("★", "; ".join(bad))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
