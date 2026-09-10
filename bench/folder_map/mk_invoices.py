# -*- coding: utf-8 -*-
"""請求書 40 ファイルの検体を作る（需要地図の上位① を実物で測るため）。

★ 揃いすぎにしない。記憶の「崩れる 3 点」を意図的に入れる:
    ① スキーマが無い   … 見出しの語・位置・シート名が file ごとに違う
    ② 正当な例外       … 赤伝(マイナス)・請求ゼロ・複数明細
    ③ 書式が意味を運ぶ … 数値が文字列・合計行・結合セル・単位つき

★ 正解（期待値）を同時に JSON で吐く。**採点は検体を作った側が持つ**
  （道具の出力から期待値を作ると恒真になる）。
"""
import json
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent / "invoices40"
OUT.mkdir(parents=True, exist_ok=True)

random.seed(20260910)

TORIHIKISAKI = [
    "あかね商事", "いろは工業", "うえだ物産", "エバラ機械", "大久保商店",
    "화林" if False else "花菱建材", "菊池電機", "くまがい精工", "ケイアイ物流", "小坂運輸",
    "佐々木製作所", "しなの化成", "鈴村テック", "瀬戸内海運", "曽根田工務店",
    "高梨産業", "千葉見本市", "つばめ交通", "寺岡食品", "戸田金属",
]

# 見出しの語ゆれ（★ スキーマが無い）
AMOUNT_LABEL = ["金額", "請求額", "合計金額", "ご請求金額", "税込金額"]
DATE_LABEL = ["日付", "請求日", "発行日", "年月日"]
NAME_LABEL = ["取引先", "御中", "宛先", "顧客名"]
SHEET_NAME = ["請求書", "Sheet1", "請求", "invoice", "御請求書"]

answers = []

for i in range(40):
    name = TORIHIKISAKI[i % len(TORIHIKISAKI)]
    if i >= len(TORIHIKISAKI):
        name = name + "（2期）"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME[i % len(SHEET_NAME)]

    amount_label = AMOUNT_LABEL[i % len(AMOUNT_LABEL)]
    date_label = DATE_LABEL[i % len(DATE_LABEL)]
    name_label = NAME_LABEL[i % len(NAME_LABEL)]

    # ★ 見出しの「上」に飾りを入れるファイルと入れないファイル（＝見出し行の位置が違う）
    title_rows = i % 3          # 0/1/2 行
    for t in range(title_rows):
        ws.append(["請求書" if t == 0 else f"発行番号 INV-{1000+i}"])
        if t == 0:
            ws.cell(row=1, column=1).font = Font(bold=True, size=14)

    # ★ 列順もバラす（実物の請求書は揃っていない）。位置で当たってしまう検体にしない。
    order = [0, 1, 2, 3]
    if i % 5 == 1:
        order = [1, 0, 2, 3]        # 日付が先頭
    elif i % 5 == 2:
        order = [0, 2, 1, 3]        # 金額と日付が入れ替わる
    elif i % 5 == 3:
        order = [3, 0, 1, 2]        # 備考が先頭
    labels = [name_label, date_label, amount_label, "備考"]
    ws.append([labels[k] for k in order])

    n_rows = 1 + (i % 3)        # 1〜3 明細（★ 複数明細は正当な例外）
    total = 0
    for r in range(n_rows):
        amt = (i + 1) * 1000 + r * 250
        if i == 7:              # ★ 赤伝（マイナス）
            amt = -amt
        if i == 13:             # ★ 請求ゼロ
            amt = 0
        # ★ 書式が意味を運ぶ: 一部は数値でなく文字列（「1,000円」）
        cell_amt = f"{amt:,}円" if i % 9 == 4 else amt
        cells = [name, f"2026-08-{(i % 28) + 1:02d}", cell_amt, ""]
        ws.append([cells[k] for k in order])
        total += amt

    # ★ 合計行があるファイルとないファイル（あると「明細の和」と二重計上しうる）
    has_total = (i % 4 == 0)
    if has_total:
        tcells = ["合計", "", total, ""]
        ws.append([tcells[k] for k in order])

    # ★ 結合セル（書式が意味を運ぶ／読み取りを壊しやすい）
    if i % 11 == 3:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

    path = OUT / f"請求書_{i+1:02d}_{name}.xlsx"
    wb.save(path)

    answers.append({
        "file": path.name,
        "取引先": name,
        "請求額": total,          # ★ 明細の和（合計行は二重計上しない）
        "明細数": n_rows,
        "見出し行": title_rows + 1,
        "見出しの語": {"取引先": name_label, "日付": date_label, "金額": amount_label},
        "合計行あり": has_total,
        "列順": order,
        "金額が文字列": i % 9 == 4,
        "備考": ("赤伝" if i == 7 else "請求ゼロ" if i == 13 else ""),
    })

(OUT.parent / "_答え.json").write_text(
    json.dumps(answers, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")

print(f"作った: {len(answers)} ファイル → {OUT}")
print(f"  合計請求額（明細の和）: {sum(a['請求額'] for a in answers):,}")
print(f"  見出し行が 1 行目でない: {sum(1 for a in answers if a['見出し行'] != 1)} 件")
print(f"  金額が文字列: {sum(1 for a in answers if a['金額が文字列'])} 件")
print(f"  合計行あり: {sum(1 for a in answers if a['合計行あり'])} 件")
print(f"  明細が 2 行以上: {sum(1 for a in answers if a['明細数'] > 1)} 件")
print(f"  正当な例外: {[a['file'] for a in answers if a['備考']]}")
