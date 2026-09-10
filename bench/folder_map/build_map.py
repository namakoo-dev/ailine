# -*- coding: utf-8 -*-
"""40 冊を先に読み込んで、**その場で語のマップを作る**（Namakoo 案・2026-09-10）。

★ なぜこの形か:
    出荷する対応表（取引先←宛先/御中/顧客名…）は、この 20 社この語彙でしか効かない。
    次の客は「得意先」「客先」と書く。**マップを配るのでなく、その場のフォルダから作る**。

★ 2 段構え:
    ① 語を 1 つも読まず、**中身の輪郭**で列の役割を決める（1 冊の中 + 冊をまたぐ証拠）
    ② 同じ役割に落ちた見出しの語を集めて **語のマップ**にする
       → 以後は語で速く引ける。人が目で見て直せる。attr に覚えさせられる。

★ 冊をまたぐ証拠が効く所（1 冊では出ない情報）:
    備考   … 全冊で空          → 一意率では取引先と区別できないが、空率で分かれる
    取引先 … 冊ごとに値が違う   → 明細 1 行でも「冊間で散る」ことで分かる
    日付   … 全冊で同じ書式
"""
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
import openpyxl

import ailine as A

HERE = Path(__file__).resolve().parent
FOLDER = HERE / "invoices40"
ANSWERS = {a["file"]: a for a in json.loads((HERE / "_答え.json").read_text(encoding="utf-8"))}

DATE_RE = re.compile(r"^\s*\d{4}[-/年]\d{1,2}[-/月]\d{1,2}日?\s*$")


def as_number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    s = str(v).strip().replace(",", "").replace("円", "").replace("¥", "")
    try:
        return float(s)
    except ValueError:
        return None


def read_one(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    end = min(ws.max_row or 1, A.MAX_ROWS, A.STRUCT_HEADER_SCAN_ROWS)
    stats = A._row_char_stats(ws, 1, end, 1, min(ws.max_column or 1, A.MAX_COLS))
    row, confident = A.detect_header_row({"rows": stats})
    hr = row if confident else 1
    last = ws.max_row or hr
    ncol = min(ws.max_column or 1, 12)
    cols = []
    for c in range(1, ncol + 1):
        head = ws.cell(row=hr, column=c).value
        vals = [ws.cell(row=r, column=c).value for r in range(hr + 1, last + 1)]
        nonempty = [v for v in vals if v not in (None, "")]
        cols.append({
            "col": c,
            "見出し": (str(head).strip() if head not in (None, "") else ""),
            "値": vals,
            "空率": 1 - len(nonempty) / (len(vals) or 1),
            "日付らしさ": sum(1 for v in nonempty if DATE_RE.match(str(v))) / max(len(nonempty), 1),
            "数らしさ": sum(1 for v in nonempty if as_number(v) is not None) / max(len(nonempty), 1),
            "文字値": {str(v) for v in nonempty
                       if as_number(v) is None and not DATE_RE.match(str(v))},
        })
    wb.close()
    return {"file": path.name, "見出し行": hr, "列": cols}


# ── ① 全冊を読む（インポート） ─────────────────────────────
books = [read_one(p) for p in sorted(FOLDER.glob("*.xlsx"))]
print(f"読み込んだ: {len(books)} 冊\n")

# ── ② 冊をまたぐ証拠を集める（見出しの語ごとに輪郭を平均） ──────
by_word = defaultdict(lambda: {"空率": [], "日付らしさ": [], "数らしさ": [],
                                "文字値の集合": [], "冊数": 0})
for b in books:
    for c in b["列"]:
        w = c["見出し"] or f"(無題:{c['col']})"
        e = by_word[w]
        e["空率"].append(c["空率"])
        e["日付らしさ"].append(c["日付らしさ"])
        e["数らしさ"].append(c["数らしさ"])
        e["文字値の集合"].append(c["文字値"])
        e["冊数"] += 1

def avg(xs):
    return sum(xs) / len(xs) if xs else 0.0

def role_of(e):
    """語を読まずに役割を決める。★ 冊をまたぐ散らばりを使う。"""
    if avg(e["空率"]) > 0.8:
        return "備考"                       # 全冊ほぼ空
    if avg(e["日付らしさ"]) > 0.5:
        return "日付"
    if avg(e["数らしさ"]) > 0.5:
        return "金額"
    sets = [s for s in e["文字値の集合"] if s]
    if not sets:
        return None
    union = set().union(*sets)
    # ★ 冊ごとに違う値 → union が冊数に比例して増える（取引先）
    #   全冊同じ値 → union は小さい（固定ラベル）
    spread = len(union) / max(len(sets), 1)
    return "取引先" if spread > 0.5 else "固定ラベル"

# ★ A: 少数派を捨てる。40 冊のうち 1 冊だけに出る語は「語」でなく「データ」
#   （見出し行を読み違えた冊が、そのデータ値を見出しとして持ち込む）。
#   ★ 分母（何冊か）が在って初めて言える判断 ── 先に全冊を読む形の効き所。
MIN_BOOKS = max(2, len(books) // 20)
word_map = defaultdict(list)
dropped = []
for w, e in sorted(by_word.items(), key=lambda kv: -kv[1]["冊数"]):
    r = role_of(e)
    if not r:
        continue
    if e["冊数"] < MIN_BOOKS:
        dropped.append((w, e["冊数"], r))
        continue
    word_map[r].append((w, e["冊数"]))
if dropped:
    print(f"★ 少数派として捨てた語（{MIN_BOOKS} 冊未満）: "
          + "、".join(f"{w}({n}冊→{r})" for w, n, r in dropped))

print("=== その場で作った語のマップ（語は 1 つも事前に知らない）===")
for role in ("取引先", "日付", "金額", "備考", "固定ラベル"):
    if word_map.get(role):
        items = "、".join(f"{w}({n}冊)" for w, n in sorted(word_map[role], key=lambda x: -x[1]))
        print(f"  {role:6s} ← {items}")

# ── ③ 作ったマップで抽出し、答えと突き合わせる ────────────────
lookup = {w: role for role, ws_ in word_map.items() for w, _ in ws_}
ok_name = ok_amt = both = fallback = 0
misses = []
for b in books:
    ans = ANSWERS[b["file"]]
    picked = {}
    for c in b["列"]:
        r = lookup.get(c["見出し"])
        if r in ("取引先", "日付", "金額") and r not in picked:
            picked[r] = c
    used_fallback = not {"取引先", "金額"} <= set(picked)
    if used_fallback:
        fallback += 1
        # ★ B: マップで引けなかったら輪郭に落ちる（見出しが空の列はここでしか拾えない）
        taken = {c["col"] for c in picked.values()}
        if "金額" not in picked:
            cand = [c for c in b["列"] if c["数らしさ"] > 0.5 and c["col"] not in taken]
            if cand:
                picked["金額"] = max(cand, key=lambda c: c["数らしさ"])
                taken.add(picked["金額"]["col"])
        if "取引先" not in picked:
            cand = [c for c in b["列"] if c["文字値"] and c["空率"] < 0.5
                    and c["col"] not in taken and c["日付らしさ"] < 0.5]
            if cand:
                picked["取引先"] = min(cand, key=lambda c: len(c["文字値"]))
    name = amt = None
    if "取引先" in picked:
        vs = [str(v) for v in picked["取引先"]["値"]
              if v not in (None, "") and str(v) != "合計"]
        name = vs[0] if vs else None
    if "金額" in picked:
        nm = picked.get("取引先", {}).get("値", [])
        s = 0.0
        for k, v in enumerate(picked["金額"]["値"]):
            if k < len(nm) and str(nm[k]).strip() == "合計":
                continue
            nv = as_number(v)
            if nv is not None:
                s += nv
        amt = s
    n_ok = (name == ans["取引先"])
    a_ok = (amt is not None and abs(amt - ans["請求額"]) < 0.5)
    ok_name += n_ok
    ok_amt += a_ok
    both += (n_ok and a_ok)
    if not (n_ok and a_ok):
        misses.append((b["file"], (name, amt), (ans["取引先"], ans["請求額"]),
                       b["見出し行"], ans["見出し行"],
                       [c["見出し"] for c in b["列"]]))

n = len(books)
print(f"\n=== 作ったマップで 40 冊から抜き出す ===")
print(f"取引先が当たった   {ok_name}/{n}")
print(f"請求額が当たった   {ok_amt}/{n}")
print(f"★ 両方当たった     {both}/{n}")
print(f"マップで引けず輪郭に落ちた {fallback}/{n}")
if misses:
    print("\n=== 外した分（★ 名指し）===")
    for f, got, want, hr, hr_ans, heads in misses:
        print(f"  {f}\n     得た={got} 答え={want} 見出し行={hr}(正:{hr_ans}) 見出し={heads}")
