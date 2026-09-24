"""行と列の「位置」の語と解決 ── 「〇〇の下に」「3行目」「売上の右に」を、表のどこかへ決める。

★ 2026-09-24 に src/ailine/__init__.py から**移しただけ**（本文は 1 文字も変えていない）。
★ 行と列の位置の語は**組**として同じ冊に置く（試験が両方の軸を並べて縛る ──
  tests/test_the_edge_of_a_table_is_a_position.py）。本体に残る位置の兄弟が後から入る先もここ。
★ ailine を import しない（持ち出せる部品）。本体は from ailine_core.anchor import ... で束ね直す。
"""
from __future__ import annotations

import re
from ailine_core.book_view import BookView
from ailine_core.row_words import _re_row_number_word
from ailine_core.table_scan import data_extent
from openpyxl.utils import column_index_from_string
from pathlib import Path
from ailine_core.quotes import _task_outside_quotes


def _digit_candidates(raw: str, headers: list) -> list:
    """数字表記の列参照を 0 起点/1 起点の両解釈で実在列名に変換した候補（重複除去・順序維持）。
       数字表記でなければ空リスト。"""
    s = str(raw).strip()
    if not re.fullmatch(r"\d+", s):
        return []
    n = int(s)
    cands = []
    for idx in (n, n - 1):        # 0起点読み(n) と 1起点読み(n-1) の両方を試す
        if 0 <= idx < len(headers) and headers[idx] not in cands:
            cands.append(headers[idx])
    return cands


def resolve_col_ref(raw, headers: list) -> tuple:
    """(実在列名 or None, 推定だったか, エラー文 or None)。
       ★ 列名を正とする。実在すればそのまま。数字表記なら 0/1 起点の両候補を試し、
       一意に決まればそれを『推定』として解決、決まらなければ CLARIFY 相当のエラーを返す。"""
    s = str(raw)
    if s in headers:
        return s, False, None
    cands = _digit_candidates(s, headers)
    if len(cands) == 1:
        return cands[0], True, None
    if len(cands) > 1:
        return None, False, f"列『{s}』は複数の解釈が可能で一意に決まりません: {cands}"
    # ★★ 2026-08-28（Namakoo が請求書のデモで実測）: 「F列に…」が断られていた。
    #   人は表計算の座標（A1 方式の列文字）で当たり前に指す。**断るのでなく解ける。**
    #   ★ 見出しに同じ名前が在れば**そちらが勝つ**（上の s in headers が先に返る）ので、
    #     『URL』のような英字の見出しを列文字と誤読する事故は起きない。
    #   ★ 解けたら「推定」として返す（呼び側が解釈行に (推定) を付ける＝黙って決めない）。
    letter = s[:-1] if s.endswith("列") else s
    if letter and re.fullmatch(r"[A-Za-z]{1,2}", letter):
        try:
            idx = column_index_from_string(letter.upper())
        except ValueError:
            idx = 0
        if 1 <= idx <= len(headers):
            return headers[idx - 1], True, None
    known = ", ".join(headers) if headers else "(無し)"
    return None, False, f"列『{s}』がありません。ある列: {known}"


_re_between = re.compile(r"([^\s、。]+?)\s*と\s*([^\s、。]+?)\s*の\s*間")


def _re_anchor(suffix: str):
    return re.compile(r"([^\s、。]+?)" + re.escape(suffix))


_ANCHOR_AFTER = ("の下に", "の下へ", "の後に", "の後ろに", "の次に")


_ANCHOR_BEFORE = ("の上に", "の上へ", "の前に")


# ★★ 2026-09-09（Namakoo「言い間違えを除けば位置語の指定が精度に直結する」）:
#   位置の語彙を軸ごとに並べたら**列にだけ端が在り、行には無かった**:
#       概念        列                    行
#       隣（後）    6 語                  5 語
#       隣（前）    5 語                  3 語
#       2 つの間    _re_between（共有）   同左        ← ここだけ対称だった
#       端          _COL_HEAD/_COL_TAIL   ★ 無い
#   ★ 09-08 に列側へ端を足したので、非対称が**広がって**いた。到達率の器が
#     名指しした 4 件（「一番下に行を入れて」等）はここに落ちていた。
#   ★ 行と列の非対称は、この repo で 3 度目（間・端・…）。列と**同じ形**で持つ。
_ANCHOR_TOP = ("一番上", "いちばん上", "先頭", "最初", "上端", "一番最初")


_ANCHOR_BOTTOM = ("一番下", "いちばん下", "末尾", "最後", "下端", "一番最後", "最終行")


def _row_word_number(word: str) -> int:
    """「4行目」「４行」→ 4（全角も受ける）。"""
    digits = "".join(ch for ch in word if ch.isdigit() or ch in "０１２３４５６７８９")
    return int(digits.translate(_ZENKAKU_DIGITS))


def row_number_anchor(task: str) -> tuple:
    """「4行目の下に」「2行目の前に」── **行番号と向き**だけから位置を出す。

    戻り値: (入れる行, 依頼文が言った行番号, 説明) ── 当たらなければ (None, None, "")。
    ★ 表に訊く必要が無い（純関数）。だから**位置を決める側と、その位置を審査する側の
      両方が同じここを通る**（片配線を作らない ── この repo が何度も踏んだ形）。
    """
    text = (task or "").replace("　", " ")
    # ★★ 2026-08-29（Namakoo）:「4行目と5行目は両方ともヤマノ食品。取引先で指定は
    #   出来ない」── 中身で指せない表では、人は番号でしか言えない。ならば
    #   「4行目と5行目の間に」も同じ引き算で出す（ここも表に訊く必要が無い）。
    m = _re_between.search(text)
    if m:
        a_, b_ = m.group(1).strip(), m.group(2).strip()
        if _re_row_number_word.fullmatch(a_) and _re_row_number_word.fullmatch(b_):
            na, nb = _row_word_number(a_), _row_word_number(b_)
            if nb - na != 1:
                return None, None, ""        # 隣り合っていない ── 決めない
            return nb, na, f"{na}行目と{nb}行目の間＝{nb}行目"
        return None, None, ""                # 片方でも名前なら、表に訊く側の仕事
    for sufs, after in ((_ANCHOR_AFTER, True), (_ANCHOR_BEFORE, False)):
        for suf in sufs:
            m = _re_anchor(suf).search(text)
            if not m:
                continue
            name = m.group(1).strip()
            if not _re_row_number_word.fullmatch(name):
                return None, None, ""      # 名前で指している（表に訊く側の仕事）
            n = _row_word_number(name)
            at = n + 1 if after else n
            return at, n, f"{n}行目の{chr(0x4E0B) if after else chr(0x4E0A)}に入れる＝{at}行目"
    return None, None, ""


# ★★ 2026-08-31（Namakoo の提案した通しを俺が先に走らせて出た・1 幕目が全滅）:
#   「**8行目に**丸山工業の行を作って」で、名前として『8行目に丸山工業』を丸ごと
#   切り出していた（区切りが空白と読点しか無く、**行番号をまたいで飲み込む**）。
#   ★ そのとき task_names_a_row_number は正しく 8 を返していた ── **行番号が
#     分かっているのに、名前の切り出しがそれを無視していた**。
#   ★ 助詞と行番号の語は名前に含まれない ── そこで切る（語彙ではなく文法の線）。
_re_row_of = re.compile(r"([^\s、。をにへはがでとのも]+?)\s*の\s*行")


def _is_number_like(s) -> bool:
    """数字だけの文字列か（依頼文に出る数と、行の名前を混同しないため）。"""
    try:
        float(str(s).replace(",", ""))
        return True
    except (TypeError, ValueError):
        return False


def _cell_row_name_for(book_meta: dict, sheet, row: int, header_row: int = 1):
    """その行を人が呼ぶときの名前（1 列目の値）。分からなければ None。"""
    rows, _h = _table_rows_for_anchor(book_meta, sheet, header_row)
    vals = rows.get(row) or []
    return vals[0] if vals else None


# ★ 「7 行目」「7行」「第7行」── 人は行を**番号**でも指す。
_re_row_number_in_task = re.compile(r"(?:第)?\s*([0-9０-９]{1,4})\s*行(?:目)?")


def task_names_a_row_number(task: str) -> int | None:
    """依頼文が指している行番号（1 起点）。無い/複数あって決まらないなら None。"""
    nums = {int(m.translate(_ZENKAKU_DIGITS))
             for m in _re_row_number_in_task.findall(task or "")}
    return nums.pop() if len(nums) == 1 else None


def _resolve_named_row(book_meta: dict, sheet: str | None, name: str) -> tuple:
    """行の名前 → 行番号（1 起点）。決められなければ (None, 断りの文)。

    ★ 2026-08-27: 住所の解決はここ 1 箇所に集める（resolve_row_anchor もこれを使う形へ
      寄せていく）。★ 探す範囲は**物理の使用範囲**（走査が最初の空で止まる穴を避ける）。
    ★ 見つからない・複数ある時は**決めない** ── 推測で別の行に書くのが一番こわい。
    """
    path = book_meta.get("path")
    if not path:
        return None, "表を読めないため、どの行かを決められません"
    hr = int((book_meta.get("header_rows") or {}).get(sheet, 1) or 1)
    # ★★ 2026-08-31（Namakoo が実測・「この基本操作ができない」）:
    #   「6行目と5行目を入れ替えて」が CLARIFY に落ちていた ── **行番号で指すと黙る**、
    #   08-29 に追加・削除で直したのと同じ非対称が、**入れ替えには残っていた**。
    #   ★ ここは住所の解決を集めている 1 箇所なので、ここに足すと全部の op に効く
    #     （入れ替え専用の判定を作らない）。
    if _re_row_number_word.fullmatch(str(name or "").strip()):
        _n = _row_word_number(name)
        if _n > hr:
            return _n, f"{_n}行目（依頼文の行番号）"
        return None, f"{_n}行目は見出し行（{hr}行目）またはその上です"
    # ★★ 2026-09-09: LLM は位置を**語**で返すことがある（挿入位置=「最後」）。
    #   旧版はそれを**行の名前**として実表に探しに行き、「『最後』という行が
    #   見つかりません」と断っていた ── 位置語を値として扱っていた形。
    #   ★ ここは住所の解決を集めている 1 箇所なので、足せば全部の op に効く
    #     （上の行番号と同じ理由・専用の判定を作らない）。
    _pos = str(name or "").strip()
    if _pos and any(w in _pos for w in _ANCHOR_TOP + _ANCHOR_BOTTOM):
        try:
            with BookView(Path(path)) as _bv0:
                _last0, _ = data_extent(_bv0.sheet(sheet), hr)
        except Exception:
            return None, None
        if any(w in _pos for w in _ANCHOR_TOP):
            return hr + 1, f"『{_pos}』＝{hr + 1}行目（見出しの次）"
        return _last0 + 1, f"『{_pos}』＝{_last0 + 1}行目（表の終わりの次）"
    try:
        with BookView(Path(path)) as bv:
            ws = bv.sheet(sheet)
            last, last_col = data_extent(ws, hr)
            hits = [r for r in range(hr + 1, last + 1)
                     if any(str(ws.cell(row=r, column=c).value or "").strip() == name
                             for c in range(1, last_col + 1))]
    except Exception as e:
        return None, f"表を読めませんでした（{type(e).__name__}）"
    if not hits:
        return None, f"『{name}』という行が見つかりません"
    if len(hits) > 1:
        return None, (f"『{name}』が {len(hits)} 行あります"
                       f"（{'、'.join(str(h) for h in hits)}行目）── どれか決められません")
    return hits[0], f"『{name}』の行＝{hits[0]}行目"


def _table_rows_for_anchor(book_meta: dict, sheet, header_row: int) -> tuple:
    """位置解決のために実表を読む（行番号 → 値の並び、と見出しの並び）。読めなければ空。"""
    path = book_meta.get("path")
    if not path:
        return {}, []
    try:
        with BookView(Path(path)) as bv:
            ws = bv.sheet(sheet)
            last, last_col = data_extent(ws, header_row)
            rows = {r: [str(ws.cell(row=r, column=c).value or "").strip()
                         for c in range(1, last_col + 1)]
                     for r in range(header_row + 1, last + 1)}
            heads = [str(ws.cell(row=header_row, column=c).value or "").strip()
                      for c in range(1, last_col + 1)]
            return rows, heads
    except Exception:
        return {}, []


def _row_named_anywhere_in_task(task: str, rows: dict, headers: list,
                                 require_possessive: bool = False):
    """依頼文に literal で現れる**実在の値**が、ちょうど 1 行にしか無いならその行。

    ★★ 2026-08-28（Namakoo「行の削除もできない」）: 「ナットを削除して」のように、
      人は「〜の行」と言わないことがある。言い回しを足すのではなく**表に訊く**。
    ★ 見出しの語は除く（列名を行の名前と読み違えない）。
    ★ 2 行に当たったら決めない（推測で別の行を消すのが一番こわい）。
    ★★ 2026-08-30（Namakoo「セル指定しているのに値を上書きできない」）:
      「7行B列を『{{合計:税込金額}}』に上書き」で、**引用符の中の『合計』**が表の
      合計行に当たり、そこを狙った操作に読み替えられていた。
      ★ 引用符の中は**値**であって、対象の名指しではない ── ここが 4 つの呼び出しの
        合流点なので、**この 1 行**で全部に効く（呼び出し側に配らない）。
    """
    text = _task_outside_quotes(task)
    heads = {h for h in headers if h}
    best = None
    for r, vals in (rows or {}).items():
        for v in vals:
            # ★★ 2026-08-31（Namakoo「LLM の揺れが一番厄介だ」→ 追ったら半分は機械の責任）:
            #   「金額が**60000**以上の行を抜き出して」で、機械が『60000』を**行の名前**
            #   として解き（金額列に 60000 が在る）、「『60000』の行＝3行目」と確信して
            #   行追加に読み替えていた。★ LLM が揺れた回に、**機械がその揺れを
            #   『確信をもって間違った操作』に育てていた**。
            #   ★ 揺れは消せないが、**増幅しないことはできる** ── 依頼文に出る数は
            #     ほぼ常に閾値や個数で、行の名前ではない。
            #   ★ 判定は既にある `_is_number_like`（「依頼文に出る数と、行の名前を
            #     混同しないため」）を借りる ── 1 箇所でしか使われていなかった。
            if not v or v in heads or len(v) < 2 or v not in text or _is_number_like(v):
                continue
            # ★★ 2026-08-29（Namakoo が実測）: 「丸山工業の担当に『佐藤』を入れて」で
            #   **書き込む値『佐藤』**が別の行の担当欄にも在るため、行の候補が 2 つに
            #   なって「決められない」に落ちていた ── 値を行の名前と読んでいた。
            #   ★ 人が行を指すときは「**〜の**」と言う。セルを指す経路ではそれを要求する
            #     （「ナットを削除して」のように の が無い経路は今までどおり）。
            if require_possessive and f"{v}の" not in text:
                continue
            if best is None:
                best = (r, v)
            elif best[0] != r:
                return None          # 2 行以上に当たる ── 決めない
            elif len(v) > len(best[1]):
                best = (r, v)
    return best


def resolve_row_anchor(task: str, book_meta: dict, sheet: str | None,
                        header_row: int = 1, anchor_out: dict | None = None) -> tuple:
    """依頼文の「**みかんの下に**」「**みかんとぶどうの間に**」から行番号を決める。

    ★ 2026-08-27（Namakoo が実測）: ADD_ROW は位置を**行番号**でしか受け取れないのに、
      人は相対で言う。LLM に数えさせると外し、空行だけの INSERT_ROWS に落ちていた。
    ★ 分担を変える: **LLM は「誰の隣か」を言うだけ／行番号は機械が実表を数えて決める**
      （列名の解決を機械 3 段でやっているのと同じ形）。
    ★ 見つからない・複数ある時は**決めない**（推測で行を挿すと、静かに別の場所へ入る）。
    戻り値: (行番号 or None, 説明 or 断りの理由 or None)
    """
    # ★★ 2026-08-30（Namakoo「セル指定しているのに値を上書きできない」）:
    #   「7行B列を『{{合計:税込金額}}』に上書き」で、**引用符の中の『合計』**を位置の
    #   目印として拾い、『合計』の行＝9行目 と解いていた。そのせいで一段目が行の挿入を
    #   返した回に「行追加として読み直しました」が発火し、頼んでいない行が挿さりかけた。
    #   ★ 列では既に塞いだ穴（_task_names_single_real_column）が、行では開いていた
    #     ── **行と列の非対称**、この repo が何度も踏んだ形。
    #   ★ 引用符の中は**値**であって、対象の名指しではない（Namakoo の決めた約束）。
    #     だから位置を探す時は引用符の中を見ない ── 「『みかん』の行を削除して」の
    #     ように名前を引用する書き方は、引用符なしで書いてもらう（列と同じ扱い）。
    text = _task_outside_quotes(task).replace("　", " ")
    want_after, name, second = None, None, None
    m = _re_between.search(text)
    if m:
        want_after, name, second = True, m.group(1).strip(), m.group(2).strip()
    else:
        # ★★ 2026-09-09（盲検 D が打った言い方で実測）: 「北斗精機**の行**の下に」で
        #   掴むのが『北斗精機の行』になり、実表に無いので落ちて、後段の「<X>の行」
        #   規則が**その行そのもの**（6行目）を返していた ── 「の下に」が丸ごと消え、
        #   新しい行が**上**に入る。★ 事務の人が最も自然に言う形。
        #   ★ 掴んだ語から**構造の語**（の行／の列）を落とす。09-08 に列側でやった
        #     「掴んだ語を実表の見出しで切り直す」と同じ形（軸が違うだけ）。
        def _strip_structure(s: str) -> str:
            t = (s or "").strip()
            for suffix in ("の行", "の列", "行", "列"):
                if len(t) > len(suffix) and t.endswith(suffix):
                    return t[: -len(suffix)]
            return t

        for suf in _ANCHOR_AFTER:
            m = _re_anchor(suf).search(text)
            if m:
                want_after, name = True, _strip_structure(m.group(1))
                break
        if name is None:
            for suf in _ANCHOR_BEFORE:
                m = _re_anchor(suf).search(text)
                if m:
                    want_after, name = False, _strip_structure(m.group(1))
                    break
    if not name:
        # ★ 2026-08-27（実測）:「りんごの行を削除して」── 人は行を**中身**で指す。
        #   相対の言い回しが無くても、「<X>の行」なら X を実表で探す。
        m2 = _re_row_of.search(text)
        if m2:
            want_after, name = None, m2.group(1).strip()
    # ★ 2026-08-28: 言い回しが 1 つも当たらない回も、**表に訊いてから**諦める。
    if not name:
        rows_h, heads_h = _table_rows_for_anchor(book_meta, sheet, header_row)
        alt = _row_named_anywhere_in_task(task, rows_h, heads_h)
        if alt:
            return alt[0], f"『{alt[1]}』の行＝{alt[0]}行目"
        # ★★ 2026-09-09: 端の指定は**隣の指定より後**に見る（「みかんの下に」の方が
        #   具体的なので、両方書いてあったら隣を採る ── 列側と同じ順序）。
        #   ★ 実表を数えて決める（「一番下」は見出しでも 1 行目でもない）。
        _last = header_row + len(rows_h)
        for w in _ANCHOR_TOP:
            if w in text:
                return header_row + 1, f"『{w}』＝{header_row + 1}行目（見出しの次）"
        for w in _ANCHOR_BOTTOM:
            if w in text and rows_h:
                return _last + 1, f"『{w}』＝{_last + 1}行目（表の終わりの次）"
        return None, None
    # ★ 2026-08-27（自分で入れた誤爆・既存の検体が捕まえた）:
    #   「**2行目の前に**1行挿入して」の「2行目」を中身の名前として探し、
    #   見つからず断っていた。**行番号は名前ではない** ── 表を探しに行かない。
    # ★★ 2026-08-29（Namakoo が実測・「行の追加が出来なくなってる」）:
    #   そのとき「探さない」を「**決めない**」と書いてしまった。結果:
    #     「ヤマノ食品の下に丸山工業の行を作って」→ 5行目・値も入る（✓）
    #     「4行目の下に丸山工業の行を作って」  → 機械が黙る → LLM の 4 がそのまま通り、
    #                                            **上に空行**が挿さった（✗）
    #   ★ 同じ「下に」なのに、**指し方が名前か番号かで結果が変わっていた**。
    #     ここは表に訊く必要すらない ── 番号と向きが揃っているのだから**引き算で出る**。
    #   ★ 位置が出れば、`insert_rows_should_have_been_add_row` の証拠①も立つので、
    #     値つきの行（ADD_ROW）へ回る ── 空行に落ちる道が同時に塞がる。
    if _re_row_number_word.fullmatch(name):
        at, _n, note = row_number_anchor(task)
        return (at, note) if at is not None else (None, None)
    path = book_meta.get("path")
    if not path:
        return None, None
    try:
        with BookView(Path(path)) as bv:
            ws = bv.sheet(sheet)
            # ★★ 2026-08-27（Namakoo が実測・俺が新しい所で開けた同じ穴）:
            #   `_scan_last_row` は 1 列目を上から見て**最初の空で止まる**。
            #   下書きに空行が 1 本あると、その下の「みかん」を探せず、位置解決が黙って
            #   失敗して LLM の行番号がそのまま通っていた。
            #   ★ **探す範囲は物理の使用範囲から取る**（今週この repo が 3 度直した形）。
            last, last_col = data_extent(ws, header_row)
            ws_rows = {r: [str(ws.cell(row=r, column=c).value or "").strip()
                            for c in range(1, last_col + 1)]
                        for r in range(header_row + 1, last + 1)}
            headers_here = [str(ws.cell(row=header_row, column=c).value or "").strip()
                             for c in range(1, last_col + 1)]
            hits = [r for r, vals in ws_rows.items() if name in vals]
            # ★★ 2026-09-07: 当たった行を**呼び出し側へ渡す口**（判断はここ 1 箇所のまま）。
            #   削除の道は「2 行に当たったら断る」でなく「**聞いてから両方消す**」に変えた
            #   ── 買い手の目には「消してと言ったのに何も起きない」は失敗だから。
            #   ★ 一覧を作る所を 2 箇所に書かない（片配線を作らない）ので、out 引数にする。
            if anchor_out is not None:
                anchor_out["name"], anchor_out["rows"] = name, list(hits)
    except Exception:
        return None, None
    if not hits:
        # ★★ 2026-08-28（Namakoo「行の削除もできない」）: 「ナット**を**削除して」が
        #   『1行目は見出し行です』で断られていた。人は「〜の行」と言わないこともある。
        #   ★ 言い回しを足すのではなく、**表に訊く**: 依頼文に literal で現れる値が
        #     この表のちょうど 1 行にしか無いなら、それがその行。
        #     （列名も見出しも除く ── 「数量が100未満の行」のような条件文は当たらない）
        alt = _row_named_anywhere_in_task(task, ws_rows, headers_here)
        if alt:
            return alt[0], f"『{alt[1]}』の行＝{alt[0]}行目"
        # ★★ 2026-08-31（通しの 1 幕目で全滅した形）:「8行目に丸山工業の行を作って」
        #   ── **これから置く**行なので、名前が表に無いのは当たり前。
        #   ★ 依頼文が行番号を名指ししているなら、それが場所（表に無いことは断りの
        #     理由にならない）。実測では task_names_a_row_number が 8 を返せていたのに、
        #     名前が見つからないほうで先に断っていた。
        _n_here = task_names_a_row_number(task)
        if _n_here and _n_here > header_row:
            return _n_here, f"{_n_here}行目（依頼文の行番号）"
        # ★★ 2026-09-09: 端の語（一番下/最後/…）は**位置**であって行の名前ではない。
        #   ★ 住所を解く所は 2 つ在るので、両方に同じ語彙を持たせる
        #     （片方だけだと「一番下に足して」は通るのに「最後の行を削除して」が
        #     断られる ── 実測でそうなった）。
        _pos2 = str(name or "").strip()
        if _pos2 and any(w in _pos2 for w in _ANCHOR_TOP + _ANCHOR_BOTTOM):
            if any(w in _pos2 for w in _ANCHOR_TOP):
                return header_row + 1, f"『{_pos2}』＝{header_row + 1}行目（見出しの次）"
            _last2 = header_row + len(ws_rows)
            return _last2, f"『{_pos2}』＝{_last2}行目（表の最後）"
        return None, (f"『{name}』という行が見つかりません"
                       "（この表に在る値で指してください・行番号でも指せます）")
    if len(hits) > 1:
        # ★★ 2026-08-29（Namakoo）:「どうしても中身でさせない場面が出てくる。例えば
        #   4行目と5行目は両方ともヤマノ食品。取引先で指定は出来ない」── そのとおりで、
        #   ここは**断って終わる場所ではなく、行番号の道へ渡す場所**。
        #   ★ 候補の行番号は機械がもう知っている ── そのまま言う（人に数え直させない）。
        _rows = "、".join(str(h) for h in hits)
        _ex = f"{hits[0]}行目"
        return None, (f"『{name}』が {len(hits)} 行あります（{_rows}行目） ── どれか決められません。"
                       f"行番号で指してください（例:「{_ex}の下に…」「{_ex}を削除して」）")
    row = hits[0]
    if want_after is None:          # 「<X>の行」＝ その行そのもの
        return row, f"『{name}』の行＝{row}行目"
    at = row + 1 if want_after else row
    where = "下" if want_after else "上"
    note = f"『{name}』（{row}行目）の{where}＝{at}行目"
    if second:
        note += f"（『{second}』との間）"
    return at, note


# ★ 2026-08-27: 列の相対位置。行（_ANCHOR_AFTER/_BEFORE）と**同じ形**で持つ ──
#   「位置は op に依らず位置」なので、片方だけ賢くしない。
_COL_AFTER = ("の右に", "の右へ", "の右側に", "の後ろに", "のうしろに", "の次に")


_COL_BEFORE = ("の左に", "の左へ", "の左側に", "の前に", "の手前に")


# ★★ 2026-09-08（盲検の検品が「取引先の列を**一番左に**持ってきて」で挙げた）:
#   端を指す言い方を 1 つも持っていなかった ── 隣（誰かの右/左）しか読めない。
#   ★ 列を作る時にも動かす時にも同じ言い回しが来るので、**位置の層に置く**
#     （op ごとに if を書かない ── この層が在る理由そのもの）。
_COL_HEAD = ("一番左", "いちばん左", "先頭", "最初", "左端", "いちばん前", "一番前")


_COL_TAIL = ("一番右", "いちばん右", "末尾", "最後", "右端", "いちばん後ろ", "一番後ろ")


# 「原価と売上の右側に」＝ 2 つのうち右の方の隣（Namakoo が挙げた実例）。
_re_col_pair = re.compile(r"([^\s、。]+?)\s*と\s*([^\s、。]+?)\s*の\s*(右|左)")


def _header_index(headers: list, name: str) -> tuple:
    """列名 → (1 起点の位置, 実際の見出し名)。『原価列』のように「列」が付いた言い方も受ける。
       決まらなければ (None, name)（推測しない）。
       ★ 実際の見出し名も返すのは、解釈行に**表に在る名前**を出すため
         （『原価列』と書かれても『原価』と表示する ── 人が突き合わせられる形にする）。"""
    names = [str(h) for h in headers]
    if name in names:
        return names.index(name) + 1, name
    if name.endswith("列") and name[:-1] in names:
        return names.index(name[:-1]) + 1, name[:-1]
    return None, name


_ZENKAKU_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")


def resolve_col_anchor(task: str, headers: list) -> tuple:
    """依頼文の「**原価の右に**」「**原価と売上の右側に**」から、新しい列が入る位置
       （1 起点）を決める。

    ★ 分担は行と同じ: **LLM は「誰の隣か」を言うだけ／位置は機械が実表の見出しから決める。**
    ★ 見つからない・決められない時は**決めない**（黙って末尾に付けない ── 静かに
      違う場所へ入るのが一番こわい、を列でも同じに扱う）。
    戻り値: (位置 or None, 説明 or 断りの理由 or None)。
            (None, None) = 位置の言い回しが**そもそも無い**（呼び側が末尾を選べる）
    """
    text = (task or "").replace("　", " ")
    names = [str(h) for h in headers]
    # ★★ 2026-08-29（84 件の効果検体で 3 表とも同じ形で落ちた）:
    #   「料理と主材料の**間に**区分の列を追加して」が解けず、黙って末尾に付いていた。
    #   行は `_re_between`（「AとBの間」）を持っているのに、列は「右／左」しか
    #   見ていなかった ── **行と列の非対称**。Namakoo が名指しした所そのもの。
    #   ★ 同じ正規表現を列にも通す（軸が違うだけで、位置の言い回しは同じ）。
    mb = _re_between.search(text)
    if mb:
        a, c = mb.group(1).strip(), mb.group(2).strip()
        ia, a = _header_index(names, a)
        ic, c = _header_index(names, c)
        if ia is not None and ic is not None:
            hi = max(ia, ic)
            return hi, f"『{a}』と『{c}』の間＝{hi}列目"
        # ★ 見出しに無いなら、それは列の話ではない（行の「間」かもしれない）── 触らない。
    m = _re_col_pair.search(text)
    if m:
        a, c, side = m.group(1).strip(), m.group(2).strip(), m.group(3)
        ia, a = _header_index(names, a)
        ic, c = _header_index(names, c)
        if ia is None or ic is None:
            missing = [x for x, i in ((a, ia), (c, ic)) if i is None]
            return None, (f"『{"』『".join(missing)}』という列がありません"
                           f"（ある列: {"、".join(names)}）")
        lo, hi = min(ia, ic), max(ia, ic)
        at = hi + 1 if side == "右" else lo
        return at, f"『{a}』と『{c}』の{side}＝{at}列目"
    for suf in _COL_AFTER + _COL_BEFORE:
        m = _re_anchor(suf).search(text)
        if not m:
            continue
        # ★★ 2026-09-08（列移動を入れて実測）: この正規表現は**いちばん早い開始位置**
        #   から伸びるので、「締め日を金額の右に」で『締め日を金額』を丸ごと掴み、
        #   「そんな列はありません」と嘘の診断を出していた（列追加でも同じ形）。
        #   ★ 掴んだ文字列を**実表の見出しで切り直す**（推測でなく実表で決める）。
        _grabbed = m.group(1).strip()
        if _grabbed not in names:
            _tails = [h for h in names if h and _grabbed.endswith(h)]
            if len(_tails) >= 1:
                _grabbed = max(_tails, key=len)
        idx, name = _header_index(names, _grabbed)
        if idx is None:
            return None, (f"『{name}』という列がありません"
                           f"（ある列: {"、".join(names)}）")
        after = suf in _COL_AFTER
        at = idx + 1 if after else idx
        return at, f"『{name}』（{idx}列目）の{"右" if after else "左"}＝{at}列目"
    # ★ 端の指定は**隣の指定より後**に見る ── 「原価の右に」の方が具体的なので、
    #   両方書いてあったら隣を採る（具体が一般に勝つ）。
    for w in _COL_HEAD:
        if w in text:
            return 1, f"『{w}』＝1列目"
    for w in _COL_TAIL:
        if w in text:
            return len(names) + 1, f"『{w}』＝{len(names) + 1}列目"
    return None, None
