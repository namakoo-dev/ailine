"""④ 決定論 codegen ── op と解決済みの引数から LibreOffice Basic を組む（LLM は一切使わない）。

★ 2026-09-23 に src/ailine/__init__.py から**移しただけ**（本文は 1 文字も変えていない）。
  入口は `codegen_dsl`、op ごとの組み立ては `CODEGEN_BY_OP` の表（足すのは表の 1 行）。
★ ailine を import しない（持ち出せる部品）。本体は `from ailine_core.codegen import ...` で束ね直すので、
  ailine.X の名前（公開面・試験）はそのまま残る。
"""
from __future__ import annotations

import re
from ailine_core.chart_range import chart_data_last_row
from ailine_core.postconditions._shared import COLOR_MAP
from ailine_core.primitives import is_number as _is_number
from ailine_core.target_sheet import wrap_basic_for_sheet
from openpyxl.utils import column_index_from_string, get_column_letter
from ailine_core.compare_codes import _EXTRACT_CMP_CODE
from ailine_core.row_words import _re_row_number_word


def _wrap_basic(body: str) -> str:
    """CONTRACT と同じ骨格（Option 2行 + Sub Run(oDoc As Object) 1つ）で包む。"""
    return "Option VBASupport 1\nOption Explicit\n\nSub Run(oDoc As Object)\n" + body + "End Sub\n"


def _wrap_basic_for_sheet(body: str, book_meta: dict, target_sheet: str | None) -> str:
    """★ 挙動変更#2: 薄い配線。本体（対象シートの一時的な並べ替えロジック）は
       ailine_core.target_sheet.wrap_basic_for_sheet に置く（移植可能性の番人 —
       tests/ailine_py_line_budget.txt 参照）。_wrap_basic を「対象シートが1枚目のときの
       既定ラップ」としてコールバックで渡す（ailine_core → ailine の逆流を避けるため）。"""
    return wrap_basic_for_sheet(body, _wrap_basic, book_meta.get("sheets") or [], target_sheet)


def _scan_last_row_basic(var: str = "oSheet", key_col: str = "0",
                          start_row: str = "1", min_ok: str | None = None) -> str:
    """走査ループの定型（refs の作法どおり：A列を上から走査して最終データ行を探す）。
       ★ W3: start_row（Basic 0起点の走査開始行＝見出し行の直下）を渡すと、見出しが
       物理1行目でない帳票でも正しい行から走査する（既定 "1" は旧挙動と同一）。
       min_ok（"lastRow がこの値未満なら Exit Sub"の閾値）を渡すと保存境界を調整できる
       （既定は start_row と同じ＝『データ0行なら何もしない』。見出しを含めて範囲を
       スタイリングする操作は min_ok に start_row-1 相当を渡す＝データ0行でも見出しは扱える）。

    ★★ 2026-09-16（掃き出しの続き・その日のうちに踏んだ穴）:
      同じ「表の終わり」の判断を helpers/*.bas の **11 か所**から `TableLastRow` へ畳んだが、
      **生成側が組み立てて埋め込むこちら**を数えていなかった（ここ 1 本から 7 箇所へ吐く）。
      当たる op は 6 つ ── APPEND_TOTAL / BOLD / CENTER_ALIGN / COMPUTE_COLUMN /
      FILL_COLOR / SET_COLUMN_VALUE。実測した症状:

          合計行（左端が空）のある表で「取引先の列を太字にして」
            → 見出しとデータ行だけ太字になり、合計行に届かず
              「列『取引先』に太字でないセルがある」で ×（原本は無傷だが操作できない）

      ★ 朝に「測定がデータを落とす形を 1 つ見つけたら、他の形を列挙してから閉じる」と
        書いた当人が、同じ日に 1 つ直して満足した。★ 畳む時は **生成する側も数える**。
    ★ ここは既に 1 本に畳まれていたので、直しはこの関数の中だけで 6 op 全部に届く。
    ★ key_col を渡された回（A 列でない列で数える指定）は従来どおり自前で走査する ──
      TableLastRow は「見出しの幅のどこかに値が在るか」で数えるので、意味が違う。
    """
    if min_ok is None:
        min_ok = start_row
    if key_col == "0":
        # ★ 見出しの幅で数える（左端が空の行にも届く）。判断は helpers の 1 本に任せる。
        return (f"    lastRow = TableLastRow({var}, {int(start_row) - 1})\n"
                f"    If lastRow < {min_ok} Then Exit Sub\n")
    return (f"    lastRow = {start_row}\n"
            f"    Do While {var}.getCellByPosition({key_col}, lastRow).getString() <> \"\"\n"
            f"        lastRow = lastRow + 1\n"
            f"    Loop\n"
            f"    lastRow = lastRow - 1\n"
            f"    If lastRow < {min_ok} Then Exit Sub\n")


def _codegen_sort(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """SORT の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    col_idx = headers[first_sheet].index(resolved_args["col"])
    asc = "True" if resolved_args["order"] == "asc" else "False"
    last_col = len(headers[first_sheet]) - 1
    # ★ 合計行が在る回だけ、終わりの行を足して渡す。
    #   ★ 挙動が変わらない回は**生成する文字列も変えない**（凍結した検体・
    #     ゴールデンを理由なく動かさない ── 動かすと差分の意味が薄まる）。
    #   ★ 引数を増やさず**別の腕**を呼ぶ（既存の目録・README・凍結した検体を動かさない）。
    _end = resolved_args.get("_sort_end_row")
    if _end:
        return wrap(f"    Call SortByColumnUpTo(oDoc, {hr0}, {last_col}, {col_idx}, "
                     f"{asc}, {int(_end) - 1})\n")
    return wrap(f"    Call SortByColumn(oDoc, {hr0}, {last_col}, {col_idx}, {asc})\n")


def _codegen_lookup_fill(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """LOOKUP_FILL の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    theaders = headers[resolved_args["target_sheet"]]
    key_idx = theaders.index(resolved_args["key_col"])
    src = resolved_args["source_sheet"].replace('"', '""')
    target_col_name = resolved_args["target_col"]
    # ★ W10c 致命2: target_col が対象シートに実在しない場合（verify_dsl_args が依頼文に
    #   同じ列名があると確認済み）は、COMPUTE_COLUMN の新規列作成と同じ考え方で
    #   末尾に新しい列を作ってから転記する（無関係な既存列を上書きしない）。
    if target_col_name in theaders:
        tgt_idx = theaders.index(target_col_name)
        header_write = ""
    else:
        tgt_idx = len(theaders)   # 0起点・次の空き列
        header_name = str(target_col_name).replace('"', '""')
        header_write = (f'    oDoc.Sheets.getByIndex(0).getCellByPosition({tgt_idx}, {hr0})'
                         f'.setString("{header_name}")\n')
    # ★★ 2026-09-21（盲検 5 体目）: 参照表のどの列を読むかを**名前で**決めて渡す。
    #   長らく「1 列目がキー・2 列目が値」の決め打ちで、買い手の 6 列の発注記録では
    #   転記が使えなかった（道具は「2 列だけの表を用意してください」と言っていた）。
    #   ★★ 位置を割り出す式は**すぐ上の警告（verify_dsl_args）が既に書いていた** ──
    #     正しい列を計算しておきながら、それで書きに行かず「意図と違うかもしれません」と
    #     言うだけだった（器官は在るが配線が無い）。
    #   ★ 見つからなければ従来どおり 0/1 に落ちる（挙動は変えない・警告はそこだけ残す）。
    _sheaders = list(headers.get(resolved_args["source_sheet"], []))
    _lk = _sheaders.index(resolved_args["key_col"]) if resolved_args["key_col"] in _sheaders else 0
    _lv = _sheaders.index(target_col_name) if target_col_name in _sheaders else 1
    return wrap(header_write +
                        f'    Call VLookupFromTable(oDoc, {hr0}, {key_idx}, {tgt_idx}, '
                        f'"{src}", {_lk}, {_lv})\n')


def _codegen_aggregate(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """AGGREGATE の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    g_idx = headers[first_sheet].index(resolved_args["group_col"])
    v_idx = headers[first_sheet].index(resolved_args["value_col"])
    # ★★ 2026-09-07: 合計行を分類に混ぜない ── 除外の宣言（入口で 1 度だけ作る
    #   `_skip_rows`）をそのまま渡す。条件つき書換が既にこの形で受けている。
    _skip = ",".join(str(int(r) - 1) for r in (resolved_args.get("_skip_rows") or []))
    if _skip:
        return wrap(f'    Call SummaryTable(oDoc, {hr0}, {g_idx}, {v_idx}, "{_skip}")\n')
    return wrap(f"    Call SummaryTable(oDoc, {hr0}, {g_idx}, {v_idx})\n")


def _codegen_number_format(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """NUMBER_FORMAT の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    if resolved_args.get("_row_index"):
        _nrow0 = int(resolved_args["_row_index"]) - 1
        _lastc = max(0, len(headers.get(first_sheet) or []) - 1)
        return wrap(f"    Call FormatThousandsRow(oDoc, {_nrow0}, {_lastc})\n")
    col_idx = headers[first_sheet].index(resolved_args["col"])
    return wrap(f"    Call FormatThousands(oDoc, {hr0}, {col_idx})\n")


def _codegen_merge(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """MERGE の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    c1s, r1s, c2s, r2s = re.match(
        r"([A-Za-z]{1,3})(\d+):([A-Za-z]{1,3})(\d+)", resolved_args["range"]).groups()
    col1 = column_index_from_string(c1s.upper()) - 1
    col2 = column_index_from_string(c2s.upper()) - 1
    row1, row2 = int(r1s) - 1, int(r2s) - 1
    return wrap(f"    Call MergeCells(oDoc, {col1}, {row1}, {col2}, {row2})\n")


def _codegen_chart(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """CHART の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    v_idx = headers[first_sheet].index(resolved_args["value_col"])
    # ★ グラフ段: category_col/kind は verify_dsl_args が既定を確定させるが、codegen_dsl
    #   を単体で直接呼ぶ(ゴールデン等)呼び出し元もあるため、ここでも既定を持つ
    #   （value_col と同様の作法・省略時は1列目/"bar"）。
    cat_name = resolved_args.get("category_col") or headers[first_sheet][0]
    cat_idx = headers[first_sheet].index(cat_name)
    kind = resolved_args.get("kind") or "bar"
    # ★ operator10 ①: 合計行をグラフ範囲から除く（片配線の解消）。book_meta にファイル
    #   パスがある（実行時）場合だけ実ファイルを読める ── 手組みの book_meta（単体テスト・
    #   ゴールデン等）では従来どおり最終引数を付けない＝InsertChart が自前走査する旧挙動。
    book_path = book_meta.get("path")
    max_row_arg = ""
    if book_path is not None:
        try:
            last_row = chart_data_last_row(book_path, first_sheet, header_row)
            max_row_arg = f", {last_row - 1}"   # Basic 0起点
        except Exception:
            max_row_arg = ""
    return wrap(f'    Call InsertChart(oDoc, {hr0}, {cat_idx}, {v_idx}, "{kind}"{max_row_arg})\n')


def _codegen_append_total(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """APPEND_TOTAL の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ W6: ヘルパは無し（罫線・カンマ等の見栄えまでは踏み込まない・素の SUM 式だけ）。
    #   データ最終行の直下に [ラベル文字列 | 合計式] を書く。
    #   ラベルは対象列の左隣に置く（既存の帳票では『合計』の文字が金額の左に来るのが自然
    #   で、対象列自体を上書きしない＝既存構造を壊さない置き方）。対象列が表の最左端
    #   （col_idx=0）の場合は左隣が無いため、値のみを書きラベルは省略する。
    # ★ B: 挿入耐性式（bench/formula_spike_work2 で実測）。SUM(D2:INDEX(D:D;ROW()-1))
    #   型で書く。INDEX(D:D;ROW()-1) は「この式自身の1行上」を指すので、後でデータ行を
    #   1本挿入しても SUM 範囲が自動で追従する（静的な "D2:D5" 型は追従しない）。
    #   ★ setFormula は LO 方言＝INDEX の2引数区切りはセミコロン(;)。カンマ(,)で
    #   書くと #VALUE!/#NAME? になる（実測: formula_spike_RESULTS.md 追記分）。
    #   保存後は自動でカンマ形 "=SUM(D2:INDEX(D:D,ROW()-1))*1.1" に変換される
    #   （check_append_total はこの保存後カンマ形と照合する）。
    col_idx = headers[first_sheet].index(resolved_args["col"])
    label = str(resolved_args.get("label", "合計")).replace('"', '""')
    factor = float(resolved_args.get("factor", 1) or 1)
    col_letter = get_column_letter(col_idx + 1)
    start_excel_row = hr0 + 2   # データ先頭行（Basic 0起点 hr0+1）の Excel(1起点) 行
    factor_tail = "" if factor == 1 else f"*{factor:g}"
    _fixed_row = resolved_args.get("_at_row")
    if _fixed_row:
        # ★ 既にある合計行に書く（行は増やさない・ラベルはその行に既に在る）。
        body = ("    Dim oSheet As Object, totalRow As Long\n"
                 "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                 "    totalRow = " + str(int(_fixed_row) - 1) + "\n")
    else:
        body = ("    Dim oSheet As Object, lastRow As Long, totalRow As Long\n"
                "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                + _scan_last_row_basic(start_row=str(hr0 + 1))
                + "    totalRow = lastRow + 1\n")
    if col_idx > 0 and not _fixed_row:
        body += f'    oSheet.getCellByPosition(0, totalRow).setString("{label}")\n'
    body += (f'    oSheet.getCellByPosition({col_idx}, totalRow).setFormula('
             f'"=SUM(" & "{col_letter}" & {start_excel_row} & ":INDEX(" & "{col_letter}" & '
             f'":" & "{col_letter}" & ";ROW()-1))" & "{factor_tail}")\n')
    # ★ 2026-09-13（買い手役 2 回とも）: 合計のセルだけ `General` で、上の #,##0 の列の下に
    #   `1719200` が来ていた ── 上司に出す紙で一番目が行く 1 マス。表示形式は 1 つ上の行から写す。
    body += (f'    oSheet.getCellByPosition({col_idx}, totalRow).NumberFormat = '
             f'oSheet.getCellByPosition({col_idx}, totalRow - 1).NumberFormat\n')
    return wrap(body)


def _codegen_center_align(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """CENTER_ALIGN の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    if str(resolved_args["target"]).startswith("cell:"):
        _r, _c = (int(x) for x in str(resolved_args["target"])[5:].split(","))
        return wrap('    oDoc.Sheets.getByIndex(0).getCellByPosition(%d, %d)'
                     '.HoriJustify = com.sun.star.table.CellHoriJustify.CENTER%s'
                     % (_c - 1, _r - 1, chr(10)))
    if resolved_args["target"] == "all":
        last_col = len(headers[first_sheet]) - 1
        return wrap(f"    Call AlignCenter(oDoc, {hr0}, {last_col})\n")
    # col:NAME はヘルパ無し → refs の作法（走査して範囲を求め HoriJustify）でテンプレを書く。
    col_idx = headers[first_sheet].index(resolved_args["target"][4:])
    body = ("    Dim oSheet As Object, oRange As Object, lastRow As Long\n"
            "    oSheet = oDoc.Sheets.getByIndex(0)\n"
            + _scan_last_row_basic(start_row=str(hr0 + 1), min_ok=str(hr0))
            + f"    oRange = oSheet.getCellRangeByPosition({col_idx}, {hr0}, {col_idx}, lastRow)\n"
            "    oRange.HoriJustify = com.sun.star.table.CellHoriJustify.CENTER\n")
    return wrap(body)


def _codegen_bold(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """BOLD の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    target = resolved_args["target"]
    if target.startswith("cell:"):
        _r, _c = (int(x) for x in target[5:].split(","))
        return wrap("    Call StyleBold(oDoc, %d, %d, %d, %d)%s"
                     % (_c - 1, _r - 1, _c - 1, _r - 1, chr(10)))
    if target.startswith("row:"):
        row_idx = int(target[4:]) - 1
        # ★ W3: 列幅は Basic で走査せず、接地済みの見出し列数(headers)から決定論的に
        #   決める（多段見出しで先頭列が空欄のケースで走査が誤って -1 になるのを回避）。
        last_col = len(headers[first_sheet]) - 1
        body = (f"    Call StyleBold(oDoc, 0, {row_idx}, {last_col}, {row_idx})\n")
    else:
        col_idx = headers[first_sheet].index(target[4:])
        body = ("    Dim oSheet As Object, lastRow As Long\n"
                "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                + _scan_last_row_basic(start_row=str(hr0 + 1), min_ok=str(hr0))
                + f"    Call StyleBold(oDoc, {col_idx}, {hr0}, {col_idx}, lastRow)\n")
    return wrap(body)


def _codegen_fill_color(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """FILL_COLOR の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    target = resolved_args["target"]
    hexcolor = COLOR_MAP[resolved_args["color"]]
    if target.startswith("cell:"):
        _r, _c = (int(x) for x in target[5:].split(","))
        return wrap('    oDoc.Sheets.getByIndex(0).getCellByPosition(%d, %d)'
                     '.CellBackColor = &H%s&%s' % (_c - 1, _r - 1, hexcolor, chr(10)))
    if target.startswith("row:"):
        row_idx = int(target[4:]) - 1
        # ★ W3: 列幅は Basic で走査せず、接地済みの見出し列数(headers)から決定論的に
        #   決める（多段見出しで先頭列が空欄のケースで走査が誤って -1 になるのを回避）。
        last_col = len(headers[first_sheet]) - 1
        body = ("    Dim oSheet As Object, c As Integer\n"
                "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                f"    For c = 0 To {last_col}\n"
                f"        oSheet.getCellByPosition(c, {row_idx}).CellBackColor = &H{hexcolor}&\n"
                "    Next c\n")
    else:
        col_idx = headers[first_sheet].index(target[4:])
        body = ("    Dim oSheet As Object, lastRow As Long, r As Long\n"
                "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                + _scan_last_row_basic(start_row=str(hr0 + 1), min_ok=str(hr0))
                + f"    For r = {hr0} To lastRow\n"
                f"        oSheet.getCellByPosition({col_idx}, r).CellBackColor = &H{hexcolor}&\n"
                "    Next r\n")
    return wrap(body)


def _codegen_compute_column(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """COMPUTE_COLUMN の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    operands = resolved_args["operands"]
    operator = resolved_args["operator"]
    target = resolved_args.get("target")

    if len(operands) == 1:
        # ★ W10b 項目3: 税込み/税抜き等「列 × 率」パターン（1列 + 機械確定した factor）。
        #   codegen は既存の2列版と同じ書き方（式=セル参照、値ベタ書きの選択も同じ）。
        op1 = operands[0]
        i1 = headers[first_sheet].index(op1)
        factor = float(resolved_args["factor"])
        if target:
            new_col = headers[first_sheet].index(target)
            header_write = ""
        else:
            new_col = len(headers[first_sheet])
            # ★ W10c 中: verify_dsl_args が税込み/税抜きと判定できていれば自然な
            #   日本語見出し（例:「税込金額」）を使う。無ければ従来どおりの数式風見出し。
            header_name = str(resolved_args.get("_new_col_label")
                               or f"{op1}{operator}{factor:g}").replace('"', '""')
            header_write = f"    oSheet.getCellByPosition({new_col}, {hr0}).setString(\"{header_name}\")\n"
        if use_formula:
            col1_letter = get_column_letter(i1 + 1)
            write_line = (f'        oSheet.getCellByPosition({new_col}, i).setFormula('
                          f'"=" & "{col1_letter}" & (i + 1) & "{operator}{factor:g}")\n')
        else:
            write_line = (f"        oSheet.getCellByPosition({new_col}, i).setValue("
                          f"oSheet.getCellByPosition({i1}, i).getValue() {operator} {factor:g})\n")
        body = ("    Dim oSheet As Object, lastRow As Long, i As Long\n"
                "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                + _scan_last_row_basic(start_row=str(hr0 + 1))
                + header_write
                + f"    For i = {hr0 + 1} To lastRow\n"
                + write_line
                + "    Next i\n")
        return wrap(body)

    op1, op2 = operands
    i1 = headers[first_sheet].index(op1)
    i2 = headers[first_sheet].index(op2)
    # ★ M2c: target(実在列名) 指定時はその列に書く（新規列を作らない）。
    #   無指定なら従来どおり次の空き列に新規列を作る。
    if target:
        new_col = headers[first_sheet].index(target)
        header_write = ""   # 既存の見出しはそのまま（上書きしない）
    else:
        new_col = len(headers[first_sheet])   # 0起点で次の空き列
        # ★ W3 改定(2026-08-20): verify_dsl_args が「target は実在しないが依頼文に
        #   実在する語」と判定していれば _new_col_label に利用者の指名が入っている
        #   （単列×率パターンの税込/税抜ラベルと同じ仕組み）。無ければ従来どおりの
        #   数式風見出し。
        header_name = str(resolved_args.get("_new_col_label")
                           or f"{op1}{operator}{op2}").replace('"', '""')
        header_write = f"    oSheet.getCellByPosition({new_col}, {hr0}).setString(\"{header_name}\")\n"
    # ★ W3 Part3: 既定は式（setFormula）。formula_spike の実測どおり、単純な行内の
    #   二項演算（=B2*C2 型）は区切り記号もシート参照も無いため LO 方言(;/.)の
    #   気遣いが不要（多引数関数・シート跨ぎ参照だけが方言の対象＝ここでは無縁）。
    #   行番号だけが Basic 側ループ変数(i)で動くので、その部分だけ実行時に連結する。
    if use_formula:
        col1_letter = get_column_letter(i1 + 1)
        col2_letter = get_column_letter(i2 + 1)
        write_line = (f'        oSheet.getCellByPosition({new_col}, i).setFormula('
                      f'"=" & "{col1_letter}" & (i + 1) & "{operator}" & "{col2_letter}" & (i + 1))\n')
    else:
        write_line = (f"        oSheet.getCellByPosition({new_col}, i).setValue("
                      f"oSheet.getCellByPosition({i1}, i).getValue() {operator} "
                      f"oSheet.getCellByPosition({i2}, i).getValue())\n")
    body = ("    Dim oSheet As Object, lastRow As Long, i As Long\n"
            "    oSheet = oDoc.Sheets.getByIndex(0)\n"
            + _scan_last_row_basic(start_row=str(hr0 + 1))
            + header_write
            + f"    For i = {hr0 + 1} To lastRow\n"
            + write_line
            + "    Next i\n")
    return wrap(body)


    # --- ★ W9: 検証済みヘルパ4種の語彙昇格。いずれも helpers/*.bas のヘルパは headerRow
    #   引数を取らない（物理1行目を前提に自前走査する既存実装・ここでは変更しない）ため、
    #   codegen 側も hr0 を渡さずそのまま Call するだけ。
def _codegen_insert_rows(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """INSERT_ROWS の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    at0 = int(resolved_args["at"]) - 1   # 1起点(Excel行番号) → 0起点(Basic)
    count = int(resolved_args.get("count", 1) or 1)
    return wrap(f"    Call InsertRows(oDoc, {at0}, {count})\n")


    # ★ 2026-08-26: 表の基本操作 3 種（Namakoo が GUI を触って欠けを実測）。
    #   ★ insertByIndex / removeByIndex を使う ── 挿入は**途中でも押し下げ**、
    #     削除は**詰める**（clearContents だと空行が残る）。
def _codegen_add_row(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """ADD_ROW の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    at0 = int(resolved_args["at"]) - 1
    values = resolved_args.get("values") or {}
    headers = list(resolved_args.get("_headers") or [])
    idx, vals, kinds = [], [], []
    for name, v in values.items():
        if name not in headers:
            continue
        idx.append(str(headers.index(name)))
        # ★ 型は書く値そのものから決める。数値を文字列で書くと下流の SUM が壊れる
        #   （この repo が何度も測ってきた「静かに壊れる」形）。
        if isinstance(v, bool):
            vals.append("TRUE" if v else "FALSE"); kinds.append("s")
        elif isinstance(v, (int, float)):
            vals.append(repr(v)); kinds.append("n")
        else:
            vals.append(str(v)); kinds.append("s")
    sep = chr(1)
    body = (
        f'    Call AddRowWithValues(oDoc, {at0}, "{",".join(idx)}", '
        f'"{sep.join(vals)}", "{",".join(kinds)}")\n')
    # ★★ 2026-09-02: 既存の行が式で出している列は、新しい行にも式を写す。
    #   ★ 値は作らない（A' 原則）── **隣の行から写すだけ**。
    #     参照の付け替えは LibreOffice にやらせる（自前で式を書き換えない）。
    _inh = list(resolved_args.get("_inherit_cols") or [])
    _src = int(resolved_args.get("_inherit_from") or 0)
    if _inh and _src:
        body += ('    Call FillFormulasFromNeighbour(oDoc, %d, %d, "%s")%s'
                  % (at0, _src - 1, ",".join(str(c) for c in _inh), chr(10)))
    return wrap(body)


def _codegen_delete_rows(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """DELETE_ROWS の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★★ 2026-09-07: 名前が複数行に当たった回は、その行を**全部**消す。
    #   ★ **下から順に**呼ぶ ── 上から消すと、その下の行番号がずれる。
    #   ★ Basic 側は変えない（DeleteRows を複数回呼ぶだけ）── 新しいヘルパを作らない。
    if (_rows := list(resolved_args.get("_delete_rows") or [])):
        body = "".join(f"    Call DeleteRows(oDoc, {int(r) - 1}, 1)" + chr(10)
                       for r in sorted(_rows, reverse=True))
        return wrap(body)
    at0 = int(resolved_args["at"]) - 1
    count = int(resolved_args.get("count", 1) or 1)
    return wrap(f"    Call DeleteRows(oDoc, {at0}, {count})\n")


def _codegen_dedup_delete(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """DEDUP_DELETE の Basic ── 中身は DELETE_ROWS と**同じ**（`_delete_rows` を下から順に消す）。

    ★ 実装は 1 本（`_codegen_delete_rows`）。ここは名簿の作法（`_codegen_<op>`）と
      引数の形を守るためだけの薄い配線 ── 「下から順に消す」を書き写さない。
    ★ 引数を `**kw` で受けない: 番人が「かつてローカルだったものをグローバルで取りに
      行っていないか」を署名で見ている（test_codegen_is_a_table）。
    """
    return _codegen_delete_rows(
        op=op, resolved_args=resolved_args, book_meta=book_meta, use_formula=use_formula,
        headers=headers, first_sheet=first_sheet, header_row=header_row, hr0=hr0, wrap=wrap)


def _codegen_delete_column(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """DELETE_COLUMN の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    headers = list(resolved_args.get("_headers") or [])
    col0 = headers.index(resolved_args["col"]) if resolved_args["col"] in headers else 0
    return wrap(f"    Call DeleteColumn(oDoc, {col0})\n")


def _codegen_move_column(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """MOVE_COLUMN の Basic を組む。★ 列を動かす手は**横断層に 1 本だけ**在り、
       新しい列の配置もこの op も同じ口を通る（列を動かす実装を 2 つ持たない）。"""
    # ★ 本文は空 ── 列を動かす Call は横断層（wrap）が 1 箇所で足す。
    #   ここで書くと「移動の配線が 2 箇所」になり、番人が赤くなる（それが正しい）。
    return wrap("")


def _codegen_extract_columns(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """EXTRACT_COLUMNS の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    hdrs = list(resolved_args.get("_headers") or [])
    idx = ",".join(str(hdrs.index(c)) for c in resolved_args["cols"] if c in hdrs)
    hr0x = int(resolved_args.get("_header_row", 1)) - 1
    dst = str(resolved_args["_new_sheet"]).replace(chr(34), chr(34) * 2)
    return wrap('    Call ExtractColumns(oDoc, %d, "%s", "%s")%s'
                 % (hr0x, idx, dst, chr(10)))


def _codegen_set_where(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """SET_WHERE の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    headers = list(resolved_args.get("_headers") or [])
    wcol = headers.index(resolved_args["col"]) if resolved_args["col"] in headers else 0
    ccol = (headers.index(resolved_args["cond_col"])
             if resolved_args["cond_col"] in headers else 0)
    hr0 = int(resolved_args.get("_header_row", 1)) - 1
    code = _EXTRACT_CMP_CODE[resolved_args["cmp"]]
    thr = resolved_args["cond_value"]
    # ★ 2026-08-27（実測・生 traceback を出した）: 閾値が数値かどうかは **cmp ではなく
    #   値そのもの**で決まる。eq は「金額が 100 と等しい」にも「チェックが『◎』と
    #   等しい」にも使う ── cmp で分けると、置き換えの『◎』を float() に渡して落ちる。
    # ★ 2026-09-05: nin/in は**値の一覧**。EXTRACT と同じ Chr(2) 区切りで渡す
    #   （Basic の RowMatches Case 6/7 が区切りごとの丸ごと一致で見る）。
    if isinstance(thr, (list, tuple, set)):
        _joined = chr(2).join(str(v) for v in thr)
        thr_lit = ('"' + _joined.replace(chr(34), chr(34) * 2)
                          .replace(chr(2), '" & Chr(2) & "') + '"')
    else:
        _num = _is_number(thr) and not isinstance(thr, bool)
        thr_lit = (repr(float(thr)) if _num
                    else '"%s"' % str(thr).replace(chr(34), chr(34) * 2))
    val = str(resolved_args["value"]).replace(chr(34), chr(34) * 2)
    # ★ 外す行は**構造の事実**なので Python が渡す（条件の判定は Basic が自分で行う ──
    #   そこを渡すと事後条件が独立した検算でなくなる）。
    skip = ",".join(str(int(r) - 1) for r in (resolved_args.get("_skip_rows") or []))
    # ★★ 2026-09-06: 2 組目の条件（AND）。無ければ従来どおり引数を足さない ──
    #   Basic 側は Optional なので、既存の呼び出しは 1 文字も変わらない。
    #   ★ 述語（RowMatches）は無傷のまま **2 回呼ぶ**だけ（凍結した真理表を触らない）。
    c2 = resolved_args.get("cond2_col")
    if not c2:
        return wrap('    Call SetColumnValueWhere(oDoc, %d, %d, %d, %d, %s, "%s", "%s")%s'
                     % (hr0, wcol, ccol, code, thr_lit, val, skip, chr(10)))
    # ★ 列の番号は 1 組目と**同じ引き方**にする（headers はこの op ではその表の列名の並び）
    c2col = headers.index(c2) if c2 in headers else 0
    code2 = _EXTRACT_CMP_CODE[resolved_args["cond2_cmp"]]
    thr2 = resolved_args["cond2_value"]
    if isinstance(thr2, (list, tuple, set)):
        _j2 = chr(2).join(str(v) for v in thr2)
        thr2_lit = ('"' + _j2.replace(chr(34), chr(34) * 2)
                            .replace(chr(2), '" & Chr(2) & "') + '"')
    else:
        _n2 = _is_number(thr2) and not isinstance(thr2, bool)
        thr2_lit = (repr(float(thr2)) if _n2
                     else '"%s"' % str(thr2).replace(chr(34), chr(34) * 2))
    return wrap('    Call SetColumnValueWhere(oDoc, %d, %d, %d, %d, %s, "%s", "%s", %d, %d, %s)%s'
                 % (hr0, wcol, ccol, code, thr_lit, val, skip,
                    c2col, code2, thr2_lit, chr(10)))


def _codegen_add_column(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """ADD_COLUMN の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ 位置は Python が見出しから決めた 1 起点 → Basic は 0 起点。
    #   名前は Basic 側で書く（空なら見出しも空のまま）。
    at0 = int(resolved_args["_at_col"]) - 1
    hr0 = int(resolved_args.get("_header_row", 1)) - 1
    nm = str(resolved_args.get("name") or "").replace(chr(34), chr(34) * 2)
    return wrap('    Call InsertColumnAt(oDoc, %d, "%s", %d)%s' % (at0, nm, hr0, chr(10)))


def _codegen_swap(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """SWAP の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ 位置は Basic 側にも**名前で**渡す（Python が数えた番号を渡さない）。
    #   Basic は実文書を走査して自分で見つける ── Python の解決と食い違えば、
    #   事後条件が「宣言した位置が入れ替わっていない」と落とす（独立な 2 実装）。
    hr0 = int(resolved_args.get("_header_row", 1)) - 1
    _a = str(resolved_args["a"]).replace(chr(34), chr(34) * 2)
    _b = str(resolved_args["b"]).replace(chr(34), chr(34) * 2)
    if resolved_args.get("_axis") == "cell":
        # ★ 2 セルの入れ替え。中身は verify が実表から読んだもの（LLM 由来ではない）。
        #   ★ 式のセルは setFormula で書く（setString だと文字列になる）。
        _cs = resolved_args["_cells"]
        _vs = resolved_args["_cell_values"]
        _body = ""
        for (r, c), v in zip(_cs, reversed(_vs)):     # ★ 入れ替え＝相手の値を書く
            if isinstance(v, str) and v.startswith("="):
                _esc = formula_for_basic(v).replace(chr(34), chr(34) * 2)
                _body += ('    Call SetFormulaAt(oDoc, %d, %d, "%s")%s'
                           % (r - 1, c - 1, _esc, chr(10)))
            elif _is_number(v):
                _body += ('    Call SetCellAt(oDoc, %d, %d, "%s", "n")%s'
                           % (r - 1, c - 1, repr(float(v)), chr(10)))
            else:
                _esc = str("" if v is None else v).replace(chr(34), chr(34) * 2)
                _body += ('    Call SetCellAt(oDoc, %d, %d, "%s", "s")%s'
                           % (r - 1, c - 1, _esc, chr(10)))
        return wrap(_body)
    if resolved_args.get("_axis") == "column":
        _body = ('    Call SwapColumnsByName(oDoc, "%s", "%s", %d)%s'
                  % (_a, _b, hr0, chr(10)))
    else:
        # ★★ 2026-08-31: 行番号で指された回は**座標で**渡す。名前で渡すと Basic は
        #   表から『6行目』という名前を探して見つけられない（実測で別の行が動いた）。
        #   ★ 名前で指された回は今までどおり名前を渡す ── あちらは Basic が自分で
        #     位置を見つけるので、事後条件が独立な検算になる。
        if (_re_row_number_word.fullmatch(_a.strip())
                and _re_row_number_word.fullmatch(_b.strip())):
            _body = ('    Call SwapRowsAt(oDoc, %d, %d)%s'
                      % (resolved_args["_a_pos"] - 1, resolved_args["_b_pos"] - 1,
                         chr(10)))
        else:
            _body = ('    Call SwapRowsByName(oDoc, "%s", "%s", 0, %d)%s'
                      % (_a, _b, hr0, chr(10)))
    # ★ 入れ替えの**あと**に、写像を通した式を書き戻す（順序が意味を持つ）。
    for _r, _c, _f in resolved_args.get("_formula_rewrites") or ():
        _esc = formula_for_basic(_f).replace(chr(34), chr(34) * 2)
        _body += ('    Call SetFormulaAt(oDoc, %d, %d, "%s")%s'
                   % (_r - 1, _c - 1, _esc, chr(10)))
    return wrap(_body)


def _codegen_set_cell_value(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """SET_CELL_VALUE の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    headers = list(resolved_args.get("_headers") or [])
    col0 = headers.index(resolved_args["col"]) if resolved_args["col"] in headers else 0
    v = resolved_args["value"]
    kind = "n" if resolved_args.get("_write_numeric") else "s"
    val = (repr(float(resolved_args["_write_numeric_value"]))
            if kind == "n" else str(v))
    hr0 = int(resolved_args.get("_header_row", 1)) - 1
    # ★ 2026-08-28: 行番号で指された回は**座標で**書く（探し直す相手が無い・同名の
    #   行が 2 つある表でも狙いが定まる）。名前で指された回は今までどおり名前を渡す
    #   ── あちらは Basic が自分で位置を見つけるので、事後条件が独立な検算になる。
    if resolved_args.get("row_number"):
        return wrap('    Call SetCellAt(oDoc, %d, %d, "%s", "%s")%s'
                     % (int(resolved_args["row_number"]) - 1, col0, val, kind, chr(10)))
    return wrap(
        f'    Call SetCellByName(oDoc, "{resolved_args["row"]}", 0, {col0}, '
        f'"{val}", "{kind}", {hr0})\n')


def _codegen_draw_borders(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """DRAW_BORDERS の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    return wrap("    Call DrawTableBorders(oDoc)\n")


def _codegen_autofit(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """AUTOFIT の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    return wrap("    Call AutoFitColumns(oDoc)\n")


def _codegen_pivot(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """PIVOT の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    g_idx = headers[first_sheet].index(resolved_args["group_col"])
    v_idx = headers[first_sheet].index(resolved_args["value_col"])
    return wrap(f"    Call PivotSum(oDoc, {g_idx}, {v_idx})\n")


def _codegen_set_column_value(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """SET_COLUMN_VALUE の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ 致命3(W10e): ヘルパ無し・既存列のデータ行全部に同じ値を書く
    #   （CENTER_ALIGN の col: 分岐と同じ「走査してヘッダ直下から最終行まで」の作法）。
    # ★ operator10 ④: verify_dsl_args が列の実体（書換前）から機械決定した型
    #   （resolved_args["_write_numeric"]）に従い、数値列には setValue で数値を書く
    #   （既定/判定できない場合は従来どおり setString）。
    col_idx = headers[first_sheet].index(resolved_args["col"])
    body_head = ("    Dim oSheet As Object, lastRow As Long, r As Long\n"
                 "    oSheet = oDoc.Sheets.getByIndex(0)\n"
                 + _scan_last_row_basic(start_row=str(hr0 + 1)))
    if resolved_args.get("_write_numeric"):
        num = float(resolved_args["_write_numeric_value"])
        num_lit = str(int(num)) if num.is_integer() else repr(num)
        body = (body_head
                + f"    For r = {hr0 + 1} To lastRow\n"
                f"        oSheet.getCellByPosition({col_idx}, r).setValue({num_lit})\n"
                "    Next r\n")
    else:
        value = str(resolved_args["value"]).replace('"', '""')
        body = (body_head
                + f"    For r = {hr0 + 1} To lastRow\n"
                f"        oSheet.getCellByPosition({col_idx}, r).setString(\"{value}\")\n"
                "    Next r\n")
    return wrap(body)


def _codegen_extract(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """EXTRACT の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ ヘルパへの Call 1行だけ（helpers/AiLineHelpers.bas:ExtractRows）。
    #   型を保つコピー（getValue/setValue・getString/setString の分岐）は helper 側。
    col_idx = headers[first_sheet].index(resolved_args["col"])
    cmp_code = _EXTRACT_CMP_CODE[resolved_args["cmp"]]
    # ★ 日付比較のときだけ、表示用の文字列でなくシリアル値を Basic へ渡す
    #   （resolved["value"] は解釈行と出力シート名のために元の文字列で残してある）。
    value = resolved_args.get("_value_serial", resolved_args["value"])
    if isinstance(value, (list, tuple)):
        # ★「どれか」の一覧は Chr(2) 区切り（値にカンマ・タブが入りうる）。
        #   Basic 側 RowMatches Case 6 が同じ区切りで**丸ごと一致**を見る。
        joined = chr(2).join(str(v) for v in value)
        value_lit = '"' + joined.replace('"', '""').replace(chr(2), '" & Chr(2) & "') + '"'
    elif isinstance(value, str):
        value_lit = '"' + value.replace('"', '""') + '"'
    else:
        value_lit = f"{float(value):g}"
    dst_name = str(resolved_args["_new_sheet"]).replace('"', '""')
    # ★ 外す行は**構造の事実**なので Python が渡す（条件の判定は Basic が自分で行う
    #   ── そこを渡すと事後条件が独立した検算でなくなる）。SET_WHERE と同じ作法。
    # ★ Basic のループは **0 起点**（他のヘルパと同じ）── 1 起点の行番号を渡すと
    #   1 行ずれて効かない（実測: 「8」を渡したのに 8 行目が抜き出された）。
    #   SetColumnValueWhere も同じく -1 して渡している。
    _x_skip = ",".join(str(int(r) - 1) for r in (resolved_args.get("_skip_rows") or []))
    return wrap(f'    Call ExtractRows(oDoc, {hr0}, {col_idx}, {cmp_code}, '
                 f'{value_lit}, "{dst_name}", "{_x_skip}")\n')


def _codegen_split_cell(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """SPLIT_CELL の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ ヘルパへの Call 1行だけ（helpers/AiLineHelpers.bas:SplitColumn）。
    #   新しい見出し名は Python 側で決めて渡す（Basic 側で名前を作らない・帳票段と同じ作法）。
    col_idx = headers[first_sheet].index(resolved_args["col"])
    sep_lit = '"' + str(resolved_args["sep"]).replace('"', '""').replace(chr(10), '" & Chr(10) & "') + '"'
    names_csv = ",".join(str(n).replace(",", "_") for n in resolved_args["_new_cols"])
    return wrap(f'    Call SplitColumn(oDoc, {hr0}, {col_idx}, {sep_lit}, "{names_csv}")' + chr(10))


def _codegen_dedup(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """DEDUP の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ ヘルパへの Call 1行だけ（helpers/AiLineHelpers.bas:DedupRows・EXTRACT と同じ作法）。
    #   キー列は複数ありうるため、0起点の列インデックスをカンマ区切りの文字列で渡す
    #   （Basic の Call 引数に配列リテラルを直接書けないため・ヘルパ側で Split する）。
    key_idxs = [headers[first_sheet].index(k) for k in resolved_args["keys"]]
    key_idx_csv = ",".join(str(i) for i in key_idxs)
    dst_name = str(resolved_args["_new_sheet"]).replace('"', '""')
    return wrap(f'    Call DedupRows(oDoc, {hr0}, "{key_idx_csv}", "{dst_name}")\n')


def _codegen_report_per_row(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """REPORT_PER_ROW の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ ヘルパへの Call を行数ぶん（helpers/AiLineHelpers.bas:FillReportSheet）。
    #   ★ B: 呼ぶ前に一意名を解決済み（verify_dsl_args の unique_sheet_name）── 失敗
    #   しうる名前を copyByName に渡さない。ループは Python 側（Basic 側でシート名を
    #   作らない・設計文書の指示どおり「名前は Python 側が全部決めてから渡す」）。
    template_sheet = str(resolved_args["template_sheet"]).replace('"', '""')
    src_sheet = str(first_sheet).replace('"', '""')
    lines = []
    groups = resolved_args.get("_groups")
    if groups:
        # ★★ まとめ版（2026-08-28）: 1 グループ = 1 枚。明細行が件数ぶん増える。
        #   行を増やすと下がずれる ── ずれの数え方は report_group.output_rows_for に
        #   1 箇所だけ置き、確かめる側もそこを使う（埋める側と数え方が割れない）。
        det0 = resolved_args.get("_detail_row")
        det0 = -1 if det0 is None else int(det0) - 1
        for g in groups:
            gname = str(g["sheet"]).replace('"', '""')
            rows_csv = ",".join(str(int(r) - 1) for r in g["rows"])
            lines.append(
                f'    Call FillGroupReportSheet(oDoc, "{template_sheet}", "{gname}", '
                f'"{src_sheet}", "{rows_csv}", {hr0}, {det0})' + chr(10))
    for rr in ([] if groups else resolved_args.get("_report_rows", [])):
        new_name = str(rr["sheet"]).replace('"', '""')
        src_row0 = int(rr["row"]) - 1   # Excel(1起点) → Basic(0起点)
        lines.append(
            f'    Call FillReportSheet(oDoc, "{template_sheet}", "{new_name}", '
            f'"{src_sheet}", {src_row0}, {hr0})\n'
        )
    # ★ 2026-08-24（土台固め）: 検分シートも**同じ Basic の中で**書く。
    #   旧実装は生成後に openpyxl でブックを開き直して足しており、その往復が
    #   xl/drawings の中の図形（描かれた角印・社判）を全部捨てていた（実測）。
    #   LO 側で書けば往復そのものが無くなる ── 追加の LO 起動も要らない。
    insp = resolved_args.get("_inspection_sheet")
    report_rows = resolved_args.get("_report_rows") or []
    if insp and groups:
        # ★ まとめた回は「元の行」が複数 ── 何行をまとめたかまで検分に出す
        #   （どの発注が 1 枚に入ったかを、後から人が追えること）。
        lines.append(inspection_sheet_basic_call(
            insp, ["シート名", "元の行", "まとめた件数"],
            [[g["sheet"], ",".join(str(r) for r in g["rows"]), len(g["rows"])]
             for g in groups], "ssn"))
    elif insp and report_rows:
        n_ph = len(resolved_args.get("_placeholders") or [])
        lines.append(inspection_sheet_basic_call(
            insp, ["シート名", "元の行", "埋めた印の数"],
            [[rr["sheet"], rr["row"], n_ph] for rr in report_rows], "snn"))
    return wrap("".join(lines))


def _codegen_format_map(*, op, resolved_args, book_meta, use_formula, headers, first_sheet,
                 header_row, hr0, wrap):
    """FORMAT_MAP の Basic を組む（★ codegen_dsl から**本文をそのまま**移した）。"""
    # ★ ヘルパへの Call 1行だけ（helpers/AiLineHelpers.bas:FillFormatMapSheet）。
    #   REPORT_PER_ROW と同じ理由: 出力シート名は呼ぶ前に一意名を解決済み。印の実在は
    #   verify_dsl_args が検証済み・行の一覧(_data_rows)も Python 側で決め切る
    #   （Basic 側でシート名・行の取捨選択をしない）。
    template_sheet = str(resolved_args["template_sheet"]).replace('"', '""')
    src_sheet = str(first_sheet).replace('"', '""')
    dst_sheet = str(resolved_args["_output_sheet"]).replace('"', '""')
    header_tpl_row0 = int(resolved_args["_header_tpl_row"]) - 1
    ph_tpl_row0 = int(resolved_args["_placeholder_tpl_row"]) - 1
    src_rows_csv = ",".join(str(int(r) - 1) for r in resolved_args.get("_data_rows", []))
    # ★ 2026-08-24（土台固め）: 検分シートも同じ Basic の中で書く（帳票段と同じ処置・
    #   同じ関数）。openpyxl の往復が図形を捨てるので、往復自体を無くす。
    fm_calls = (
        f'    Call FillFormatMapSheet(oDoc, "{template_sheet}", "{src_sheet}", "{dst_sheet}", '
        f'{header_tpl_row0}, {ph_tpl_row0}, {hr0}, "{src_rows_csv}")' + chr(10))
    insp = resolved_args.get("_inspection_sheet")
    data_rows = resolved_args.get("_data_rows") or []
    if insp and data_rows:
        n_ph = len(resolved_args.get("_placeholders") or [])
        out_sheet = resolved_args.get("_output_sheet")
        fm_calls += inspection_sheet_basic_call(
            insp, ["出力シート", "出力行", "元の行", "埋めた印の数"],
            [[out_sheet, k + 1, src_row, n_ph]
             for k, src_row in enumerate(data_rows, start=1)], "snnn")
    return wrap(fm_calls)


#: op → その op の Basic を組む関数（★ 足すのはここ 1 行 ── 分岐を書き足さない）
#: ★★ 2026-09-05: codegen_dsl は 607 行のうち 536 行が `op == "…"` の 29 分岐だった。
#:   `verify_dsl_args` を 641 件凍結して 1,735 → 227 行に割ったのと**同じ形**。
#:   こちらは**順序に意味が無い**（op で振り分けるだけ）ので、切り出しでなく
#:   **表**まで畳める。ゴールデン 78 本（29 op 全部）が 1 文字の変化も許さない。
CODEGEN_BY_OP = {
    "SORT": _codegen_sort,
    "LOOKUP_FILL": _codegen_lookup_fill,
    "AGGREGATE": _codegen_aggregate,
    "NUMBER_FORMAT": _codegen_number_format,
    "MERGE": _codegen_merge,
    "CHART": _codegen_chart,
    "APPEND_TOTAL": _codegen_append_total,
    "CENTER_ALIGN": _codegen_center_align,
    "BOLD": _codegen_bold,
    "FILL_COLOR": _codegen_fill_color,
    "COMPUTE_COLUMN": _codegen_compute_column,
    "INSERT_ROWS": _codegen_insert_rows,
    "ADD_ROW": _codegen_add_row,
    "DELETE_ROWS": _codegen_delete_rows,
    "DELETE_COLUMN": _codegen_delete_column,
    "MOVE_COLUMN": _codegen_move_column,
    "EXTRACT_COLUMNS": _codegen_extract_columns,
    "SET_WHERE": _codegen_set_where,
    "ADD_COLUMN": _codegen_add_column,
    "SWAP": _codegen_swap,
    "SET_CELL_VALUE": _codegen_set_cell_value,
    "DRAW_BORDERS": _codegen_draw_borders,
    "AUTOFIT": _codegen_autofit,
    "PIVOT": _codegen_pivot,
    "SET_COLUMN_VALUE": _codegen_set_column_value,
    "EXTRACT": _codegen_extract,
    "SPLIT_CELL": _codegen_split_cell,
    "DEDUP": _codegen_dedup,
    # ★ 削除の生成は DELETE_ROWS と**同じ関数**（`_delete_rows` を下から順に消す）。
    #   op ごとに書き写すと、片方だけ直る日が来る。
    "DEDUP_DELETE": _codegen_dedup_delete,
    "REPORT_PER_ROW": _codegen_report_per_row,
    "FORMAT_MAP": _codegen_format_map,
}


def codegen_dsl(op: str, resolved_args: dict, book_meta: dict, use_formula: bool = True) -> str:
    """④ 決定論 codegen。既存ヘルパへの Call を最優先し、無い操作だけテンプレ Basic を書く。
       ★ W3: book_meta["header_rows"] があれば対象シートの見出し行(1起点)をそこから読み
       （★ 挙動変更#2 より前は対象シート＝常に1枚目だった。下記参照）、
       0起点(hr0)に変換して全 op に一貫して渡す（『三層全部が同じ見出し推定を使う』の codegen 側）。
       header_rows が無い/キーが無い book_meta（_SAMPLE_META 等の旧テスト値）は既定1行目
       （hr0=0）＝旧挙動と完全に同一の Basic を生成する。
       use_formula: COMPUTE_COLUMN の既定を式（=B2*C2 等）にする（★ W3 Part3）。False で
       従来の値ベタ書きに戻す（--values）。
       ★ 挙動変更#2: 対象シートは resolved_args["_target_sheet"]（verify_dsl_args が一箇所で
       決めた値）を読む。無ければ従来どおり1枚目（後方互換 — codegen_dsl を直接呼ぶ既存の
       単体テスト・golden はこのキーを持たない args を渡しており、その挙動は変えない）。
       wrap() は _wrap_basic の対象シート対応版（_wrap_basic_for_sheet）に book_meta/
       first_sheet を閉じ込めたショートハンド — 対象シートの決定を codegen 側でも
       一箇所（_wrap_basic_for_sheet）に寄せるための配線。"""
    headers = book_meta["headers"]
    first_sheet = resolved_args.get("_target_sheet") or book_meta["sheets"][0]
    header_row = book_meta.get("header_rows", {}).get(first_sheet, 1)
    hr0 = header_row - 1   # Basic 0起点の見出し行

    def wrap(body: str) -> str:
        # ★★ 2026-08-27（Namakoo「位置の言い回しは頻出だから全ての操作で有効に」）:
        #   新しい列を作る op は、既定でデータの**右端**に作る。依頼文が位置を言っていて
        #   機械がそれを解けたなら、作った**あとで動かす**（MoveColumnTo）。
        #   ★ op ごとに codegen を書き換えない ── ここ 1 箇所で、宣言（WRITE_NEW_COLUMN）
        #     を持つ op すべてに同じ手が付く。op が増えても配線が要らない。
        #   ★ 実測（bench/swap_formula_spike_RESULTS.md）: insertByIndex も moveRange も
        #     参照を自動で付け替えるので、動かしても式は壊れない。
        move_to = resolved_args.get("_move_new_col_to")
        if move_to is not None:
            src0 = int(resolved_args.get("_new_col_from", 0))
            body = body + ('    Call MoveColumnTo(oDoc, %d, %d)%s'
                            % (src0, int(move_to), chr(10)))
        return _wrap_basic_for_sheet(body, book_meta, first_sheet)

    gen = CODEGEN_BY_OP.get(op)
    if gen is None:
        raise ValueError(f"未対応の op: {op}")
    return gen(op=op, resolved_args=resolved_args, book_meta=book_meta,
               use_formula=use_formula, headers=headers, first_sheet=first_sheet,
               header_row=header_row, hr0=hr0, wrap=wrap)


def formula_for_basic(formula: str) -> str:
    """xlsx の式（引数の区切りが `,`）を、Basic の `setFormula` が読める形（`;`）にする。

    ★★ 2026-08-29（実測で分かった環境の事実・推測では出ない）:
      同じ式を 3 通り書いて読み戻した:
        =SUM(E2:INDEX(E:E,ROW()-1))  → **#VALUE!**
        =SUM(E2:INDEX(E:E;ROW()-1))  → 通る
        =SUM(E2:E8)                  → 通る（区切りが無いので影響なし）
      ★ `setFormula` の引数区切りは **`;`**。カンマのまま渡すと式は入るが計算できない
        （文字は正しく見えるのに値だけ壊れる ── 一番たちが悪い形）。
    ★ 文字列の中のカンマは触らない（`=IF(A1;"a,b";"c")` の中身を壊さない）。
    """
    out, in_q = [], False
    for ch in str(formula or ""):
        if ch == chr(34):
            in_q = not in_q
        out.append(";" if (ch == "," and not in_q) else ch)
    return "".join(out)


def inspection_sheet_basic_call(sheet_name: str, header: list, rows: list,
                                 types: str) -> str:
    """検分シートを LibreOffice 側で書く Basic の 1 行を組む。

    ★ 2026-08-24（土台固め）: 旧実装は openpyxl でブックを開き直して検分シートを足して
    いた。openpyxl の往復は xl/drawings の**中身の図形**（描かれた角印・社判・
    テキストボックス）を捨てる ── ファイル名は残るので、忠実度ゲートの
    ファイル名比較にも掛からなかった。実測: 雛形に角印のある請求書で、LO が正しく
    N 枚へ複製した角印を最後の openpyxl 往復が全部消し、✓ が出ていた。
    ★ LO 経路は図形を保つと実測済みなので、書き手を LO へ寄せて往復ごと無くす。
    ★ 実装は 1 つ（帳票段と様式写像段が同じ関数を呼ぶ ── 書き写さない）。

    types は列ごとの型（"s"=文字列 / "n"=数値）。Excel のシート名・セル値は制御文字を
    含めないので、レコード Chr(30) / フィールド Chr(31) の区切りは安全。
    """
    # ★ 区切りは **Basic の Chr(30)/Chr(31) 式**として書く ── 生の制御文字を .bas に
    #   埋めない（この repo は生成物に混ざった制御文字で一度事故を起こしている）。
    def q(v):
        return '"' + ("" if v is None else str(v)).replace('"', '""') + '"'

    def rec(values):
        return " & Chr(31) & ".join(q(v) for v in values)

    parts = [rec(header)] + [rec(r) for r in rows]
    payload_expr = " & Chr(30) & ".join(parts) if parts else '""'
    name = str(sheet_name).replace('"', '""')
    return (f'    Call WriteInspectionSheet(oDoc, "{name}", '
            f'{payload_expr}, "{types}")' + chr(10))
