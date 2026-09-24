"""引数の検査（op ごとの _verify_*）── 解決済みの引数が実表の上で成り立つかを、書く前に確かめる。

★ 2026-09-24 に src/ailine/__init__.py から**移しただけ**（本文は 1 文字も変えていない）。
  入口の verify_dsl_args（op の台帳を読む）と、台帳を直に読む _verify_bold・_verify_set_column_value、
  EXTRACT / SET_WHERE の 2 本は本体に残る（台帳を core へ持ち込まないため）。
★ ailine を import しない（持ち出せる部品）。本体は from ailine_core.argcheck import ... で束ね直す。
★ 名前は verify_* を避けた（ailine_core/verify.py＝独立検算と別物）。
"""
from __future__ import annotations

import openpyxl
import re
from ailine_core import cellmap, inspection, intent as intent_mismatch, report_group, split_cell, threshold, total_row
from ailine_core.anchor import _COL_AFTER, _COL_BEFORE, _cell_row_name_for, _digit_candidates, _resolve_named_row, resolve_col_anchor, resolve_col_ref, resolve_row_anchor, task_names_a_row_number
from ailine_core.book_view import BookView
from ailine_core.column_type import column_is_all_numeric
from ailine_core.dedup_key import _dedup_normalize_key_part
from ailine_core.quotes import _task_outside_quotes, extract_quoted_literal
from ailine_core.report_per_row import cells_with_multiple_placeholders, scan_placeholders, unique_sheet_name
from ailine_core.subject import name_matches_task
from ailine_core.table_scan import _col_index_by_header, _scan_last_row, data_extent
from pathlib import Path


_RATE_PCT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*[%％]")


_RATE_BAI_RE = re.compile(r"(\d+(?:\.\d+)?)\s*倍")


# ★ W10b 項目3: 「1.1を掛けた/掛ける」「1.1で割った/割る」型（税込み/税抜きの言い換えで
#   「倍」を伴わない場合の実測ギャップ・battery v5 #503 で発覚）。掛けるは n そのまま、
#   割るは 1/n（税抜き＝税込み金額から逆算する倍率）。
_RATE_KAKE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:を|に)?\s*掛け")


_RATE_WARI_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:で|に)?\s*割っ")


# ★★ 2026-09-08（盲検の検品が「経理の常用語が通らない」として挙げた・**測って戻した**）:
#   「単価を全部**1割**値上げして」は断られる。『1割』を率（=10%）として読む regex を
#   書いて実機で確かめたところ、断りが**別のことをして △**に化けた ── 単価は上がらず、
#   新しい列『単価*1.1』が増える（COMPUTE_COLUMN は**新しい列**しか書けないため）。
#   ★ 率が読めないことは本当の欠けではない。欠けているのは**既存列をその場で書き換える
#     計算**で、率だけ読めるようにすると「もっともらしく違うこと」をして通ってしまう。
#     正直な断りの方が良いので戻した。実装する時は率と一緒に入れる。
_RATE_KEYWORD_RE = re.compile(r"税|倍率")


_RATE_BARE_NUM_RE = re.compile(r"(\d+(?:\.\d+)?)")


# ★ W10c 高: 依頼文に「率らしい語」が一切無いのに COMPUTE_COLUMN の「1列×率」パターン
#   （税込み/税抜き専用）へ誤分類された時に分類そのものを疑うための、より広い信号語。
#   上の各 _RATE_*_RE より緩い（数値を伴わなくても良い）— 「倍率を求める話かどうか」
#   だけを見る。
_RATE_SIGNAL_RE = re.compile(r"[%％]|倍|掛け|割っ|税")


# ★ W10c 中: 新規列の見出しを自然な日本語にするための語（A' 原則: LLM を使わず正規表現の
#   有無判定だけで決める。査定で名指しされた「金額*1.1」という数式風の見出しの対応）。
_TAX_INCLUSIVE_RE = re.compile(r"税込")


_TAX_EXCLUSIVE_RE = re.compile(r"税抜")


def extract_rate_factor(text: str) -> tuple:
    """依頼文から明示の倍率を抽出する。戻り値は (factor, 出典スニペット) か (None, None)。
       ①「10%」「8 ％」型 → 1+n/100 ②「1.1倍」型 → n そのまま ②'「1.1を掛けた」型 → n
       そのまま ②''「1.1で割った」型 → 1/n（税抜き等の逆算） ③「税」「倍率」という語の
       前後8文字だけにある裸の小数（例:「税率0.1」）→ 1未満なら 1+n・1以上ならそのまま
       （無関係な数値の誤爆を避けるため、③だけは税/倍率の語の近傍に絞る）。
       複数の異なる値が見つかった場合は断定しない（None, None・CLARIFY に委ねる）。"""
    if not text:
        return None, None
    candidates: dict = {}   # factor -> 出典スニペット（最初に見つかったもの）
    for m in _RATE_PCT_RE.finditer(text):
        f = round(1 + float(m.group(1)) / 100, 6)
        candidates.setdefault(f, m.group(0))
    for m in _RATE_BAI_RE.finditer(text):
        f = round(float(m.group(1)), 6)
        candidates.setdefault(f, m.group(0))
    for m in _RATE_KAKE_RE.finditer(text):
        f = round(float(m.group(1)), 6)
        candidates.setdefault(f, m.group(0))
    for m in _RATE_WARI_RE.finditer(text):
        n = float(m.group(1))
        if n > 0:
            f = round(1 / n, 6)
            candidates.setdefault(f, m.group(0))
    if not candidates:
        for km in _RATE_KEYWORD_RE.finditer(text):
            window = text[max(0, km.start() - 8): km.end() + 8]
            nm = _RATE_BARE_NUM_RE.search(window)
            if nm:
                n = float(nm.group(1))
                f = round((1 + n) if n < 1 else n, 6)
                candidates.setdefault(f, nm.group(0))
    if len(candidates) == 1:
        f, snippet = next(iter(candidates.items()))
        return f, snippet
    return None, None


def lookup_vocab_factor(text: str, vocab: dict) -> tuple:
    """依頼文に用語集の語が部分一致で含まれるかを見る。戻り値は (factor, 用語) か
       (None, None)。複数の異なる語（異なる値）がヒットした場合は断定しない。"""
    if not text or not vocab:
        return None, None
    hits: dict = {}
    for term, value in vocab.items():
        if term and term in text:
            hits.setdefault(value, term)
    if len(hits) == 1:
        value, term = next(iter(hits.items()))
        return value, term
    return None, None


def lookup_vocab_tax_factor(vocab: dict) -> tuple:
    """★ operator8 ②: 恒真式の番人が CLARIFY に倒す**直前**の敗者復活。第一照合
       （lookup_vocab_factor・依頼文の字面部分一致）と依頼文の率抽出（extract_rate_factor）
       の優先順は一切変えない ―― これはその両方が外れた（label が税/込を含むのに倍率が
       確定できない）場合だけ呼ばれる最後の一手。
       label『税込み合計』は語彙 key『消費税』を字面に含まないため第一照合は当たらないが、
       「税込み/税抜き」の依頼で使う倍率は用語集の中でも key に「税」を含む語である
       可能性が高い ―― そこだけ緩めて拾う（A' 原則は維持: 実在する用語集エントリの
       値だけを使い、LLM は使わない）。
       戻り値: (factor, term, candidates)。
         ・相異なる値がちょうど1つ → (その値, 名前, ())。
         ・相異なる値が2つ以上 → (None, None, ((value, term), ...))
           （呼び出し側が候補を名指しした CLARIFY にする ―― 「登録してください」とは言わない。
           登録は既に済んだ手だから）。
         ・0件 → (None, None, ())（呼び出し側は従来どおりの登録案内）。"""
    if not vocab:
        return None, None, ()
    seen: dict = {}   # value -> term（最初に見つかった名前。表示・一意判定の両方に使う）
    for term, value in vocab.items():
        if term and "税" in term:
            seen.setdefault(value, term)
    if len(seen) == 1:
        value, term = next(iter(seen.items()))
        return value, term, ()
    if len(seen) >= 2:
        return None, None, tuple(seen.items())
    return None, None, ()


def _resolve_tax_rescue(context_word: str, context_text: str, vocab: dict | None) -> tuple:
    """★ 致命③(2026-08-23レビュー): lookup_vocab_tax_factor の敗者復活を APPEND_TOTAL・
       COMPUTE_COLUMN の両方から呼ぶ共有実装（レビュー所見: 逐語コピー2箇所のうち
       COMPUTE_COLUMN 側だけ配線が届いていなかった片配線 ── 登録済みの税語彙があるのに
       「登録してください」と嘘をつく）。context_text（APPEND_TOTAL は label、
       COMPUTE_COLUMN は task）に「税」か「込」が無ければ rescue 対象外。
       戻り値: (factor, term, error_message)。
         ・rescue 対象外（税/込を含まない） → (None, None, None)
         ・rescue 成立 → (factor, term, None)
         ・rescue 失敗（候補複数/0件） → (None, None, "エラー文言")。"""
    if not any(k in context_text for k in ("税", "込")):
        return None, None, None
    tax_factor, tax_term, tax_candidates = lookup_vocab_tax_factor(vocab or {})
    if tax_factor is not None:
        return tax_factor, tax_term, None
    if tax_candidates:
        listed = "・".join(f"{term}={value:g}" for value, term in tax_candidates)
        return None, None, (
            f"{context_word}『{context_text}』は税/込を含みますが、用語集に候補が複数あります"
            f"（{listed}）。どちらを使うか依頼文に書いてください（例:「消費税10%」）"
        )
    return None, None, (
        f"{context_word}『{context_text}』は税/込を含みますが倍率が分かりません。"
        "依頼文に税率を書く（例:「消費税10%」）か、用語集に登録してください"
        "（例: ailine vocab add 消費税 1.1）"
    )


# ★ 単位B 照合の断片ガード（呼び出し側で・単位B 本体＝ailine_core/subject.py は変更しない）。
#   name_matches_task の呼び出し元は「実在するかどうか未確認の raw_target」を渡すことがあり
#   （下の COMPUTE_COLUMN target 経路）、_standalone_occurrence は「他の“実在名”の一部でしか
#   ない出現」しか除外しない ── raw_target が「実在しない・別の複合語」の断片（例:『小計』の
#   『計』）であっても素通りする。ailine_core/subject.py の _MIN_FRAGMENT=2（「1文字の漢字は
#   偶然一致しすぎる」）と同じ理由で、①長さ2未満は最初から証拠にしない、②2文字以上でも
#   依頼文中の全出現が「より長い連続した漢字の内部」（＝別の複合語の内側）でしかないなら
#   証拠にしない。
from ailine_core.word_boundary import CJK_KANJI_RE as _CJK_KANJI_RE, stands_alone  # noqa: E402,F401 ── _CJK_KANJI_RE は再輸出。★ 2026-09-24: 判定は word_boundary の 1 本（写しの範囲に異文字が紛れていた）


def _raw_target_not_embedded_in_task(raw_target: str, task: str) -> bool:
    """raw_target の依頼文中の出現のうち、少なくとも1つが「より長い連続した漢字の内部」
       ではない（＝独立した語としての出現がある）なら True。ひらがな/カタカナ/記号は
       日本語の語境界（助詞など）として扱う ―― 漢字が両隣にも続く場合だけ『内部』とみなす。
       出現が無ければ False（そもそも証拠が無い）。"""
    return stands_alone(raw_target, task)


def _verify_sort(resolved: dict, inferred: set, first_sheet: str, book_meta: dict,
                  resolve_in) -> tuple | None:
    """SORT の引数を確かめる（★ verify_dsl_args から切り出した・挙動は 1 ビットも変えていない）。

    ★ 切り出しの形（2026-09-04）: 元は `if op == "SORT":` の分岐だった。分岐は
      **早期 return するか、何も返さずチェーンを抜ける**の 2 通りなので、ここでは
      **返すべき tuple か None（＝続行）** を返す。呼び出し側が None を見て続ける。
    ★ `resolved` は辞書、`inferred` は set なので**参照が渡り、副作用はそのまま伝わる**。
      `resolve_in` は verify_dsl_args の内部関数（クロージャで resolved を書き換える）。

    ★ 挙動不変の確かめ方: bench/verify_golden.json（641 件の入出力）と突き合わせて
      **1 件でも動いたら赤**。SORT はそのうち 149 件を占めるので、壊せば必ず鳴る。
    """
    if (err := resolve_in("col", first_sheet)):
        return False, resolved, inferred, err
    if resolved.get("order") not in ("asc", "desc"):
        return False, resolved, inferred, f"順序『{resolved.get('order')}』は asc/desc のどちらでもありません"
    # ★★ 2026-08-29（Namakoo が実測）: 合計行まで並べ替えの範囲に入れていたので、
    #   降順にすると合計（一番大きい）が**先頭へ飛び**、その式が
    #   `=SUM(#REF!:INDEX(E:E,ROW()-1))` に壊れた。番人は止めたが、人は並べ替えられない。
    #   ★ 合計行は「データ行ではない」── 並べ替えの対象から外し、最下行に残す。
    #     判定は既存の凍結規則を借りる（total_rows_in → row_has_total_word）。
    #   ★ 見つけたら**必ず画面に出す**（黙って行を外さない）。
    _s_sheet = resolved.get("_target_sheet") or first_sheet
    _s_hr = int((book_meta.get("header_rows") or {}).get(_s_sheet, 1) or 1)
    _s_tot = total_rows_in(book_meta, _s_sheet, _s_hr)
    if _s_tot:
        _s_end = min(_s_tot) - 1
        if _s_end >= _s_hr + 1:
            resolved["_sort_end_row"] = _s_end
            # ★ 開示は**解釈行**に出す（警告ではない）。SET_WHERE が合計行を外す時と
            #   同じ口を使う ── 警告にすると決裁③で ✓ が △ に落ち、合計行のある表を
            #   並べ替えるたびに「確かめきれていない」と言うことになる。
            #   ★ 宣言どおりに動いて検算も通っているのだから、それは ✓ でよい。
            resolved["_skip_label"] = ("合計行 " + "、".join(
                f"{r}行目" for r in resolved["_skip_rows"]) + "（データ行でないため並べ替えません）")
    # ★ 並べ替えで「指す先の中身が変わる式」を名指しする（★ 付き＝決裁③で ✓→△）。
    #   ここは疑いなので警告でよい ── 合計行の除外（開示）とは性質が違う。
    _s_last = resolved.get("_sort_end_row") or 10 ** 7
    if (_dw := reference_drift_warning(book_meta, _s_sheet,
                                        row_lo=_s_hr + 1, row_hi=_s_last)):
        resolved["_warnings"] = resolved.get("_warnings", []) + [_dw]
    return None


def _verify_compute_column(resolved, inferred, first_sheet, task, vocab, headers, op):
    """COMPUTE_COLUMN の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    operands = resolved.get("operands")
    # ★ W10b 項目3: 税込み/税抜き等「1列 × 率」パターン。operands が列名1つだけの
    #   配列なら『既存列×倍率』とみなす（2列の四則演算とは別モード）。倍率(factor)は
    #   APPEND_TOTAL と同じ A' 原則で LLM から受け取らず機械確定する
    #   （extract_rate_factor/lookup_vocab_factor・regex のみ）。
    single_factor_mode = isinstance(operands, list) and len(operands) == 1
    if not single_factor_mode and not (isinstance(operands, list) and len(operands) == 2):
        return False, resolved, inferred, "演算対象が2つの列名になっていません"

    if single_factor_mode:
        v, was_inferred, err = resolve_col_ref(operands[0], headers.get(first_sheet, []))
        if err:
            return False, resolved, inferred, err
        resolved["operands"] = [v]
        if was_inferred:
            inferred.add("operands")
        if resolved.get("operator") not in ("*", "/"):
            return False, resolved, inferred, (
                f"演算子『{resolved.get('operator')}』は列1つの計算（税込み/税抜き等）"
                "では * か / のみ対応です")

        llm_factor_raw = resolved.pop("factor", None)
        text_factor, text_snippet = extract_rate_factor(task)
        vocab_factor, vocab_term = (None, None)
        if text_factor is None:
            vocab_factor, vocab_term = lookup_vocab_factor(task, vocab or {})

        sources: dict = {}
        if text_factor is not None:
            resolved["factor"] = text_factor
            sources["factor"] = f"依頼文: {text_snippet}"
        elif vocab_factor is not None:
            resolved["factor"] = vocab_factor
            sources["factor"] = f"用語集: {vocab_term}"
        else:
            # ★ W10c 高: 依頼文に率らしい語が一切無いのに「1列×率」（税込み/税抜き専用）へ
            #   分類されているのは、分類そのものが誤っている可能性が高い（実測: 「氏名の
            #   列を全部『退職済み』に書き換えて」のような値の一括書き換え依頼が、税率の
            #   話と誤認されて COMPUTE_COLUMN の単列モードに落ちることがあった）。
            #   その場合は「倍率が分からない」でなく、分類自体を疑う文言に変える
            #   （率を要求する op に分類されたのに率の手がかりが無い＝CLARIFY の理由を
            #   正直に言い換える。指示は意図・保証は機械＝プロンプト側だけに頼らない）。
            if not _RATE_SIGNAL_RE.search(task or ""):
                return False, resolved, inferred, (
                    f"依頼「{task}」は『{v}』列に何らかの倍率（税率等）を掛ける操作として"
                    "解釈しましたが、依頼文に倍率らしき手がかりが見当たりません。"
                    "列の値をそのまま書き換える操作は今のところ対応していません。"
                    "倍率を掛ける処理であれば、依頼文に率を書く（例:「消費税10%」）か、"
                    "用語集に登録してください（例: ailine vocab add 消費税 1.1）"
                )
            # ★ 致命③(2026-08-23レビュー): 敗者復活（_resolve_tax_rescue・APPEND_TOTAL と
            #   共有）。第一照合（上の text_factor/vocab_factor）の優先順は変えない ──
            #   ここに来るのはその両方が外れ、かつ依頼文に率らしき語（税/込を含む）が
            #   ある場合だけ（片配線の解消: 登録済みの税語彙で「登録してください」と
            #   嘘をつかない）。
            tax_factor, tax_term, tax_err = _resolve_tax_rescue("依頼", task or "", vocab)
            if tax_factor is not None:
                resolved["factor"] = tax_factor
                sources["factor"] = f"用語集: {tax_term}（依頼『{task}』の税に適用）"
            elif tax_err:
                return False, resolved, inferred, tax_err
            else:
                return False, resolved, inferred, (
                    "倍率（税率等）が分かりません。依頼文に率を書く（例:「消費税10%」）か、"
                    "用語集に登録してください（例: ailine vocab add 消費税 1.1）"
                )
        if resolved["factor"] <= 0:
            return False, resolved, inferred, f"倍率『{resolved['factor']}』は正の数でなければなりません"
        if sources:
            resolved["_sources"] = sources
        # ★ W10c 中: 新規列の見出しの自然化。旧実装は見出しを f"{op1}{operator}{factor:g}"
        #   （例:「金額*1.1」）という数式風の文字列にしていた（査定で名指し）。target
        #   無指定(新規列作成)かつ依頼文が税込み/税抜きと分かる言い方の場合だけ、
        #   その日本語ラベルを見出しに使う（A' 原則: LLM を使わず正規表現の有無のみで
        #   決める。手がかりが無ければ従来どおりの数式風見出しにフォールバック）。
        if not resolved.get("target"):
            # ★★ 2026-08-30: **依頼文に名前が在るなら、それが名前**（作らない）。
            #   下の「税込〜」は、人が名前を書かなかった時だけの間に合わせ。
            _asked = new_column_name_from_task(task, headers.get(first_sheet, []))
            if _asked:
                resolved["_new_col_label"] = _asked
            elif _TAX_INCLUSIVE_RE.search(task or ""):
                resolved["_new_col_label"] = f"税込{v}"
            elif _TAX_EXCLUSIVE_RE.search(task or ""):
                resolved["_new_col_label"] = f"税抜{v}"
            elif _RATE_KEYWORD_RE.search(task or ""):
                resolved["_new_col_label"] = f"税込{v}" if resolved["operator"] == "*" else f"税抜{v}"
        if llm_factor_raw not in (None, ""):
            try:
                llm_factor = float(llm_factor_raw)
            except (TypeError, ValueError):
                llm_factor = None
            if llm_factor is not None and abs(llm_factor - resolved["factor"]) > 1e-9:
                mfactor = resolved["factor"]
                resolved["_warnings"] = [
                    f"LLM が返した倍率({llm_factor:g})と機械抽出の倍率({mfactor:g})が"
                    f"食い違うため機械抽出({mfactor:g})を採用しました"
                ]
    else:
        new_operands = []
        for o in operands:
            v, was_inferred, err = resolve_col_ref(o, headers.get(first_sheet, []))
            if err:
                return False, resolved, inferred, err
            new_operands.append(v)
            if was_inferred:
                inferred.add("operands")
        resolved["operands"] = new_operands
        if resolved.get("operator") not in ("+", "-", "*", "/"):
            return False, resolved, inferred, f"演算子『{resolved.get('operator')}』が不明です"
        # ★★ 2026-09-14（盲検・誤配の家系③）: 「出勤と退勤の時刻から実働時間を計算する列を」で
        #   `出勤 − 退勤`（符号が逆）が黙って通っていた。引き算と割り算は**向きで答えが変わる**
        #   ── 依頼文が向きを言っていなければ聞き返す（並び順をそのまま演算の順にしない）。
        #   ★ 足し算・掛け算は向きが無いので触らない（陰性対照）。
        # ★ 向きの関所は**並び順だけ**の話 ── 依頼文に出てこない列が混じっている回は
        #   「どの列か」の食い違いで、既存の関所（subject_mismatch → 確認）の仕事。
        #   ここで先に断ると、聞けば済む回を断りに変えてしまう（実測: 凍結検体が 7→3 になった）。
        if (task and resolved["operator"] in ("-", "/")
                and all(str(o) in task for o in new_operands)):
            _dir = threshold.direction_of(task, new_operands, resolved["operator"])
            if _dir is None:
                _x, _y = new_operands
                _verb = "引く" if resolved["operator"] == "-" else "割る"
                _ex = "引いた" if resolved["operator"] == "-" else "割った"
                return False, resolved, inferred, (
                    f"『{_x}』と『{_y}』のどちらから{_verb}のかが依頼文から決まりません"
                    f"（例:「{_y}から{_x}を{_ex}列を作って」のように書いてください）")
            if _dir != new_operands:
                resolved["_warnings"] = resolved.get("_warnings", []) + [
                    f"依頼文の向き（{_dir[0]} {resolved['operator']} {_dir[1]}）を採用しました"]
                resolved["operands"] = _dir
    # ★ M2c: target(任意) — 依頼が既存列を名指し（「小計に」等）した場合はその列に書く。
    #   無指定なら従来どおり新規列（codegen_dsl 側で分岐）。
    # ★ W3: target が実在しない場合、翻訳が「新しい列の名前」（例:「利益列を作って」の
    #   『利益』）を target と誤って埋めていることが多いと実測された（qwen2.5-coder:7b が
    #   『既存列に書く/新規に作る』の区別を安定して守らない）。実在しない＝一意に決まらない
    #   （複数解釈で曖昧）のとは別の理由なので、その場合だけ target を無指定として扱い
    #   新規列作成にフォールバックする（推測で断定しない CLARIFY の原則は、真に曖昧な
    #   ケース＝digit_candidates の複数一致にだけ残す）。
    # ★ W3 改定(2026-08-20): 上の前提（実在しない target＝ほぼ捏造）が古くなった実測が
    #   出た。「金額を数量×単価で埋めて」（金額列がまだ実在しない構成）で翻訳は正しく
    #   target:"金額" を返す。だが無条件に捨てる旧実装はそれも落とし、新規列が
    #   『数量*単価』に自動命名されていた（利用者が名前を言っているのに無視される・
    #   2026-08-19 のデモ制作で3回踏んだ実害）。★ 依頼文を判定者にする: raw_target が
    #   依頼文に実在する語として機械照合できるなら（単位B の name_matches_task を再利用
    #   ── 素朴な in 判定はしない。「税込金額を…」+target「金額」のような片方向の部分
    #   文字列の穴は単位B が塞いだ形そのものなので、同じ判定を2箇所に書かず呼ぶ）、
    #   捏造ではなく利用者の指名とみなして新規列の名前として使う。依頼文に語が無ければ
    #   従来どおり捏造とみなして捨てる（W3 本来の防御は生きている）。
    if resolved.get("target"):
        raw_target = resolved["target"]
        v, was_inferred, err = resolve_col_ref(raw_target, headers.get(first_sheet, []))
        if err:
            if "一意に決まりません" in err:
                return False, resolved, inferred, err
            # ★ 単位B 照合の断片ガード（呼び出し側・上の _raw_target_not_embedded_in_task
            #   docstring 参照）: raw_target は「実在するか未確認」の生の文字列なので、
            #   name_matches_task を素通しに使う前に (1) 1文字を弾き (2) 依頼文中の全出現が
            #   「他の複合語（実在列とは限らない）の内部」でしかないなら弾く。
            if (len(raw_target) >= 2
                    and _raw_target_not_embedded_in_task(raw_target, task)
                    and name_matches_task(raw_target, task, others=headers.get(first_sheet, []))):
                resolved["_new_col_label"] = raw_target
            del resolved["target"]
        elif (task_asks_to_add_a_column(task)
                and not name_matches_task(v, task,
                                           others=headers.get(first_sheet, []))):
            # ★★ 2026-09-02（130 件の器を広げて初めて見えた・実測）:
            #   「単価の右に、数量と単価をかけた**列を作って**」で、一段目が
            #   target='メモ'（実在するが**空**の列）を返し、道具は新しい列を作らずに
            #   **その列へ書いて ✓ を出していた**。頼んでいない場所に書いている。
            #   ★ W3 は「**実在しない** target ＝ ほぼ捏造」を捨てる。抜けていたのは
            #     「**実在するが、依頼文に無い**」列 ── そこだけ素通りだった。
            #   ★ 判定に語彙の一覧は要らない: 依頼文が「作る」と言っているか
            #     （閉じた文法）と、その名前が依頼文と機械照合できるか（既存の
            #     provenance 層）の 2 つだけ。新しい言い回しが来ても足すものは無い。
            #   ★ 道具は既に気づいていた（★で開示していた）── **止めていなかった**だけ。
            del resolved["target"]          # → 新しい列を作る（自動命名 or 依頼文の名前）
        else:
            resolved["target"] = v
            if was_inferred:
                inferred.add("target")
                # ★ W10a 項目3: 数字指定→列名解決の元の表記を残す（解釈要約の表示用・
                #   例:「列5」→「在庫」列と解決した時、確認行の直後にその経緯を見せる）。
                resolved["_target_raw"] = raw_target
    # ★★ 2026-09-02: 2 項の演算（売上 − 原価）でも、依頼文の名前を拾う。
    #   名前の抽出は**倍率の枝（税込/税抜）にしかなかった**ので、
    #   「売上から原価を引いた**利益**の列を作って」の見出しが
    #   『売上-原価』（式そのもの）になっていた ── A' 原則が抜けた形。
    #   ★ 人が名前を書いていない時だけ従来どおり式風の見出しに落ちる。
    if not resolved.get("target") and not resolved.get("_new_col_label"):
        # ★ 見出しの一覧を**渡さずに**呼ぶ ── 「既に在る名前」も受け取りたいから。
        _asked_c = new_column_name_from_task(task, [], require_position=False)
        _heads_c = [str(h) for h in (headers.get(first_sheet) or [])]
        if _asked_c and _asked_c in _heads_c:
            # ★★ 2026-09-02（実測で捕まえた実害）: 依頼した名前が**既に在る**時、
            #   「新しい列の名前ではない」として捨てて自動命名に落ちていた。すると
            #   「売上から原価を引いた利益の列を作って」を 2 回実行すると、
            #   1 回目『利益』・2 回目『売上-原価』になり、**見出しが違うので
            #   「見出しも値も同一の列を作りました」の関所が鳴らない** ──
            #   値がそっくり同じ列が静かに 2 本目として増え、✓ まで出た。
            #   （盲検 operator 査定が見つけた事故「不安でもう一回実行」の再来）
            #   ★ 意味で考えても、これは「作る」ではなく「**もう在る**」。
            #     その列を計算し直す依頼と読み、既存の**上書きの関所**に載せる
            #     ── 新しい関所も新しい終了コードも作らない。
            resolved["target"] = _asked_c
        elif _asked_c:
            resolved["_new_col_label"] = _asked_c
    return None


def _verify_lookup_fill(resolved, inferred, first_sheet, book_meta, resolve_in, check_sheet, task, headers, op):
    """LOOKUP_FILL の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := check_sheet("target_sheet")):
        return False, resolved, inferred, err
    if (err := check_sheet("source_sheet")):
        return False, resolved, inferred, err
    # ★ 挙動変更#2: 旧実装はここで「対象シートは1枚目のみ対応しています」と拒否していた
    #   （散在した『1枚目固定』の一つ・査定の致命そのもの）。LOOKUP_FILL は元々
    #   target_sheet を自分の必須 slot として名前で受け取り check_sheet で実在確認まで
    #   済ませているので、この制限を外すだけで対応できる。resolved["_target_sheet"] は
    #   LOOKUP_FILL 自身の target_sheet を正とする（他 op 用の一般解決 first_sheet より
    #   こちらを優先 — 依頼文に転記先/参照元の2シート名が両方出て一般解決が曖昧に
    #   フォールバックしていても、LOOKUP_FILL のここでの解決は影響を受けない）。
    resolved["_target_sheet"] = resolved["target_sheet"]
    # ★★ 塊③(2/2)・中核 op 致命2（2026-08-24 の盲検）:
    #   書き手（VLookupFromTable）は参照表の**列1（2 番目の列）を値**と決め打ちし、
    #   検算（check_lookup_fill）も同じ決め打ちで期待値を作る ── やる側と見る側が
    #   同じ思い込みを共有しているので必ず一致する（恒真）。
    #   実測: マスタ = 商品/区分/**単価**（3 列目）で「単価を転記して」と頼むと、
    #   **単価の列に「果物」が入って ✓** が出た。数値であるべき列に文字列が入る。
    #   ★ 商品コード/商品名/単価 のような 3 列マスタは実務でごく普通。
    #   → 書く**前**に前提を照合する: 頼まれた列名が参照表の 2 列目でなければ断る。
    #   ★ 見出しが読めない参照表では断らない（根拠が無い時に止めない）。
    #   ★ 断らずに**開示する**理由（実測で 1 度誤爆した）:
    #     事故の形   マスタ=[商品,区分,単価] → 2 列目は「区分」
    #     正しい依頼 明細  =[商品,数量,単価] → 2 列目は「数量」
    #     どちらも「2 列目 ≠ 頼まれた列」で、**列の位置だけでは区別できない**。
    #     断ると正しい依頼まで止める（既存検体で実証）。判定は変えず、
    #     何が書かれるかを名指しして ✓ を △ に降ろす（決裁③の機構に乗せる）。
    # ★★ 2026-09-21（盲検 5 体目）: 参照表の値の列を**名前で**読むようにしたので、
    #   名前が見つかる限り 3 列以上のマスタでも正しく転記できる（買い手の発注記録は 6 列で、
    #   以前は「2 列だけの表を用意してください」と言うしかなかった）。
    #   ★ 残す警告は **名前で引けなかった時だけ** ── そこは従来どおり 2 列目に落ちるので、
    #     何が書かれるかを名指しして ✓ を △ に降ろす（決裁③の機構に乗せる）。
    _src_headers = list((book_meta.get("headers") or {}).get(resolved["source_sheet"], []))
    _want = resolved.get("target_col")
    if len(_src_headers) > 2 and isinstance(_want, str) and _want not in _src_headers:
        resolved["_warnings"] = resolved.get("_warnings", []) + [
            f"参照表『{resolved['source_sheet']}』に『{_want}』という列が見つかりません。"
            f"この転記は列の名前で値を探しますが、見つからないときは 2 列目に落ちます ── "
            f"実際に書き込まれるのは『{_src_headers[1]}』の値です。"
            f"意図と違う場合は、参照表の見出しを確かめてください"]
    # ★ W10c 致命2: target_col は COMPUTE_COLUMN の target と違い OP_SCHEMA 上は必須
    #   slot なので、LLM は「存在しないなら空にする」を選べない。実測（監査再現）:
    #   対象シートに『単価』列がまだ無いのに転記を頼むと、LLM がそれと無関係な
    #   *実在する*既存列（例:「数量」）の名前を代わりに返すことがある。resolve_col_ref
    #   は実在列名なら無条件で素通しするため、これだけでは見分けられない（そのまま
    #   進めると「数量」が確認なしで上書きされる事故になる）。
    #   ここでは「実在するから信用する」をやめ、根拠を要求する:
    #   ①依頼文にその列名が書かれている ②転記元（source_sheet）の値列
    #   （VLookupFromTable ヘルパの仕様どおり常に列1＝2番目の列）と同じ名前
    #   のどちらかが無いと、実在列であっても信用しない。
    target_headers = headers.get(resolved["target_sheet"], [])
    source_headers = headers.get(resolved["source_sheet"], [])
    value_col_hint = source_headers[1] if len(source_headers) > 1 else None
    raw_target_col = resolved.get("target_col")
    raw_str = str(raw_target_col) if raw_target_col not in (None, "") else ""
    exists = raw_str in target_headers
    mentioned = bool(raw_str) and raw_str in task
    matches_value_col = value_col_hint is not None and raw_str == value_col_hint

    if exists and (mentioned or matches_value_col):
        pass   # 根拠つきで実在列を指名＝そのまま使う（上書き注意は破壊の関所が別途担当）
    elif not exists:
        cands = _digit_candidates(raw_str, target_headers)
        if len(cands) == 1:
            resolved["target_col"] = cands[0]   # 数字表記の推定は従来どおり許容
            inferred.add("target_col")
        elif mentioned:
            # ★ 依頼文にも同じ列名が書かれている＝新規作成が正しい解釈（COMPUTE_COLUMN の
            #   target 無指定＝新規列と同じ考え方）。target_col はそのまま残し、
            #   codegen_dsl 側で新規列として作る。
            pass
        else:
            known = ", ".join(target_headers) if target_headers else "(無し)"
            return False, resolved, inferred, (
                f"転記先の列『{raw_target_col}』が『{resolved['target_sheet']}』シートに"
                f"見つかりません。ある列: {known}。新しい列として作る場合は、依頼文に"
                f"その列名を書いてください（例:「{raw_target_col}という列を作って転記して」）"
            )
    else:
        # ★ 実測の事故そのもの: exists=True だが根拠が無い（依頼文にも書かれておらず、
        #   転記元の値列とも一致しない）＝上書き対象を取り違えている可能性が高い。
        hint = f"（参照表『{resolved['source_sheet']}』の値の列は『{value_col_hint}』です）" \
            if value_col_hint else ""
        return False, resolved, inferred, (
            f"転記先の列『{raw_target_col}』は実在しますが、依頼文にその列名が見当たらず、"
            f"転記元の値とも対応が確認できません{hint}。上書き対象を取り違えている"
            "可能性があるため、意図した列名を依頼文に明記してください"
        )
    if (err := resolve_in("key_col", resolved["target_sheet"])):
        return False, resolved, inferred, err
    # ★★ キー列 == 対象列 は転記として意味を成さない（2026-09-11・PENDING-20260910 ②）。
    #   実測: 「商品表から商品名を注文シートに転記して」で 7B が キー列:商品名／対象列:商品名 を
    #   返し、**空の列を空の列で引く**ので 1 件も一致せず、文書は無変化のまま
    #   事後条件が「検証対象が 0 件」と × を出していた（battery の名前の台帳 注文/lookup）。
    #   × は正しいが、**なぜかを人に教えない**。ここで書く前に断り、原因を名指しする。
    #   ★ 依頼文も語彙も読まない ── 見るのは「2 つの列名が同じか」だけ
    #     （detect_new_row_missing_key と同じ、構造だけの判定）。
    #   ★ 手がかりとして参照表の 1 列目を名指す ── VLookupFromTable は参照表の
    #     **1 列目をキー**と決め打ちして読む（ヘルパの契約・推測ではない）。
    #   ★ 数字（× → 断り）が良くなる方向の変更なので、決め手は「親切か」に置いた:
    #     「検証対象が 0 件」より「キー列と対象列が同じです。参照表の 1 列目は『コード』です」
    #     の方が、人は次に何を言えばいいか分かる。
    if str(resolved.get("key_col")) == str(resolved.get("target_col")):
        _src = list((book_meta.get("headers") or {}).get(resolved["source_sheet"], []))
        _hint = (f"参照表『{resolved['source_sheet']}』の 1 列目は『{_src[0]}』です ── "
                 f"ふつうはそれがキー列です" if _src else
                 f"参照表『{resolved['source_sheet']}』の 1 列目がキーになります")
        # ★★ 2026-09-20（盲検 5 体目 ⑤）: ここの例文は「引く側」に参照表の 1 列目を
        #   そのまま入れていた。買い手の冊は 1 列目も転記先も『商品名』だったので、
        #   **例文がそのまま同じ誤りを再生産した**（`path_fails` ── 盤で一番重い失敗）。
        #   ★ キーと同じ名前になる例は作らない。人が次に言うべきは「**何を**転記するか」
        #     なので、参照表の**他の列**を候補として名指しする。
        # ★★ 除くのは**引く側に使う列**（参照表の 1 列目）── 初版は `key_col` で除いたので、
        #   1 列目が `key_col` と違う冊では 1 列目自身が候補に残り、例文がまた
        #   「コードで引いてコードを転記して」になった（既存の番人が掴んだ・2026-09-20）。
        _puller = str(_src[0]) if _src else ""
        _others = [c for c in _src if str(c) != _puller]
        if _others:
            _how = (f"引くのは『{_puller}』のままで構いません ── "
                    f"転記したい列を選んでください（『{_others[0]}』など"
                    + ("、他に " + "／".join(f"『{c}』" for c in _others[1:4]) if len(_others) > 1 else "")
                    + f"）。依頼文にそう書いてください。例:「{_puller}で引いて{_others[0]}を転記して」")
        else:
            _how = (f"参照表『{resolved['source_sheet']}』には『{resolved['key_col']}』しか"
                    "列がありません ── 転記できる中身がその表に入っていません")
        return False, resolved, inferred, (
            f"転記のキー列と対象列がどちらも『{resolved['key_col']}』になっています。"
            f"転記はキーとは別の列を持ってくる操作です（{_hint}）。{_how}"
        )
    return None


def _verify_append_total(resolved, inferred, first_sheet, book_meta, resolve_in, args, task, vocab, headers):
    """APPEND_TOTAL の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := resolve_in("col", first_sheet)):
        return False, resolved, inferred, err
    # ★ W6: label は既定値を持つ任意項目。ここで確定させ、codegen/事後条件/
    #   確認行の全部に同じ既定解決を一貫して渡す。
    resolved["label"] = str(resolved.get("label") or "合計")
    label = resolved["label"]

    # ★★ 2026-08-29（Namakoo が実測）: 合計行が**既に在る**表で「単価列の合計行に
    #   単価の合計を書いて」と頼むと、10 行目に『単価合計』という**別の行**が増えた。
    #   ★ 真因: 合計行を「データ行」と数えて、その下に足していた。
    #   ★ 合計行が 1 つに決まり、その列がまだ空なら、**その行に書く**（行は増やさない）。
    #     判定は既存の凍結規則を借りる（total_rows_in → row_has_total_word）──
    #     ここで新しい規則を書かない。同じことを 2 箇所が決めると必ずずれる。
    _tot_sheet = resolved.get("_target_sheet") or first_sheet
    _tot_hr = int((book_meta.get("header_rows") or {}).get(_tot_sheet, 1) or 1)
    _tot_rows = total_rows_in(book_meta, _tot_sheet, _tot_hr)
    if len(_tot_rows) == 1:
        _tr = _tot_rows[0]
        _theads = [str(h) for h in ((book_meta.get("headers") or {}).get(_tot_sheet) or [])]
        _tidx = _theads.index(resolved["col"]) + 1 if resolved["col"] in _theads else 0
        _cur = None
        if _tidx:
            try:
                with BookView(Path(book_meta["path"])) as _bv:
                    _cur = _bv.sheet(_tot_sheet).cell(row=_tr, column=_tidx).value
            except Exception:
                _tidx = 0
        if _tidx and (_cur in (None, "") or str(_cur).startswith("=SUM(")):
            resolved["_at_row"] = _tr
            resolved["_at_basis"] = f"既にある合計行＝{_tr}行目（行は増やしません）"
            # ★ ラベルは**その行に既に在る物**が正（LLM の案『単価合計』で検算しない）。
            try:
                with BookView(Path(book_meta["path"])) as _bv2:
                    _lbl = _bv2.sheet(_tot_sheet).cell(row=_tr, column=1).value
                if _lbl not in (None, ""):
                    resolved["label"] = str(_lbl)
                    label = resolved["label"]
            except Exception:
                pass
        # ★★ 2026-08-29: ここで「既に値が入っています」と**断るのはやめた**。
        #   既存の番人（事後条件の算術の検算＝二重計上に ✓ を出さない／単位F の関所）が
        #   同じ事故を既に止めていて、断りを重ねると**その番人の出番が消える**
        #   ── 過去の事故を守っている検体が通らなくなる（実測で 3 本落ちた）。
        #   ★ 埋められる時だけ埋め、それ以外は今までどおり深い番人に任せる。

    # ★ A': factor は LLM から受け取らない。LLM が返した値(あれば)はいったん取り出して
    #   おき、機械抽出/用語集の結果と食い違う場合だけ WARN として記録する（常に機械が勝つ）。
    llm_factor_raw = resolved.pop("factor", None)

    text_factor, text_snippet = extract_rate_factor(task)
    vocab_factor, vocab_term = (None, None)
    if text_factor is None:
        vocab_factor, vocab_term = lookup_vocab_factor(task, vocab or {})

    sources: dict = {}
    if text_factor is not None:
        resolved["factor"] = text_factor
        sources["factor"] = f"依頼文: {text_snippet}"
    elif vocab_factor is not None:
        resolved["factor"] = vocab_factor
        sources["factor"] = f"用語集: {vocab_term}"
    else:
        resolved["factor"] = 1.0

    if resolved["factor"] <= 0:
        return False, resolved, inferred, f"倍率『{resolved['factor']}』は正の数でなければなりません"

    # ★ 恒真式の番人（最優先）: label が「税」/「込」を含むのに倍率が確定できず既定
    #   1.0 のままだと、税抜き金額に「税込み」ラベルが付いた恒真の誤りを事後条件が
    #   pass にしてしまう（args 基準の検証だから）。ここで機械的に CLARIFY へ倒す
    #   （語リストは 税/込 の2語で凍結・むやみに増やさない）。
    # ★ operator8 ②: CLARIFY に倒す直前に敗者復活（lookup_vocab_tax_factor・
    #   docstring 参照）。第一照合（上の text_factor/vocab_factor）の優先順は変えない
    #   ―― ここに来るのはその両方が外れた場合だけ。
    if resolved["factor"] == 1.0 and any(k in label for k in ("税", "込")):
        tax_factor, tax_term, tax_err = _resolve_tax_rescue("ラベル", label, vocab)
        if tax_factor is not None:
            resolved["factor"] = tax_factor
            sources["factor"] = f"用語集: {tax_term}（ラベル『{label}』の税に適用）"
        elif tax_err:
            return False, resolved, inferred, tax_err

    if sources:
        resolved["_sources"] = sources
    if llm_factor_raw not in (None, ""):
        try:
            llm_factor = float(llm_factor_raw)
        except (TypeError, ValueError):
            llm_factor = None
        if llm_factor is not None and abs(llm_factor - resolved["factor"]) > 1e-9:
            mfactor = resolved["factor"]
            resolved["_warnings"] = [
                f"LLM が返した倍率({llm_factor:g})と機械抽出の倍率({mfactor:g})が"
                f"食い違うため機械抽出({mfactor:g})を採用しました"
            ]

# --- ★ 2026-08-26: 表の基本操作 3 種（追加・行削除・列削除）---------------
    return None


def _verify_report_per_row(resolved, inferred, first_sheet, book_meta, resolve_in, check_sheet, sheets, headers):
    """REPORT_PER_ROW の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := check_sheet("template_sheet")):
        return False, resolved, inferred, err
    template_sheet = resolved["template_sheet"]
    if template_sheet == first_sheet:
        return False, resolved, inferred, (
            f"雛形シートとデータシートが同じ『{template_sheet}』です。"
            "雛形は別のシートに用意してください")
    if (err := resolve_in("name_col", first_sheet)):
        return False, resolved, inferred, err

    book_path = book_meta.get("path")
    if book_path is None:
        return False, resolved, inferred, (
            "帳票段はファイルの実体が無いと検証できません（book_meta に path が無い）")
    data_headers = headers.get(first_sheet, [])
    header_row_here = book_meta.get("header_rows", {}).get(first_sheet, 1)

    try:
        wb_tpl = openpyxl.load_workbook(book_path)
    except Exception as e:
        return False, resolved, inferred, f"雛形の読み込みに失敗しました: {e}"
    try:
        tpl_ws = wb_tpl[template_sheet]
        placeholders = scan_placeholders(tpl_ws, tpl_ws.max_row or 1, tpl_ws.max_column or 1)
        # ★ 縦の結合セルは、明細行を増やすと崩れる（値は合うので事後条件は通ってしまう）。
        #   日本の請求書の雛形は結合だらけなので、起きる方に賭けるべき事象（設計査読）。
        tpl_vmerges = [(m.min_row, m.max_row, m.coord)
                        for m in tpl_ws.merged_cells.ranges if m.min_row != m.max_row]
    finally:
        wb_tpl.close()

    if not placeholders:
        return False, resolved, inferred, (
            f"雛形『{template_sheet}』に印（{{{{列名}}}}）が見つかりません。"
            "転記したいセルに {{列名}} の形で印を置いてください")

    # ★ 2026-08-24: 1 セルに印が 2 つ以上あるなら、埋めずに断る。埋めると 1 セルに
    #   2 回書くことになり後の値が前を消す ── 「それらしく埋まって片方が生で残る」
    #   （盲検の査定で名指しされた事故）より、雛形を直してくださいと言う方が正しい。
    if (dupes := cells_with_multiple_placeholders(placeholders)):
        cell, names = dupes[0]
        return False, resolved, inferred, (
            f"雛形『{template_sheet}』の {cell} に印が {len(names)} つあります"
            f"（{chr(12539).join(names)}）。1 つのセルに置ける印は 1 つまでです ── "
            f"別々のセルに分けてください")

    # ★★ 2026-08-28（Namakoo「同名の取引先から複数の発注があるケースでは
    #   請求書を一枚にまとめないといけない」）: 印を 3 種類に仕分ける。
    #   {{列名}} / {{明細:列名}} / {{合計:列名}}。**雛形が形を決める**ので、
    #   依頼文にも一段目の語彙（OPS_DOC）にも 1 文字も足さない。
    mark_layout, layout_err = report_group.classify_placeholders(placeholders)
    if layout_err:
        return False, resolved, inferred, f"雛形『{template_sheet}』: {layout_err}"
    if mark_layout.detail_row is not None:
        crossing = [c for lo, hi, c in tpl_vmerges if lo <= mark_layout.detail_row <= hi]
        if crossing:
            return False, resolved, inferred, (
                f"雛形『{template_sheet}』の明細行（{mark_layout.detail_row}行目）を、"
                f"縦に結合したセルが横切っています（{'・'.join(crossing[:3])}）。"
                "明細行は件数ぶん増えるので、この結合は崩れます ── "
                "結合を解くか、明細行の外へずらしてください")

    resolved_placeholders = []
    for ph in placeholders:
        ph_kind, ph_col = report_group.mark_kind(ph.column_name)
        if ph_col not in data_headers:
            return False, resolved, inferred, (
                f"雛形『{template_sheet}』の印『{{{{{ph.column_name}}}}}』"
                f"（{ph.cell}）が指す列『{ph_col}』は、データシート"
                f"『{first_sheet}』に見つかりません。実在する列名を印にしてください"
            )
        col_idx = data_headers.index(ph_col) + 1
        if ph_kind == "total" and not ph.whole:
            return False, resolved, inferred, (
                f"雛形『{template_sheet}』の合計の印『{{{{{ph.column_name}}}}}』"
                f"（{ph.cell}）は、セル全体を印にしてください（合計は数値です）")
        if not ph.whole:
            # ★ 訂正3: 部分一致の印は原理的に文字列にしかなれない ── 数値列には使わせない
            #   （検体には無いが自分の検体で固定する境界。設計文書の指示どおり）。
            try:
                is_numeric = column_is_all_numeric(book_path, first_sheet, col_idx,
                                                    header_row_here)
            except Exception:
                is_numeric = False
            if is_numeric:
                return False, resolved, inferred, (
                    f"雛形『{template_sheet}』の印『{{{{{ph.column_name}}}}}』"
                    f"（{ph.cell}）はセルの一部分（部分一致）ですが、列『{ph_col}』は"
                    "数値です。数値列には部分一致の印を使えません"
                    "（セル全体を印にしてください: 例 " + "{{" + ph.column_name + "}}）"
                )
        resolved_placeholders.append({
            "cell": ph.cell, "row": ph.row, "col": ph.col,
            "column_name": ph_col, "kind": ph_kind, "mark": ph.column_name,
            "whole": ph.whole, "raw": ph.raw, "col_idx": col_idx,
        })
    resolved["_placeholders"] = resolved_placeholders

    try:
        wb_data = openpyxl.load_workbook(book_path, data_only=True)
    except Exception as e:
        return False, resolved, inferred, f"データシートの読み込みに失敗しました: {e}"
    try:
        src_ws = wb_data[first_sheet]
        last_row = _scan_last_row(src_ws, header_row=header_row_here)
        rows_in = []
        for r in range(header_row_here + 1, last_row + 1):
            label_val = src_ws.cell(row=r, column=1).value
            vals = {h: src_ws.cell(row=r, column=i + 1).value
                    for i, h in enumerate(data_headers)}
            rows_in.append((r, label_val, vals))
    finally:
        wb_data.close()
    verdict = total_row.split_total_rows_multi(rows_in) if rows_in else total_row.TotalRowVerdict(
        excluded=[], adopted_rows=[], mismatches=[])
    row_values = {r: v for r, _l, v in rows_in}

    # ★★ まとめるか、1 行 1 枚か ── **雛形と実表の両方**が決める（人に選ばせない）:
    #   ・雛形に明細/合計の印が在る → まとめる（1 件でも同じ道を通る）
    #   ・印は無いが同じ名前が 2 行以上ある → **断る**（2 枚に割れた紙は仕事にならない）
    name_col_here = resolved["name_col"]
    name_idx = data_headers.index(name_col_here) + 1
    groups = report_group.build_groups(
        [(r, [row_values[r].get(h) for h in data_headers]) for r in verdict.adopted_rows],
        name_idx)
    grouped = mark_layout.detail_row is not None or bool(mark_layout.total)
    # ★★ 2026-08-28（設計査読で名指しされた・自分で開けかけた穴）:
    #   ここで**断って**はいけない。同名が 2 行あっても正しい帳票がある ──
    #   領収書・納品書は取引ごとに 1 枚だし、締め日違いの月別請求も同じ形
    #   （OPS_DOC 自身が REPORT_PER_ROW の用途に領収書を挙げている）。
    #   既に在る処方は「断ること」ではなく「✓ を出さないこと」だった（2026-08-24）。
    #   ★ 反転させずに、△ の警告文へ**まとめ方への道**を足すだけにする。

    used = set(sheets) | {template_sheet}
    report_rows = []
    if grouped:
        for g in groups:
            sheet_name = unique_sheet_name(str(g.name), used)
            used.add(sheet_name)
            report_rows.append({"row": g.rows[0], "sheet": sheet_name,
                                 "name": g.name, "rows": list(g.rows)})
    else:
        for r in verdict.adopted_rows:
            raw_name = row_values[r].get(name_col_here)
            sheet_name = unique_sheet_name(str(raw_name), used)
            used.add(sheet_name)
            report_rows.append({"row": r, "sheet": sheet_name})
    if not report_rows:
        return False, resolved, inferred, (
            "帳票にするデータ行がありません（表が空か、全行が合計行と判定されました）"
        )
    inspection_sheet = unique_sheet_name(inspection.SHEET_NAME, used)
    used.add(inspection_sheet)

    if grouped:
        # ★ 1 枚に 1 つしか書けない欄が、グループの中で食い違っていないか。
        #   食い違ったら**埋めずに断る** ── 推測で選ぶと、別の担当者の名前が客に届く。
        by_name = {g.name: g for g in groups}
        for rr in report_rows:
            g = by_name[rr["name"]]
            for ph in resolved_placeholders:
                if ph["kind"] == "value":
                    vals = report_group.value_conflicts(g, row_values, ph["column_name"])
                    if vals:
                        # ★ 「足すなら {{合計:担当}}」は、担当のような文字列の列では
                        #   意味を成さない ── 出せる道だけを名指しする。
                        _way = f"『{{{{明細:{ph['column_name']}}}}}』にしてください"
                        if all(report_group.is_numeric(v) for v in vals):
                            _way = (f"『{{{{明細:{ph['column_name']}}}}}』、"
                                     f"足すなら『{{{{合計:{ph['column_name']}}}}}』"
                                     "にしてください")
                        return False, resolved, inferred, (
                            f"『{g.name}』の {list(g.rows)}行目で"
                            f"『{ph['column_name']}』が食い違っています（{vals}）。"
                            f"1 枚の紙には 1 つしか書けません ── 明細に出すなら{_way}")
                elif ph["kind"] == "total":
                    _s, serr = report_group.sum_for(g, row_values, ph["column_name"])
                    if serr:
                        return False, resolved, inferred, f"『{g.name}』: {serr}"
        resolved["_groups"] = [{"sheet": rr["sheet"], "name": rr["name"],
                                 "rows": rr["rows"]} for rr in report_rows]
        resolved["_detail_row"] = mark_layout.detail_row
    else:
        # ★ 2026-08-24: 重複を知った瞬間に言う（`_2` を付けたのがその瞬間）。
        #   実測: 3 社の売上表（4 行）で請求書 4 枚・同じ取引先が 2 枚に分かれて ✓ が出た。
        #   ★ 付きなので count_suspicious_advisories が拾い、決裁③で ✓→△ に降格する。
        if (dup := duplicate_name_warning(
                name_col_here,
                [row_values[r].get(name_col_here) for r in verdict.adopted_rows])):
            resolved["_warnings"] = resolved.get("_warnings", []) + [dup]
    resolved["_report_rows"] = report_rows
    resolved["_report_sheet_names"] = [rr["sheet"] for rr in report_rows]
    resolved["_inspection_sheet"] = inspection_sheet
    resolved["_source_headers"] = tuple(data_headers)
    return None


def _verify_set_cell_value(resolved, inferred, book_meta, task, sheets, headers, op):
    """SET_CELL_VALUE の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    # ★ 2026-08-27（Namakoo「梨の売上にピンポイントで入れたい」）:
    #   SET_COLUMN_VALUE は**列を丸ごと**同じ値にする op で、1 セルを狙えなかった。
    _sheet_c = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _headers_c = [str(h) for h in
                   ((book_meta.get("headers") or {}).get(_sheet_c) or [])]
    _row_name = str(resolved.get("row", "")).strip()
    _col_name = str(resolved.get("col", "")).strip()
    _row_no = resolved.get("row_number")
    # ★★ 2026-08-30（Namakoo「行と列による一意の指定も出来た方がいい」→ 実測）:
    #   「1行F列を「税込金額(10%)」にして」で、第二段は col に**書き込む値**を入れて
    #   きた（col=『税込金額(10%)』）。列の名前と値が入れ替わっている。
    #   ★ 依頼文が英字で列を名指ししているなら、それが正 ── 機械が実表から決める
    #     （行番号を機械が決めるのと同じ分担・LLM の欄の中身に頼らない）。
    _letters = {m.group(0) for m in _re_a1_col_word.finditer(_task_outside_quotes(task))}
    if len(_letters) == 1:
        _cand = next(iter(_letters)).replace("列", "").strip()
        _v2, _inf2, _err2 = resolve_col_ref(_cand, _headers_c)
        if not _err2:
            _col_name = _v2
    # ★ 2026-08-28: 列は**列文字でも**指せる（「F列に」）── resolve_col_ref が解く。
    if _col_name not in _headers_c:
        _v, _inf, _err = resolve_col_ref(_col_name, _headers_c)
        if _err:
            return False, resolved, inferred, (
                f"列『{_col_name}』がこの表にありません"
                f"（ある列: {"、".join(_headers_c)}）")
        if _inf:
            inferred.add("col")
        _col_name = _v
    if not _row_name and not _row_no:
        return False, resolved, inferred, (
            "どの行かが読み取れません（行の名前か行番号で指してください）")
    # ★★ 2026-08-28: 行番号で指された時は**番号を正**にする（人が数えて言っている）。
    #   ★ 名前も同時に在るなら、その行に本当にその名前が在るかを確かめる ──
    #     三項（依頼・宣言・実体）。食い違ったら書かずに断る。
    if _row_no:
        _hr_c = int((book_meta.get("header_rows") or {}).get(_sheet_c, 1) or 1)
        _path_c = book_meta.get("path")
        try:
            with BookView(Path(_path_c)) as _bvc:
                _wsc = _bvc.sheet(_sheet_c)
                _lastc, _colsc = data_extent(_wsc, _hr_c)
                _rowvals = [str(_wsc.cell(row=int(_row_no), column=c).value or "").strip()
                             for c in range(1, _colsc + 1)] if int(_row_no) <= _lastc else None
        except Exception as e:
            return False, resolved, inferred, f"表を読めませんでした（{type(e).__name__}）"
        # ★★ 2026-08-30（Namakoo「行と列による一意の指定も出来た方がいい。
        #   ピンポイントに操作できるようになる」）: それまで見出し行は一律で断って
        #   いたので、**列の名前を直す手段が 1 つも無かった**（実測: 計算列の見出しが
        #   「金額*1.1」に化けた表を、人が直せない）。
        #   ★ 人が**行番号と列を書いて名指しした**のは、いちばん強い証拠 ──
        #     見出しでも書かせる。ただし黙って書かない（解釈行で必ず言う）。
        #   ★ LLM が推した行・名前から解いた行では、この道は開けない（下の else 側）。
        if _rowvals is None and int(_row_no) > _hr_c:
            return False, resolved, inferred, (
                f"{_row_no}行目はこの表の範囲外です（見出しは{_hr_c}行目・データは"
                f"{_hr_c + 1}〜{_lastc}行目）")
        if int(_row_no) < _hr_c:
            return False, resolved, inferred, (
                f"{_row_no}行目は見出し行（{_hr_c}行目）より上です ── "
                "表の外には書けません")
        if int(_row_no) == _hr_c:
            if task_names_a_row_number(task) != int(_row_no):
                return False, resolved, inferred, (
                    f"{_row_no}行目は見出し行です ── 見出しの名前を変えるなら、"
                    f"行番号と列で名指ししてください"
                    f"（例:「{_hr_c}行G列を「新しい名前」にして」）")
            resolved["_writes_header"] = True
        # ★ 実測: 第二段は row に**行番号そのもの**を入れてくることがある（"7"）。
        #   それは名前ではないので、名前としては扱わない（食い違い扱いにしない）。
        if _row_name.isdigit() or _row_name == str(_row_no):
            _row_name = ""
        if _row_name and _row_name not in _rowvals:
            return False, resolved, inferred, (
                f"{_row_no}行目に『{_row_name}』がありません"
                f"（その行: {"、".join(v for v in _rowvals if v)}）── "
                "行番号と名前が食い違っています")
        _hitrow = int(_row_no)
        _note_c = f"{_hitrow}行目（依頼文の行番号）"
        if resolved.get("_writes_header"):
            # ★ 見出しを書き換える回に「対象の行:取引先」と出ると読み手を誤らせる
            #   （実測で出た）── 何をしているのかを、その言葉で言う。
            _row_name = "見出し"
            _note_c = f"{_hitrow}行目（見出し行）── 見出しの名前を変えます"
        else:
            _row_name = _row_name or (_rowvals[0] if _rowvals and _rowvals[0] else str(_hitrow))
    else:
        # ★ 行が実在し・1 つに決まることを**適用前に**確かめる（推測で別の行に書かない）。
        _hitrow, _note_c = _resolve_named_row(book_meta, _sheet_c, _row_name)
        if _hitrow is None:
            return False, resolved, inferred, _note_c
    resolved["_row_index"] = _hitrow
    # ★ 見出しを書き換えると、**その列は元の名前で引けなくなる** ── 位置を残す
    #   （検算は名前でなく座標で見る）。実測で「列『税込み金額』が見つからない」と
    #   落ちた（書き込み自体は成功していたのに）。
    if _col_name in _headers_c:
        resolved["_col_index"] = _headers_c.index(_col_name) + 1
    resolved["row"] = _row_name
    resolved["col"] = _col_name
    resolved["_headers"] = _headers_c
    resolved["_at_basis"] = _note_c
    # ★ 値は LLM に決めさせず、依頼文から機械が取る（A' 原則・SET_COLUMN_VALUE と同じ線）。
    #   ★ ただし 1 セルなので**裸の数字も受ける** ── 「梨の売上を2000にして」を
    #     引用符の有無で断るのは、道具の都合を人に押し付けている（実測の困りごと）。
    _lit = extract_quoted_literal(task)
    if _lit is None:
        # ★★ 2026-08-29（Namakoo が実測・直した先で出た穴）:
        #   「丸山工業の締め日を**2026/08/31**にして」で **31** が書かれて ✓ が出た。
        #   裸の数字を拾う正規表現が、日付の**末尾だけ**を掴んでいた。
        #   ★ 先に機械の引き算（依頼文から、既に分かっている物を引く）を通す ──
        #     こちらは値を**丸ごと**取るので、途中で切れない。
        _lit = bare_value_from_task(task, _row_name, _col_name, _headers_c)
    if _lit is None:
        _m = _re_bare_number.search(task or "")
        _lit = _m.group(1) if _m else None
    if _lit is None:
        # ★★ 2026-08-29（Namakoo が実測）: 「丸山重工の右にPCパーツ」が
        #   『文字なら「」で囲んで』で断られていた。**引用符は道具の都合**であって、
        #   人の書き方の問題ではない（この repo が何度も自分に言ってきた線）。
        #   ★ A' 原則の芯は「引用符が在ること」ではなく「**依頼文に在る値**であること」。
        #     だから条件をそちらへ置き直す: 依頼文に literal で在り・見出しの語でなく・
        #     行の名前でもない値なら、引用符が無くても受ける。
        #   ★ それでも**画面に出してから書く**（「こう読みました」に値が出る）。
        _cand = str((resolved.get("value") if resolved.get("value") is not None else ""))
        _cand = _cand.strip()
        _bad = {str(h) for h in _headers_c} | {str(_row_name), str(_col_name)}
        if _cand and _cand in (task or "") and _cand not in _bad:
            _lit = _cand
    if _lit is None:
        return False, resolved, inferred, (
            "書き込む値が依頼文から読み取れません"
            "（依頼文に書かれている値をそのまま使います ── "
            "紛らわしいときは「」で囲んでください）")
    resolved["value"] = _lit
    # ★ 2026-08-27（実測）: `_is_number` は**型**で見るので、文字列 "2000" は False。
    #   ここへ来る値は必ず文字列なので、**数字として読めるか**で判定する。
    #   ★ これを外すと `'2000'` が文字列でセルに入り、下流の SUM が静かに壊れる
    #     （この repo が何度も測ってきた形）。
    try:
        resolved["_write_numeric_value"] = float(str(_lit).replace(",", ""))
        resolved["_write_numeric"] = True
    except ValueError:
        pass
    return None


def _verify_format_map(resolved, inferred, first_sheet, book_meta, check_sheet, sheets, headers):
    """FORMAT_MAP の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := check_sheet("template_sheet")):
        return False, resolved, inferred, err
    template_sheet = resolved["template_sheet"]
    if template_sheet == first_sheet:
        return False, resolved, inferred, (
            f"雛形シートとデータシートが同じ『{template_sheet}』です。"
            "雛形は別のシートに用意してください")

    book_path = book_meta.get("path")
    if book_path is None:
        return False, resolved, inferred, (
            "様式写像段はファイルの実体が無いと検証できません（book_meta に path が無い）")
    data_headers = headers.get(first_sheet, [])
    header_row_here = book_meta.get("header_rows", {}).get(first_sheet, 1)

    try:
        wb_tpl = openpyxl.load_workbook(book_path)
    except Exception as e:
        return False, resolved, inferred, f"雛形の読み込みに失敗しました: {e}"
    try:
        tpl_ws = wb_tpl[template_sheet]
        placeholders = scan_placeholders(tpl_ws, tpl_ws.max_row or 1, tpl_ws.max_column or 1)
        if not placeholders:
            return False, resolved, inferred, (
                f"雛形『{template_sheet}』に印（{{{{列名}}}}）が見つかりません。"
                "出力したい列の直下のセルに {{列名}} の形で印を置いてください")

        # ★ 2026-08-24: 1 セルに印が 2 つ以上あるなら、埋めずに断る。埋めると 1 セルに
        #   2 回書くことになり後の値が前を消す ── 「それらしく埋まって片方が生で残る」
        #   （盲検の査定で名指しされた事故）より、雛形を直してくださいと言う方が正しい。
        if (dupes := cells_with_multiple_placeholders(placeholders)):
            cell, names = dupes[0]
            return False, resolved, inferred, (
                f"雛形『{template_sheet}』の {cell} に印が {len(names)} つあります"
                f"（{chr(12539).join(names)}）。1 つのセルに置ける印は 1 つまでです ── "
                f"別々のセルに分けてください")

        # ★ 第一波: 印は全部1つの行に置かれている前提（設計文書「見出し行 + 直下1行に印」）。
        #   最初に見つかった印の行を「印行」、その直上を「見出し行」とみなす。
        ph_row = placeholders[0].row
        header_tpl_row = ph_row - 1
        if header_tpl_row < 1:
            return False, resolved, inferred, (
                f"雛形『{template_sheet}』の印（{ph_row}行目）の上に見出し行がありません。"
                "印の1つ上の行に出力したい列名を書いてください")

        row_placeholders = sorted(
            (ph for ph in placeholders if ph.row == ph_row), key=lambda p: p.col)
        resolved_placeholders = []
        header_texts = []
        for ph in row_placeholders:
            if ph.column_name not in data_headers:
                return False, resolved, inferred, (
                    f"雛形『{template_sheet}』の印『{{{{{ph.column_name}}}}}』"
                    f"（{ph.cell}）が指す列『{ph.column_name}』は、データシート"
                    f"『{first_sheet}』に見つかりません。実在する列名を印にしてください"
                )
            col_idx = data_headers.index(ph.column_name) + 1
            if not ph.whole:
                # ★ REPORT_PER_ROW と同じ境界: 部分一致の印は原理的に文字列にしかなれない。
                try:
                    is_numeric = column_is_all_numeric(book_path, first_sheet, col_idx,
                                                        header_row_here)
                except Exception:
                    is_numeric = False
                if is_numeric:
                    return False, resolved, inferred, (
                        f"雛形『{template_sheet}』の印『{{{{{ph.column_name}}}}}』"
                        f"（{ph.cell}）はセルの一部分（部分一致）ですが、列『{ph.column_name}』は"
                        "数値です。数値列には部分一致の印を使えません"
                        "（セル全体を印にしてください: 例 " + "{{" + ph.column_name + "}}）"
                    )
            header_texts.append(tpl_ws.cell(row=header_tpl_row, column=ph.col).value)
            resolved_placeholders.append({
                "cell": ph.cell, "row": ph.row, "col": ph.col,
                "column_name": ph.column_name, "whole": ph.whole, "raw": ph.raw,
                "col_idx": col_idx, "out_col": len(resolved_placeholders) + 1,
            })
    finally:
        wb_tpl.close()
    resolved["_placeholders"] = resolved_placeholders
    resolved["_header_texts"] = header_texts
    resolved["_header_tpl_row"] = header_tpl_row
    resolved["_placeholder_tpl_row"] = ph_row

    try:
        wb_data = openpyxl.load_workbook(book_path, data_only=True)
    except Exception as e:
        return False, resolved, inferred, f"データシートの読み込みに失敗しました: {e}"
    try:
        src_ws = wb_data[first_sheet]
        last_row = _scan_last_row(src_ws, header_row=header_row_here)
        rows_in = []
        for r in range(header_row_here + 1, last_row + 1):
            label_val = src_ws.cell(row=r, column=1).value
            vals = {h: src_ws.cell(row=r, column=i + 1).value
                    for i, h in enumerate(data_headers)}
            rows_in.append((r, label_val, vals))
    finally:
        wb_data.close()
    verdict = total_row.split_total_rows_multi(rows_in) if rows_in else total_row.TotalRowVerdict(
        excluded=[], adopted_rows=[], mismatches=[])
    if not verdict.adopted_rows:
        return False, resolved, inferred, (
            "写す行がありません（表が空か、全行が合計行と判定されました）"
        )

    used = set(sheets) | {template_sheet}
    output_sheet = unique_sheet_name(str(template_sheet) + "_出力", used)
    used.add(output_sheet)
    inspection_sheet = unique_sheet_name(inspection.SHEET_NAME, used)
    used.add(inspection_sheet)

    resolved["_data_rows"] = list(verdict.adopted_rows)
    resolved["_output_sheet"] = output_sheet
    resolved["_inspection_sheet"] = inspection_sheet
    resolved["_source_headers"] = tuple(data_headers)
    return None


def _verify_swap(resolved, inferred, first_sheet, book_meta, task, sheets, headers):
    """SWAP の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    # ★★ 2026-08-31: セルの入れ替えは **a/b を要求する前**に見る。
    #   実測: 一段目が OUT_OF_VOCAB を返す言い方（「みどり建設の単価と丸和物流の
    #   単価を入れ替えて」）では a/b が空で、ここで先に落ちていた。
    #   ★ 座標は依頼文と実表だけで解ける ── LLM の返事に依存させない。
    _sheet_s0 = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _hr_s0 = int((book_meta.get("header_rows") or {}).get(_sheet_s0, 1) or 1)
    _cells0 = swap_targets_are_cells(task, book_meta, _sheet_s0, _hr_s0)
    _a = str(resolved.get("a", "")).strip()
    _b = str(resolved.get("b", "")).strip()
    if not _cells0 and (not _a or not _b):
        return False, resolved, inferred, (
            "入れ替える 2 つを取り出せませんでした"
            "（『みかんとぶどうを入れ替えて』のように 2 つの名前を書いてください）")
    if not _cells0 and _a == _b:
        return False, resolved, inferred, (
            f"『{_a}』と『{_b}』が同じものです（入れ替えになりません）")
    _sheet_s = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _hr_s = int((book_meta.get("header_rows") or {}).get(_sheet_s, 1) or 1)
    _headers_s = [str(h) for h in ((book_meta.get("headers") or {}).get(_sheet_s) or [])]
    resolved["_headers"] = _headers_s
    resolved["_header_row"] = _hr_s
    resolved["a"], resolved["b"] = _a, _b
    # ★★ 2026-08-31: 行/列を決める**前に**、セルの入れ替えでないかを見る。
    #   実測: 「丸和物流の単価とみどり建設の単価を入れ替えて」で行を丸ごと
    #   入れ替えて ✓ を出していた（頼んだのは 2 セル）。
    if (_cells := _cells0):
        # ★ 中身は**実表から読む**（LLM に値を作らせない・A' 原則）。
        _p_s = book_meta.get("path")
        try:
            with BookView(Path(_p_s)) as _bv_s:
                _ws_s = _bv_s.sheet(_sheet_s)
                _vals = [_ws_s.cell(row=r, column=c).value for r, c in _cells]
        except Exception as _e:
            return False, resolved, inferred, f"表を読めませんでした（{type(_e).__name__}）"
        if _vals[0] == _vals[1]:
            return False, resolved, inferred, (
                "入れ替える 2 つのセルの中身が同じです（入れ替えになりません）")
        resolved["_axis"] = "cell"
        # ★ 2026-08-31: a/b が空の経路（一段目が OUT_OF_VOCAB だった回）だと
        #   解釈行に「入れ替える一方: もう一方:」と**空欄**が出ていた。
        #   嘘の空欄を見せない ── その行を人が呼ぶ名前（1 列目）で埋める。
        if not _a or not _b:
            _n0 = _cell_row_name_for(book_meta, _sheet_s, _cells[0][0], _hr_s)
            _n1 = _cell_row_name_for(book_meta, _sheet_s, _cells[1][0], _hr_s)
            resolved["a"] = _a or (str(_n0) if _n0 else f"{_cells[0][0]}行目")
            resolved["b"] = _b or (str(_n1) if _n1 else f"{_cells[1][0]}行目")
        resolved["_cells"] = [list(c) for c in _cells]
        resolved["_cell_values"] = list(_vals)
        resolved["_a_pos"], resolved["_b_pos"] = _cells[0][0], _cells[1][0]
        resolved["_axis_label"] = (
            f"セル（{_headers_s[_cells[0][1] - 1]} の {_cells[0][0]}行目 と "
            f"{_headers_s[_cells[1][1] - 1]} の {_cells[1][0]}行目）")
        return True, resolved, inferred, None
    as_col = _a in _headers_s and _b in _headers_s
    _ra, _note_a = _resolve_named_row(book_meta, _sheet_s, _a)
    _rb, _note_b = _resolve_named_row(book_meta, _sheet_s, _b)
    as_row = _ra is not None and _rb is not None
    hint = _swap_axis_hint(task)
    # ★ 三項（依頼・宣言・実体）: 依頼文の「行/列」という語と、LLM が挙げた 2 つの名前と、
    #   実際の表。どれか 2 つだけで決めると、欠けた項を代用して恒真になる。
    if as_col and as_row:
        if hint is None:
            return False, resolved, inferred, (
                f"『{_a}』『{_b}』は列の見出しにも、行の中身にも両方あります ── "
                "どちらを入れ替えるのか決められません"
                "（『〜の列を入れ替えて』『〜の行を入れ替えて』と書いてください）")
        as_col, as_row = (hint == "column"), (hint == "row")
    if hint == "row" and not as_row:
        return False, resolved, inferred, f"行として決められません（{_note_a}／{_note_b}）"
    if hint == "column" and not as_col:
        return False, resolved, inferred, (
            f"列として決められません（ある列: {"、".join(_headers_s)}）")
    if not as_col and not as_row:
        return False, resolved, inferred, (
            f"入れ替える対象を決められません（{_note_a}／{_note_b}／"
            f"ある列: {"、".join(_headers_s)}）")
    resolved["_headers"] = _headers_s
    resolved["_header_row"] = _hr_s
    resolved["a"], resolved["b"] = _a, _b
    if as_col:
        resolved["_axis"] = "column"
        resolved["_axis_label"] = "列（見出しで一致）"
        resolved["_a_pos"] = _headers_s.index(_a) + 1
        resolved["_b_pos"] = _headers_s.index(_b) + 1
    else:
        resolved["_axis"] = "row"
        resolved["_axis_label"] = f"行（{_note_a}／{_note_b}）"
        resolved["_a_pos"] = _ra
        resolved["_b_pos"] = _rb
        if min(_ra, _rb) <= _hr_s:
            return False, resolved, inferred, (
                f"見出し行（{_hr_s}行目）を巻き込む入れ替えは受け付けません")
    # ★★ 2026-08-29: 入れ替えは「表に写像 π を掛ける」ことで、式もその対象。
    #   LibreOffice の自動付け替えに任せず、**π を通した式を自分で書き戻す**。
    _sh = (cellmap.swap_cols(resolved["_a_pos"], resolved["_b_pos"]) if as_col
            else cellmap.swap_rows(resolved["_a_pos"], resolved["_b_pos"]))
    _rw, _rw_why = formula_rewrites_for_shift(
        book_meta, resolved.get("_target_sheet") or first_sheet, _sh)
    if _rw_why and "実行しません" in _rw_why:
        return False, resolved, inferred, _rw_why
    if _rw:
        resolved["_formula_rewrites"] = sorted(
            (r, c, f) for (r, c), f in _rw.items())
        resolved["_formula_rewrites_label"] = (
            f"{len(_rw)} 個の式を、入れ替え後の位置に合わせて書き直します"
            "（操作前と同じ計算結果に戻ることを、適用後に読み戻して確かめます）")
    elif _rw_why:
        resolved["_warnings"] = resolved.get("_warnings", []) + [_rw_why]
    # ★ 入れ替えでも「指す先の中身が変わる式」を名指しする（並べ替えと同じ目）。
    #   ★ 軸で区画が変わるだけ ── 行なら 2 行、列なら 2 列（行と列を同じ形で書く）。
    _sw_sheet = resolved.get("_target_sheet") or first_sheet
    _lo, _hi = min(resolved["_a_pos"], resolved["_b_pos"]), max(resolved["_a_pos"],
                                                                 resolved["_b_pos"])
    _kw = ({"col_lo": _lo, "col_hi": _hi} if as_col
            else {"row_lo": _lo, "row_hi": _hi})
    # ★★ 2026-08-31（Namakoo「この指す中身が変わるとはどういうこと？」→ 実測）:
    #   「金額と単価を入れ替えて」で ⚠ が出るが、**中身が 3 つとも事実に反していた**:
    #     ・「指す先の中身が変わる」→ 変わらない（税込金額は金額を指し続けた）
    #     ・「**行**が入れ替わる」   → 入れ替えたのは列
    #     ・「直していません」       → **直している**（=E2*1.1 → =D2*1.1）
    #   そして嘘の ⚠ のせいで、正しく動いた操作の ✓ が △ に落ちていた。
    #   ★ 片配線の**逆**: 警告は並べ替え（式を直さない）用に作って入れ替えにも配線し、
    #     そのあと入れ替えだけ式を直すようになったのに、警告は昔の前提のまま残った。
    #   ★ 書き直す式は名指しから外す（別シートから指す式は書き直さないので残す）。
    if (_dw := reference_drift_warning(book_meta, _sw_sheet,
                                        rewritten=set(_rw),
                                        unit=("列" if as_col else "行"), **_kw)):
        resolved["_warnings"] = resolved.get("_warnings", []) + [_dw]
    return None


def _verify_add_row(resolved, inferred, book_meta, task, sheets, headers, op):
    """ADD_ROW の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    # ★★ 2026-08-27（Namakoo が実測）: 「みかんの下に梨を追加して」が動かなかった。
    #   位置を**行番号**でしか受け取れないのに、人は相対で言う。LLM に数えさせると
    #   外し、空行だけの INSERT_ROWS に落ちていた。
    #   ★ 分担を変える: LLM は「誰の隣か」を言うだけ／**行番号は機械が実表を数えて決める**
    #     （列名の解決を機械 3 段でやっているのと同じ形）。
    #   ★ 機械が決めた位置は LLM の数字より優先する ── 実表を見た側が正しい。
    # ★★ 2026-09-16（盲検の買い手役 2 体目・再現 2/2・EXIT=0 で通っていた）:
    #   「納期が…より後の行のチェック列に★を入れて」が**行追加**に化け、受注台帳に
    #   取引先も金額も空の行が入った。★ 買い手の値段の回答は 0 円で、理由はこの 1 件だった。
    #   ★ 「その op にしてよいか」を依頼文に問い返す器官（OP_META の requires_word）は在るのに、
    #     宣言していたのは PIVOT ただ 1 つ ── 行を足す op は「足す意図」を何も要求していなかった。
    #   ★ 判定は ailine_core/intent.py に 1 つだけ置き、ここは材料を渡すだけ（既存の作法）。
    _sheet0 = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _hr0 = int((book_meta.get("header_rows") or {}).get(_sheet0, 1) or 1)
    _anchor0: dict = {}
    _at_anchor, _anchor_note = resolve_row_anchor(
        task, book_meta, _sheet0, header_row=_hr0, anchor_out=_anchor0)
    if _anchor_note and _at_anchor is None:
        # ★★ 2026-09-07（外部の検品・Namakoo 決裁「削除は聞く」）: 名前が**複数行**に
        #   当たった回、旧版は断って「行番号で指してください」と言っていた。断り方は
        #   親切だが、買い手の目には**「消してと言ったのに何も起きない」＝失敗**だ。
        #   ★ 決められないのではなく **聞けば決まる** ── 行を全部挙げて確認を取る。
        #   ★ 削除は取り返しがつかないので**必ず聞く**（既存の破壊の関所に載せる ──
        #     新しい関所も新しい exit code も作らない）。
        _rows0 = list(_anchor0.get("rows") or [])
        if op == "DELETE_ROWS" and len(_rows0) > 1:
            resolved["_delete_rows"] = _rows0
            # ★ 宣言を実体に合わせる ── 2 行消すのに「行数:1」と出したら、
            #   今日ずっと潰している「宣言と実体のずれ」を自分で作ることになる。
            resolved["at"], resolved["count"] = _rows0[0], len(_rows0)
            resolved["_at_basis"] = (
                f"『{_anchor0.get('name')}』の行＝"
                + "、".join(f"{r}行目" for r in _rows0))
            resolved["_confirm_delete"] = (
                f"『{_anchor0.get('name')}』に当てはまる {len(_rows0)} 行"
                f"（{'、'.join(str(r) for r in _rows0)}行目）を削除します")
        else:
            return False, resolved, inferred, _anchor_note
    if _at_anchor is not None:
        resolved["at"] = _at_anchor
        resolved["_at_basis"] = _anchor_note
    at_raw = str(resolved.get("at", "")).strip()
    if not (at_raw.isdigit() and int(at_raw) >= 1):
        return False, resolved, inferred, f"行番号『{resolved.get('at')}』が不正です（1以上の整数）"
    resolved["at"] = int(at_raw)
    _sheet = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _hr = int((book_meta.get("header_rows") or {}).get(_sheet, 1) or 1)
    # ★ 見出し行より上を触らせない（表の骨格を壊す操作は受け付けない）。
    if int(at_raw) <= _hr:
        return False, resolved, inferred, (
            f"{at_raw}行目は見出し行（{_hr}行目）またはその上です ── "
            "見出しを壊す操作は受け付けません")
    if op == "DELETE_ROWS":
        c = resolved.get("count")
        if c in (None, ""):
            resolved["count"] = 1
            inferred.add("count")
        else:
            cs = str(c).strip()
            if not (cs.isdigit() and int(cs) >= 1):
                return False, resolved, inferred, f"削除行数『{c}』が不正です（1以上の整数）"
            resolved["count"] = int(cs)
        # ★★ 2026-09-17（盲検 3 体目から辿った・同じ家系の片割れ）: 上の「名前が複数行に
        #   当たった」回だけが関所に載っていて、**行番号で指した削除は素通り**していた
        #   （実測:「3行目を削除して」で値 11 個が消えて exit 0）。
        #   ★ 2026-09-07 の決裁「削除は取り返しがつかないので必ず聞く」を、宣言
        #     （WRITE_REMOVE）の家系ぜんぶに配線する ── 片方だけ直すと次の片配線を作る。
        #   ★ 鳴る条件は上書き・列削除と対称: **消えるものが在る時だけ**（空行は黙って消す）。
        if not resolved.get("_confirm_delete"):
            _lost = _rows_existing_value_count(
                book_meta.get("path"), _sheet, int(resolved["at"]), int(resolved["count"]))
            if _lost > 0:
                _n = int(resolved["count"])
                _where = (f"{resolved['at']}行目"
                          if _n == 1 else
                          f"{resolved['at']}行目から {_n} 行")
                resolved["_confirm_delete"] = (
                    f"{_where}には値が {_lost} 件あります（その行ごと削除します）")
    else:
        # ★ 値は**列名で**受ける。実在しない列名はここで弾く（幻覚の封鎖）。
        vals = resolved.get("values")
        # ★ 2026-08-27（実測）: LLM は values を**並び**で返すことがある
        #   （['梨', 600, 300]）。列名の対応は**機械が付けられる** ── 左から順に
        #   当てる。多すぎる時だけ断る（推測で余りを捨てない）。
        #   ★ 決めた対応は解釈行に出す（_values_label）── 黙って割り当てない。
        _headers_now = (book_meta.get("headers") or {}).get(
            resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]) or []
        if isinstance(vals, (list, tuple)):
            if len(vals) > len(_headers_now):
                return False, resolved, inferred, (
                    f"入れる値が {len(vals)} 個ありますが、列は {len(_headers_now)} 本です"
                    f"（ある列: {"、".join(map(str, _headers_now))}）")
            # ★ 2026-08-27（実測・俺が入れた壊し方）: LLM は埋まらない列を None で
            #   返すことがある（['梨', None, None]）。そのまま渡すと codegen が
            #   `str(None)` を書き、セルに**文字列 "None"** が入った。
            #   ★ 指定の無い列には**何も書かない**（空欄のままにする）。
            #   ★ 事後条件はこの壊れ方を捕まえていた（rc=1）── 番人は効いていたが、
            #     壊れた物を作ってから気づく形だったので、入口で落とす。
            vals = {str(h): v for h, v in zip(_headers_now, vals)
                     if v is not None and v != ""}
            resolved["values"] = vals
            inferred.add("values")
        if not isinstance(vals, dict) or not vals:
            return False, resolved, inferred, (
                "入れる値が読み取れません（列名と値の組で書いてください）")
        headers = (book_meta.get("headers") or {}).get(
            resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]) or []
        unknown = [k for k in vals if str(k) not in [str(h) for h in headers]]
        if unknown:
            return False, resolved, inferred, (
                f"列『{"、".join(map(str, unknown))}』がこの表にありません"
                f"（ある列: {"、".join(map(str, headers))}）")
        # ★ 同上: 値が空の列は書かない（"None" という文字列を作らない）。
        resolved["values"] = {str(k): v for k, v in vals.items()
                               if v is not None and v != ""}
        # ★★ 2026-09-10（出荷前の実機テストが落ちて分かった・実測 11 回中 7 回）:
        #   「5行目に丸山工業の行を作って」に対し、モデルが件数 1・単価 1000・金額 1000 を
        #   **発明して**書き、`✓` が出ていた。依頼に値の指定はどこにも無い。
        #   ★ 重いのは 2 次被害の方 ── 発明した値は下の継承の除外集合に入るので、
        #     **その列の式が入らなくなる**（金額が =B5*C5 でなく直値 1000 になり、
        #     件数を直しても追随しない）。判定は ailine_core/intent.py に 1 つだけ。
        #   ★ 落とすのは「依頼文に接地しない値」だけ。落とした列は必ず名指しで出す
        #     （黙って空にするのは別の嘘になる）。空欄は誤値より安い。
        _ungrounded = intent_mismatch.values_not_grounded_in_the_request(
            task, resolved["values"])
        if _ungrounded:
            resolved["_dropped_label"] = (
                "／".join(_ungrounded) + "（依頼に無いので空のままにします）")
            resolved["values"] = {k: v for k, v in resolved["values"].items()
                                   if k not in set(_ungrounded)}
        if not resolved["values"]:
            return False, resolved, inferred, (
                "依頼文に入れる値が見当たりません（頼まれていない値は書きません）"
                " ── 空の行を入れるだけなら「"
                f"{resolved.get('at')}行目に空行を入れて」と頼めます")
        resolved["_headers"] = [str(h) for h in headers]
        resolved["_values_label"] = "／".join(
            f"{k}={v}" for k, v in resolved["values"].items())
        # ★★ 2026-09-02（README の「既知の問題」に自分で書いていた）:
        #   追加した行に既存の式が引き継がれず、利益列が**空のまま**だった。
        #   宣言した値は正しいので ✓ は正しいが、人の期待とは違う。
        #   ★ 引き継ぐのは「全データ行が式を持つ列」だけ ── 形で決める（列挙しない）。
        #     合計列は E2..E7 が直値なので自然に外れる。
        #   ★ **黙ってやらない。**解釈行に出す（_inherit_label）。
        if op == "ADD_ROW" and resolved.get("at"):
            _ih_sheet = resolved.get("_target_sheet")
            _ih_hr = int(resolved.get("_header_row") or 1)
            _ih_cols, _ih_from = formula_columns_to_inherit(
                book_meta, _ih_sheet, _ih_hr, int(resolved["at"]),
                set(resolved["values"].keys()))
            if _ih_cols and _ih_from:
                resolved["_inherit_cols"] = _ih_cols
                resolved["_inherit_from"] = _ih_from
                _hd = resolved["_headers"]
                resolved["_inherit_label"] = "／".join(
                    (_hd[c] if 0 <= c < len(_hd) else f"{c + 1}列目")
                    for c in _ih_cols) + f"（{_ih_from}行目から）"
    return None


def _verify_add_column(resolved, inferred, book_meta, task, sheets, headers):
    """ADD_COLUMN の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    _sheet_c = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _hr_c = int((book_meta.get("header_rows") or {}).get(_sheet_c, 1) or 1)
    _headers_c = [str(h) for h in ((book_meta.get("headers") or {}).get(_sheet_c) or [])]
    _name_c = str(resolved.get("name") or "").strip()
    # ★★ 2026-08-27（実測）: 名前を言っていない依頼に対し、LLM が「新しい列」という
    #   **依頼文に無い名前を作って**返す回があった（3 回中 1 回）。A' 原則の違反 ──
    #   値は LLM に確定させない。**依頼文に現れない名前は採らない**（空欄に倒す）。
    #   ★ 空欄は誤った名前より安い: 見出しが空なら △ になり、人が気づける。
    #     もっともらしい名前が付くと、人は「自分がそう言った」と思ってしまう。
    if _name_c and _name_c not in (task or ""):
        resolved["_name_dropped"] = _name_c
        _name_c = ""
    # ★ 同名の列が既に在るなら断る（黙って 2 本目を作らない ── 後で列名の解決が
    #   「2 つあります」で詰まる形を、作る側で防ぐ）。
    if _name_c and _name_c in _headers_c:
        return False, resolved, inferred, (
            f"列『{_name_c}』は既にあります（{_headers_c.index(_name_c) + 1}列目）")
    _at_c, _note_c = resolve_col_anchor(task, _headers_c)
    if _at_c is None and _note_c:
        return False, resolved, inferred, _note_c
    if _at_c is None:
        # 位置の言い回しが無い＝末尾。**黙って決めない**ので根拠を必ず出す。
        _at_c = len(_headers_c) + 1
        _note_c = f"末尾＝{_at_c}列目（依頼文に位置の指定が無いため）"
    # ★★ 2026-08-27（Namakoo が GUI で実測）: 見出しも値も無い列を**末尾**に足すと、
    #   セルは 1 つも増えないので機械には**何も変わって見えない**（物理の使用範囲は
    #   値のあるセルで測るため）。事後条件は正しく「列数が合わない」で × を出すが、
    #   利用者には「動かなかった」としか見えない ── **やる前に断って理由を言う**。
    #   ★ 途中に挿す場合は右の列がずれるので見える（そちらは通す）。
    if not _name_c and int(_at_c) > len(_headers_c):
        return False, resolved, inferred, (
            "見出しも値も無い列を末尾に足しても、ファイルの中身は何も変わりません"
            "（空の列はセルを持たないので機械にも見えません）── "
            "見出しの名前を言ってください（例: 「原価の右にチェックという列を追加して」）")
    resolved["name"] = _name_c
    resolved["_at_col"] = int(_at_c)
    resolved["_at_basis"] = _note_c
    resolved["_headers"] = _headers_c
    resolved["_header_row"] = _hr_c
    # ★ 名前が無いなら「空のまま」と画面に書く（黙って空欄を作らない）。
    # ★ 見出しが空の列を作ると、**その右にある列も走査できなくなる**
    #   （走査は見出し行の最初の空で止まる）── 作る前に言う。判定は △ に落ちる。
    _dropped = resolved.get("_name_dropped")
    resolved["_name_label"] = _name_c or (
        ("（依頼文に無い名前『%s』は採りませんでした・" % _dropped if _dropped
          else "（名前なし・")
        + "見出しは空のまま ── 右にある列も走査できなくなり、判定は △ になります）")
    return None


def _verify_chart(resolved, inferred, first_sheet, resolve_in, task, headers):
    """CHART の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := resolve_in("value_col", first_sheet)):
        return False, resolved, inferred, err
    # ★ グラフ段①: kind も A' 原則の中へ（cmp と同じ作法）。依頼文からの機械抽出が
    #   非 None かつ LLM の kind と食い違えば機械が勝つ（EXTRACT の cmp と同じ形）。
    llm_kind_raw = str(resolved.get("kind") or "").strip().lower()
    mechanical_kind = extract_chart_kind_from_task(task)
    if mechanical_kind is not None and mechanical_kind != llm_kind_raw:
        kind = mechanical_kind
        resolved["_warnings"] = resolved.get("_warnings", []) + [
            f"LLM が返した種類({llm_kind_raw or '(空)'})と依頼文の機械抽出({mechanical_kind})が"
            f"食い違うため機械抽出({mechanical_kind})を採用しました"
        ]
    else:
        kind = llm_kind_raw or "bar"
    if kind not in _CHART_KINDS:
        return False, resolved, inferred, (
            f"グラフ種類『{resolved.get('kind')}』は {'/'.join(_CHART_KINDS)} のどれでもありません"
        )
    resolved["kind"] = kind
    # ★ グラフ段②: category_col(省略可・既定は先頭列)。指定があれば実在列検証。
    raw_cat = resolved.get("category_col")
    if raw_cat in (None, ""):
        first_col = (headers.get(first_sheet) or [None])[0]
        if first_col is None:
            return False, resolved, inferred, f"シート『{first_sheet}』に列がありません"
        resolved["category_col"] = first_col
        inferred.add("category_col")
    elif (err := resolve_in("category_col", first_sheet)):
        return False, resolved, inferred, err
    return None


def _verify_insert_rows(resolved, inferred, book_meta, task, sheets, op):
    """INSERT_ROWS の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    # ★ 2026-08-27（実測）:「みかんの下に空行を入れて」で LLM が 3 行目と言った
    #   （みかんが 3 行目なので、下は 4 行目）。**位置は op に関係なく位置** ──
    #   同じ機械の解決を通す（片配線を作らない）。
    _sheet_i = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    _hr_i = int((book_meta.get("header_rows") or {}).get(_sheet_i, 1) or 1)
    _at_i, _note_i = resolve_row_anchor(task, book_meta, _sheet_i, header_row=_hr_i)
    # ★ 2026-08-27（実測）: 見つからなかった時に**黙って LLM の行番号へ落ちて**いた。
    #   位置を名指しした依頼で場所が特定できないなら、推測で挿さずに断る
    #   （静かに別の場所へ入るのが一番こわい ── ADD_ROW と同じ線に揃える）。
    if _note_i and _at_i is None:
        return False, resolved, inferred, _note_i
    if _at_i is not None:
        resolved["at"] = _at_i
        resolved["_at_basis"] = _note_i
    at_raw = str(resolved.get("at", "")).strip()
    if not (at_raw.isdigit() and int(at_raw) >= 1):
        return False, resolved, inferred, f"行番号『{resolved.get('at')}』が不正です（1以上の整数）"
    resolved["at"] = int(at_raw)
    count_raw = resolved.get("count")
    if count_raw in (None, ""):
        resolved["count"] = 1
        inferred.add("count")
    else:
        count_str = str(count_raw).strip()
        if not (count_str.isdigit() and int(count_str) >= 1):
            return False, resolved, inferred, f"挿入行数『{count_raw}』が不正です（1以上の整数）"
        resolved["count"] = int(count_str)
    return None


def _verify_extract_columns(resolved, inferred, first_sheet, book_meta, task, headers):
    """EXTRACT_COLUMNS の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    _sheet_x = resolved.get("_target_sheet") or first_sheet
    _hdrs_x = [str(h) for h in (headers.get(_sheet_x) or [])]
    # ★ 値は LLM に作らせない: **依頼文に現れる実在の列名**を、出現順に機械が拾う。
    #   LLM の cols は「候補の当たり」としてだけ使い、実在照合を通ったものだけ採る。
    _asked = [c for c in _hdrs_x if c and c in (task or "")]
    _llm_cols = resolved.get("cols")
    if isinstance(_llm_cols, str):
        _llm_cols = [x.strip() for x in _llm_cols.split(",") if x.strip()]
    _llm_cols = [str(c) for c in (_llm_cols or []) if str(c) in _hdrs_x]
    _cols_x = _asked or _llm_cols
    if not _cols_x:
        return False, resolved, inferred, (
            "残す列が依頼文から読み取れません"
            f"（ある列: {'、'.join(_hdrs_x)}）── 列名をそのまま書いてください")
    if len(_cols_x) >= len(_hdrs_x):
        return False, resolved, inferred, (
            "全部の列が指定されています（抜き出す意味がありません）")
    # ★ 依頼文の出現順に並べる（人が書いた順で出す ── 表の順に勝手に直さない）
    _cols_x = sorted(set(_cols_x), key=lambda c: (task or "").find(c))
    resolved["cols"] = _cols_x
    resolved["_cols_label"] = "・".join(_cols_x)
    resolved["_headers"] = _hdrs_x
    resolved["_header_row"] = int((book_meta.get("header_rows") or {}).get(_sheet_x, 1) or 1)
    resolved["_source_headers"] = tuple(_hdrs_x)
    resolved["_new_sheet"] = _EXTRACT_SHEET_NAME_FORBIDDEN_RE.sub(
        "_", "・".join(_cols_x) + "だけ")[:31]
    return None


def _verify_dedup(resolved, inferred, first_sheet, headers):
    """DEDUP の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    raw_keys = resolved.get("keys")
    if not isinstance(raw_keys, list) or not raw_keys:
        return False, resolved, inferred, (
            "重複を判定する列が依頼文から読み取れません。どの列が同じなら重複とみなすか、"
            "依頼文に列名を書いてください（例:「取引先が同じ行を重複として除いて」）"
        )
    resolved_keys = []
    for raw_key in raw_keys:
        v, was_inferred, err = resolve_col_ref(raw_key, headers.get(first_sheet, []))
        if err:
            return False, resolved, inferred, err
        resolved_keys.append(v)
        if was_inferred:
            inferred.add("keys")
    resolved["keys"] = resolved_keys
    # ★ 単位H: EXTRACT と同じ作法（出力シートの見出し署名の材料を resolved に積む）。
    resolved["_source_headers"] = tuple(headers.get(first_sheet, []))
    resolved["_new_sheet"] = _dedup_output_sheet_name(resolved_keys)
    return None


def duplicate_rows(path, sheet: str, keys, header_row: int = 1) -> tuple:
    """実物の表を読み、**2 件目以降**の重複行（1 起点の行番号）を上から順に返す。

    ★★ 2026-09-21（盲検 5 体目・Namakoo 決裁「消す側を作る」）: 削除する行を
      **Python 側が決める**。Basic 側で探索させない理由は 2 つ:
      ① 確認の文と実際に消す行が**同じ材料**から出る（見せた数と違う数を消さない）
      ② 「何が重複か」の判断を 2 つ書かない（Basic にもう 1 つ書けば、片方だけ直る日が来る）
    ★ 鍵の規則は事後条件と同じ `_dedup_normalize_key_part`（前後空白だけ落とす・型が違えば別）。
    ★ 読めない・列が無いなら空 ── 根拠が無い時に消しに行かない。
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:   # noqa: BLE001
        return ()
    try:
        if sheet not in wb.sheetnames:
            return ()
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) <= header_row:
            return ()
        head = [str(c) if c is not None else "" for c in rows[header_row - 1]]
        idxs = [head.index(k) for k in keys if k in head]
        if len(idxs) != len(list(keys)):
            return ()
        seen, dup = set(), []
        for n, row in enumerate(rows[header_row:], start=header_row + 1):
            key = tuple(_dedup_normalize_key_part(row[i] if i < len(row) else None) for i in idxs)
            if key in seen:
                dup.append(n)
            else:
                seen.add(key)
        return tuple(dup)
    except Exception:   # noqa: BLE001
        return ()
    finally:
        wb.close()


def _verify_dedup_delete(resolved, inferred, first_sheet, headers, book_meta):
    """DEDUP_DELETE の引数を確かめ、**消える行を数えて確認に積む**。

    ★★ 削除は取り返しがつかない（2026-09-07 の決裁「削除は必ず聞く」）。
      `_confirm_delete` を積むと、既存の破壊の関所がそのまま「削除しますか？」で聞く
      ── 新しい関所も新しい終了コードも作らない。
    ★ 鳴る条件は上書き・列削除と**対称**: **消えるものが在る時だけ**（重複が 0 なら黙る）。
    """
    r = _verify_dedup(resolved, inferred, first_sheet, headers)
    if r is not None:
        return r
    # ★ 非破壊版が積む出力シート名はこの op には無い（別シートを作らない）。
    resolved.pop("_new_sheet", None)
    sheet = resolved.get("_target_sheet") or first_sheet
    hr = (book_meta.get("header_rows") or {}).get(sheet, 1)
    dup = duplicate_rows(book_meta.get("path"), sheet, resolved["keys"], header_row=hr)
    if not dup:
        # ★ 破壊する op で**無言の no-op** を作らない ── 消すものが無いなら、そう言って止まる。
        return False, resolved, inferred, (
            f"重複している行はありません（判定キー: {'・'.join(resolved['keys'])}）"
            " ── 消すものが無いので、何もしていません")
    # ★ 鍵の名前は既存の DELETE_ROWS と**同じ** ── 生成部（_codegen_delete_rows）を
    #   そのまま使い回す。「下から順に消す」を 2 箇所に書かない。
    resolved["_delete_rows"] = list(dup)
    _shown = "・".join(str(n) for n in dup[:5]) + ("…" if len(dup) > 5 else "")
    resolved["_confirm_delete"] = (
        f"重複が {len(dup)} 行あります（{_shown} 行目 ── その行ごと削除します。"
        f"判定キー: {'・'.join(resolved['keys'])}）")
    return None


def _verify_split_cell(resolved, inferred, first_sheet, book_meta, resolve_in, task):
    """SPLIT_CELL の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := resolve_in("col", first_sheet)):
        return False, resolved, inferred, err
    # ★★ 2026-09-04: 区切りは**依頼文から機械が取る**（A' 原則 ── SET_COLUMN_VALUE の
    #   value と同じ作法）。実測で「読点で分けて」に対し模型が sep="," を返し、
    #   「カンマが見つからない」で止まっていた（断り方は正しいが依頼は読めている）。
    #   ★ 引用の中身は**それが区切りとして意味を持つ時だけ**採る（「『済』の行を…」の
    #     ような、区切りと無関係な引用に引きずられないため）。
    llm_sep_raw = resolved.get("sep")
    _quoted = extract_quoted_literal(task)
    text_sep = (split_cell.SEPARATOR_ALIASES.get(_quoted)
                if _quoted in split_cell.SEPARATOR_ALIASES else None)
    _basis = f"依頼文の引用: 「{_quoted}」" if text_sep else None
    if text_sep is None:
        text_sep = split_cell.separator_named_in(task)
        _basis = "依頼文の名指し" if text_sep else None
    sep = text_sep or split_cell.normalize_separator(llm_sep_raw)
    if text_sep is not None:
        resolved["_sources"] = {**resolved.get("_sources", {}), "sep": _basis}
        _llm_norm = split_cell.normalize_separator(llm_sep_raw)
        if _llm_norm is not None and _llm_norm != text_sep:
            _a = split_cell.describe_separator(_llm_norm)
            _b = split_cell.describe_separator(text_sep)
            resolved["_warnings"] = resolved.get("_warnings", []) + [
                f"LLM が返した区切り('{_a}')と依頼文('{_b}')が食い違うため"
                f"依頼文側('{_b}')を採用しました"]
    if not sep:
        return False, resolved, inferred, (
            "区切り(sep)が読み取れません（改行/カンマ/、/スペース などで指定してください）"
        )
    resolved["sep"] = sep
    # ★ 何列必要かは**実データ**が決める（LLM に数えさせない）。
    values = _column_values(book_meta, first_sheet, resolved["col"])
    parts = split_cell.max_parts(values, sep)
    if parts < 2:
        return False, resolved, inferred, (
            f"列『{resolved['col']}』に区切り『{split_cell.describe_separator(sep)}』が"
            f"見つからないため、分けられません"
        )
    resolved["_parts"] = parts
    resolved["_new_cols"] = [f"{resolved['col']}_{k}" for k in range(1, parts + 1)]
    return None


def _verify_number_format(resolved, inferred, first_sheet, book_meta, resolve_in, task=""):
    """NUMBER_FORMAT の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    # ★ 2026-08-29: 行にも掛けられるようにした（行と列は軸違い）。
    #   行が指定されている回は列を要求しない。
    if resolved.get("row_number"):
        _nf_sheet = resolved.get("_target_sheet") or first_sheet
        _nf_hr = int((book_meta.get("header_rows") or {}).get(_nf_sheet, 1) or 1)
        _nf_row = int(resolved["row_number"])
        if _nf_row <= _nf_hr:
            return False, resolved, inferred, (
                f"{_nf_row}行目は見出し行（{_nf_hr}行目）またはその上です")
        resolved["_row_index"] = _nf_row
        resolved.pop("col", None)
        resolved["_at_basis"] = f"{_nf_row}行目"
    elif (err := resolve_in("col", first_sheet)):
        return False, resolved, inferred, err
    if resolved.get("style") != "thousands":
        return False, resolved, inferred, f"書式『{resolved.get('style')}』は未対応です（対応: thousands）"
    # ★★ 2026-09-07: 上の検査は**宣言**しか見ていない。LLM は「円マーク」を、持っている
    #   書式（thousands）へ**正規化して**返すので素通りし、¥ が付かないまま ✓ が出ていた。
    #   ★ 依頼文の側を見る（三項のうち依頼が落ちる形は、今日これで 4 つ目）。
    if (_fmt := intent_mismatch.format_asked_but_not_supported(
            task, (book_meta.get("headers") or {}).get(first_sheet) or [])):
        return False, resolved, inferred, (
            f"{_fmt}はこの道具の数値書式では扱えません（扱えるのは桁区切りだけです）"
            "── 桁区切りだけでよければ「桁区切りを付けて」と頼んでください。"
            "要望として記録します")
    return None


def _verify_delete_column(resolved, inferred, book_meta, sheets, headers):
    """DELETE_COLUMN の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    name = str(resolved.get("col", "")).strip()
    headers = (book_meta.get("headers") or {}).get(
        resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]) or []
    if name not in [str(h) for h in headers]:
        return False, resolved, inferred, (
            f"列『{name}』がこの表にありません（ある列: {"、".join(map(str, headers))}）")
    if len([h for h in headers if str(h) != ""]) <= 1:
        return False, resolved, inferred, "列が 1 本しかないので削除できません"
    resolved["col"] = name
    resolved["_headers"] = [str(h) for h in headers]
    # ★★ 2026-09-17（盲検 3 体目・製造業の購買）: 「検収日の列を消して」で、値の入った
    #   15 件が**確認なしに** exit 0 の ✓ で消えた。買い手の言葉:
    #   「上書きより削除の方が怖いのに、厳しい方が緩い」。
    #   ★ 関所（破壊の関所）も、削除用の聞き文（「削除しますか？」）も**既に在った** ──
    #     配線されていたのは DELETE_ROWS の「名前が複数行に当たった」1 ケースだけで、
    #     列の削除はどれだけ値が入っていても素通りしていた（器官は在るが配線が無い）。
    #   ★ 2026-09-07 の決裁は「削除は取り返しがつかないので**必ず聞く**」。目の前に在った
    #     1 ケースにだけ配線された形なので、同じ数え方（上書きの関所が使う件数）を借りる。
    #   ★ 鳴る条件は上書きと対称にする ── **消えるものが在る時だけ**（空の列は黙って消す）。
    _sheet_d = resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]
    if book_meta.get("path") and _sheet_d:
        _lost = _column_existing_value_count(
            Path(book_meta["path"]), _sheet_d, name,
            header_row=int((book_meta.get("header_rows") or {}).get(_sheet_d, 1) or 1))
        if _lost > 0:
            resolved["_confirm_delete"] = (
                f"列『{name}』には値が {_lost} 件あります（この列ごと削除します）")
    return None


def _verify_move_column(resolved, inferred, book_meta, task):
    """MOVE_COLUMN の引数を確かめ、**目的地を機械が決める**。

    ★ 分担は列追加と同じ ── LLM は「どの列を」だけ言う。「どこへ」は依頼文の
      言い回し（一番左／金額の右／…）から `resolve_col_anchor` が実表の見出しで解く。
    ★ 位置が読み取れなければ**動かさない**（黙って端へ寄せない ── 静かに違う所へ
      入るのが一番こわい、を列でも同じに扱う）。
    """
    name = str(resolved.get("col", "")).strip()
    headers = [str(h) for h in ((book_meta.get("headers") or {}).get(
        resolved.get("_target_sheet") or (book_meta.get("sheets") or [None])[0]) or [])]
    if name not in headers:
        return False, resolved, inferred, (
            f"列『{name}』がこの表にありません（ある列: {"、".join(headers)}）")
    if len([h for h in headers if h != ""]) <= 1:
        return False, resolved, inferred, "列が 1 本しかないので動かせません"
    at, note = resolve_col_anchor(task, headers)
    if at is None:
        return False, resolved, inferred, (
            note or "どこへ動かすかが依頼文から読み取れません"
            "（例:「取引先の列を一番左に」「担当を金額の右に」）")
    src0 = headers.index(name)
    # ★★ `at` は「今の見出しの何番目の**手前**に入れるか」（1起点）。動かす列は
    #   自分もその並びに居るので、自分より右を指したら**自分が抜けた分**だけ詰まる。
    #   ★ ここが持つのは**動いた後の最終位置**（事後条件が突き合わせるのもこれ）。
    #     ヘルパ(MoveColumnTo)が要る引数は別物なので、codegen 側で換算する
    #     ── 「約束する形」と「道具に渡す形」を混ぜない。
    to0 = int(at) - 1
    if src0 < to0:
        to0 -= 1
    to0 = max(0, min(to0, len(headers) - 1))
    if to0 == src0:
        return False, resolved, inferred, (
            f"列『{name}』はもうそこに在ります（{src0 + 1}列目）── 動かす必要がありません")
    resolved["col"] = name
    resolved["_headers"] = headers
    resolved["_move_from"] = src0
    resolved["_move_to"] = to0
    resolved["_at_basis"] = note
    resolved["_to_label"] = f"{to0 + 1}列目"
    # ★★ 移動を出す口は repo に**1 つだけ**（codegen の wrap）── 2 本目を書くと
    #   片配線になるので、番人が機械で止める（test_column_placement）。
    #   ★ ここは横断層と同じ鍵に材料を置くだけ。ヘルパの第2引数は「1本広げた
    #     途中の並び」での位置なので、右へ動かす回だけ最終位置より 1 大きい。
    resolved["_new_col_from"] = src0
    resolved["_move_new_col_to"] = to0 + 1 if to0 > src0 else to0
    # ★ 動いた後の見出しの並びを**先に**確定して宣言に載せる（事後条件が突き合わせる）。
    rest = headers[:src0] + headers[src0 + 1:]
    resolved["_headers_after"] = rest[:to0] + [name] + rest[to0:]
    return None


def _verify_pivot(resolved, inferred, first_sheet, resolve_in):
    """PIVOT の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    # ★ AGGREGATE と同じ2 slot（group_col/value_col）。列名の実在確認だけ共有する。
    if (err := resolve_in("group_col", first_sheet)):
        return False, resolved, inferred, err
    if (err := resolve_in("value_col", first_sheet)):
        return False, resolved, inferred, err
    return None


#: ★ 集計は**合計しか無い** ── 「何件」を数量の合計で代用しない（盲検 #50 の誤配）。
_COUNT_WORDS = ("件数", "何件", "回数", "人数", "何人", "何回")


def _verify_aggregate(resolved, inferred, first_sheet, resolve_in, task="", sheet_headers=()):
    """AGGREGATE の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if (err := resolve_in("group_col", first_sheet)):
        return False, resolved, inferred, err
    # ★★ 2026-09-14（盲検 #50「担当者ごとに何件受注したか件数も出して」→ 数量の**合計**）:
    #   件数を数える器がこの道具に無いのに、近い操作で代用して別の数字を件数として出していた。
    #   ★ 語彙の穴は穴と言う（代用しない）── ただし合計する列を依頼文が名指ししていれば通す。
    if task and any(w in task for w in _COUNT_WORDS):
        _named = [h for h in (sheet_headers or ()) if h and h in task]
        if str(resolved.get("value_col") or "") not in _named:
            return False, resolved, inferred, (
                "件数は数えられません（この道具の集計は合計だけです）── "
                "合計なら「担当者ごとに金額を集計して」のように、合計する列を書いてください")
    if (err := resolve_in("value_col", first_sheet)):
        return False, resolved, inferred, err
    return None


def _verify_merge(resolved, inferred):
    """MERGE の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    if not re.fullmatch(r"[A-Za-z]{1,3}\d+:[A-Za-z]{1,3}\d+", str(resolved.get("range", ""))):
        return False, resolved, inferred, f"範囲『{resolved.get('range')}』の形式が不正です（例: A1:C1）"
    return None


def _verify_draw_borders():
    """DRAW_BORDERS の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    pass   # 引数無し・表全体が対象（検証することが無い）
    return None


def _verify_autofit():
    """AUTOFIT の引数を確かめる（★ verify_dsl_args から切り出した・挙動不変）。

    ★ 返り値は **返すべき tuple か None（＝続行）**。op 分岐は「早期 return するか、
      何も返さずチェーンを抜ける」の 2 通りしかないので、この形で 1 つずつ剥がせる。
    ★ resolved(dict) / inferred(set) は参照が渡り、副作用はそのまま伝わる。
    ★ 挙動不変は bench/verify_golden.json（641 件）との突き合わせで確かめている。
    """
    pass   # 引数無し・全列が対象（検証することが無い）
    return None


_EXTRACT_SHEET_NAME_FORBIDDEN_RE = re.compile(r'[:\\/?*\[\]]')


# ★ グラフ段①: kind の機械抽出（cmp と同じ作法・extract_cmp_from_task の兄弟）。
#   折れ線/推移→line・円/構成比/割合/内訳→pie・棒→bar・手掛かりなし→None。
_CHART_KINDS = ("bar", "line", "pie")


_CHART_KIND_LABELS = {"bar": "棒", "line": "折れ線", "pie": "円"}


_CHART_KIND_WORDS = (
    ("line", ("折れ線", "推移")),
    ("pie", ("構成比", "割合", "内訳", "円")),
    # ★ 断片ガード①: 「棒」単独は「相棒」等の複合語と衝突するため、単独の「棒」でなく
    #   「棒グラフ」全体を語にする（凍結検体はどれも「棒グラフ」表記のみで単独「棒」を要らない）。
    ("bar", ("棒グラフ",)),
)


# ★ 断片ガード②: 「円」は単独では通貨表記（「500円」）と衝突するため、直前の文字が
#   数字（半角/全角）なら採用しない（extract_cmp_from_task の gte/lte 数字近傍ガードと同じ考え方）。
_CHART_KIND_YEN_GUARD = frozenset({"円"})


_CHART_KIND_NUM_RE = re.compile(r'[0-9０-９]')


def extract_chart_kind_from_task(task: str) -> str | None:
    """依頼文からグラフ種別を機械抽出する（extract_cmp_from_task と同じ作法）。
       一致が無ければ None（機械は断定しない）。複数の種別語が現れても、依頼文中で
       最初に出現した(かつ断片ガードを通った)ものを採る。"""
    if not task:
        return None
    best = None   # (出現位置, kind名)
    for kind_name, words in _CHART_KIND_WORDS:
        for w in words:
            idx = task.find(w)
            while idx >= 0:
                if (w in _CHART_KIND_YEN_GUARD and idx > 0
                        and _CHART_KIND_NUM_RE.match(task[idx - 1])):
                    idx = task.find(w, idx + 1)
                    continue
                if best is None or idx < best[0]:
                    best = (idx, kind_name)
                break
    return best[1] if best else None


def _column_values(meta: dict, sheet: str, col: str, limit: int = 500) -> list:
    """対象列の中身を実際に読む（宣言でなく実体を見る側の共通の入口）。
       読めない時は空リスト ── ここで例外を投げると無関係な入力まで巻き添えで壊れる。"""
    try:
        book_path = meta.get("path")
        headers = meta.get("headers", {}).get(sheet) or []
        if not book_path or col not in headers:
            return []
        col_idx1 = headers.index(col) + 1
        header_row = int(meta.get("header_rows", {}).get(sheet, 1))
        wb = openpyxl.load_workbook(book_path, read_only=True, data_only=True)
        try:
            ws = wb[sheet]
            return [row[0] for row in ws.iter_rows(
                min_row=header_row + 1, min_col=col_idx1, max_col=col_idx1,
                max_row=header_row + limit, values_only=True)]
        finally:
            wb.close()
    except Exception:
        return []


def _dedup_output_sheet_name(keys: list) -> str:
    """★ A'（EXTRACT の兄弟）: 出力シート名は機械が決め打ちで組む（LLM に名前を決めさせない）。
       例: keys=["取引先"] → 『取引先の重複除去』。禁止文字の置換・31文字上限は
       _extract_output_sheet_name と同じ規則（_EXTRACT_SHEET_NAME_FORBIDDEN_RE を共有）。"""
    label = "・".join(keys)
    name = f"{label}の重複除去"
    return _EXTRACT_SHEET_NAME_FORBIDDEN_RE.sub("_", name)[:31]


def duplicate_name_warning(col: str, values) -> str | None:
    """シート名の元になる列に同じ値が複数あるなら、その事実を名指しする 1 文（無ければ None）。

    ★ 実測（盲検の使い勝手レビュー・2026-08-24）: 取引先 3 社の売上表（4 行）に
      「取引先ごとに請求書を作って」と頼むと**請求書が 4 枚**でき、同じ取引先が
      『あかつき商事』(120,000) と『あかつき商事_2』(64,000) の 2 枚に分かれた。
      それでも「データ4行 → 出力4枚」で ✓ が出る。帳票段は **1 行 1 枚**の op なので
      機械としては正しいが、依頼者の「ごとに」は**まとめて 1 枚**を指している。
    ★ 機械は `_2` を付けたその瞬間に重複を知っている。知っていて黙るのが一番悪い。
      1 枚にまとめるかどうかは人が決めることなので、**名指しして人に返す**
      （★ 付き → 決裁③で ✓→△）。
    """
    seen = {}
    for v in values:
        if v is None or str(v).strip() == "":
            continue
        key = str(v).strip()
        seen[key] = seen.get(key, 0) + 1
    dupes = sorted(((k, n) for k, n in seen.items() if n > 1), key=lambda t: (-t[1], t[0]))
    if not dupes:
        return None
    head = "・".join(f"『{k}』{n} 行" for k, n in dupes[:3])
    more = f" ほか {len(dupes) - 3} 件" if len(dupes) > 3 else ""
    # ★ 印（⚠）は _warnings を印字する側が付ける ── ここで ★ を足すと「⚠ ★」と二重になる。
    # ★ 2026-08-28: 「先に集計してから」は**行き止まり**だった（この道具の中に道が無い）。
    #   まとめ方は在る ── 雛形に明細の印を置けば 1 枚にまとまる。出口を名指しする。
    #   ★ 断らないのは、取引ごとに 1 枚が正しい帳票（領収書・納品書）があるから。
    return (f"列『{col}』に同じ値が複数あります（{head}{more}）。"
             f"1 行につき 1 枚ずつ作るので、同じ相手の書類が別々の枚に分かれます"
             f"（1 枚にまとめるなら、雛形の明細行に『{{{{明細:列名}}}}』の印を、"
             f"合計欄に『{{{{合計:金額}}}}』の印を置いてください。"
             f"取引ごとに 1 枚が正しい書類 ── 領収書・納品書 ── ならこのままで大丈夫です）")


# ★ 2026-08-27: 入れ替えの依頼を見分ける（第二段翻訳へ回すための**証拠**であって、
#   名前をここから取るためのものではない ── 名前は LLM が言い、機械が実表で確かめる）。
_re_swap_ask = re.compile(
    r"(?:入れ?替え|入替|交換|逆に\s*し|前後を\s*入れ)")


# 依頼文が軸（行/列）をはっきり書いている場合だけ、その語を採る。
_re_swap_axis_row = re.compile(r"行\s*(?:同士\s*)?(?:を|の|で)?\s*[^。]{0,6}?(?:入れ?替え|入替|交換)")


_re_swap_axis_col = re.compile(r"列\s*(?:同士\s*)?(?:を|の|で)?\s*[^。]{0,6}?(?:入れ?替え|入替|交換)")


# 「AとBの〈列〉」── 2 つの行の、同じ 1 列を指す言い方（末尾の「の」まで含めて当てる）。
_re_between_and = re.compile(r"^.*?([^\s、。との]+?)\s*と\s*([^\s、。との]+?)\s*の$")


def swap_targets_are_cells(task: str, book_meta: dict, sheet: str | None,
                            header_row: int = 1) -> list | None:
    """依頼文が **2 つのセル**の入れ替えを指しているなら [(行,列), (行,列)]（でなければ None）。

    ★★ 2026-08-31（Namakoo が実測・✓ が出たのに操作が違った）:
      「丸和物流の**単価**とみどり建設の**単価**を入れ替えて」で、
      **行を丸ごと入れ替えて ✓ を出していた**（16 セルが動いた・頼んだのは 2 セル）。
      番人は「宣言どおり行が入れ替わったか」を見るので、**宣言そのものが違えば通る**
      ── 三項（依頼・宣言・実体）の「依頼」が抜けた形。
    ★ 依頼文には証拠が在る: **両側とも「〜の〈列名〉」と列を名指ししている**。
      行の入れ替えなら列は出てこない。
    ★ 言い回しを数え上げない ── 見るのは**実表に在る見出し**と**実表に在る行**だけ。
      「〈何か〉の〈実在する列〉」がちょうど 2 つ在り、その〈何か〉が行として一意に
      決まるときだけ、セルの入れ替えと読む。決まらなければ None（推測しない）。
    """
    text = _task_outside_quotes(task)
    headers = [str(h) for h in ((book_meta.get("headers") or {}).get(sheet) or [])]
    if not headers:
        return None
    found = []
    for h in sorted(headers, key=len, reverse=True):   # 長い見出しから（部分一致よけ）
        start = 0
        while True:
            i = text.find("の" + h, start)
            if i < 0:
                break
            start = i + 1
            name = text[:i].rsplit("と", 1)[-1].strip(" 　、。")
            for w in ("を", "は", "が", "の"):
                name = name.split(w)[-1] if name.endswith(w) else name
            if name:
                found.append((i, name, h))
    # ★★ 2026-08-31（Namakoo が実測・2 つ目の形）:
    #   「丸和物流**と**近江スチール**の項目**を入れ替えて」は、列名が **1 回しか出ない**。
    #   前の形（〈A〉の〈列〉と〈B〉の〈列〉）だけを見ていたので拾えず、
    #   行を丸ごと入れ替えていた（人は 2 セルのつもり）。
    #   ★ 「〈A〉と〈B〉の〈列〉」＝ **2 つの行の、同じ 1 列**。実表に照らして解ける
    #     ときだけ（列は実在の見出し・A と B は行として一意）── 推測しない。
    if len(found) == 1:
        _i, _name, _h = found[0]
        _pair = _re_between_and.match(text[:_i] + "の")
        if _pair:
            _rows = []
            for _nm in (_pair.group(1).strip(), _pair.group(2).strip()):
                _r, _ = _resolve_named_row(book_meta, sheet, _nm)
                if _r is None:
                    return None
                _rows.append(_r)
            if _rows[0] != _rows[1] and _h in headers:
                _c = headers.index(_h) + 1
                return [(_rows[0], _c), (_rows[1], _c)]
        return None
    if len(found) != 2:
        return None
    cells = []
    for _i, name, h in sorted(found):
        row, note = _resolve_named_row(book_meta, sheet, name)
        if row is None or h not in headers:
            return None
        cells.append((row, headers.index(h) + 1))
    return cells if cells[0] != cells[1] else None


def _swap_axis_hint(task: str) -> str | None:
    """依頼文が『行を』『列を』とはっきり書いているときだけ "row"/"column" を返す。
       ★ 両方書いてある／どちらも無いなら None ── 推測しない（機械が実表で決める側へ回す）。"""
    t = task or ""
    row = bool(_re_swap_axis_row.search(t))
    col = bool(_re_swap_axis_col.search(t))
    if row == col:
        return None
    return "row" if row else "column"


_re_row_unit = re.compile(r"[0-9０-９]*\s*行\s*(?:を|も)?\s*(?:足|追加|入れ|挿入)")


_re_row_word = re.compile(r"(?:第)?\s*[0-9０-９]{1,4}\s*行(?:目)?")


_re_a1_col_word = re.compile(r"[A-Za-z]{1,2}\s*列")


# ★ 助詞と語尾は**閉じた文法の集合**（業務語彙の列挙ではない）。落としても意味は減らない。
_TAIL_WORDS = ("にして", "にする", "に変えて", "に変える", "と入れて", "と書いて",
                "を入れて", "を書いて", "を追加して", "を追加", "を記入して", "を記入",
                "にセット", "入れて", "書いて", "変えて", "して", "ください", "です")


_PARTICLES = "をにへはとがのでも、。 　"


# ★ 末尾の「を＋動詞」（を作って／をつくる／を新設して…）。**語でなく形**で書く。
#   ★ 語尾（て・た・る…）は**必須**にする。省略可にすると「担当を佐藤」の
#     『を佐藤』まで食って、値そのものを消してしまった（実測）。
_re_verb_tail = re.compile(
    r"を[^\sをにへはがでとのも、。]{1,6}(?:て|た|る|ます|ください|下さい)\s*$")


def bare_value_from_task(task: str, row_name: str | None, col_name: str | None,
                          headers=None) -> str | None:
    """依頼文から、機械が**引き算で**書き込む値を切り出す。

    ★★ 2026-08-29（Namakoo が実測）: 「丸山重工の右にPCパーツ」で、第二段は
      row/col しか返さず **value を返さなかった**（qwen も gemma4 も）。
      LLM が値を出さないなら、機械が出す ── 機械は既に「誰の行か」「どの列か」を
      知っているので、依頼文からそれらを**引く**だけでいい。
    ★ 引くのは: 行の名前・列の名前・見出しの語・「N行目」「F列」・位置の語・助詞と語尾。
      どれも閉じた集合（業務語彙の列挙ではない）。
    ★ 残りが**依頼文の中に連続した文字列として在る**ことを最後に確かめる
      ── 切れ端を継ぎ足した幽霊の値を作らないため。
    """
    text = (task or "")
    out = text
    # ★ row_name は 1 つとは限らない（「AとBの間に」は目印が 2 つ）── 並びも受ける。
    names = list(row_name) if isinstance(row_name, (list, tuple)) else [row_name]
    for w in names + [col_name] + [str(h) for h in (headers or [])]:
        if w:
            out = out.replace(str(w), " ")
    out = _re_row_word.sub(" ", out)
    out = _re_a1_col_word.sub(" ", out)
    # ★ 位置の語（列の左右だけでなく、**行の上下と間**も落とす）。
    #   ★ 2026-08-29: 「味噌汁**の上に**新品を入れて」で『上に新品』が値になっていた
    #     ── 列の語だけ落として行の語を落としていなかった（また行と列の非対称）。
    for w in ("の右隣", "の左隣", "のとなり", "のあいだ", "の間",
               "の右", "の左", "の隣", "の上", "の下", "の前", "の後ろ", "列"):
        out = out.replace(w, " ")
    for w in _TAIL_WORDS:
        out = out.replace(w, " ")
    # ★★ 2026-08-30（1B の検体で 6 件・7B でも同じ形で落ちていた）:
    #   「ボルトとナットの間に**新品を作って**」で値が取れず、空行の挿入に落ちていた。
    #   「を追加して」は _TAIL_WORDS に在り、「を作って」は無い ── **列挙の穴そのもの**。
    #   ★ 動詞を数え上げると必ず漏れる（この repo が何度も踏んだ形）。
    #     語尾は**閉じた文法**なので、語ではなく**形**で書く:
    #       「を」＋（助詞を含まない短い語）＋（て／た／る／…）が末尾に付いていたら落とす。
    #   ★ セルに書く値は名詞なので、この形が値の一部になることはまずない。
    out = _re_verb_tail.sub(" ", out)
    #   ★ 「1行足して」のように**行そのものを足す**依頼は値ではない（凍結済みの
    #     述語を借りる ── 新しい語を数え上げない）。実測で『足』が値になりかけた。
    if _re_row_unit.search(task or ""):
        return None
    out = out.strip(_PARTICLES).strip()
    while out and out[0] in _PARTICLES:
        out = out[1:]
    while out and out[-1] in _PARTICLES:
        out = out[:-1]
    if not out or " " in out or "　" in out:
        return None                       # 2 つ以上に割れた ── 決めない
    if out not in text:
        return None                       # 連続していない ── 継ぎ足した値は使わない
    # ★★ 2026-08-29: 「スプリング**を作って**」がそのまま値になった。動詞の語尾を
    #   数え上げても必ず漏れる（今日 3 度目）── **文法の線**で弾く:
    #   セルに書く値の中に助詞は入らない。残っていたら、それは文がまだ切れていない証拠。
    if any(ch in out for ch in "をにへはがでとのも"):
        return None
    bad = {str(h) for h in (headers or [])} | {str(row_name or ""), str(col_name or "")}
    if out in bad:
        return None
    return out


# ★ 式を書き直す本数の上限。これを超えたら書き直さない（黙って諦めない・下で言う）。
FORMULA_REWRITE_LIMIT = 200


def formula_rewrites_for_shift(book_meta: dict, sheet: str | None, shift) -> tuple:
    """操作前の式を写像に通して、**操作後に在るべき式**を並べる。

    戻り値: (書き直し {(行, 列): 式}, 断りの理由 or None)

    ★★ 2026-08-29（Namakoo「合計行ごと参照を変えずに追記したいってことじゃないの？」）:
      「税込み金額と金額を入れ替えて」が × になった。断り自体は正しかった（実測で
      合計式が二列にまたがり、両方 1,000,440 ＝ 金額＋税込み金額 になっていた）。
      ★ だが利用者が欲しいのは**意味を保ったまま位置だけ入れ替わった表**で、
        それは機械が全部言える ── 操作前の式と π が分かっているのだから、
        操作後の式は π(操作前) でしかない。
      ★ だから「後から直す」のではなく **最初から正しく書く**。LibreOffice の
        自動付け替えを当てにしない（範囲の片側だけ動かすことがある・実測）。
      ★ 合っているセルに同じ内容を書き戻すのは無害 ── 見分けるより確実。
    ★ 消える参照が 1 つでもあれば**書き直さずに断る**（壊れた式を書かない）。
    ★ 本数が多すぎる回は書き直さない ── ただし黙らない（呼び出し側が理由を出す）。
    """
    path = book_meta.get("path")
    if not path or not sheet:
        return {}, None
    try:
        cells = {}
        with BookView(Path(path)) as bv:
            ws = bv.sheet(sheet)
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        cells[(cell.row, cell.column)] = v
    except Exception:
        return {}, None                    # 読めない回は何もしない（断定しない）
    if not cells:
        return {}, None
    if len(cells) > FORMULA_REWRITE_LIMIT:
        return {}, (f"式が {len(cells)} 個あるため、参照の書き直しは行いません"
                     f"（上限 {FORMULA_REWRITE_LIMIT} 個）")
    out, lost = cellmap.formulas_after(cells, shift)
    if lost:
        where = "、".join(f"{r}行{c}列" for r, c in lost[:3])
        return {}, f"この操作で参照が消える式があります（{where}）── 実行しません"
    return out, None


def formula_columns_to_inherit(book_meta: dict, sheet: str | None, header_row: int,
                                at: int, declared: set) -> tuple:
    """新しい行に式を引き継ぐ列（0 起点の列番号の一覧）と、**写す元の行**（1 起点）。

    ★★ 2026-09-02（README の「既知の問題」に自分で書いていた）:
      「みかんの下に梨を追加して」の後、梨の行の利益列は**空のまま**だった。
      宣言した値だけを書くので `✓` は正しいが、**人が期待するものとは違う**。
    ★ 式は発明ではない ── **隣の行から写す**（依頼文にも実表にも無い値は作らない＝A' 原則）。
      参照の付け替えは LibreOffice にやらせる（自分で式の文字列を書き換えると、
      それは 2 つ目の参照解決の実装になる ── SwapRowsByName が moveRange を使うのと同じ線）。

    ★ 引き継ぐ条件は **全データ行が式を持っていること**。これで合計列が自然に外れる:
      金額列は E2..E7 が直値・E8 だけ =SUM なので「全部が式」ではない。
      逆に 税込金額 は全行 =E*1.1 なので引き継ぐ。**列挙ではなく形で決める。**
    ★ 合計行は写す元にしない（=SUM を新しい行に配ると壊れる）。判定は既存の
      凍結規則（total_rows_in）を借りる ── 同じことを 2 箇所で決めない。
    ★ 人が値を指定した列は触らない（**人の指定が勝つ**）。

    返り値: (0 起点の列番号のリスト, 写す元の行番号) / 引き継ぐものが無ければ ([], 0)
    """
    path = (book_meta or {}).get("path")
    if not path:
        return [], 0
    totals = set(total_rows_in(book_meta, sheet, header_row))
    try:
        with BookView(Path(path)) as bv:
            ws = bv.sheet(sheet)
            last, wide = data_extent(ws, header_row)
            rows = [r for r in range(header_row + 1, last + 1) if r not in totals]
            if not rows or wide < 1:
                return [], 0
            heads = [str(ws.cell(row=header_row, column=c).value or "")
                      for c in range(1, wide + 1)]
            cols = []
            for c in range(1, wide + 1):
                if heads[c - 1] in declared:
                    continue                       # ★ 人が指定した列は触らない
                if all(bv.cell_formula(r, c, sheet) is not None for r in rows):
                    cols.append(c - 1)
    except Exception:
        return [], 0                               # 読めない回は黙る（断定しない）
    if not cols:
        return [], 0
    # ★ 写す元は「新しい行のすぐ上のデータ行」。無ければ下から取る。
    above = [r for r in rows if r < at]
    if above:
        return cols, max(above)
    below = [r for r in rows if r >= at]
    if not below:
        return [], 0
    return cols, min(below) + 1                    # ★ 挿入で 1 行ずれた後の位置


def reference_drift_warning(book_meta: dict, sheet: str | None, *,
                             row_lo: int = 1, row_hi: int = 10 ** 7,
                             col_lo: int = 1, col_hi: int = 10 ** 4,
                             rewritten=None, unit: str = "行") -> str | None:
    """動かす区画を**外から**指している式を見つけ、1 行にして返す（無ければ None）。

    ★★ 2026-08-29（Namakoo の指摘 → 実測で裏取り）:
      並べ替えると、範囲の外から特定の 1 行を指している式は**追従しない**。
      実測: `=B3`（ラベルは「ぶどうの金額」）が、並べ替え後に みかん の 200 を指した。
      別シートからの `=売上!B2` も同じ（りんご 100 → ぶどう 300）。
      ★ **式は 1 文字も壊れていない**ので、値でも文字列でも捕まらない ── 参照を読むしかない。
      ★ そして ailine は ✓ を出していた（並べ替え自体は宣言どおりだから）。
        「静かに壊れて合格が出る」── この製品が一番嫌う形に、ぴったり当てはまっていた。
    ★ 直さない: Excel も LibreOffice も、範囲の外から特定の行を指す式は並べ替えで
      追従させない（アドレスに留まるのが既定の意味）。
      「ぶどうの金額 = B3」は行に追従してほしいが「3行目の値 = B3」は留まってほしい
      ── 機械には区別できない。**名指しして人に返す**（補正は人が決めてから）。
    ★ 範囲（SUM(B2:B4)）は鳴らさない ── そちらは領域を指すので正しく追従する。
    """
    path = book_meta.get("path")
    if not path or not sheet:
        return None
    try:
        hits = cellmap.refs_pointing_into(Path(path), sheet, row_lo, row_hi, col_lo, col_hi)
    except Exception:
        return None                      # 読めない回は黙る（断定しない）
    # ★ 2026-08-31: **こちらで書き直す式**は、もう「ずれる式」ではないので名指ししない
    #   （別シートから指している式は書き直していないので、そちらは残す）。
    return cellmap.reference_drift_note(
        cellmap.drop_rewritten(hits, rewritten, sheet), unit=unit)


_re_bare_number = re.compile(r"([0-9][0-9,]*(?:\.[0-9]+)?)\s*(?:に|へ|と)?\s*(?:し|する|して|に)")


# 依頼文が「列を追加/足す/挿入」と言っているか（第二段へ回すための証拠）。
# ★ 2026-08-27（自分で開けた穴・実機の検体が捕まえた）: 「列を**入れ替え**て」が
#   「入れ」に当たって列追加として横取りされ、入れ替えが動かなくなった。
#   ★ 語の一部が別の語の一部でありうる ── 部分文字列の穴は、この repo で 2 度目。
# ★ 2026-08-27（Namakoo「◎を入れて では動作しない」）: 「列**に**『◎』を入れて」まで
#   列追加として拾っていた。★ 列の直後の助詞で分かれる ── 「列**を**追加/入れる」は
#   列そのものが対象、「列**に**…を入れる」は列が**行き先**。助詞は意味を運んでいる。
_re_add_col_ask = re.compile(r"列\s*(?:を|の)\s*[^。]{0,4}?(?:追加|足し|足す|挿入|入れ(?!替)|作)")


def task_asks_to_add_a_column(task: str) -> bool:
    """依頼文が「列を追加」を求めているか（位置も名前もここでは決めない）。"""
    return bool(_re_add_col_ask.search(task or ""))


def total_rows_in(book_meta: dict, sheet: str | None, header_row: int = 1) -> list:
    """データ行ではない「合計行」の行番号（1 起点）。

    ★★ 2026-08-28（Namakoo が請求書のデモで実測）: 「金額が10万以上の行に印を付けて」が
      **合計行にも印を付けた**。条件としては真だが、合計行は請求の行ではない ── 意味が違う。
    ★ 判定は既存の凍結規則を借りる（ailine_core.total_row.row_has_total_word:
      合計/小計/総計は部分一致・『計』は完全一致・『設計部』等は誤爆しない断片ガードつき）。
      ここで新しい規則を書かない ── 同じことを 2 箇所が別々に決めると必ずずれる。
    ★ 見つけたら**必ず画面に出す**（黙って行を外さない）。
    """
    path = book_meta.get("path")
    if not path:
        return []
    try:
        with BookView(Path(path)) as bv:
            ws = bv.sheet(sheet)
            last, cols = data_extent(ws, header_row)
            out = []
            for r in range(header_row + 1, last + 1):
                vals = [ws.cell(row=r, column=c).value for c in range(1, cols + 1)]
                if total_row.row_has_total_word(vals):
                    out.append(r)
    except Exception:
        return []
    return out


# 「〜の右に」等の位置の言い回し（列版・_COL_AFTER/_COL_BEFORE と同じ語彙）。
_re_after_position = re.compile("(?:" + "|".join(
    re.escape(w) for w in (*_COL_AFTER, *_COL_BEFORE)) + ")")


# 名前のうしろに付く「列を追加して」等。★ 語尾は閉じた文法の集合（業務語彙ではない）。
# 修飾節の終わり ── 助詞、または活用語尾（引い**た** / 掛け**て** / 足し**た**）。
# ★ 「の」は入れない（「粗利の列」の の は名前の側）。★ 業務語彙ではなく閉じた文法。
_re_clause_end = re.compile(r"[をにへはがでとも]|[ぁ-ん](?:た|て|だ)")


# ★ 位置語が無い回の入口 ── 「**作る**」と言っている語尾だけ（裸の「列」は入れない）。
_NEW_COL_MAKE_TAILS = ("という列を追加して", "という列を作って", "という列を追加",
                        "という列を作る", "の列を追加して", "の列を作って",
                        "の列を追加", "の列を作る", "列を追加して", "列を作って",
                        "列を追加", "列を作る")


_NEW_COL_TAILS = ("という列を追加して", "という列を作って", "という列を追加", "という列を作る",
                   "という列", "の列を追加して", "の列を作って", "の列を追加", "の列を作る",
                   "列を追加して", "列を作って", "列を追加", "列を作る", "列",
                   "を追加して", "を作って", "を追加", "を作る", "を入れて", "を足して")


def new_column_name_from_task(task: str, headers=None, *,
                               require_position: bool = True) -> str | None:
    """依頼文が名指ししている、**新しい列の名前**（決まらなければ None）。

    ★★ 2026-08-30（Namakoo が実測・下書きに 2 本できた）:
      「金額の右に税込み金額を追加」を 2 回頼んで、見出しが
        1 回目「税込金額」（**「み」が落ちた**）／2 回目「金額*1.1」（**式が名前になった**）
      になった。前者は道具が `f"税込{列名}"` と**作った**名前、後者は式そのもの。
      ★ どちらも A' 原則（値も名前も依頼文から取る）が抜けていた ── **人が書いた
        名前がそこに在るのに、機械が別の名前を発明していた**。
      ★ しかも解釈行に名前が出ていなかったので、間違いに気づく手がかりが無かった。
    ★ 引き算で切り出す: 位置の言い回し（「〜の右に」）の**うしろ**から、語尾を落とす。
      全体を置換しない ── 「金額の右に税込み金額」で『金額』を全部消すと『税込み』になる。
    ★ 実在する見出しと同じ名前なら None（それは新しい列ではない）。
    """
    text = _task_outside_quotes(task).replace(chr(12288), " ")
    m = None
    for m2 in _re_after_position.finditer(text):
        m = m2                       # 最後の位置語のうしろを見る（「AとBの右に X」）
    if m:
        name = text[m.end():]
    elif require_position:
        # ★★ 2026-09-02: 既定は**位置語が在る時だけ**（従来どおり）。
        #   実測で分かったこと: 位置語なしを無条件に許すと、税の枝（W10c で設計）まで
        #   書き換わる ──「税込みの列を追加して」で見出しが『税込み』になった。
        #   『税込み』は名前ではなく**修飾語**で、機械が組む『税込金額』のほうが良い。
        #   ★ 測っていない所まで直しを広げない（断る範囲を広げるのと同じ失敗）。
        #     欠けていたのは**2 項の演算**の枝だけなので、そこだけ明示的に呼ぶ。
        return None
    else:
        # ★★ 2026-09-02（A の確認中に見つけた）: 位置を言わない依頼では、ここに
        #   入る前に空文字になっていた ── 「売上から原価を引いた**利益**の列を作って」で
        #   見出しが『売上-原価』（式そのもの）になっていた。A' 原則が抜けた形。
        #   ★ 位置語が無い時は、**語尾の手前まで**を候補にして、その中の
        #     **修飾節の終わり**から始める（引き算は位置語の時と同じ考え方）。
        #     節の終わり = 助詞（を に へ は が で と も）か、活用語尾（〜た/て/だ）。
        #     ★ 「の」は**入れない** ── 「粗利の列」の の は名前側に属する。
        #   ★★ 初版は語尾を先に切らずに走査したので、「作**って**」自身が節の終わりに
        #     当たり、名前ごと飲み込んで空になっていた（実測で捕まえた）。
        # ★★ 実測で捕まえた誤爆（既存の検体が赤くした）: 語尾に **裸の「列」** を
        #   許すと、「A行G列を『税込み金額』に上書き」で『A行G』を新しい列の名前として
        #   拾った。★ 位置語が無い回は、**作る**と言っている語尾だけを入口にする
        #   （裸の「列」は「〜の右に 利益列」のような位置語つきの回のためのもの）。
        _end = min((text.find(w) for w in _NEW_COL_MAKE_TAILS if text.find(w) > 0),
                    default=-1)
        if _end < 0:
            return None                  # 「〜の列を作って」の形 が無い＝名指しでない
        _head = text[:_end]
        _cut = 0
        for _mb in _re_clause_end.finditer(_head):
            _cut = _mb.end()
        name = _head[_cut:]
    if not name.strip():
        return None
    for w in _NEW_COL_TAILS:         # 長い語尾から落とす（並びが長さ順）
        i = name.find(w)
        if i > 0:
            name = name[:i]
            break
    name = name.strip().strip("、。 ")
    while name and name[0] in "をにへはがでとのも 　":
        name = name[1:]
    if len(name) < 2 or " " in name:
        return None
    if any(ch in name for ch in "をにへはがでとも"):
        return None                  # 文がまだ切れていない（助詞が残っている）
    if name in {str(h) for h in (headers or [])}:
        return None                  # 既にある列 ── 新しい名前ではない
    return name if name in text else None


def _column_existing_value_count(book_path: Path, sheet_name: str, col_name: str,
                                  header_row: int = 1) -> int:
    """★ M2c / W10a: target(既存列指定)列に、見出し行を除いて値が入っているセルの件数。
       上書き検知の明示（確認行の注意書き）と W10a の破壊の関所（確認メッセージの件数）が
       共有する。読めない/列やシートが見つからない場合は 0（保守的に『無い』扱い＝誤って
       止めない）。★ W3: header_row(1起点)で見出しの実位置を受け取る
       （省略時は物理1行目・旧挙動と同一）。"""
    try:
        wb = openpyxl.load_workbook(book_path, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return 0
        ws = wb[sheet_name]
        idx = _col_index_by_header(ws, col_name, header_row=header_row)
        if idx is None:
            wb.close()
            return 0
        last = _scan_last_row(ws, header_row=header_row)
        count = sum(1 for r in range(header_row + 1, last + 1)
                    if ws.cell(row=r, column=idx).value not in (None, ""))
        wb.close()
        return count
    except Exception:
        return 0


def _rows_existing_value_count(book_path, sheet_name: str, at: int, count: int = 1) -> int:
    """消そうとしている行に、値の入っているセルが何個あるか（削除の関所の件数）。

    ★★ 2026-09-17（盲検 3 体目）: 列の上書きには関所が在り、**削除には無かった**。
      買い手の言葉:「上書きより削除の方が怖いのに、厳しい方が緩い」。
    ★ 列側（_column_existing_value_count）と**同じ作法**で書く ── 読めなければ 0 を返し、
      誤って止めない。鳴るのは「消えるものが在る」と数えられた時だけ。
    ★ 数えるのは**値の個数**であって行数ではない ── 事後の表示が「消した中身（15 行）」と
      空行まで数えて出していたのを、関所の側では繰り返さない。
    """
    if not book_path or not sheet_name:
        return 0
    try:
        wb = openpyxl.load_workbook(book_path, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return 0
        ws = wb[sheet_name]
        n = 0
        for r in range(int(at), int(at) + max(1, int(count))):
            for c in range(1, (ws.max_column or 0) + 1):
                if ws.cell(row=r, column=c).value not in (None, ""):
                    n += 1
        wb.close()
        return n
    except Exception:
        return 0
